#!/usr/bin/env python3
"""Build a teacher-facing brief from immutable project results; no model calls."""
from pathlib import Path
import csv
import hashlib
import html
import json
import math
import re
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/teacher_brief_20261027'
ASSETS = OUT/'assets'
sys.path.insert(0, str(ROOT/'.cache/teacher_reporting_vendor'))
sys.path.insert(0, str(ROOT/'.tools/pdf_bundle'))
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from fontTools.ttLib import TTFont as FontFile
from fontTools.varLib.instancer import instantiateVariableFont
from fontTools import subset
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.colors import HexColor

NAVY = '#122D45'
TEAL = '#087E8B'
GOLD = '#D59A42'
RED = '#B74F45'
GREEN = '#28795E'
INK = '#20374B'
MUTED = '#647887'
BG = '#F4F7F9'
LINE = '#DCE5EB'
PALE = '#E7F2F2'
ARMS = ['SFT_clean', 'SFT_aug', 'DPO_clean', 'DPO_aug']
SOURCES = {}


def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()


def readcsv(relative):
    p = ROOT/'reports'/relative
    SOURCES[str(p.relative_to(ROOT))] = sha(p)
    with p.open(encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        for k, v in r.items():
            if v == '': r[k] = None
            elif v in ['True', 'False']: r[k] = v == 'True'
            else:
                try:
                    x = float(v)
                    r[k] = int(x) if x.is_integer() else x
                except (ValueError, TypeError): pass
    return rows


def pct(x): return f'{x*100:.2f}%'
def pp(x): return f'{x*100:+.2f}'
def ci(r, a='low', b='high'): return f'[{r[a]*100:+.2f}, {r[b]*100:+.2f}]'


def data():
    d = {}
    for key, path in {
        'T':'week_final_20261026/metrics.csv',
        'Tci':'week_final_20261026/paired_G_differences.csv',
        'Gtrain':'week_final_20261026/training_lineage.csv',
        'rationale':'week_final_20261026/rationale_metrics.csv',
        'control':'random_control_20261024/metrics.csv',
        'control_ci':'random_control_20261024/paired_deltas.csv',
        'prefix':'prefix_full_panel_20261026/net_rates.csv',
        'reverse':'prefix_full_panel_20261026/reverse_removal.csv',
        'cross':'cross_round_prefix_20261025/flip_rates.csv',
        'cross_mean':'cross_round_prefix_20261025/uniform_trigger_means.csv',
        'gaps':'closing_diagnostics_20261025/symmetric_round_gaps.csv',
        'chain':'closing_diagnostics_20261025/chain_question_matched.csv',
        'Vmetrics':'verifier_robustness_clean_20261026/metrics.csv',
        'repair':'verifier_robustness_clean_20261026/repair_damage.csv',
        'contrast':'verifier_robustness_clean_20261026/contrast_repair_damage.csv',
        'stability':'verifier_robustness_clean_20261026/prefix_stability.csv',
    }.items(): d[key] = readcsv(path)
    d['T0'] = [r for r in d['T'] if r['V']=='V0']
    d['control'] = [r for r in d['control'] if r['scenario']=='C95_exclude_reference_suspect']
    d['E'] = [r for r in d['Vmetrics'] if r['cohort']=='E_new_final']
    d['Erepair'] = [r for r in d['repair'] if r['cohort']=='E_new_final']
    d['main'] = [r for r in d['Erepair'] if r['view']=='G3_vLLM_newE' and r['variant']=='we_can']
    d['Estability'] = [r for r in d['stability'] if r['cohort']=='E_new_final' and r['view']=='G3_vLLM_newE' and r['variant']=='we_can']
    assert len(d['T'])==12 and len(d['T0'])==4 and len(d['prefix'])==8 and len(d['main'])==4
    assert all(r['TP']+r['FN']+r['FP']+r['TN']+r['unknown']==1196 for r in d['T'])
    for r in d['T']:
        assert abs(r['FPR']-r['FP']/(r['FP']+r['TN']))<1e-12
        assert abs(r['FNR']-r['FN']/(r['FN']+r['TP']))<1e-12
    for r in d['Erepair']:
        assert r['net_repaired']==r['repaired_FP']+r['repaired_FN']-r['new_FP']-r['new_FN']
    assert [(r['repaired_FP'],r['repaired_FN'],r['new_FP'],r['new_FN']) for r in d['main']]==[(16,0,6,11),(5,1,23,6),(13,0,0,8),(12,0,0,7)]
    return d


def fonts():
    cache = ROOT/'.cache/teacher_reporting_fonts'
    # Keep portable embedded subsets small; all summary and deck text is included.
    text = (OUT/'研究总结.md').read_text()+Path(__file__).read_text()
    for p in (ROOT/'reports').glob('**/*.md'):
        if 'prefix_sources' not in str(p): text += p.read_text()
    text += ''.join(chr(i) for i in range(32,127))+'−→←Δβαμ≥≤×Σ'
    character_key=hashlib.sha256(''.join(sorted(set(text))).encode()).hexdigest()
    paths=[]
    for weight, name in [(400,'BriefRegular'),(700,'BriefBold')]:
        path=cache/(name+'.ttf')
        marker=cache/(name+'.characters.sha256')
        if not path.exists() or not marker.exists() or marker.read_text()!=character_key:
            f=FontFile(cache/'NotoSansSC.ttf')
            instantiateVariableFont(f, {'wght':weight}, inplace=True)
            opts=subset.Options();opts.name_IDs=['*'];opts.layout_features=['*']
            sub=subset.Subsetter(options=opts);sub.populate(text=text);sub.subset(f);f.save(path)
            marker.write_text(character_key)
        pdfmetrics.registerFont(TTFont(name,str(path)))
        paths.append(path)
    pdfmetrics.registerFontFamily('BriefRegular',normal='BriefRegular',bold='BriefBold',italic='BriefRegular',boldItalic='BriefBold')
    plt.rcParams.update({'font.family': FontProperties(fname=paths[0]).get_name(), 'font.size':11,
                         'axes.spines.top':False,'axes.spines.right':False,'axes.edgecolor':LINE,
                         'axes.labelcolor':INK,'text.color':INK,'xtick.color':MUTED,'ytick.color':MUTED,
                         'figure.facecolor':'white','axes.facecolor':'white','axes.unicode_minus':False})
    from matplotlib import font_manager
    for p in paths: font_manager.fontManager.addfont(str(p))
    return paths


def savefig(fig, name):
    fig.savefig(ASSETS/(name+'.png'),dpi=180,bbox_inches='tight',facecolor='white')
    plt.close(fig)


def figures(d):
    from matplotlib.patches import FancyBboxPatch
    fig,ax=plt.subplots(figsize=(11.8,2.1));ax.set(xlim=(0,12),ylim=(0,2));ax.axis('off')
    cards=[('生成器 Gt','A池每题生成4条'),('固定验证器 V0','评分与筛选'),('SFT 更新 Gt+1','每题至多1条'),('统一评价','C开发 / T留出')]
    for i,(title,sub) in enumerate(cards):
        x=.15+i*3
        ax.add_patch(FancyBboxPatch((x,.43),2.55,1.2,boxstyle='round,pad=.08',facecolor=PALE if i<3 else BG,edgecolor=LINE))
        ax.text(x+1.275,1.18,title,ha='center',fontsize=13,color=NAVY,weight='bold')
        ax.text(x+1.275,.78,sub,ha='center',fontsize=10,color=MUTED)
        if i<3: ax.annotate('',xy=(x+2.87,1.03),xytext=(x+2.65,1.03),arrowprops={'arrowstyle':'->','color':TEAL,'lw':2})
    savefig(fig,'research_flow')
    fig,axs=plt.subplots(1,2,figsize=(11.4,3.5),layout='constrained')
    x=np.arange(4)
    for key,label,color in [('FPR','误接受 FPR',RED),('FNR','误拒 FNR',GOLD)]:
        y=np.array([r[key]*100 for r in d['T0']]);axs[0].plot(x,y,'o-',lw=2.4,color=color,label=label)
        for a,b in zip(x,y):axs[0].annotate(f'{b:.2f}',(a,b),xytext=(0,8),textcoords='offset points',ha='center',fontsize=10)
    axs[0].set(xticks=x,xticklabels=['G0','G1','G2','G3*'],ylim=(0,61),ylabel='错误率 / %');axs[0].legend(frameon=False,ncol=2,loc='upper left')
    y=np.array([r['balanced_accuracy']*100 for r in d['T0']]);axs[1].plot(x,y,'o-',lw=2.5,color=TEAL)
    for a,b in zip(x,y):axs[1].annotate(f'{b:.2f}%',(a,b),xytext=(0,9),textcoords='offset points',ha='center',fontsize=11)
    axs[1].set(xticks=x,xticklabels=['G0','G1','G2','G3*'],ylim=(64,77),ylabel='平衡准确率 BA / %')
    for ax in axs:ax.grid(axis='y',color=LINE,alpha=.6)
    savefig(fig,'T_cross_round')
    rr=d['prefix'];labels=['原文','中性控制','we need to','we can','we will','it is essential to','逗号+换行','reason step by step']
    fig,axs=plt.subplots(1,2,figsize=(11.6,4.3),layout='constrained')
    for ax,key,color in zip(axs,['FPR','FNR'],[RED,TEAL]):
        vals=[r[key]*100 for r in rr];ax.barh(np.arange(8),vals,color=[NAVY if i==0 else color for i in range(8)],alpha=.9)
        for i,v in enumerate(vals):ax.text(v+.5,i,f'{v:.2f}%',va='center',fontsize=9)
        ax.set(yticks=np.arange(8),yticklabels=labels,xlim=(0,77 if key=='FPR' else 23),xlabel=('误接受率 FPR' if key=='FPR' else '误拒率 FNR')+' / %');ax.invert_yaxis()
    savefig(fig,'prefix_full_panel')
    rr=[r for r in d['gaps'] if r['split']=='A' and r['metric']=='FPR']
    fig,ax=plt.subplots(figsize=(8.5,3.5),layout='constrained')
    for i,r in enumerate(rr):
        val=r['G3_minus_G2']*100
        ax.barh(i,val,color=[NAVY,TEAL,GOLD][i],height=.55)
        ax.text(val+.12,i,f'{val:.2f} pp',va='center',fontsize=13)
    ax.set(yticks=range(3),yticklabels=['自然原文','删除自然开场','自然近义替换'],xlim=(0,9.2),xlabel='A池 G3−G2 的FPR差 / 百分点');ax.invert_yaxis()
    savefig(fig,'natural_gap')
    fig,ax=plt.subplots(figsize=(10.5,3.7),layout='constrained')
    xs=np.arange(4)
    for j,(field,label,color) in enumerate([('repaired_FP','修复 FP',GREEN),('repaired_FN','修复 FN','#8AB8A2'),('new_FP','新增 FP',RED),('new_FN','新增 FN',GOLD)]):
        vals=[r[field] for r in d['main']];bars=ax.bar(xs+(j-1.5)*.18,vals,width=.18,color=color,label=label)
        for b,v in zip(bars,vals):
            if v:ax.text(b.get_x()+.09,v+.35,str(v),ha='center',fontsize=11)
    ax.set(xticks=xs,xticklabels=ARMS,ylim=(0,29),ylabel='回答条数');ax.legend(ncol=4,frameon=False,loc='upper left');ax.grid(axis='y',color=LINE,alpha=.4);ax.set_axisbelow(True)
    savefig(fig,'repair_damage')
    fig,ax=plt.subplots(figsize=(9.6,3.3),layout='constrained')
    for split,color in [('A',TEAL),('C',GOLD)]:
        rr=[r for r in d['cross_mean'] if r['split']==split]
        ax.plot(range(4),[r['mean_flip_rate']*100 for r in rr],'o-',color=color,label=f'{split}池',lw=2.2)
    ax.set(xticks=range(4),xticklabels=['G0','G1','G2','G3'],ylim=(0,17),ylabel='原TN注入翻转率 / %');ax.legend(frameon=False);ax.grid(axis='y',color=LINE,alpha=.5)
    ax.axvline(1.5,color=MUTED,ls='--',lw=1);ax.text(1.55,1.4,'生成引擎切换',fontsize=9,color=MUTED)
    savefig(fig,'cross_prefix')


FLOWS=[
 ('Day 1','环境与初测','模型部署、镜像测速、生成/评分/训练冒烟检查'),
 ('Day 2','解析与任务诊断','GSM8K / MATH自然错误覆盖比较'),
 ('Day 3','冻结DAPO协议','坏题、长度、标签、题池与工作点'),
 ('Day 4—5','固定V0更新G','G0→G1→G2；G3为后续扩展'),
 ('配套对照','监督质量与缓存','Vmatched / Oracle / Random、逐题监督权重'),
 ('Day 6','更新V的原对照','Vrand / Vtarget，记录接受偏移与负结果'),
 ('Day 7','T留出评价','T299统一引擎、混淆矩阵、区间与案例'),
 ('后续补充','前缀诊断与修复','内容保持干预、机制对照、SFT × DPO'),
]
POOLS=[('A',600,600,'G反馈训练题池','生成4候选，V0筛选后每题至多1条更新G'),('B',200,200,'V训练题池','G监督选样不使用B回答；新修复训练取其中64题'),('C',100,95,'开发与前缀诊断','初排4题后96，后续再隔离1道参考疑问题，主分析95'),('T',300,299,'七天流程最终留出','所有G/V整题排除1道缺图题；每G1196回答'),('E',64,60,'新V修复评价','所有G/V/前缀整题隔离4题；每G240回答')]
MECHANISMS=[
 ('学术性','20对AWL/表外词；A64、C48','A 2.11% vs 2.81%；C 2.81% vs 3.54%','未显示AWL优势','awl_comparison_20261024/AWL_comparison_report.md'),
 ('长度','38前缀；字符/词数/token','字符Spearman A−0.414、C−0.470；token−0.121、−0.202','表述与格式仍是重要因素','prefix_length_20261025/prefix_length_report.md'),
 ('题型与符号','输入特征分层','多组风险上升，组间交互未通过校正','未定位独有题型退化','fpr_patterns_20261024/EFP_SITE_report.md'),
 ('训练链富集','匹配511道实际SFT题','we need to 15.51%→15.85%→16.24%→14.63%','未见一致放大','closing_diagnostics_20261025/supplementary_evidence_report.md'),
 ('自然开场贡献','G2/G3对称编辑，A598','FPR差7.69pp→删除7.12pp / 替换6.90pp','主要差距仍保留','closing_diagnostics_20261025/supplementary_evidence_report.md'),
 ('跨G注入变化','同A64/C29，四G，各1条原TN','A 5.99/8.85/9.03/8.68%；C 5.36/12.84/11.49/12.84%','有阶段差异，未见一致单调性','cross_round_prefix_20261025/cross_round_prefix_report.md'),
 ('自然删除','741条原判对的自然TP/TN','A删除新增FN5/340、FP5/306；C为1/34、1/61','自然表达编辑也有局部损伤','natural_invariance_20261024/natural_invariance_report.md'),
]
LITERATURE=[
 ('R1','Zheng et al. (2023)','Judging LLM-as-a-Judge with MT-Bench and Chatbot Arena','评审偏差与推理局限的研究背景','https://arxiv.org/abs/2306.05685'),
 ('R2','AdvJudge-Zero v3','Binary Decision Flips in LLM-as-a-Judge via Adversarial Control Tokens','短控制字符串、监督加固；本项目修复数据思路的主要借鉴','https://arxiv.org/html/2512.17375v3'),
 ('R3','Yang et al. (2026)','Turning Bias into Bugs: Bandit-Guided Style Manipulation Attacks on LLM Judges','语义保持风格编辑与评审操控的相关工作','https://arxiv.org/abs/2605.26156'),
 ('R4','Rafailov et al. (2023)','Direct Preference Optimization: Your Language Model is Secretly a Reward Model','本项目DPO分支采用的标准参考相对目标','https://arxiv.org/abs/2305.18290'),
]


def workbook(d):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.chart import LineChart,BarChart,Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb=Workbook();nav=wb.active;nav.title='00_阅读导航'
    nav.append(['阶段研究结果总览','2026-10-27｜10.20—10.27']);nav.append(['研究主题','迭代数学数据中的验证器可靠性、前缀敏感性与修复'])
    nav.append(['主结论1','T299：G0→G2，固定V0 BA下降3.24个百分点'])
    nav.append(['主结论2','C95/G3：we can使FPR+10.55pp、FNR−6.12pp'])
    nav.append(['主结论3','E60/G3/we can：DPO两组净修复各5条，同时新增FN；增强未显示额外净收益'])
    nav.append(['数字说明','率为数值百分比格式；Δ及区间标题注明“百分点”时按100倍导出'])
    nav.append(['标签与区间','独立最终答案核验；未知单列；同题多回答按题配对重采样'])
    nav.append(['阅读顺序','T跨轮 → 前缀面板 → 机制对照 → 修复新增 → 直接增强对照'])
    nav.append(['工作表','内容与口径'])
    table_n=0
    def sheet(name,headers,rows,notes='',pctcols=(),ppcols=()):
        nonlocal table_n
        ws=wb.create_sheet(name);ws.append([notes]);ws.append(headers)
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
        for row in rows:
            vals=[row.get(k) for k in headers] if isinstance(row,dict) else list(row)
            ws.append(vals)
        ws.freeze_panes='C3';ws.auto_filter.ref=f'A2:{get_column_letter(len(headers))}{ws.max_row}'
        if rows:
            table_n+=1;t=Table(displayName=f'BriefTable{table_n}',ref=ws.auto_filter.ref)
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
        for j,h in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(j)].width=min(48,max(13,len(h)*1.35+3))
        for j in pctcols:
            for col in ws.iter_cols(min_col=j,max_col=j,min_row=3):
                for c in col:c.number_format='0.00%'
        for j in ppcols:
            for col in ws.iter_cols(min_col=j,max_col=j,min_row=3):
                for c in col:c.number_format='+0.00;-0.00;0.00'
        ws.sheet_view.showGridLines=False
        nav.append([name,notes]);nav.cell(nav.max_row,1).hyperlink=f"#'{name}'!A1";nav.cell(nav.max_row,1).style='Hyperlink'
        return ws
    sheet('01_七天与后续流程',['阶段','主题','完成工作'],FLOWS,'原七天计划的完成节点；G3和前缀修复为授权扩展')
    sheet('02_题池与指标',['题池','原始题数','主要统计题数','用途','说明'],POOLS,'A/B/C/T按题及可识别来源分离；E为新修复题池')
    wsT=sheet('03_T固定V0',['G','题目','正确','错误','未知','G正确率','FP','TN','FN','TP','FPR','FNR','BA','未知率','AUROC'],
       [[r['G'],299,r['correct'],r['wrong'],r['unknown'],r['G_accuracy_verified'],r['FP'],r['TN'],r['FN'],r['TP'],r['FPR'],r['FNR'],r['balanced_accuracy'],r['unknown_fraction'],r['AUROC']] for r in d['T0']],
       'T299，每G1196回答；四个G统一vLLM，V0原评分与0.5阈值',pctcols=[6,11,12,13,14])
    sheet('04_T跨轮配对区间',['比较','指标','差值_百分点','95%下限_百分点','95%上限_百分点'],
       [[r['contrast'],r['metric'],r['delta']*100,r['low']*100,r['high']*100] for r in d['Tci'] if r['V']=='V0' and r['metric'] in ['FPR','FNR','balanced_accuracy']],
       'T299按题配对bootstrap10000；包含G3扩展，单训练轨迹',ppcols=[3,4,5])
    sheet('05_G实际训练',['模型','父模型','回答','步数','回答loss_token','正确','错误','未知','有标签错误率'],
       [[r['model'],r['parent'],r['examples'],r['steps'],r['loss_tokens'],r['selected_correct_by_original_reference'],r['selected_wrong_by_original_reference'],r['selected_unknown'],r['selected_verified_error_fraction']] for r in d['Gtrain'] if r['model'].startswith('G')],
       '主路线A选样，独立标签仅记录质量；G3单列扩展',pctcols=[9])
    wsC=sheet('06_G监督质量对照',['模型','G正确率','正确','错误','未知','FPR','FNR','BA','监督token'],
       [[r['model'],r['G_accuracy_verified'],r['correct'],r['wrong'],r['unknown'],r['FPR'],r['FNR'],r['balanced_accuracy'],{'G0':0,'Random_SFT':224951,'Vmatched_SFT':225565,'Oracle_SFT':225565}[r['model']]] for r in d['control']],
       'C95/380回答；三个训练分支共同289道A题，37步，原Transformers协议',pctcols=[2,6,7,8])
    sheet('07_G对照配对区间',['比较','指标','差值_百分点','95%下限_百分点','95%上限_百分点'],
       [[r['contrast'],r['metric'],r['estimate']*100,r['low']*100,r['high']*100] for r in d['control_ci'] if r['scenario']=='C95_exclude_reference_suspect' and r['metric'] in ['G_accuracy_verified','FPR','FNR']],
       'C95按题配对；原对照区间沿用2000次原结果',ppcols=[3,4,5])
    wsP=sheet('08_前缀完整面板',['前缀名称','实际字符串','FPR','FNR','BA','ΔFPR_pp','FPR下限_pp','FPR上限_pp','ΔFNR_pp','FNR下限_pp','FNR上限_pp','FPR_BH_q','FNR_BH_q'],
       [[r['variant'],json.dumps(r['prefix'],ensure_ascii=False),r['FPR'],r['FNR'],r['BA'],r['FPR_delta']*100,r['FPR_low']*100,r['FPR_high']*100,r['FNR_delta']*100,r['FNR_low']*100,r['FNR_high']*100,r['FPR_q_BH12'],r['FNR_q_BH12']] for r in d['prefix']],
       'C95/G3完整380回答：正确147、错误218、未知15；3040逻辑评分，C为开发集',pctcols=[3,4,5],ppcols=range(6,12))
    sheet('09_前缀双向翻转',['前缀','TN→FP','FP→TN','FN→TP','TP→FN','错误分母','正确分母','未知'],
       [[r['variant'],r['TN_to_FP'],r['FP_to_TN'],r['FN_to_TP'],r['TP_to_FN'],r['wrong_n'],r['correct_n'],r['unknown_n']] for r in d['prefix']],
       '同一回答原文→注入；两方向同时统计，分母区别于原TN翻转率')
    sheet('10_跨G同前缀',['池','G','独立题','前缀','实际字符串','翻转数','原TN翻转率','引擎'],
       [[r['split'],r['G_round'],r['questions'],r['variant'],json.dumps(r['prefix'],ensure_ascii=False),r['flips'],r['flip_rate'],r['source_engine']] for r in d['cross']],
       'A64/C29，每题每G一条原TN；21设置，G0/G1 TF，G2/G3 vLLM。不是全池FPR',pctcols=[7])
    sheet('11_机制对照',['假设','实验覆盖','观察结果','结论','原报告'],MECHANISMS,'既有机制诊断汇总；条件样本各有分母，不合并估计总体发生率')
    sheet('12_自然编辑跨轮差',['池','编辑','指标','G3−G2_pp','下限_pp','上限_pp','相对原文缩小_pp'],
       [[r['split'],r['variant'],r['metric'],r['G3_minus_G2']*100,r['low']*100,r['high']*100,r['gap_reduction_after_edit']*100] for r in d['gaps'] if r['metric']=='FPR'],
       'A598/C95，自然开场对称编辑；同引擎G2/G3；这是编辑后差距的描述',ppcols=range(4,8))
    sheet('13_原V更新T结果',['G','V','TP','FN','FP','TN','FPR','FNR','BA','接受率','AUROC'],
       [[r['G'],r['V'],r['TP'],r['FN'],r['FP'],r['TN'],r['FPR'],r['FNR'],r['balanced_accuracy'],r['retention_all'],r['AUROC']] for r in d['T']],
       'T299同一回答比较V0/Vrand/Vtarget；原更新与本次修复分支分开',pctcols=range(7,11))
    training=[]
    for arm in ARMS:
        p=ROOT/'reports/verifier_robustness_clean_20261026'/(arm+'_training_report.json');r=json.loads(p.read_text());SOURCES[str(p.relative_to(ROOT))]=sha(p)
        training.append([arm,64,128,640,r['steps'],r['input_tokens'],r['chosen_loss_tokens'],r['rejected_tokens_used'],2e-5,8,16,r['seconds_this_invocation'],r['peak_memory_bytes']/2**30])
    sheet('14_修复训练配置',['分支','题目','独立回答','呈现次数','步数','输入token','监督标签token','rejected_token','学习率','LoRA_r','alpha','训练保存重载秒','峰值显存GiB'],training,
       '四组各自从V0初始化；clean原文重复5次，aug为原文+4前缀；同题同前缀正负配对')
    sheet('15_新E所有指标',['G视图','前缀','V','TP','FN','FP','TN','FPR','FNR','BA','未知','接受集错误率'],
       [[r['view'],r['variant'],r['V'],r['TP'],r['FN'],r['FP'],r['TN'],r['FPR'],r['FNR'],r['BA'],r['unknown_accepted']+r['unknown_rejected'],r['accepted_error_risk']] for r in d['E']],
       'E60每G240回答；G0为96正确129错误15未知，G3为91正确129错误20未知',pctcols=[8,9,10,12])
    wsR=sheet('16_修复与新增错误',['G视图','前缀','V','修复FP','修复FN','新增FP','新增FN','净修复','已知分母','净修复率','净修复率下限','净修复率上限','ΔFPR_pp','ΔFNR_pp','未知改判数','含未知净修复下界','含未知净修复上界'],
       [[r['view'],r['variant'],r['V'],r['repaired_FP'],r['repaired_FN'],r['new_FP'],r['new_FN'],r['net_repaired'],r['known_n'],r['net_repair_rate'],r['net_repair_low'],r['net_repair_high'],r['FPR_delta']*100,r['FNR_delta']*100,r['unknown_decisions_changed'],r['net_repaired_all_lower'],r['net_repaired_all_upper']] for r in d['Erepair']],
       '每个新V相对同输入版本V0；净修复=修复FP+修复FN−新增FP−新增FN；未知界与抽样区间分开',pctcols=[10,11,12],ppcols=[13,14])
    sheet('17_增强直接对照',['G视图','前缀','比较_后减前','修复FP','修复FN','新增FP','新增FN','净修复','净修复率下限','净修复率上限'],
       [[r['view'],r['variant'],r['contrast'],r['repaired_FP'],r['repaired_FN'],r['new_FP'],r['new_FN'],r['net_repaired'],r['net_repair_low'],r['net_repair_high']] for r in d['contrast'] if r['cohort']=='E_new_final'],
       '起点是比较名中的前模型，非一律V0；同目标aug-clean用于检查增强增益',pctcols=[9,10])
    sheet('18_局部修复和新翻转',['V','已知回答','原文→注入TN→FP','原文→注入FP→TN','原文→注入FN→TP','原文→注入TP→FN','V0注入诱发FP','两版本均已修复FP'],
       [[r['V'],r['known'],r.get('TN_to_FP',0),r.get('FP_to_TN',0),r.get('FN_to_TP',0),r.get('TP_to_FN',0),r.get('V0_injection_induced_FP',0),r.get('new_stably_correct_FP',0)] for r in d['Estability']],
       'E60/G3/we can；同V原文→注入；DPO均修复7个V0局部案例，但仍有其他新增注入FP')
    sheet('19_参考文献',['编号','作者或名称','原论文','借鉴与关系','链接'],LITERATURE,'已核对原论文；训练方法引用来源，项目贡献定位于问题、数据、对照与分析')
    sheet('20_数据来源与哈希',['原始表格路径','SHA256'],sorted(SOURCES.items()),'汇总读取既有结果，无新增模型训练或评分；完整精度保留在来源CSV')
    for ws in wb:
        for row in ws:
            for c in row:
                c.font=Font(name='Microsoft YaHei',size=10,color=INK.lstrip('#'))
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.row==1:
                    c.fill=PatternFill('solid',fgColor=BG.lstrip('#'));c.font=Font(name='Microsoft YaHei',size=11,bold=True,color=TEAL.lstrip('#'))
                elif c.row==2 and ws!=nav:
                    c.fill=PatternFill('solid',fgColor=NAVY.lstrip('#'));c.font=Font(name='Microsoft YaHei',size=10,bold=True,color='FFFFFF')
                if c.data_type=='f': raise AssertionError('Unexpected Excel formula')
        ws.row_dimensions[1].height=40
        for i in range(2,ws.max_row+1):
            estimated=1
            for c in ws[i]:
                value=str(c.value or '')
                width=ws.column_dimensions[c.column_letter].width
                units=sum(2 if ord(ch)>255 else 1 for ch in value)
                estimated=max(estimated,math.ceil(units/max(width-2,8)))
            ws.row_dimensions[i].height=min(82,max(32,estimated*14+8))
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:2'
    nav.column_dimensions['A'].width=28;nav.column_dimensions['B'].width=105;nav.freeze_panes='A10';nav.sheet_view.showGridLines=False
    for i in [1,9]:
        for c in nav[i]:
            c.fill=PatternFill('solid',fgColor=NAVY.lstrip('#'));c.font=Font(name='Microsoft YaHei',size=11,bold=True,color='FFFFFF')
    for row in nav.iter_rows(min_row=10,max_col=1):
        for c in row:c.font=Font(name='Microsoft YaHei',size=11,color=TEAL.lstrip('#'),underline='single')
    for row in wb['19_参考文献'].iter_rows(min_row=3,min_col=5,max_col=5):
        for c in row:
            c.hyperlink=c.value;c.font=Font(name='Microsoft YaHei',size=10,color=TEAL.lstrip('#'),underline='single')
    charts=wb.create_sheet('21_可编辑汇总图表');charts.sheet_view.showGridLines=False
    line=LineChart();line.title='T299：固定V0的FPR、FNR与BA';line.y_axis.title='比例';line.add_data(Reference(wsT,min_col=11,max_col=13,min_row=2,max_row=6),titles_from_data=True);line.set_categories(Reference(wsT,min_col=1,min_row=3,max_row=6));line.height=10;line.width=23;charts.add_chart(line,'A1')
    bar=BarChart();bar.title='C95/G3：前缀对FPR、FNR的影响';bar.add_data(Reference(wsP,min_col=3,max_col=4,min_row=2,max_row=10),titles_from_data=True);bar.set_categories(Reference(wsP,min_col=1,min_row=3,max_row=10));bar.height=11;bar.width=26;charts.add_chart(bar,'A23')
    nav.append(['21_可编辑汇总图表','原生Excel图表：T跨轮与完整前缀面板']);nav.cell(nav.max_row,1).hyperlink="#'21_可编辑汇总图表'!A1"
    for ws in [wsT,wsP]:
        ws.conditional_formatting.add(f'C3:N{ws.max_row}',ColorScaleRule(start_type='min',start_color='F6FAFB',end_type='max',end_color='C8E1E4'))
    wb.save(OUT/'关键实验结果.xlsx')
    return len(wb.sheetnames)


class Deck:
    """Identical page coordinates for an editable PPTX and a vector PDF."""
    W,H=960,540
    def __init__(self):
        from pptx import Presentation
        from pptx.util import Inches
        self.prs=Presentation();self.prs.slide_width=Inches(13+1/3);self.prs.slide_height=Inches(7.5)
        self.pdf=Canvas(str(OUT/'导师汇报.pdf'),pagesize=(self.W,self.H))
        self.pdf.setTitle('验证器可靠性：跨轮漂移、前缀敏感性与修复');self.pdf.setAuthor('阶段研究汇报')
        self.page=0;self.textlog=[];self.issues=[]

    def rect(self,x,y,w,h,fill=BG,stroke=None,radius=0):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(fill.lstrip('#'))
        if stroke:shape.line.color.rgb=RGBColor.from_string(stroke.lstrip('#'));shape.line.width=Pt(.6)
        else:shape.line.fill.background()
        if radius:
            try:shape.adjustments[0]=.12
            except (AttributeError,IndexError):pass
        self.pdf.setFillColor(HexColor(fill));self.pdf.setStrokeColor(HexColor(stroke or fill))
        self.pdf.roundRect(x,self.H-y-h,w,h,radius,stroke=bool(stroke),fill=1)

    @staticmethod
    def wrap(text,w,size,font):
        result=[]
        for original in str(text).split('\n'):
            if not original:result.append('');continue
            line=''
            # Preserve English words when practical, wrap CJK at character boundaries.
            parts=re.findall(r'[A-Za-z0-9_./:+%−→<>-]+\s*|.',original)
            for part in parts:
                if line and pdfmetrics.stringWidth(line+part,font,size)>w:
                    result.append(line.rstrip());line=part.lstrip()
                else:line+=part
            result.append(line.rstrip())
        return result

    def text(self,x,y,w,h,text,size=18,color=INK,bold=False,align='left'):
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
        from pptx.oxml.xmlchemy import OxmlElement
        font='BriefBold' if bold else 'BriefRegular'
        lines=self.wrap(text,w,size,font)
        while len(lines)*size*1.35>h and size>10.5:
            size-=.5;lines=self.wrap(text,w,size,font)
        if len(lines)*size*1.35>h+1:self.issues.append({'page':self.page,'text':text,'height':h,'size':size})
        self.textlog.append({'page':self.page,'text':str(text),'size':size,'x':x,'y':y,'width':w,'height':h})
        box=self.slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=box.text_frame
        tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.35)
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            run=p.add_run();run.text=line;run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
            ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
            self.pdf.setFont(font,size);self.pdf.setFillColor(HexColor(color))
            yy=self.H-y-size*1.02-i*size*1.35
            if align=='center':self.pdf.drawCentredString(x+w/2,yy,line)
            elif align=='right':self.pdf.drawRightString(x+w,yy,line)
            else:self.pdf.drawString(x,yy,line)

    def image(self,path,x,y,w,h):
        from PIL import Image
        from pptx.util import Pt
        im=Image.open(path);ratio=min(w/im.width,h/im.height);ww,hh=im.width*ratio,im.height*ratio
        xx,yy=x+(w-ww)/2,y+(h-hh)/2
        self.slide.shapes.add_picture(str(path),Pt(xx),Pt(yy),width=Pt(ww),height=Pt(hh))
        self.pdf.drawImage(str(path),xx,self.H-yy-hh,ww,hh,mask='auto')

    def new(self,title,section,subtitle='',source=''):
        if self.page:self.pdf.showPage()
        self.page+=1;self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0,0,self.W,self.H,'#FFFFFF');self.rect(0,0,10,self.H,TEAL)
        self.text(42,21,800,20,section.upper(),10.5,TEAL,bold=True)
        self.text(42,51,875,55,title,28,NAVY,bold=True)
        if subtitle:self.text(43,106,870,42,subtitle,13,MUTED)
        self.rect(42,505,875,1,LINE)
        self.text(43,513,835,19,source or '迭代合成数学数据中的验证器可靠性｜阶段汇报 2026.10.27',8.5,MUTED)
        self.text(880,511,35,20,f'{self.page:02d}',11,MUTED,align='right')

    def card(self,x,y,w,h,label,value,desc,color=TEAL):
        self.rect(x,y,w,h,BG,radius=9);self.rect(x,y,4,h,color)
        self.text(x+18,y+13,w-35,28,label,13,color,bold=True)
        self.text(x+18,y+49,w-35,65,value,31,NAVY,bold=True)
        self.text(x+18,y+121,w-35,h-131,desc,16,INK)

    def callout(self,text,y=441,color=PALE):
        self.rect(43,y,873,51,color,radius=7);self.text(58,y+10,843,36,text,16,NAVY,bold=True)

    def table(self,x,y,widths,headers,rows,rowh=42,size=14):
        total=sum(widths);self.rect(x,y,total,rowh,NAVY)
        xx=x
        for w,s in zip(widths,headers):self.text(xx+10,y+9,w-18,rowh-10,s,size,'#FFFFFF',True);xx+=w
        for i,row in enumerate(rows):
            yy=y+(i+1)*rowh;self.rect(x,yy,total,rowh,BG if i%2==0 else '#FFFFFF')
            xx=x
            for w,s in zip(widths,row):self.text(xx+10,yy+8,w-18,rowh-8,str(s),size,INK);xx+=w

    def finish(self):
        self.pdf.save();self.prs.save(OUT/'导师汇报.pptx')
        (OUT/'slide_text_audit.json').write_text(json.dumps({'pages':self.page,'textboxes':self.textlog,'overflow':self.issues},ensure_ascii=False,indent=2))
        assert not self.issues,self.issues
        return self.page


def slides(d):
    s=Deck()
    # 1. Cover, intentionally sparse.
    s.new('迭代合成数学数据中的验证器可靠性','RESEARCH BRIEF')
    s.rect(43,179,873,244,NAVY,radius=10)
    s.text(69,208,815,59,'跨轮漂移 · 前缀敏感性 · 修复探索',29,'#FFFFFF',True)
    s.text(70,285,801,80,'从固定V的迭代实验出发，追踪表述变化怎样影响判断，\n再检验修复旧错误与保留原有能力的取舍。',20,'#D5E5EC')
    s.text(45,451,870,38,'2026.10.20—10.27｜七天流程与后续补充｜导师交流材料',17,MUTED)
    # 2. Three findings.
    s.new('已完成一条“观察—诊断—修复”的研究链','01 / 研究概览')
    s.card(43,158,277,263,'跨轮可靠性','−3.24 pp','T299：G0→G2\n固定V0的BA下降\n95% CI [−6.06,−0.39]',TEAL)
    s.card(338,158,277,263,'表述敏感性','+10.55 pp','C95/G3：we can\n提高FPR，同时降低FNR\n数学正文与答案不变',RED)
    s.card(633,158,283,263,'修复与代价','+5 条','E60/G3/we can：\nDPO两组均净修复5条\n同时新增7或8个FN',GOLD)
    s.callout('研究焦点：既看是否修好了错误，也看是否破坏了原本正确的判断。')
    # 3. Research logic.
    s.new('固定V不变，回答分布会随G更新而变化','02 / 研究问题','三个问题：跨轮是否稳定？表述是否影响判断？更新V怎样衡量得失？')
    s.image(ASSETS/'research_flow.png',38,167,885,165)
    s.rect(44,350,872,76,BG,radius=8)
    s.text(62,365,836,54,'G0是原始模型；G1/G2/G3依次更新1/2/3次。\nA/B/C/T是题池；每个G阶段可以在多个题池上生成与评价。',18,INK)
    s.callout('主线固定V0更新G；修复实验另从V0开分支，保留原始轨迹与缓存。')
    # 4. Progress.
    s.new('原七天核心流程完成，后续围绕现象补对照','03 / 工作进展')
    s.table(43,150,[103,181,589],['阶段','主题','完成内容'],FLOWS,rowh=36,size=12.5)
    # 5. Data and metrics.
    s.new('用有独立答案标签的数学任务研究V','04 / 模型、数据与指标','G：Qwen2.5-Math-1.5B-Instruct　｜　V0：Qwen3-4B-Instruct-2507')
    s.table(43,156,[66,97,188,522],['池','主分析题数','用途','说明'],[[a,c,e,f] for a,b,c,e,f in POOLS],rowh=42,size=13)
    s.text(44,424,872,66,'FPR：错误回答中被接受的比例　　FNR：正确回答中被拒绝的比例\nBA = 1−(FPR+FNR)/2；每题4候选，未知单列，完整回答输入V。',17,NAVY,True)
    # 6. Training quality.
    s.new('G更新使用了怎样的监督？','05 / 固定V0的反馈训练','A池4候选 → V0接受 → 每题至多随机1条 → 最多512条 / 1 epoch')
    rr=[r for r in d['Gtrain'] if r['model'].startswith('G')]
    s.table(43,162,[85,106,202,159,197,124],['模型','样本','正确/错误/未知','有标签错误率','回答监督token','步数'],[[r['model'],r['examples'],f"{r['selected_correct_by_original_reference']}/{r['selected_wrong_by_original_reference']}/{r['selected_unknown']}",pct(r['selected_verified_error_fraction']),f"{r['loss_tokens']:,}",r['steps']] for r in rr],rowh=54,size=14)
    s.text(45,391,869,43,'独立标签在选样后记录质量；主路线不按标准答案过滤。',18,INK)
    s.callout('训练数据含较多错误：这是含噪验证反馈下的更新，G正确率需独立评价。')
    # 7. Main independent evaluation.
    s.new('T留出集上，固定V0的可靠性下降','06 / 七天主结果','T299 × 每题4回答；四个G统一vLLM评价；G3为扩展',source='来源：week_final_20261026｜T299｜按题配对bootstrap 10,000次')
    s.image(ASSETS/'T_cross_round.png',38,156,889,258)
    s.text(46,421,866,69,'G正确率约40%；G0→G2：BA −3.24pp，95% CI [−6.06,−0.39]。\nFNR +3.65pp；FPR +2.82pp，后者区间跨0。',17,NAVY,True)
    # 8. Controls.
    s.new('更可靠的监督，还未转化成明确的G优势','07 / Oracle、Vmatched与Random','同G0初始化、共同289道A题、37步；监督预算约22.5万回答token')
    s.table(43,158,[182,230,149,156,156],['分支','入训正确/错误/未知','G正确率','V0 FPR','V0 FNR'],[[r['model'].replace('_SFT',''),{'G0':'—','Vmatched_SFT':'223/63/3','Oracle_SFT':'289/0/0','Random_SFT':'198/80/11'}[r['model']],pct(r['G_accuracy_verified']),pct(r['FPR']),pct(r['FNR'])] for r in d['control']],rowh=48,size=14)
    s.text(45,417,867,72,'C95上，三组未检出明确的G正确率优势；Random输出的FPR更低。\n缓存审计提示：错误可识别性与逐题监督权重值得进一步分析。',17,INK)
    # 9. Intervention design.
    s.new('从“什么变化了”转向可控的文本干预','08 / 前缀问题的提出','由自然输出中的误判出发，检验常见表达能否改变验证器决定。')
    s.rect(43,157,419,199,BG,radius=9);s.rect(482,157,434,199,PALE,radius=9)
    s.text(62,177,380,31,'原始输入',18,TEAL,True);s.text(62,223,379,104,'同一题目\n完整数学解答\n最终答案',22,NAVY)
    s.text(502,177,393,31,'干预输入',18,TEAL,True);s.text(502,222,390,111,'同一题目\nwe can + 原完整解答\n相同最终答案',22,NAVY)
    s.text(46,379,868,57,'先确认原TN能否翻转，再补齐全部错误与正确回答，\n同时记录两方向决定变化与连续分数。',18,INK)
    s.callout('数学正文与独立正确性标签保持不变，变化来自指定的表述干预。')
    # 10. Full panel.
    s.new('部分前缀使V更愿接受：FPR升、FNR降','09 / 完整前缀面板','C95/G3：380回答 = 正确147 + 错误218 + 未知15；8设置，3,040逻辑评分',source='来源：prefix_full_panel_20261026｜固定V0、提示、A/B评分与0.5阈值')
    s.image(ASSETS/'prefix_full_panel.png',32,150,900,286)
    s.callout('we can：23个TN→FP；10个FN→TP、1个TP→FN。影响同时涉及两类回答。',y=444)
    # 11. Explanation tests.
    s.new('用系列对照逐步收束解释','10 / 机制诊断')
    rows=[('词汇学术性','20对AWL内外词','未显示AWL优势'),('前缀长度','字符、词数、token','纯长度不足以概括效果'),('输入题型','题型/符号条件FPR','未定位独有退化题型'),('训练链表达','候选→接受→SFT→输出','未见一致富集'),('自然删除/替换','同引擎G2/G3对称编辑','主要跨轮差距仍保留'),('四阶段同前缀','同题G0—G3、21设置','未见一致单调上涨')]
    s.table(43,155,[170,288,415],['问题','做了什么','得到什么'],rows,rowh=45,size=14)
    # 12. Natural attribution.
    s.new('局部漏洞，与自然跨轮差距分别验证','11 / 自然表达与跨轮变化','A598，同vLLM引擎的G2/G3；改掉自然开场后重新评分')
    s.image(ASSETS/'natural_gap.png',42,160,550,251)
    s.rect(616,172,300,242,BG,radius=9)
    s.text(636,192,261,34,'主要差距仍保留',20,TEAL,True)
    s.text(636,247,260,144,'原差距 7.69pp\n删除后 7.12pp\n替换后 6.90pp\n\n继续研究更广的回答与错误组成。',17,INK)
    s.callout('可触发的局部误判已确认；自然漂移还需要更广的机制解释。')
    # 13. Contribution.
    s.new('本人的研究切入与项目贡献','12 / 研究增量与方法来源')
    s.rect(43,155,425,273,BG,radius=8);s.rect(486,155,430,273,PALE,radius=8)
    s.text(62,173,387,35,'本人提出并推进',20,NAVY,True)
    s.text(62,227,385,183,'从跨轮误判提出表达敏感性假设\n设计近义、AWL、长度、自然编辑对照\n同时跟踪FP、FN与新增错误\n将局部干预联系到真实SFT轨迹',17,INK)
    s.text(506,173,388,35,'借鉴并清楚引用',20,NAVY,True)
    s.text(506,227,386,184,'评审偏差：MT-Bench、BITE\n短字符串与监督加固：AdvJudge-Zero\n偏好优化目标：标准DPO\n工程、缓存与图表：工具辅助',17,INK)
    s.callout('贡献定位：具体问题与对照设计、完整回答数据，以及跨轮和修复损伤的分析。')
    # 14. Repair design.
    s.new('从V0独立开始：SFT × DPO × 前缀增强','13 / 验证器修复设计','B/G0：64题、128完整回答、正负各64；每组640次呈现、80步')
    s.table(43,158,[174,358,341],['分支','输入构造','只学习判断标签'],[
        ('SFT_clean','原文 × 5','正确A/B的交叉熵'),('SFT_aug','原文 + 4前缀','同前缀覆盖正确与错误'),('DPO_clean','原文 × 5','正确标签chosen，错误标签rejected'),('DPO_aug','原文 + 4前缀','同上，参考模型为冻结V0')],rowh=47,size=14)
    s.text(46,410,867,73,'原Vrand/Vtarget曾出现过强接受倾向；本轮仅监督1个A/B token。\nLoRA r8 / alpha16，LR 2e−5；不训练G正文，0.5阈值保持。',16,INK)
    # 15. Main repair results.
    s.new('DPO修复部分FP，也新增FN','14 / 新题修复结果','E60/G3/we can：240回答，已知220、未知20；均相对同一输入上的V0',source='来源：verifier_robustness_clean_20261026｜四组从V0独立初始化，原数据和模型保留')
    s.image(ASSETS/'repair_damage.png',43,153,872,257)
    s.text(49,416,860,72,'净修复：SFT_clean −1；SFT_aug −23；DPO两组均 +5。\nDPO_aug：FPR 34.11%→24.81%，FNR 8.79%→16.48%。',18,NAVY,True)
    # 16. Repair tradeoff.
    s.new('局部修好，与整体净改善是两个层次','15 / 如何评价修复')
    s.card(43,159,277,253,'局部修复','7 / 7','两组DPO都修复V0因we can注入新增的7个FP，且两版本均判对。',TEAL)
    s.card(338,159,277,253,'新的问题','7 / 3','对各自原文再注入，DPO_clean / aug仍分别新增7 / 3个FP。',RED)
    s.card(633,159,283,253,'增强的净增益','0 条','DPO_aug−clean主比较净差0；95%区间为[−1.82,+1.81]pp。',GOLD)
    s.callout('当前DPO收益主要体现为更保守的判断；下一步需保留正确接受、提高净收益。')
    # 17. Conclusions.
    s.new('目前已形成什么研究判断？','16 / 阶段结论')
    items=[('01','固定V面对更新后的G输出，可靠性确实需要重新检验。','T上BA下降；G正确率未明显提高。'),('02','一些常见前缀能够改变验证器决定。','完整面板同时量化误接受增加和误拒减少。'),('03','简单删除或小规模增强训练都存在取舍。','应以修复FP/FN与新增FP/FN共同评价。')]
    for i,(num,title,desc) in enumerate(items):
        y=157+i*99;s.rect(43,y,873,86,BG,radius=8);s.text(61,y+15,62,49,num,29,TEAL,True);s.text(137,y+13,755,32,title,19,NAVY,True);s.text(137,y+47,755,30,desc,16,MUTED)
    # 18. Next research.
    s.new('下一步：让判定稳定，同时保持内容区分','17 / 后续研究计划')
    for i,(title,desc) in enumerate([
        ('工作点与判别能力分开','比较风险—保留率曲线和相同FNR/保留率；检查是否只是更保守。'),
        ('提高修复数据覆盖','增加正确接受保持样本；按表达家族留出，跟踪新增FN来源。'),
        ('复现后进入新一轮反馈','少量种子与人工数学核验；冻结配置后新题确认，再将修复V放回G循环。')]):
        y=154+i*96;s.rect(43,y,873,83,PALE if i==0 else BG,radius=8);s.text(62,y+12,831,29,title,20,NAVY,True);s.text(62,y+48,831,28,desc,16,INK)
    s.text(47,457,865,32,'现有基础：可追溯的G/V模型、完整回答、独立标签、连续分数、题池和训练记录。',16,TEAL,True)
    # 19. Compact references and method note.
    s.new('数据口径与代表性文献','附录 / 阅读与复现')
    s.text(45,151,869,91,'最终答案标签；未知单列；坏题整题隔离；区间按题配对。\nC用于开发，T299为七天留出，E60为新修复题池；结果分别呈现。\n单条G训练轨迹、单种子修复；后续先复现，再扩大。',16,INK)
    y=266
    for rid,authors,title,role,url in LITERATURE:
        s.text(45,y,868,29,f'[{rid}] {authors}  ·  {title}',12.5,NAVY,True)
        s.text(45,y+27,868,22,url,11,TEAL)
        self_link=(45,s.H-(y+49),910,s.H-y)
        s.pdf.linkURL(url,self_link,relative=0)
        y+=54
    return s.finish()


def report_pdf():
    """Readable A4 companion, generated from the same complete Markdown summary."""
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,KeepTogether,PageBreak
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from PIL import Image as PILImage
    pagew,pageh=A4;available=pagew-96
    styles={
        'h1':ParagraphStyle('title',fontName='BriefBold',fontSize=24,leading=34,textColor=HexColor(NAVY),spaceAfter=16),
        'h2':ParagraphStyle('section',fontName='BriefBold',fontSize=15,leading=22,textColor=HexColor(TEAL),spaceBefore=17,spaceAfter=10,keepWithNext=True),
        'h3':ParagraphStyle('sub',fontName='BriefBold',fontSize=12,leading=18,textColor=HexColor(NAVY),spaceBefore=12,spaceAfter=8,keepWithNext=True),
        'p':ParagraphStyle('body',fontName='BriefRegular',fontSize=10,leading=16.8,textColor=HexColor(INK),spaceAfter=9,wordWrap='CJK'),
        'cell':ParagraphStyle('cell',fontName='BriefRegular',fontSize=8.0,leading=12.2,textColor=HexColor(INK),wordWrap='CJK'),
        'head':ParagraphStyle('head',fontName='BriefBold',fontSize=8.0,leading=12.2,textColor=colors.white,wordWrap='CJK')}
    def inline(t):
        t=html.escape(t)
        t=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',t)
        t=re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'<i>\1</i>',t)
        t=re.sub(r'`([^`]+)`',r'<font color="#087E8B">\1</font>',t)
        def link(m):
            label,url=m.groups()
            if url.startswith('https://'):return f'<link href="{url}" color="#087E8B">{label}</link>'
            return label
        return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',link,t)
    def footer(c,doc):
        c.setStrokeColor(HexColor(LINE));c.line(48,42,pagew-48,42)
        c.setFont('BriefRegular',8);c.setFillColor(HexColor(MUTED));c.drawString(48,27,'验证器可靠性研究｜阶段总结｜2026.10.27');c.drawRightString(pagew-48,27,str(doc.page))
    story=[];lines=(OUT/'研究总结.md').read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('!['):
            path=OUT/re.search(r'\(([^)]+)\)',line).group(1)
            im=PILImage.open(path);w=available;h=w*im.height/im.width
            story.extend([Spacer(1,5),Image(str(path),width=w,height=h),Spacer(1,10)]);continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            rows=[[c.strip() for c in r.strip('|').split('|')] for r in block]
            rows=[r for r in rows if not all(re.match(r'^:?-+:?$',c) for c in r)]
            n=len(rows[0]);weights=[]
            for j in range(n):
                mx=max(len(re.sub(r'[*`]', '',r[j])) for r in rows)
                weights.append(min(38,max(10,mx))**.68)
            widths=[available*v/sum(weights) for v in weights]
            cells=[[Paragraph(inline(c),styles['head'] if k==0 else styles['cell']) for c in r] for k,r in enumerate(rows)]
            t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor(NAVY)),('ROWBACKGROUNDS',(0,1),(-1,-1),[HexColor(BG),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,-1),(-1,-1),.5,HexColor(LINE))]))
            story.extend([t,Spacer(1,11)]);continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));style=styles['h1' if level==1 else 'h2' if level==2 else 'h3']
            if line.startswith('## 10.'): story.append(PageBreak())
            story.append(Paragraph(inline(line.lstrip('#').strip()),style));continue
        if re.match(r'^\d+\. ',line) or line.startswith('- '):
            label=line if not line.startswith('- ') else '• '+line[2:]
            story.append(Paragraph(inline(label),styles['p']));continue
        para=[line]
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|\d+\. |[-] )',lines[i].strip()):para.append(lines[i].strip());i+=1
        story.append(Paragraph(inline(' '.join(para)),styles['p']))
    doc=SimpleDocTemplate(str(OUT/'研究总结.pdf'),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=48,bottomMargin=57,title='迭代合成数学数据中的验证器可靠性：阶段研究总结',author='阶段研究总结')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def verify_and_package(d,n_sheets,n_slides):
    from pypdf import PdfReader
    from pptx import Presentation
    from openpyxl import load_workbook
    import pymupdf as fitz
    from PIL import Image,ImageOps,ImageDraw
    pdf=PdfReader(OUT/'导师汇报.pdf');paper=PdfReader(OUT/'研究总结.pdf')
    assert len(pdf.pages)==n_slides==19
    assert all(len(p.extract_text())>25 for p in pdf.pages)
    full='\n'.join(p.extract_text() for p in paper.pages)
    for snippet in ['72.79','69.55','65.14','24.81','53.73','AdvJudge-Zero']:
        assert snippet in full,snippet
    rendered_text=(OUT/'研究总结.md').read_text()+''.join(r['text'] for r in json.loads((OUT/'slide_text_audit.json').read_text())['textboxes'])
    for path in (ROOT/'.cache/teacher_reporting_fonts').glob('Brief*.ttf'):
        cmap=FontFile(path).getBestCmap()
        missing=sorted({c for c in rendered_text if ord(c)>31 and ord(c) not in cmap})
        assert not missing,(path.name,missing)
    ppt=Presentation(OUT/'导师汇报.pptx');assert len(ppt.slides)==19
    for slide in ppt.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=ppt.slide_width+1000
            assert shape.top+shape.height<=ppt.slide_height+1000
    wb=load_workbook(OUT/'关键实验结果.xlsx',read_only=True,data_only=False);assert len(wb.sheetnames)==n_sheets
    for ws in wb:
        for row in ws.iter_rows():assert all(c.data_type!='f' for c in row)
    counts={ws.title:max(0,ws.max_row-2) for ws in wb if ws.title not in ['00_阅读导航','21_可编辑汇总图表']};wb.close()
    preview=OUT/'previews';preview.mkdir(exist_ok=True)
    doc=fitz.open(OUT/'导师汇报.pdf');thumbs=[]
    for i,page in enumerate(doc):
        pix=page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False)
        pix.save(str(preview/f'slide_{i+1:02d}.png'))
        im=Image.open(preview/f'slide_{i+1:02d}.png').convert('RGB');im.thumbnail((480,270));tile=Image.new('RGB',(500,305),'#E7EDF1');tile.paste(im,((500-im.width)//2,9));ImageDraw.Draw(tile).text((12,284),f'{i+1:02d}',fill='#20374B');thumbs.append(tile)
    contact=Image.new('RGB',(1500,math.ceil(len(thumbs)/3)*305),'#E7EDF1')
    for i,im in enumerate(thumbs):contact.paste(im,((i%3)*500,(i//3)*305))
    contact.save(preview/'slides_contact_sheet.jpg',quality=90)
    doc=fitz.open(OUT/'研究总结.pdf')
    for i in [0, min(3,len(doc)-1),len(doc)-1]:doc[i].get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(str(preview/f'report_page_{i+1:02d}.png'))
    # Inputs are checked after rendering: packaging never mutates the experiments.
    assert all(sha(ROOT/p)==h for p,h in SOURCES.items())
    (OUT/'source_manifest.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2))
    (OUT/'README.md').write_text('# 导师交流材料\n\n建议先阅读研究总结PDF，交流时使用导师汇报PPT。\n\n- 研究总结.md / 研究总结.pdf：完整研究主线、结果、贡献与方法来源。\n- 导师汇报.pptx / 导师汇报.pdf：19页可编辑PPT及同版PDF。\n- 关键实验结果.xlsx：22张工作表，包含题池说明、关键计数、区间、前缀与修复对照；2张原生Excel图表。\n- assets/：与正文一致的重绘图表。\n\n材料基于既有缓存与报告整理，没有新增模型训练或评分。\n')
    main=['研究总结.md','研究总结.pdf','导师汇报.pptx','导师汇报.pdf','关键实验结果.xlsx','README.md']
    with zipfile.ZipFile(OUT/'导师交流材料.zip','w',zipfile.ZIP_DEFLATED) as z:
        for name in main:z.write(OUT/name,name)
        for p in ASSETS.glob('*.png'):z.write(p,'assets/'+p.name)
    with zipfile.ZipFile(OUT/'导师交流材料.zip') as z:assert z.testzip() is None
    audit={'status':'passed','date':'2026-10-27','slides':n_slides,'summary_pdf_pages':len(paper.pages),'excel_sheets':n_sheets,'excel_data_rows':counts,'source_files_rechecked':len(SOURCES),'source_data_unchanged':True,'model_training_or_scoring_added':False,'PPTX_text_editable':True,'PDF_Chinese_fonts_embedded':True,'all_slide_textboxes_fit':True,'raw_counts_and_FPR_FNR_recomputed':True,'zip_crc_passed':True,'artifacts':{p.name:{'bytes':p.stat().st_size,'sha256':sha(p)} for p in OUT.iterdir() if p.suffix in ['.pdf','.pptx','.xlsx','.zip'] or p.name=='研究总结.md'}}
    (OUT/'delivery_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    print(json.dumps(audit,ensure_ascii=False,indent=2))


def main():
    OUT.mkdir(exist_ok=True);ASSETS.mkdir(exist_ok=True)
    d=data();fonts();figures(d);n_sheets=workbook(d);n_slides=slides(d);report_pdf();verify_and_package(d,n_sheets,n_slides)


if __name__=='__main__':main()
