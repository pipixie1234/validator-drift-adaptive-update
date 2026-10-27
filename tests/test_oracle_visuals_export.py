from io import BytesIO
from openpyxl import Workbook, load_workbook
from pilot.oracle_visuals import excel_safe_value


def test_problem_text_control_character_survives_as_visible_escape_in_excel():
    original='Answer in the form \frac{m}{n}; literal LaTeX: \\frac{m}{n}.'
    safe=excel_safe_value(original)
    assert '\x0c' in original and '\x0c' not in safe
    assert r'\u000crac{m}{n}' in safe and r'\frac{m}{n}' in safe
    wb=Workbook();wb.active.cell(1,1,safe)
    buffer=BytesIO();wb.save(buffer);buffer.seek(0)
    assert load_workbook(buffer).active.cell(1,1).value==safe


def test_xml_controls_are_escaped_while_tabs_and_newlines_remain_readable():
    value='first\tsecond\nthird\rnext\x00\x0b\x1f\ufffe'
    escaped=excel_safe_value(value)
    assert escaped.startswith('first\tsecond\nthird\rnext')
    assert escaped.endswith(r'\u0000\u000b\u001f\ufffe')
