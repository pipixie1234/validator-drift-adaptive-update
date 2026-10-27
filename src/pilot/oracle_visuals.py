"""Matched oracle-control figures, workbook and Chinese review report."""
import json
import re
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

from .core import read_jsonl,write_json
from .oracle_control import RUN,REPORT,ARMS,CONFIG
from .formal_common import sha,utc

NAMES={'G0':'G0','Vmatched_SFT':'V matched','Oracle_SFT':'Oracle'}
COLORS={'G0':'#2379a8','Vmatched_SFT':'#d7872c','Oracle_SFT':'#458a65'}
PRIMARY='C95_exclude_reference_suspect'
EXCEL_CONTROL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]')


def excel_safe_value(value, candidate_id=''):
    """Escape unsupported XML characters visibly; JSONL originals stay untouched."""
    if isinstance(value,(dict,list)):
        value=json.dumps(value,ensure_ascii=False)
    if isinstance(value,str):
        value=EXCEL_CONTROL.sub(lambda m: '\\u'+f'{ord(m.group()):04x}',value)
        if len(value)>32000:
            return '全文见candidates_full.jsonl，candidate_id='+candidate_id
    return value


def pct(x):return 'NA' if x is None else f'{100*x:.2f}%'
def signed(x):return 'NA' if x is None else f'{100*x:+.2f}'
def table(head,rows):return '\n'.join(['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |']+['| '+' | '.join(map(str,r))+' |' for r in rows])+'\n'


def render(result,points,cis,deltas):
    cfg=result['study'];models=['G0',*ARMS]
    get=lambda model:result['metrics'][PRIMARY+'::'+model]['point']
    ci=lambda model,k:result['metrics'][PRIMARY+'::'+model]['ci95'][k]
    delta=lambda contrast,k:next(r for r in deltas if r['scenario']==PRIMARY and r['contrast']==contrast and r['metric']==k)
    figure_dir=REPORT/'figures';figure_dir.mkdir(exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'savefig.dpi':180})
    names=[];pdf=PdfPages(REPORT/'oracle_control_figures.pdf')
    def save(fig,name):
        fig.tight_layout(pad=1.5)
        for ext in ['png','pdf','svg']:fig.savefig(figure_dir/f'{name}.{ext}',bbox_inches='tight')
        pdf.savefig(fig,bbox_inches='tight');plt.close(fig);names.append(name)

    fig,axs=plt.subplots(1,3,figsize=(13.5,4.7))
    for ax,key,title in zip(axs,['G_accuracy_verified','G_accuracy_lower_all','unknown_fraction'],['Accuracy among verifiable answers','Correct / all outputs (lower bound)','Unknown labels / all outputs']):
        for i,g in enumerate(models):
            y=get(g)[key]*100;c=ci(g,key)
            ax.errorbar(i,y,yerr=[[y-c['low']*100],[c['high']*100-y]],fmt='o',color=COLORS[g],capsize=4)
            ax.annotate(f'{y:.2f}%',(i,y),xytext=(0,12),textcoords='offset points',ha='center',fontsize=9)
        ax.set_xticks(range(3),[NAMES[g] for g in models]);ax.set_ylabel('Percent');ax.set_title(title);ax.grid(axis='y',alpha=.2)
    fig.suptitle('C95: two fresh-G0 matched updates; exploratory question-cluster 95% intervals',y=1.03)
    save(fig,'01_generator_development_metrics')

    fig,ax=plt.subplots(figsize=(10,5.3))
    keys=['G_accuracy_verified','G_accuracy_lower_all','FPR','FNR','unknown_fraction']
    labels=['G accuracy, verifiable','Correct / all outputs','V0 false acceptance (FPR)','V0 false rejection (FNR)','Unknown fraction']
    rr=[delta('Oracle_SFT-Vmatched_SFT',k) for k in keys]
    for i,r in enumerate(rr):
        y=r['estimate']*100
        ax.errorbar(y,i,xerr=[[y-r['low']*100],[r['high']*100-y]],fmt='o',color=COLORS['Oracle_SFT'],capsize=4)
        ax.annotate(f'{y:+.2f}',(y,i),xytext=(0,9),textcoords='offset points',ha='center',fontsize=9)
    ax.axvline(0,color='black',lw=.8);ax.set_yticks(range(5),labels);ax.invert_yaxis();ax.set_xlabel('Oracle minus V matched, percentage points')
    ax.set_title('C95 paired differences; positive means higher metric, not always better');ax.grid(axis='x',alpha=.2)
    save(fig,'02_matched_oracle_effects')

    fig,axs=plt.subplots(1,3,figsize=(12,4.7))
    matrices=[np.array([[get(g)['TP'],get(g)['FN']],[get(g)['FP'],get(g)['TN']]]) for g in models]
    for ax,g,matrix in zip(axs,models,matrices):
        ax.imshow(matrix,cmap='Blues',vmin=0,vmax=max(m.max() for m in matrices))
        for i in range(2):
            for j in range(2):ax.text(j,i,str(matrix[i,j]),ha='center',va='center',fontsize=16,color='white' if matrix[i,j]>max(m.max() for m in matrices)/2 else 'black')
        ax.set_xticks([0,1],['V accept','V reject']);ax.set_yticks([0,1],['Correct','Wrong']);ax.set_title(NAMES[g]+f'\nunknown={get(g)["unknown"]}')
    fig.suptitle('C95 fixed V0: final-answer truth vs decision; unknowns excluded from matrices',y=1.03)
    save(fig,'03_fixed_V0_confusion')

    fig,axs=plt.subplots(1,2,figsize=(12,4.6))
    bottom=np.zeros(2)
    for key,color in [('correct','#458a65'),('wrong','#c54d46'),('unknown','#a9a9a9')]:
        values=np.array([cfg['summaries'][g][key]/289*100 for g in ARMS])
        axs[0].bar(range(2),values,bottom=bottom,color=color,label=key)
        for i,v in enumerate(values):
            if v>5:axs[0].text(i,bottom[i]+v/2,f'{v:.1f}%',ha='center',va='center',color='white')
        bottom+=values
    axs[0].set_xticks(range(2),[NAMES[g] for g in ARMS]);axs[0].set_ylabel('All selected examples, percent');axs[0].set_ylim(0,114)
    axs[0].legend(loc='upper center',ncol=3,fontsize=8,frameon=False);axs[0].set_title('289 identical questions; 225,565 loss tokens each')
    for g in ARMS:
        steps=read_jsonl(RUN/g/'train/G1/steps.jsonl');axs[1].plot([s['step'] for s in steps],[s['mean_loss'] for s in steps],color=COLORS[g],label=NAMES[g])
    axs[1].set_xlabel('Optimizer step');axs[1].set_ylabel('Mean response-only SFT loss');axs[1].legend();axs[1].set_title('Same 37 steps; training loss is not test accuracy')
    save(fig,'04_training_match_and_loss')

    fig,axs=plt.subplots(1,2,figsize=(12,4.6))
    for g in models:
        rr=[r for r in read_jsonl(REPORT/'candidates_full.jsonl') if r['G_round']==g and r['question_id']!='dapo:train:24d751b7-4429-4605-9754-d1451911fb26']
        lengths=np.sort([r['output_tokens'] for r in rr]);axs[0].plot(lengths,np.arange(1,len(lengths)+1)/len(lengths),label=NAMES[g],color=COLORS[g])
    axs[0].set_xlabel('Output tokens');axs[0].set_ylabel('Empirical CDF');axs[0].legend();axs[0].set_title('C95: identical generation length budget')
    for i,k in enumerate(['FPR','FNR']):
        for j,g in enumerate(models):
            y=get(g)[k]*100;c=ci(g,k);x=i+(j-1)*.2
            axs[1].errorbar(x,y,yerr=[[y-c['low']*100],[c['high']*100-y]],fmt='o',color=COLORS[g],capsize=4,label=NAMES[g] if i==0 else None)
    axs[1].set_xticks([0,1],['FPR','FNR']);axs[1].set_ylabel('Percent');axs[1].legend();axs[1].set_title('Frozen V0 operating point = 0.5')
    save(fig,'05_output_length_and_V0_errors')
    pdf.close()

    summary_rows=[{'arm':g,**cfg['summaries'][g]} for g in ARMS]
    wb=Workbook();wb.remove(wb.active)
    sheets={'readme':[{'item':'主比较','value':'Oracle_SFT vs Vmatched_SFT；两组G0新初始化，共同289题、225565监督token、37步。'},
                      {'item':'边界','value':'Oracle仅保证冻结最终答案核验正确，不保证全过程正确。Vmatched题集也条件化于oracle可得性。'},
                      {'item':'评估','value':'C95为新参考疑点整题排除敏感性主视图；原C96和reference=1影子视图并列，T封存。'},
                      {'item':'统计','value':'整题成对bootstrap2000次；探索性95%区间，不含训练种子变异。'}],
            'training_match':summary_rows,'training_quality':read_jsonl(RUN/'training_quality_audit.jsonl'),
            'metrics':points,'metrics_ci95':cis,'paired_deltas':deltas,
            'Vmatched_steps':read_jsonl(RUN/'Vmatched_SFT/train/G1/steps.jsonl'),
            'Oracle_steps':read_jsonl(RUN/'Oracle_SFT/train/G1/steps.jsonl'),
            'candidates':read_jsonl(REPORT/'candidates_full.jsonl')}
    escaped_fields=[]
    for name,rs in sheets.items():
        ws=wb.create_sheet(name)
        if not rs:continue
        keys=list(dict.fromkeys(k for r in rs for k in r));ws.append(keys)
        for r in rs:
            ws.append([None]*len(keys))
            for j,k in enumerate(keys,1):
                v=r.get(k)
                if isinstance(v,str) and EXCEL_CONTROL.search(v):
                    escaped_fields.append({'sheet':name,'excel_row':ws.max_row,'field':k,'candidate_id':r.get('candidate_id'),
                                           'escaped_codepoints':sorted({f'U+{ord(x):04X}' for x in EXCEL_CONTROL.findall(v)})})
                v=excel_safe_value(v,r.get('candidate_id',''))
                c=ws.cell(ws.max_row,j,v)
                if isinstance(v,str):c.data_type='s'
                c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245477')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        from openpyxl.utils import get_column_letter
        for j,k in enumerate(keys,1):ws.column_dimensions[get_column_letter(j)].width=60 if k in ['text','question'] else 24
    wb.save(REPORT/'oracle_control_review.xlsx')
    write_json(REPORT/'excel_export_notes.json',{'visible_escaping_only':True,'raw_JSONL_unchanged':True,
               'encoding':'Unsupported XML characters displayed as literal \\uXXXX, never silently corrected or removed',
               'escaped_fields':escaped_fields})

    effect=delta('Oracle_SFT-Vmatched_SFT','G_accuracy_verified')
    robust=delta('Oracle_SFT-Vmatched_SFT','G_accuracy_lower_all')
    if effect['low']>0:
        conclusion='本次共同预算下，oracle分支的可核验开发正确率高于匹配V分支，成对区间未跨0。这与“反馈质量限制了本次学习”一致；还需结合全部输出口径、覆盖变化及训练随机性，不能宣称普遍因果规律。'
    elif effect['high']<0:
        conclusion='本次oracle分支的可核验开发正确率反而低于匹配V分支。提高最终答案监督可靠性在这次覆盖和预算下没有带来更好更新，应继续检查覆盖、过程质量、更新设置和训练随机性；不能因此宣布SFT无效。'
    else:
        conclusion='本次没有检出oracle分支相对匹配V分支的明确开发正确率优势；区间跨0。即使点估计有差异，也不足以在当前题池与单种子下区分“反馈质量限制”和“覆盖/更新预算限制”，更不能据此宣布SFT无效。'
    quality=cfg['summaries']['Vmatched_SFT'];e=quality['wrong']/(quality['correct']+quality['wrong'])
    train_times=sum(r['total_seconds'] for r in result['training_audit']['arms'])
    lines=['# 从G0出发的匹配oracle监督对照', '**'+conclusion+'**',
           '研究问题：保持response-only SFT、题目覆盖及训练预算基本不变，提高最终答案监督的可靠性，能否获得更好的生成器更新？这是明确标记的oracle诊断，独立真值有意用于oracle筛选；不混入原无真值筛选主轨迹。',
           '## 匹配做到了什么',
           table(['项目','Vmatched_SFT','Oracle_SFT'],[
             ['初始化','原始G0 + 新LoRA','原始G0 + 新LoRA'],['共同A题目','289，同一题单','289，同一题单'],
             ['正确 / 错误 / 未知','223 / 63 / 3','289 / 0 / 0'],['可核验错误率',pct(e),'0.00%'],
             ['实际参与损失的回答token','225,565','225,565'],['完整输入token','264,654','264,654'],
             ['epoch / 优化器步','1 / 37','1 / 37'],['microbatch / 梯度累积','1 / 8，末组1条','1 / 8，末组1条'],
             ['学习率 / LoRA','1e-4 / r16, alpha32, dropout0.05','相同'],['样本重复 / 截断','均无','均无']]),
           '两组使用同一冻结trainer和AdamW实现（betas0.9/0.999，eps1e-8，weight_decay0.01，梯度裁剪1.0），同一题目顺序再用种子42打乱，因此每个梯度累积组的题目也相同。两组均不继承原G1/G2。实际末组只有1条，按实际组大小归一化，不把不足8条的梯度缩小。',
           '独立审计精确重建578条训练张量，并逐步核对每条样本只消费一次、74个优化器步骤、实际token总量、G0初始化和adapter保存重载。原训练器曾用B题面做工程探针；本次通过隔离包装明确替换为独立固定文本“Compute 17 + 23.”，无B/C/T题面进入训练探针。冻结原源码未修改。',
           '产物命名说明：为精确复用冻结优化逻辑，每个隔离分支内部仍叫G1阶段；它们不是原formal/G1。底层manifest继承原sample_selection文字及truth_input=False，后者指张量中没有单独参考答案字段，不能据此说oracle选样没用真值。实际选择以本study配置matching及各分支selection_semantics.json为准；Oracle的selected.jsonl明确selection_used_gold=true。原底层记录保留以便追溯。',
           '## 题目选择及无法完全匹配的地方',
           'A池原G0候选中，793条冻结独立标签正确，覆盖304道题。与原G1实际V选中500道题相交得289道；两组共同缩小到这一预算。Vmatched保留这些题原先实际选中的回答，不根据真值重选回答。Oracle在同题现有正确回答中选1条，先精确匹配总体监督token，再尽量减小逐题长度差。没有复制正确回答或补造样本。',
           f'两组有{cfg["same_candidate_in_both_arms"]}条完整回答相同，68条不同。平均每题监督token绝对差{cfg["per_question_token_abs_difference"]["mean"]:.1f}，最大{cfg["per_question_token_abs_difference"]["max"]}；**总体token精确一致不代表逐题长度、过程质量和风格完全匹配**。',
           '题集被条件化为“G0至少曾答对一次且原V路线实际选中过”，只覆盖A的48.17%，不代表完整600题。由于这一条件，匹配V分支的已核验错误率22.03%，明显低于旧500条分支的53.73%；这会降低两个新分支的噪声差异，也限制外推。旧G1的500条、414830监督token、63步只作历史背景，不能与289条oracle直接当作公平对照。',
           'Oracle的“正确”仅指冻结独立最终答案核验通过，既不是全部过程正确的保证，也不是所有公开参考答案已获人工认证。Oracle目标仍是原G0完整回答，没有把参考解、修订答案或C/T材料插入训练。Vmatched的答案没有按真值筛选，但共同题单的确定使用了oracle可得性，这一条件明确披露。',
           '## 相同开发评估协议',
           '两组分别生成C原96道有效题、每题4候选，各384条。沿用原Transformers BF16 SDPA逐题采样函数、原题目种子、温度0.7/top_p0.9、同一提示和max_new_tokens=min(3840,4096−prompt_tokens)。V0原参数、提示、完整回答输入、A/B似然评分实现和0.5阈值全部不变。没有根据C表现改训练或阈值；G0相同协议下的原384条输出及V0评分直接复用。B未用于本对照，T继续封存。',
           '前一轮发现C一题x³−y³=xy+61的源参考为6，而完整整数解给出的x之和是1。本报告以整题排除后的C95作为探索性主视图，并列原C96冻结标签及reference=1影子敏感性视图；不覆盖原标签或暗中修订主协议。未知保持未知，截断始终未知。',
           table(['C95 / 模型','正确 / 错误 / 未知','G正确率（可核验）','正确/全部输出','未知率','V0 FPR','V0 FNR'],[
             [g,f'{get(g)["correct"]}/{get(g)["wrong"]}/{get(g)["unknown"]}',pct(get(g)['G_accuracy_verified']),pct(get(g)['G_accuracy_lower_all']),pct(get(g)['unknown_fraction']),pct(get(g)['FPR']),pct(get(g)['FNR'])] for g in models]),
           '![开发指标](figures/01_generator_development_metrics.png)',
           '## 匹配分支的主要差值',
           table(['C95 / Oracle−Vmatched','差值/百分点','成对95%区间/百分点'],[
             [label,signed(delta('Oracle_SFT-Vmatched_SFT',k)['estimate']),f'[{signed(delta("Oracle_SFT-Vmatched_SFT",k)["low"])}, {signed(delta("Oracle_SFT-Vmatched_SFT",k)["high"])}]']
             for k,label in [('G_accuracy_verified','G可核验正确率'),('G_accuracy_lower_all','正确/全部输出'),('unknown_fraction','未知比例'),('FPR','V0误接受率'),('FNR','V0误拒率')]]),
           '每次重采样整道题，4候选一起进入，两模型和G0共用同一题目重采样索引；2000次百分位95%区间，种子20261022。这些是探索性点位区间，未作多重比较校正，也不含重新训练种子的随机性。未知不作为负标签；正确/全部输出是完整正确率下界，完整上界另在表中报告。',
           '**'+conclusion+'**',
           '![主要差值](figures/02_matched_oracle_effects.png)',
           table(['C95 / 相对G0','可核验正确率变化/百分点','成对95%区间/百分点'],[
             [g,signed(delta(g+'-G0','G_accuracy_verified')['estimate']),f'[{signed(delta(g+"-G0","G_accuracy_verified")["low"])}, {signed(delta(g+"-G0","G_accuracy_verified")["high"])}]'] for g in ARMS]),
           '## 原C96与参考修订敏感性',
           table(['视图','模型','可核验G正确率','正确/全部输出','未知率','FPR'],[
             [scenario,g,pct(result['metrics'][scenario+'::'+g]['point']['G_accuracy_verified']),pct(result['metrics'][scenario+'::'+g]['point']['G_accuracy_lower_all']),pct(result['metrics'][scenario+'::'+g]['point']['unknown_fraction']),pct(result['metrics'][scenario+'::'+g]['point']['FPR'])]
             for scenario in ['C96_frozen_original','C96_reference1_overlay'] for g in models]),
           '这些视图用于检查结论是否依赖一题参考或可核验分母，不挑选最有利的视图。完整计数、区间、保留率、接受集风险、AUROC及长度在Excel/CSV中给出。',
           '![V0混淆矩阵](figures/03_fixed_V0_confusion.png)',
           '## 如何解释这个对照',
           '如果oracle在完整输出和可核验口径上均有一致改善，可把它视为本次预算下监督可靠性限制学习的证据；仍需复现。如果未改善，应继续检查共享题集是否过易或覆盖不足、37步/225565回答token的预算是否足够、正确最终答案下的过程质量，以及不同训练种子的变异。无论哪一种结果，都不能从这一条小规模轨迹直接推出SFT普遍有效或无效。',
           '本对照保持题目、总体token及更新方法一致，但更换了部分完整回答，因此过程内容、局部长度和梯度方向同时会变化。它检验的是“改用更可靠的完整监督回答”这一实际干预，并未实现只更改单个正确性标签而保持所有监督文本不变的反事实。',
           '## 运行资源与复现',
           f'单张RTX4090；两组训练合计{train_times:.1f}秒。开发生成单卡4个独立原采样进程、每分支2个，已有输出逐题落盘；训练和评估均由tmux后台执行。V0评分在完整生成后使用原实现。原正式结果与源码哈希保持一致。',
           '本次评估开始前的vLLM镜像探测速度约0.25–0.45MiB/s，因此本对照保留原引擎。随后按用户指令另开后台备用环境任务，定位官方vLLM0.10.0+cu126（匹配Torch2.7.1+cu126）后，完整主包下载实际约3MiB/s，已按官方SHA256校验；其安装和GPU兼容状态见`reports/vllm_setup_20261022/`。备用准备不打断本轮评估，也不将vLLM输出混入本报告。预编译CUDA依赖兼容说明见[vLLM官方安装文档](https://docs.vllm.ai/en/v0.10.0/getting_started/installation/gpu.html)。',
           '- `oracle_control_review.xlsx`：训练匹配、578条入选质量索引、主指标、区间、差值、逐步训练日志和1152条完整回答。',
           '- `excel_export_notes.json`：题目原文中的Excel不支持控制字符仅在表格中显式转义为\\uXXXX；完整JSONL和模型输入保持原样，未据此重写题目或重算结果。',
           '- `oracle_control_figures.pdf` / `figures/`：5张图，PNG/PDF/SVG；`index.html`为浏览入口。',
           '- `results.json`、`metrics.csv`、`metrics_ci95.csv`、`paired_deltas.csv`：完整数值。',
           '- `training_audit.json`、`analysis_plan.json`、`validation.json`：张量审计、评估完成前固定的比较口径和最终验证。',
           '- 训练配置：`configs/oracle_control_20261022.json`；运行与adapter：`runs/oracle_control_20261022/`；所有控制分支均与原formal目录隔离。',
           '复现：`scripts/pilot oracle_control prepare`、`scripts/pilot oracle_control train --arm Vmatched_SFT`、`scripts/pilot oracle_control train --arm Oracle_SFT`、`scripts/pilot oracle_evaluate run`。已完成训练/输出会使用缓存，不重复消费预算。',
           '日期：北京时间2026-10-22；原始产物UTC时间保留。']
    (REPORT/'oracle_control_report.md').write_text('\n\n'.join(lines)+'\n')
    html='<!doctype html><html lang="zh"><meta charset="utf-8"><title>Oracle监督对照</title><style>body{font:17px system-ui;max-width:1200px;margin:40px auto}img{width:100%}section{margin:45px 0}</style><h1>匹配Oracle监督对照</h1><p>两组G0起步、共同289题、225565监督token、37步；T封存。</p><p><a href="oracle_control_report.md">中文报告</a> · <a href="oracle_control_review.xlsx">Excel</a> · <a href="oracle_control_figures.pdf">图册</a></p>'
    html+=''.join(f'<section><h2>{n}</h2><img src="figures/{n}.png" alt="{n}"></section>' for n in names)+'</html>'
    (REPORT/'index.html').write_text(html)
    write_json(REPORT/'visuals_manifest.json',{'created_at':utc(),'figures':names,'source_results_sha256':sha(REPORT/'results.json'),
               'visuals_code_sha256':sha(Path(__file__)),'workbook_sheets':len(wb.sheetnames)})
