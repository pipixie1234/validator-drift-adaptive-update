"""One additional G0-initialized, length-conditioned random SFT control.

Original formal and oracle artifacts are read-only. Selection sees candidate text
and lengths, never verifier scores or correctness labels. Truth is joined only
after immutable selection files have been committed.
"""
import argparse
import contextlib
import copy
import json
import math
import os
import random
import shutil
import subprocess
import time
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np

from .core import (ROOT, read_jsonl, recover_jsonl, write_json, write_jsonl,
                   ensure_run_manifest, training_example, digest, cache_identity)
from .formal_common import (RUN as FORMAL, PROTOCOL, protocol, messages, inputs,
                            sha, utc, WorkBudget, gpu_lock)
from .oracle_control import RUN as ORACLE_RUN, CONFIG as ORACLE_CONFIG
from .oracle_control import validate_plan as validate_oracle

RUN = ROOT/'runs/random_control_20261024'
REPORT = ROOT/'reports/random_control_20261024'
CONFIG = ROOT/'configs/random_control_20261024.json'
ARM = 'Random_SFT'
MODELS = ['G0', 'Random_SFT', 'Vmatched_SFT', 'Oracle_SFT']
TRAIN = RUN/ARM/'train/G1'
SEED = 20261024
WORKERS = 4
SCENARIOS = ['C95_exclude_reference_suspect', 'C96_frozen_original', 'C96_reference1_overlay']
METRICS = ['G_accuracy_verified', 'G_accuracy_lower_all', 'unknown_fraction',
           'FPR', 'FNR', 'retention_all', 'accepted_error_FDR', 'AUROC', 'mean_output_tokens']


def checkpoint(): return RUN/ARM/'checkpoints/G1_SFT'
def generated(): return RUN/ARM/'eval/generated_C.jsonl'
def scored(): return RUN/ARM/'eval/V0_C.jsonl'


def draw_length_conditioned(length_options, target, seed=SEED, tolerance=.005, max_draws=100000):
    """First uniform product draw satisfying length only; no best-result search.

Conditional on success this is uniform on admissible complete candidate tuples,
not unconditional per-question uniform sampling. If no draw qualifies, retain
the FIRST draw and report the unmatched budget; never pick the closest by truth.
    """
    if not length_options or any(not x for x in length_options):
        raise ValueError('Every question needs at least one structurally complete candidate')
    sizes = np.array([len(x) for x in length_options])
    lengths = np.zeros((len(sizes), int(sizes.max())), dtype=np.int64)
    for i, options in enumerate(length_options): lengths[i, :len(options)] = options
    rng = np.random.default_rng(seed)
    first = None
    totals = []
    for attempt in range(1, max_draws+1):
        choice = rng.integers(sizes)
        total = int(lengths[np.arange(len(sizes)), choice].sum())
        totals.append(total)
        if first is None: first = choice.copy()
        if abs(total-target) <= target*tolerance:
            return choice.tolist(), {'draws': attempt, 'length_matched': True,
                'selected_total': total, 'first_draw_total': totals[0],
                'rejected_length_totals': totals[:-1]}
    total = int(lengths[np.arange(len(sizes)), first].sum())
    return first.tolist(), {'draws': max_draws, 'length_matched': False,
        'selected_total': total, 'first_draw_total': totals[0],
        'rejected_length_totals': totals}


def prepare():
    from transformers import AutoTokenizer
    validate_oracle()
    RUN.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    if CONFIG.exists(): validate(); print('Immutable random selection cached.', flush=True); return
    oc = json.loads(ORACLE_CONFIG.read_text())
    target = oc['summaries']['Oracle_SFT']['supervised_tokens']
    # Commit the rule BEFORE candidate length computation or label/score joins.
    selection_plan = {'study': 'random_SFT_length_matched_control_v1', 'code_sha256': sha(Path(__file__)),
        'seed': SEED, 'question_source': 'oracle_control/common_A_questions.jsonl',
        'questions_conditioned_on_old_oracle_availability': True,
        'candidate_eligibility': 'Original G0_A; not generation-truncated; exact full training tensor <=4096. No answer-format, truth or V filter.',
        'draw': 'Uniform independent candidate per sorted question, repeated until FIRST complete tuple within token tolerance.',
        'target_loss_tokens': target, 'relative_tolerance': .005, 'max_draws': 100000,
        'fallback': 'Use FIRST uniform draw if no budget match; explicitly report token mismatch.',
        'conditioning_limit': 'Uniform over length-admissible tuples, not unconditional uniform candidate marginals.',
        'no_V_scores_or_labels_in_selection': True, 'no_repeated_examples_or_truncation': True,
        'training_seed': 42, 'initialization': 'Fresh original G0 plus new LoRA; no old adapters',
        'training': oc['training'], 'engineering_probe': 'Compute 17 + 23.',
        'evaluation': {'questions': 96, 'candidates_each': 4, 'engine': 'original Transformers BF16 SDPA',
            'workers': WORKERS, 'V0_threshold': .5, 'protocol_sha256': sha(PROTOCOL),
            'scenarios': SCENARIOS, 'T_read': False, 'B_used': False},
        'analysis': {'primary_contrasts': ['Vmatched_SFT-Random_SFT', 'Oracle_SFT-Random_SFT', 'Random_SFT-G0'],
            'primary_metric': 'G_accuracy_verified; report correct/all and unknown alongside',
            'bootstrap_replicates': 2000, 'bootstrap_seed': 20261022,
            'unit': 'Paired whole questions; exploratory pointwise percentile95; single training seed',
            'development_reuse_disclosed': True}}
    ensure_run_manifest(RUN/'selection_plan.json', selection_plan)
    qs = [r['question_id'] for r in read_jsonl(ORACLE_RUN/'common_A_questions.jsonl')]
    assert qs == sorted(set(qs)) and len(qs) == 289
    raw = read_jsonl(FORMAL/'G0_A.jsonl')
    tok = AutoTokenizer.from_pretrained(protocol()['models']['G']['path'], local_files_only=True)
    options = defaultdict(list); excluded = []
    for r in sorted(raw, key=lambda r: r['candidate_id']):
        if r['question_id'] not in set(qs): continue
        if r['truncated']:
            excluded.append({'candidate_id': r['candidate_id'], 'reason': 'generation_truncated'}); continue
        ex = training_example(tok, messages('G', r['question']), r['text'], 4096)
        if ex is None:
            excluded.append({'candidate_id': r['candidate_id'], 'reason': 'over_training_context'}); continue
        options[r['question_id']].append((r, ex, sum(t != -100 for t in ex['labels'][1:])))
    choice, trace = draw_length_conditioned([[x[2] for x in options[q]] for q in qs], target)
    selected = []; tensors = []
    for q, j in zip(qs, choice):
        r, ex, loss_tokens = options[q][j]
        selected.append({k: r[k] for k in ['question_id','candidate_id','question','text','G_round','prompt_tokens','output_tokens']} |
                        {'input_tokens': len(ex['input_ids']), 'supervised_tokens': loss_tokens,
                         'selection_used_gold': False, 'selection_used_verifier': False})
        tensors.append({'question_id': q, 'candidate_id': r['candidate_id'], **ex})
    TRAIN.mkdir(parents=True, exist_ok=True)
    write_jsonl(TRAIN/'selected.jsonl', selected); write_jsonl(TRAIN/'tokenized.jsonl', tensors)
    write_json(RUN/'random_draw_trace.json', trace)
    write_jsonl(RUN/'structural_exclusions.jsonl', excluded)
    ensure_run_manifest(RUN/'selection_commit.json', {'selected_sha256': sha(TRAIN/'selected.jsonl'),
        'tokenized_sha256': sha(TRAIN/'tokenized.jsonl'), 'selection_plan_sha256': sha(RUN/'selection_plan.json')})
    # Correctness is joined ONLY AFTER selection is committed and cannot affect it.
    labels = {r['candidate_id']: r for r in read_jsonl(FORMAL/'labels_G0_A.jsonl')}
    quality = [{'question_id': r['question_id'], 'candidate_id': r['candidate_id'],
                'correct': labels[r['candidate_id']]['correct'], 'supervised_tokens': r['supervised_tokens']}
               for r in selected]
    write_jsonl(RUN/'training_quality_audit.jsonl', quality)
    truths = [r['correct'] for r in quality]
    summary = {'examples': len(selected), 'questions': len(selected), 'optimizer_steps': math.ceil(len(selected)/8),
        'correct': sum(x is True for x in truths), 'wrong': sum(x is False for x in truths),
        'unknown': sum(x is None for x in truths), 'supervised_tokens': sum(r['supervised_tokens'] for r in selected),
        'input_tokens': sum(r['input_tokens'] for r in selected),
        'target_supervised_tokens': target, 'relative_token_difference': trace['selected_total']/target-1,
        'length_matched': trace['length_matched'], 'random_draws': trace['draws']}
    sources = [FORMAL/'G0_A.jsonl', FORMAL/'labels_G0_A.jsonl', PROTOCOL, ORACLE_CONFIG,
        ORACLE_RUN/'common_A_questions.jsonl', ROOT/'reports/oracle_control_20261022/candidates_full.jsonl']
    cfg = {'selection_plan': selection_plan, 'summary': summary,
        'selected_sha256': sha(TRAIN/'selected.jsonl'), 'tokenized_sha256': sha(TRAIN/'tokenized.jsonl'),
        'draw_trace_sha256': sha(RUN/'random_draw_trace.json'),
        'original_sources_sha256': {str(p.relative_to(ROOT)): sha(p) for p in sources},
        'code_sha256': sha(Path(__file__)), 'prepared_at': utc()}
    ensure_run_manifest(CONFIG, cfg)
    ensure_run_manifest(REPORT/'analysis_plan.json', selection_plan['analysis'] | {'study_sha256': sha(CONFIG)})
    (REPORT/'README.md').write_text('# 随机选样SFT补充对照\n\n同一G0、共同289道A题、每题一条完整回答。随机抽样只按总体长度预算接受，不使用V分数或真标签筛选。\n\n运行：`scripts/pilot random_control run`；断线后可在tmux恢复，缓存回答不重跑，T封存。\n')
    status('prepared', training_summary=summary)
    print(json.dumps(summary, ensure_ascii=False), flush=True)


def validate():
    validate_oracle()
    cfg = json.loads(CONFIG.read_text())
    assert cfg['code_sha256'] == sha(Path(__file__))
    assert cfg['selection_plan'] == json.loads((RUN/'selection_plan.json').read_text())
    assert sha(TRAIN/'selected.jsonl') == cfg['selected_sha256']
    assert sha(TRAIN/'tokenized.jsonl') == cfg['tokenized_sha256']
    assert sha(RUN/'random_draw_trace.json') == cfg['draw_trace_sha256']
    assert all(sha(ROOT/p) == h for p,h in cfg['original_sources_sha256'].items())
    return cfg


def status(stage, **fields):
    write_json(RUN/'progress.json', {'stage': stage, 'updated_at': utc(), 'pid': os.getpid(), **fields})


def train():
    from . import formal_train as frozen
    cfg = validate()
    ensure_run_manifest(RUN/ARM/'selection_semantics.json', {
        'arm': ARM, 'selection': cfg['selection_plan']['draw'], 'length_conditioning': True,
        'selection_used_gold': False, 'selection_used_V': False,
        'shared_question_set_was_previously_conditioned_on_oracle_availability': True,
        'inherited_manifest_sample_selection_text_is_not_operative_for_this_branch': True,
        'private_stage_G1_is_not_original_formal_G1': True, 'engineering_probe': 'Compute 17 + 23.'})
    frozen.RUN = RUN/ARM
    frozen.inputs = lambda split: [{'question': 'Compute 17 + 23.'}] if split == 'B' else (_ for _ in ()).throw(ValueError('Independent probe only'))
    frozen.progress = lambda stage, **kw: status(stage, **kw)
    class Budget(WorkBudget):
        def __init__(self, stage): super().__init__('random_control/'+stage)
    frozen.WorkBudget = Budget
    frozen.train('G1')
    report = json.loads((TRAIN/'report.json').read_text())
    assert report['source_generator'] == 'G0' and report['parent_adapter_fingerprint'] is None
    assert report['supervised_tokens'] == cfg['summary']['supervised_tokens'] and report['optimizer_steps'] == 37
    write_json(RUN/'training_complete.json', {'completed_at': utc(), 'report_sha256': sha(TRAIN/'report.json')})


def worker(assignment):
    from . import formal_run as original
    a = json.loads(Path(assignment).read_text())
    assert a['study_sha256'] == sha(CONFIG) and a['code_sha256'] == sha(Path(__file__))
    pool = {r['question_id']: r for r in inputs('C')}
    rows = [pool[q] for q in a['question_ids']]; wd = Path(a['directory'])
    original.RUN = wd
    original.inputs = lambda split: rows if split == 'C' else (_ for _ in ()).throw(ValueError('C only'))
    original.output_path = lambda *args: wd/'generated.jsonl'
    original.adapter_path = lambda name: checkpoint()
    original.gpu_lock = contextlib.nullcontext
    original.progress = lambda stage, **kw: write_json(wd/'progress.json', {'stage': stage, 'updated_at': utc(), **kw})
    class Budget:
        def __init__(self, stage): pass
        def check(self):
            if (RUN/'cancel').exists(): raise RuntimeError('Coordinator requested stop')
        def charge(self, unit): pass
    original.WorkBudget = Budget
    original.run('generate', ARM, 'C')


def generate():
    from .formal_parallel import merge_ready
    validate(); rows = inputs('C'); generated().parent.mkdir(parents=True, exist_ok=True)
    ensure_run_manifest(RUN/'evaluation.manifest.json', {'study_sha256': sha(CONFIG),
        'code_sha256': sha(Path(__file__)), 'engine': 'original Transformers BF16 SDPA',
        'sampler_sha256': sha(ROOT/'src/pilot/formal_run.py'), 'workers': WORKERS,
        'whole_question_shards': True, 'C_questions': len(rows), 'expected_candidates': len(rows)*4,
        'V0_scoring_unchanged': True, 'B_used': False, 'T_read': False})
    with gpu_lock():
        budget = WorkBudget('random_control/parallel_C_generation')
        saved = recover_jsonl(generated()); children = []; logs = []; paths = []
        try:
            for i in range(WORKERS):
                wd = RUN/ARM/f'eval/worker_{i}'; wd.mkdir(exist_ok=True); paths.append(wd/'generated.jsonl')
                assignment = wd/'assignment.json'
                ensure_run_manifest(assignment, {'question_ids': [r['question_id'] for r in rows[i::WORKERS]],
                    'directory': str(wd), 'study_sha256': sha(CONFIG), 'code_sha256': sha(Path(__file__))})
                log = (wd/'worker.log').open('ab'); logs.append(log)
                children.append(subprocess.Popen([str(ROOT/'scripts/pilot'), 'random_control', 'worker', '--assignment', str(assignment)],
                    cwd=ROOT, env=dict(os.environ, CUDA_VISIBLE_DEVICES='0'), stdout=log, stderr=subprocess.STDOUT))
            last = -1
            while True:
                merge_ready(generated(), paths, rows, saved)
                if len(saved) != last:
                    status('C_generation', candidates=len(saved), expected=384, workers=WORKERS,
                        worker_pids=[p.pid for p in children if p.poll() is None])
                    print(f'Random_SFT C: {len(saved)}/384 candidates', flush=True)
                    budget.charge(f'candidates_{len(saved)}'); last = len(saved)
                codes = [p.poll() for p in children]
                if any(c is not None and c != 0 for c in codes): raise RuntimeError(f'Worker failed: {codes}; inspect worker logs')
                if all(c is not None for c in codes): break
                budget.check(); time.sleep(1)
            merge_ready(generated(), paths, rows, saved); assert len(saved) == 384
            budget.charge('generation_finalization')
            write_json(RUN/'generation_complete.json', {'completed_at': utc(), 'candidates':384})
        finally:
            for p in children:
                if p.poll() is None: p.terminate()
            for p in children:
                try: p.wait(timeout=15)
                except subprocess.TimeoutExpired: p.kill(); p.wait()
            for log in logs: log.close()


def score():
    from . import formal_run as original
    validate(); status('V0_scoring', candidates=len(read_jsonl(generated())))
    original.RUN = RUN/ARM/'eval'
    original.output_path = lambda name, split, role='G': generated() if role == 'G' else scored()
    original.progress = lambda stage, **kw: status(stage, **kw)
    class Budget(WorkBudget):
        def __init__(self, stage): super().__init__('random_control/'+stage)
    original.WorkBudget = Budget
    original.run('score', ARM, 'C')


def audit(cfg):
    from transformers import AutoTokenizer
    from .runtime import model_identity
    from .day3_protocol import formal_generation_settings
    tok = AutoTokenizer.from_pretrained(protocol()['models']['G']['path'], local_files_only=True)
    source = {r['candidate_id']:r for r in read_jsonl(FORMAL/'G0_A.jsonl')}
    chosen = read_jsonl(TRAIN/'selected.jsonl'); tensors = read_jsonl(TRAIN/'tokenized.jsonl')
    qs = [r['question_id'] for r in read_jsonl(ORACLE_RUN/'common_A_questions.jsonl')]
    assert [r['question_id'] for r in chosen] == qs and len(chosen) == len(tensors) == 289
    # Rebuild eligible options and replay the actual random draw without labels/scores.
    options = defaultdict(list)
    for r in sorted(source.values(), key=lambda r:r['candidate_id']):
        if r['question_id'] not in set(qs) or r['truncated']: continue
        ex = training_example(tok, messages('G',r['question']),r['text'],4096)
        if ex is not None: options[r['question_id']].append((r['candidate_id'], sum(t!=-100 for t in ex['labels'][1:])))
    choices, trace = draw_length_conditioned([[x[1] for x in options[q]] for q in qs], cfg['summary']['target_supervised_tokens'])
    assert trace == json.loads((RUN/'random_draw_trace.json').read_text())
    assert [options[q][j][0] for q,j in zip(qs,choices)] == [r['candidate_id'] for r in chosen]
    for r,t in zip(chosen,tensors):
        raw = source[r['candidate_id']]
        assert r['question'] == raw['question'] and r['text'] == raw['text'] and not raw['truncated']
        assert not r['selection_used_gold'] and not r['selection_used_verifier']
        ex = training_example(tok,messages('G',r['question']),r['text'],4096)
        assert all(t[k] == ex[k] for k in ['input_ids','labels','attention_mask','prefix_length'])
        assert not any(k in t for k in ['correct','gold_answer','reference','score','accepted'])
    steps = read_jsonl(TRAIN/'steps.jsonl'); order = list(range(289)); random.Random(42).shuffle(order)
    assert [c for s in steps for c in s['candidate_ids']] == [chosen[i]['candidate_id'] for i in order]
    assert len(steps) == 37 and len(set(c for s in steps for c in s['candidate_ids'])) == 289
    tr = json.loads((TRAIN/'report.json').read_text())
    assert tr['source_generator'] == 'G0' and tr['parent_adapter_fingerprint'] is None
    assert tr['base_frozen'] and tr['save_reload_passed'] and tr['finite_losses_and_gradients']
    assert tr['supervised_tokens'] == sum(s['supervised_tokens'] for s in steps) == cfg['summary']['supervised_tokens']
    base, adapter = model_identity('G',checkpoint()); vfp,_ = model_identity('V'); p = protocol()
    assert adapter == tr['final_adapter_fingerprint']
    gs = read_jsonl(generated()); vs = {r['candidate_id']:r for r in read_jsonl(scored())}
    pool = {r['question_id']:r for r in inputs('C')}
    assert len(gs) == len(vs) == len({(r['question_id'],r['candidate_index']) for r in gs}) == 384
    for r in gs:
        assert r['question'] == pool[r['question_id']]['question'] and r['G_round'] == ARM and r['split'] == 'C'
        assert r['model_fingerprint'] == base and r['adapter_fingerprint'] == adapter
        assert r['candidate_id'] == f'{r["question_id"]}:formal:{ARM}:{r["candidate_index"]}'
        actual = formal_generation_settings(r['prompt_tokens']); seed = 42+int(digest(r['question_id'])[:8],16)%1000000
        key = cache_identity('G',base,adapter,r['question_id'],messages('G',r['question']),dict(actual,seed=seed))
        assert r['generation_seed'] == seed and r['question_cache_key'] == key
        assert r['cache_key'] == digest({'question_key':key,'candidate_index':r['candidate_index']})
        assert r['prompt_sha256'] == digest(messages('G',r['question']))
        assert r['max_new_tokens'] == actual['max_new_tokens'] and r['prompt_tokens']+r['output_tokens'] <= 4096
        assert not any(k in r for k in ['correct','gold_answer','reference'])
        vr = vs[r['candidate_id']]; msg = messages('V',r['question'],r['text'])
        assert vr['model_fingerprint'] == vfp and vr['adapter_fingerprint'] is None and vr['threshold'] == .5
        assert vr['prompt_tokens']+4 <= 8192 and vr['prompt_sha256'] == digest(msg)
        assert vr['cache_key'] == cache_identity('V',vfp,None,r['candidate_id'],msg,p['V_scoring'])
        a,b = vr['log_likelihood_A'],vr['log_likelihood_B']
        expected = 1/(1+math.exp(b-a)) if b-a<700 else 0.
        assert abs(expected-vr['score']) < 2e-7 and vr['accepted'] == (a>=b)
    result = {'passed':True,'validated_at':utc(),'random_draw_reproduced':True,
        'exact_training_tensors_rebuilt':289,'optimizer_steps':37,'same_question_order_and_groups':True,
        'full_G_and_V0_records_checked_each':384,'fresh_G0':True,'T_read':False,'B_used':False,
        'selection_used_V_or_truth':False,'question_set_previously_oracle_conditioned':True,
        'length_matched':cfg['summary']['length_matched'],'training':tr,
        'original_and_oracle_sources_unchanged':True}
    write_json(REPORT/'validation.json',result); return result


def report():
    from . import oracle_report
    from .day2_protocol import numeric
    from .explanation_probes import SUSPECT
    from .v0_drift_analysis import analyze, clean, interval, csv_write
    cfg = validate(); status('analysis')
    # Only redirect this process's generic label join; no original files modified.
    oracle_report.generated = lambda arm: generated()
    oracle_report.scored = lambda arm: scored()
    oracle_report.RUN = RUN
    new = oracle_report.load_control(ARM)
    old = read_jsonl(ROOT/'reports/oracle_control_20261022/candidates_full.jsonl')
    data = {g:[r for r in old if r['G_round']==g] for g in MODELS if g != ARM}; data[ARM]=new
    assert all(len(data[g]) == 384 for g in MODELS)
    result = {'study':cfg,'metrics':{},'paired_deltas':[],'T_read':False}
    points=[]; cis=[]; deltas=[]
    for scenario in SCENARIOS:
        view=copy.deepcopy(data)
        if scenario.startswith('C95'): view={g:[r for r in rr if r['question_id']!=SUSPECT] for g,rr in view.items()}
        elif scenario == 'C96_reference1_overlay':
            for rr in view.values():
                for r in rr:
                    if r['question_id']==SUSPECT and r['correct'] is not None: r['correct']=numeric(r['extracted_answer'])=='1'
        qs=sorted({r['question_id'] for r in view['G0']})
        samples=np.random.default_rng(20261022).integers(len(qs),size=(2000,len(qs))); boot={}
        for g in MODELS:
            point,ci,draws,_=analyze(view[g],qs,samples)
            for k in METRICS:
                if k not in ci:ci[k]=interval(draws[k])
            result['metrics'][scenario+'::'+g]={'point':point,'ci95':ci}; boot[g]=draws
            points.append({'scenario':scenario,'model':g,**{k:v for k,v in point.items() if not isinstance(v,dict)}})
            for k in METRICS:cis.append({'scenario':scenario,'model':g,'metric':k,'estimate':point[k],**ci[k]})
        for older,newer in [('G0',ARM),(ARM,'Vmatched_SFT'),(ARM,'Oracle_SFT'),('G0','Vmatched_SFT'),('G0','Oracle_SFT')]:
            for k in METRICS:
                deltas.append({'scenario':scenario,'contrast':newer+'-'+older,'metric':k,
                    'estimate':result['metrics'][scenario+'::'+newer]['point'][k]-result['metrics'][scenario+'::'+older]['point'][k],
                    **interval(boot[newer][k]-boot[older][k])})
    result['paired_deltas']=clean(deltas); result['training_audit']=audit(cfg); result['completed_at']=utc()
    write_json(REPORT/'results.json',clean(result))
    write_jsonl(REPORT/'candidates_full.jsonl',[r for g in MODELS for r in data[g]])
    for name,rs in [('metrics',points),('metrics_ci95',cis),('paired_deltas',deltas)]:csv_write(REPORT/f'{name}.csv',rs)
    render(result,points,cis,deltas)
    bundle()
    write_json(RUN/'state.json',{'status':'complete','completed_at':utc(),'C_generated':384,'V0_scored':384,
        'cached_comparison_records':1152,'T_read':False,'report':str(REPORT/'random_control_report.md')})
    status('complete',report=str(REPORT/'random_control_report.md'))
    print(json.dumps({'status':'complete','report':str(REPORT/'random_control_report.md')},ensure_ascii=False),flush=True)


def render(result,points,cis,deltas):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from openpyxl import Workbook, load_workbook
    from .oracle_visuals import excel_safe_value, table, pct
    cfg=result['study']; primary=SCENARIOS[0]
    get=lambda g: result['metrics'][primary+'::'+g]['point']
    colors=['#2379a8','#8656a3','#d7872c','#458a65']; names=['G0','Random','V matched','Oracle']
    fd=REPORT/'figures';fd.mkdir(exist_ok=True)
    with PdfPages(REPORT/'random_control_figures.pdf') as pdf:
        fig,axs=plt.subplots(1,3,figsize=(13,4.5))
        for ax,key,title in zip(axs,['G_accuracy_verified','G_accuracy_lower_all','unknown_fraction'],['Verifiable accuracy','Correct / all outputs','Unknown fraction']):
            for i,g in enumerate(MODELS):
                y=get(g)[key]*100;ci=result['metrics'][primary+'::'+g]['ci95'][key]
                ax.errorbar(i,y,yerr=[[max(0,y-ci['low']*100)],[max(0,ci['high']*100-y)]],fmt='o',capsize=4,color=colors[i])
                ax.annotate(f'{y:.2f}%',(i,y),xytext=(0,10),textcoords='offset points',ha='center',fontsize=8)
            ax.set_xticks(range(4),names,rotation=15);ax.set_title(title);ax.set_ylabel('Percent');ax.grid(axis='y',alpha=.2)
        fig.suptitle('C95: one additional G0 random-SFT control; paired question-cluster intervals');fig.tight_layout()
        for ext in ['png','svg','pdf']:fig.savefig(fd/f'01_four_arm_metrics.{ext}',bbox_inches='tight')
        pdf.savefig(fig);plt.close(fig)
        fig,axs=plt.subplots(1,4,figsize=(13,4))
        mats=[np.array([[get(g)['TP'],get(g)['FN']],[get(g)['FP'],get(g)['TN']]]) for g in MODELS]
        vmax=max(m.max() for m in mats)
        for ax,g,name,m in zip(axs,MODELS,names,mats):
            ax.imshow(m,cmap='Blues',vmin=0,vmax=vmax)
            for i in range(2):
                for j in range(2):ax.text(j,i,str(int(m[i,j])),ha='center',va='center',color='white' if m[i,j]>vmax/2 else 'black')
            ax.set_xticks([0,1],['Accept','Reject']);ax.set_yticks([0,1],['Correct','Wrong']);ax.set_title(name+f'\nunknown={int(get(g)["unknown"])}')
        fig.suptitle('Fixed V0 confusion matrices; unknown labels excluded');fig.tight_layout()
        for ext in ['png','svg','pdf']:fig.savefig(fd/f'02_confusion_matrices.{ext}',bbox_inches='tight')
        pdf.savefig(fig);plt.close(fig)
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,rows):
        ws=wb.create_sheet(name)
        if not rows:return
        keys=list(rows[0]);ws.append(keys)
        for row in rows:
            for j,k in enumerate(keys,1):
                c=ws.cell(ws.max_row+1,j) if j==1 else ws.cell(ws.max_row,j)
                value=excel_safe_value(row.get(k),row.get('candidate_id',''));c.value=value
                if isinstance(value,str):c.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=24
    sheet('Training_match',[{'model':g,**s} for g,s in training_summaries(cfg).items()])
    sheet('Metrics',points);sheet('Intervals',cis);sheet('Paired_differences',deltas)
    sheet('Random_training_quality',read_jsonl(RUN/'training_quality_audit.jsonl'))
    sheet('Training_steps',read_jsonl(TRAIN/'steps.jsonl'))
    sheet('Full_C_records',read_jsonl(REPORT/'candidates_full.jsonl'))
    wb.save(REPORT/'random_control_review.xlsx')
    check=load_workbook(REPORT/'random_control_review.xlsx');assert len(check.sheetnames)==7
    assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
    summary=cfg['summary']; tr=result['training_audit']['training']
    lines=['# 随机选样SFT补充对照','',
        '本轮仅新增Random_SFT。原G0、Vmatched_SFT、Oracle_SFT直接复用已缓存C输出；不重训、不改标签、不改变V0。',
        '', '## 固定设计与选样边界','',
        '三条SFT分支均从原始G0开始，同一批289道A题、每题一条完整回答、1 epoch、37步；学习率1e-4、LoRA r16/alpha32/dropout0.05、训练种子42及原优化实现一致。',
        '随机组候选来自已有G0_A，每题均匀随机抽一条；只对整批回答loss token总量做±0.5%长度条件接受，选择第一个合格组合。种子20261024，最多100000次；失败则保留第一次均匀抽样，明确报告预算不匹配。选样不读取V分数或正确性标签，标签仅在selected/tensors哈希提交之后关联作质量记录。',
        f'本次抽样次数{summary["random_draws"]}；回答监督token为{summary["supervised_tokens"]}，目标225565，相差{summary["relative_token_difference"]*100:+.3f}%。长度匹配是否通过：{summary["length_matched"]}。',
        '**随机性限制：这是给定总体长度预算的条件随机抽样，不是无条件逐题均匀抽样。**同一题单此前已条件化为“G0至少有一个正确候选且原V路线选中过”，不能将本对照外推到全部A，也不能称为完全不涉及真值的题目抽样。回答选择本身不使用真值。',
        '不同回答的逐题长度和过程内容仍可能不同；总体token匹配不等于梯度方向或每题监督长度相同。完整正确标签仍不保证全过程正确。独立探针Compute 17 + 23.，没有B/C/T进入训练。',
        '', '## 实际训练数据','',
        table(['模型','正确/错误/未知','回答loss token','优化器步'],[[g,f'{s["correct"]}/{s["wrong"]}/{s["unknown"]}',s['supervised_tokens'],s['optimizer_steps']] for g,s in training_summaries(cfg).items()]),
        '## C95开发结果','',
        '沿用原Transformers BF16采样器、相同题目种子和每题4候选、上下文4096/max_new≤3840。V0参数/提示/A-B似然评分及0.5阈值固定；不以vLLM替换研究引擎。C95排除既有参考疑点整题，另列原C96和reference=1影子视图，未知不当作错误。',
        table(['模型','正确/错误/未知','可核验正确率','正确/全部','未知率','V0 FPR','V0 FNR'],[[g,f'{int(get(g)["correct"])}/{int(get(g)["wrong"])}/{int(get(g)["unknown"])}',pct(get(g)['G_accuracy_verified']),pct(get(g)['G_accuracy_lower_all']),pct(get(g)['unknown_fraction']),pct(get(g)['FPR']),pct(get(g)['FNR'])] for g in MODELS]),
        '![四分支指标](figures/01_four_arm_metrics.png)','',
        '## 相对随机组与G0的差值','',
        table(['比较','可核验正确率差/百分点','成对95%区间/百分点'],[[r['contrast'],f'{100*r["estimate"]:+.2f}',f'[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]'] for r in deltas if r['scenario']==primary and r['metric']=='G_accuracy_verified']),
        '区间按整题配对重采样2000次，四候选共同进入；探索性点位95%区间、未做多重比较校正，也不包含重新训练种子的变异。C已经用于前序诊断，本轮没有使用独立新测试集，因此不能作最终测试结论。',
        '随机SFT与G0比较反映本次无质量筛选更新的结果；V/Oracle与随机组的差值用于检查筛选的增量作用。若区间跨0，不解释为等效或证明筛选无用；若点估计改善，也不归结为唯一因果机制。',
        '![V0混淆矩阵](figures/02_confusion_matrices.png)','',
        '## 完整性与资源','',
        f'289条训练张量精确重建、随机抽样独立复现、37步与题目顺序核对通过；384条新生成及384条完整V0评分逐记录审计通过。训练{tr["total_seconds_this_invocation"]:.1f}秒，峰值显存{tr["peak_allocated_bytes_this_invocation"]/1024**3:.2f}GiB。生成单4090、4个原采样进程，tmux可恢复。T继续封存。',
        '原正式和Oracle产物哈希保持一致。运行目录runs/random_control_20261024；完整原文candidates_full.jsonl（1536条），实际随机训练质量training_quality_audit.jsonl在运行目录。Excel仅显示层转义XML不支持字符，原文与模型输入不改写。',
        '',f'报告完成UTC时间：{utc()}。']
    (REPORT/'random_control_report.md').write_text('\n\n'.join(lines))
    (REPORT/'index.html').write_text('<meta charset="utf-8"><title>随机SFT对照</title><h1>随机SFT对照</h1><p><a href="random_control_report.md">中文报告</a> · <a href="random_control_review.xlsx">Excel</a> · <a href="random_control_figures.pdf">图表PDF</a></p><img width="100%" src="figures/01_four_arm_metrics.png"><img width="100%" src="figures/02_confusion_matrices.png">')


def training_summaries(cfg):
    oc=json.loads(ORACLE_CONFIG.read_text())
    return {ARM:cfg['summary'],**oc['summaries']}


def bundle():
    shutil.copy2(CONFIG,REPORT/CONFIG.name)
    shutil.copy2(Path(__file__),REPORT/Path(__file__).name)
    shutil.copy2(RUN/'training_quality_audit.jsonl',REPORT/'training_quality_audit.jsonl')
    files=sorted(p for p in REPORT.rglob('*') if p.is_file() and p.name!='bundle_manifest.json')
    hashes={str(p.relative_to(REPORT)):sha(p) for p in files}
    write_json(REPORT/'bundle_manifest.json',{'created_at':utc(),'files_sha256':hashes})
    archive=REPORT.with_suffix('.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(REPORT.rglob('*')):
            if p.is_file():z.write(p,str(Path(REPORT.name)/p.relative_to(REPORT)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        for name,h in hashes.items():
            import hashlib
            assert hashlib.sha256(z.read(REPORT.name+'/'+name)).hexdigest()==h


def pipeline():
    RUN.mkdir(parents=True,exist_ok=True)
    try:
        for stage in ['prepare','train','generate','score','report']:
            subprocess.run([str(ROOT/'scripts/pilot'),'random_control',stage],cwd=ROOT,check=True)
    except Exception as error:
        write_json(RUN/'state.json',{'status':'failed','stage':stage,'failed_at':utc(),
            'error':str(error),'resume_command':'scripts/pilot random_control run',
            'saved_candidates_and_training_checkpoints_retained':True,'T_read':False})
        raise


def main():
    p=argparse.ArgumentParser();p.add_argument('stage',choices=['prepare','train','generate','worker','score','report','run']);p.add_argument('--assignment');a=p.parse_args()
    if a.stage=='worker':worker(a.assignment)
    elif a.stage=='run':pipeline()
    else:globals()[a.stage]()


if __name__=='__main__':main()
