"""Validate the new diagnostics without changing frozen outputs or labels."""
import argparse
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

from .core import ROOT, read_jsonl, write_json, training_example
from .formal_common import RUN, sha, utc, messages, gpu_lock, WorkBudget
from .formal_multigpu_revision import verify_execution_revision
from .explanation_probes import DEST


def loss_check():
    import torch
    from .runtime import load
    target = DEST/'likelihood_loss_validation.json'
    if target.exists(): return json.loads(target.read_text())
    probes = read_jsonl(DEST/'likelihood_inputs.jsonl')
    stored = {(r['probe_id'], r['model_round']): r for r in read_jsonl(DEST/'likelihood_scores.jsonl')}
    # A selected training text and the longest observed probe text check shifting,
    # prompt masking, template suffix and chunk accumulation against the HF loss.
    chosen = [probes[0], max(probes, key=lambda r: len(r['text']))]
    checks=[]
    with gpu_lock():
        budget=WorkBudget('explanation_likelihood_implementation_check')
        model,tok=load('G')
        assert not model.training
        with torch.inference_mode():
            for r in chosen:
                ex=training_example(tok,messages('G',r['question']),r['text'],4096)
                ids=torch.tensor([ex['input_ids']],device='cuda')
                labels=torch.tensor([ex['labels']],device='cuda')
                loss=float(model(input_ids=ids,labels=labels,use_cache=False).loss)
                recorded=stored[r['probe_id'],'G0']['response_nll_per_token']
                difference=abs(loss-recorded)
                assert difference < 1e-5, (loss,recorded)
                checks.append({'probe_id':r['probe_id'],'HF_response_only_loss':loss,
                               'recorded_chunked_nll':recorded,'absolute_difference':difference,
                               'input_tokens':len(ex['input_ids'])})
        assert all(p.grad is None for p in model.parameters())
        budget.charge('official_HF_response_loss_comparison')
        del model,tok
        torch.cuda.empty_cache()
    result={'passed':True,'checked_at':utc(),'examples':checks,'tolerance':1e-5,'training':False,'gradient_created':False}
    write_json(target,result);return result


def main():
    p=argparse.ArgumentParser();p.add_argument('--gpu-loss-check',action='store_true');args=p.parse_args()
    verify_execution_revision()
    if args.gpu_loss_check: loss_check()
    a=json.loads((DEST/'analysis.json').read_text());plan=a['plan']
    assert all(sha(ROOT/name)==h for name,h in plan['formal_sources_sha256'].items())
    run=json.loads((DEST/'likelihood_run.json').read_text())
    assert run['scores_sha256']==sha(DEST/'likelihood_scores.jsonl')
    scores=read_jsonl(DEST/'likelihood_scores.jsonl');probes=read_jsonl(DEST/'likelihood_inputs.jsonl')
    assert len(probes)==158 and len(scores)==474
    ids={(r['probe_id'],g) for r in probes for g in ['G0','G1','G2']}
    assert ids=={(r['probe_id'],r['model_round']) for r in scores}
    old_reports={g:json.loads((RUN/f'train/{g}/report.json').read_text()) for g in ['G1','G2']}
    assert all(r['adapter_fingerprint']==(None if r['model_round']=='G0' else old_reports[r['model_round']]['final_adapter_fingerprint']) for r in scores)
    for split in ['C','B']:
        for scheme in ['output_length','answer_gap','question_length']:
            for contrast in ['G1-G0','G2-G0','G2-G1']:
                rr=[r for r in a['decomposition'] if r['split']==split and r['scheme']==scheme and r['contrast']==contrast]
                vals={r['component']:r['estimate'] for r in rr}
                assert abs(vals['total']-vals['composition']-vals['within_stratum'])<1e-12
    loss_validation=json.loads((DEST/'likelihood_loss_validation.json').read_text())
    assert loss_validation['passed']
    reviewed=read_jsonl(DEST/'semantic_review_initial18.jsonl')
    assert len(reviewed)==18 and all(r['literal_evidence'] in r['text'] and r['human_review_status']=='pending' and r['human_error_type'] is None for r in reviewed)
    figures=sorted((DEST/'figures').glob('*.png'));assert len(figures)==10
    for path in (DEST/'figures').glob('*.svg'):ET.parse(path)
    # Each page dictionary is a PDF leaf; /Pages is the page tree, not a page.
    pages=len(re.findall(rb'/Type\s*/Page\b',(DEST/'explanations_figures.pdf').read_bytes()));assert pages==10
    wb=load_workbook(DEST/'explanations_review.xlsx');assert len(wb.sheetnames)==16
    formula_cells=[f'{ws.title}:{c.coordinate}' for ws in wb for row in ws for c in row if c.data_type=='f'];assert not formula_cells
    testlog=(DEST/'tests.log').read_text();assert '115 passed' in testlog
    src=[ROOT/'src/pilot'/n for n in ['explanation_analysis.py','explanation_probes.py','explanation_review.py','explanation_visuals.py','explanation_validate.py']]
    result={'passed':True,'checked_at':utc(),'formal_source_files_unchanged':len(plan['formal_sources_sha256']),
            'frozen_execution_source_chain_passed':True,'T_read':False,'new_training':False,
            'likelihood_examples':len(probes),'likelihood_scores':len(scores),'likelihood_model_fingerprints_match':True,
            'likelihood_against_official_HF_loss':loss_validation,
            'decomposition_identity_verified_all18_contrasts':True,'semantic_initial_reviews':18,'human_reviews_pending':18,
            'literal_review_evidence_checked':True,'figures':len(figures),'pdf_pages':pages,'workbook_sheets':len(wb.sheetnames),
            'excel_formula_cells':formula_cells,'pytest_passed':115,
            'analysis_source_sha256':{str(path.relative_to(ROOT)):sha(path) for path in src},
            'limitations':['Semantic error taxonomy has only an assistant-reviewed18-case sample, not full-population labels.',
                           'No new training seeds or clean-feedback randomized training controls; evidence is descriptive.',
                           'Reference=1 is an explicit sensitivity overlay pending review; original labels unchanged.']}
    write_json(DEST/'validation.json',result)
    print(json.dumps({'validation':'passed','formal_sources_unchanged':len(plan['formal_sources_sha256']),'pdf_pages':pages,'tests_passed':115}))


if __name__=='__main__':main()
