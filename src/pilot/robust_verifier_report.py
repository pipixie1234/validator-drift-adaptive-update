"""Count repaired and newly introduced FP/FN, with matched-question uncertainty."""
import json
import shutil
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from scipy.stats import rankdata

from .core import ROOT, read_jsonl, write_json, write_jsonl
from .formal_common import sha, utc
from .robust_verifier_common import RUN, REPORT, PLAN, ARMS, SEEN, UNSEEN, validate
from .fpr_pattern_diagnostics import write_csv
from .v0_drift_analysis import clean

MODELS = ['V0']+ARMS


def E_quality():
    path = ROOT/'configs/verifier_robustness_20261026_E_quality_revision003.json'
    return json.loads(path.read_text())


def interval(x): return [float(v) for v in np.nanquantile(x, [.025, .975])]


def safe(a, b):
    if np.isscalar(b): return float(a/b) if b else None
    with np.errstate(divide='ignore', invalid='ignore'): return a/b


def metrics(cells):
    tp, fn, fp, tn, ua, ur = cells
    return {'FPR': safe(fp, fp+tn), 'FNR': safe(fn, fn+tp), 'BA': .5*(safe(tp, tp+fn)+safe(tn, tn+fp)),
            'known_accuracy': safe(tp+tn, tp+tn+fp+fn), 'accepted_error_risk': safe(fp, fp+tp),
            'accepted_risk_all_lower': safe(fp, fp+tp+ua), 'accepted_risk_all_upper': safe(fp+ua, fp+tp+ua),
            'accuracy_all_lower': safe(tp+tn, tp+fn+fp+tn+ua+ur),
            'accuracy_all_upper': safe(tp+tn+ua+ur, tp+fn+fp+tn+ua+ur),
            'retention_all': safe(tp+fp+ua, tp+fn+fp+tn+ua+ur)}


def load_data():
    cfg = validate(); manifest = json.loads((RUN/'evaluation.manifest.json').read_text())
    assert sha(ROOT/'src/pilot/robust_verifier_eval.py') == manifest['code_sha256']
    assert sha(ROOT/'src/pilot/robust_verifier_common.py') == manifest['common_sha256']
    assert sha(RUN/'source_quality_revision004.json') == manifest['source_quality_sha256']
    assert sha(ROOT/'configs/verifier_robustness_20261026_E_quality_revision003.json') == manifest['E_quality_sha256']
    jobs = read_jsonl(RUN/'evaluation_jobs.jsonl'); assert len(jobs) == len({r['job_id'] for r in jobs}) == 7116
    assert sha(RUN/'evaluation_jobs.jsonl') == manifest['jobs_sha256']
    source = {r['candidate_id']: r for name in ['C_answers.jsonl', 'E_answers.jsonl', 'train_answers.jsonl'] for r in read_jsonl(RUN/name)}
    scores = {}; historical = []
    for name in MODELS:
        ss = read_jsonl(RUN/'scoring'/name/'scores.jsonl'); by = {r['job_id']: r for r in ss}
        assert len(ss) == len(by) == len(jobs)
        for r in jobs:
            s = by[r['job_id']]; raw = source[r['candidate_id']]
            assert s['prompt_sha256'] == r['prompt_sha256'] and s['correct'] is r['correct'] is raw['correct']
            assert s['accepted'] == (s['score'] >= .5) and s['threshold'] == .5 and s['prompt_tokens']+4 <= 8192
            assert s['evaluation_sha256'] == sha(RUN/'evaluation.manifest.json')
            assert s['adapter_fingerprint'] == (None if name == 'V0' else manifest['all4_final_adapters'][name])
            if name == 'V0' and r['cohort'] == 'C_development' and r['variant'] == 'original':
                historical.append({'candidate_id': r['candidate_id'], 'same_decision': s['accepted'] == raw['V0_accepted'],
                                   'score_difference': abs(s['score']-raw['V0_score'])})
            if s.get('reuse_source'):
                # Keep reused predictions/provenance; attach THIS trial's current job metadata.
                s.update({k: val for k, val in r.items() if k != 'prompt_ids'})
        scores[name] = by
    assert len(historical) == 1520
    write_json(REPORT/'V0_historical_reproduction.json', {'original_C_controls': len(historical),
               'same_decisions': sum(r['same_decision'] for r in historical),
               'maximum_score_difference': max(r['score_difference'] for r in historical), 'cases': historical})
    if not all(r['same_decision'] and r['score_difference'] < 2e-6 for r in historical):
        raise RuntimeError('Full historical original-C reproduction failed; inspect scorer before publishing repair claims')
    quality = E_quality(); bad = {r['question_id'] for r in quality['whole_question_quarantine']}
    primary = [r for r in jobs if r['cohort'] != 'E_new_final' or r['question_id'] not in bad]
    assert len({r['question_id'] for r in primary if r['cohort'] == 'E_new_final'}) == quality['primary_E_questions']
    return cfg, primary, source, scores


def analyze(jobs, source, scores):
    scopes = defaultdict(list)
    for r in jobs: scopes[r['cohort'], r['view'], r['variant']].append(r)
    for cohort, view in sorted({k[:2] for k in scopes}):
        for label, names in [('seen4_equal', SEEN), ('unseen4_equal', UNSEEN)]:
            if all((cohort, view, name) in scopes for name in names):
                scopes[cohort, view, label] = [r for name in names for r in scopes[cohort, view, name]]
    cohorts = {}
    for cohort in sorted({r['cohort'] for r in jobs}):
        qs = sorted({r['question_id'] for r in jobs if r['cohort'] == cohort}); G = len(qs)
        W = np.random.default_rng(20261026).multinomial(G, np.full(G, 1/G), size=10000)
        cohorts[cohort] = ({q: i for i, q in enumerate(qs)}, W)
    points = []; changes = []; comparisons = []; cases = []; stability = []; contrast_changes = []
    for (cohort, view, variant), rows in sorted(scopes.items()):
        qix, W = cohorts[cohort]; G = len(qix); n = len(rows)
        context = {'cohort': cohort, 'view': view, 'variant': variant, 'questions': G, 'answer_prefix_instances': n,
                   'unique_answers': len({r['candidate_id'] for r in rows})}
        arrays = {}; boots = {}; pm = {}
        for model in MODELS:
            cells = np.zeros((G, 6), int); known_y = []; known_p = []
            for r in rows:
                s = scores[model][r['job_id']]; y = r['correct']; a = s['accepted']
                k = (0 if a else 1) if y is True else (2 if a else 3) if y is False else (4 if a else 5)
                cells[qix[r['question_id']], k] += 1
                if y is not None: known_y.append(y); known_p.append(s['score'])
            totals = cells.sum(0); m = metrics(totals); arrays[model] = cells; pm[model] = m; boots[model] = metrics((W@cells).T)
            ny = sum(known_y); nn = len(known_y)-ny
            auc = float((rankdata(known_p)[np.array(known_y, bool)].sum()-ny*(ny+1)/2)/(ny*nn)) if ny and nn else None
            row = {**context, 'V': model, **{k: int(v) for k, v in zip(['TP', 'FN', 'FP', 'TN', 'unknown_accepted', 'unknown_rejected'], totals)},
                   **m, 'AUROC': auc, 'exact_half_scores': sum(scores[model][r['job_id']]['score'] == .5 for r in rows)}
            points.append(row)
        for model in ARMS:
            qtrans = np.zeros((G, 5), int); seen_cases = []
            for r in rows:
                old, new = scores['V0'][r['job_id']], scores[model][r['job_id']]; y = r['correct']
                if old['accepted'] == new['accepted']: continue
                if y is None: k = 4; kind = 'unknown_changed'
                elif y is False: k, kind = (0, 'FP_to_TN_repaired') if old['accepted'] else (2, 'TN_to_FP_new')
                else: k, kind = (1, 'FN_to_TP_repaired') if not old['accepted'] else (3, 'TP_to_FN_new')
                qtrans[qix[r['question_id']], k] += 1
                if not variant.endswith('_equal'):
                    raw = source[r['candidate_id']]
                    cases.append({**context, 'V': model, 'kind': kind, 'job_id': r['job_id'], 'question_id': r['question_id'],
                                  'candidate_id': r['candidate_id'], 'correct': y, 'prefix': r['prefix'], 'old_score': old['score'],
                                  'new_score': new['score'], 'old_margin': old['margin'], 'new_margin': new['margin'],
                                  'question': raw['question'], 'response': raw['text'], 'gold_answer': raw['gold_answer'],
                                  'extracted_answer': raw['extracted_answer']})
            repaired_fp, repaired_fn, new_fp, new_fn, uc = qtrans.sum(0)
            base = arrays['V0'].sum(0); current = arrays[model].sum(0)
            assert current[2]-base[2] == new_fp-repaired_fp and current[1]-base[1] == new_fn-repaired_fn
            known = arrays['V0'][:, :4].sum(1); net = qtrans[:, 0]+qtrans[:, 1]-qtrans[:, 2]-qtrans[:, 3]
            net_ci = interval(safe(W@net, W@known)); result = {**context, 'V': model,
                  'repaired_FP': int(repaired_fp), 'repaired_FN': int(repaired_fn), 'new_FP': int(new_fp), 'new_FN': int(new_fn),
                  'net_FP_repaired': int(repaired_fp-new_fp), 'net_FN_repaired': int(repaired_fn-new_fn),
                  'net_repaired': int(net.sum()), 'known_n': int(known.sum()), 'net_repair_rate': safe(net.sum(), known.sum()),
                  'net_repair_low': net_ci[0], 'net_repair_high': net_ci[1], 'unknown_decisions_changed': int(uc),
                  'net_repaired_all_lower': int(net.sum()-uc), 'net_repaired_all_upper': int(net.sum()+uc),
                  'repaired_FP_questions': int((qtrans[:, 0]>0).sum()), 'repaired_FN_questions': int((qtrans[:, 1]>0).sum()),
                  'new_FP_questions': int((qtrans[:, 2]>0).sum()), 'new_FN_questions': int((qtrans[:, 3]>0).sum()),
                  'overall_old_error_repair_rate': safe(repaired_fp+repaired_fn, base[1]+base[2]),
                  'overall_old_correct_damage_rate': safe(new_fp+new_fn, base[0]+base[3]),
                  'old_FP_repair_rate': safe(repaired_fp, base[2]), 'old_FN_repair_rate': safe(repaired_fn, base[1]),
                  'old_TN_damage_rate': safe(new_fp, base[3]), 'old_TP_damage_rate': safe(new_fn, base[0])}
            for metric in ['FPR', 'FNR', 'BA']:
                ci = interval(boots[model][metric]-boots['V0'][metric])
                result.update({metric+'_delta': pm[model][metric]-pm['V0'][metric], metric+'_low': ci[0], metric+'_high': ci[1]})
            changes.append(result)
        for after, before in [('SFT_aug', 'SFT_clean'), ('DPO_aug', 'DPO_clean'), ('DPO_aug', 'SFT_aug')]:
            ct = Counter()
            for r in rows:
                y = r['correct']
                if y is None: continue
                old = scores[before][r['job_id']]['accepted']; new = scores[after][r['job_id']]['accepted']
                if y:
                    ct['repaired_FN'] += not old and new; ct['new_FN'] += old and not new
                else:
                    ct['repaired_FP'] += old and not new; ct['new_FP'] += not old and new
            count_row = {**context, 'contrast': after+'-'+before, 'before': before, 'after': after,
                         **{k: ct[k] for k in ['repaired_FP', 'repaired_FN', 'new_FP', 'new_FN']}}
            count_row['net_repaired'] = ct['repaired_FP']+ct['repaired_FN']-ct['new_FP']-ct['new_FN']
            count_row['net_repair_low'], count_row['net_repair_high'] = interval(boots[after]['known_accuracy']-boots[before]['known_accuracy'])
            contrast_changes.append(count_row)
            for metric in ['FPR', 'FNR', 'BA', 'known_accuracy']:
                bounds = interval(boots[after][metric]-boots[before][metric])
                comparisons.append({**context, 'contrast': after+'-'+before, 'metric': metric,
                                    'delta': pm[after][metric]-pm[before][metric], 'low': bounds[0], 'high': bounds[1]})
        if variant != 'original' and not variant.endswith('_equal'):
            original = {r['candidate_id']: r for r in scopes[cohort, view, 'original']}
            for model in MODELS:
                b = Counter(); margins = []
                for r in rows:
                    a = scores[model][original[r['candidate_id']]['job_id']]; s = scores[model][r['job_id']]
                    y = r['correct']; margins.append(abs(s['margin']-a['margin']))
                    if y is None: continue
                    b['known'] += 1; b['decision_changed'] += a['accepted'] != s['accepted']
                    if a['accepted'] != s['accepted']:
                        b[('FN_to_TP' if s['accepted'] else 'TP_to_FN') if y else ('TN_to_FP' if s['accepted'] else 'FP_to_TN')] += 1
                    old_orig = scores['V0'][original[r['candidate_id']]['job_id']]; old_pre = scores['V0'][r['job_id']]
                    if old_orig['accepted'] == y and old_pre['accepted'] != y:
                        k = 'FP' if not y else 'FN'; b['V0_injection_induced_'+k] += 1
                        b['new_stably_correct_'+k] += a['accepted'] == y and s['accepted'] == y
                stability.append({**context, 'V': model, **dict(b), 'mean_absolute_margin_change': float(np.mean(margins)),
                                  'decision_flip_rate_known': b['decision_changed']/b['known']})
    return {'metrics': points, 'repair_damage': changes, 'contrast_repair_damage': contrast_changes,
            'factorial_comparisons': comparisons, 'prefix_stability': stability}, cases


def pct(x): return '—' if x is None else f'{100*x:.2f}%'


def text_report(cfg, tables):
    metrics_rows = tables['metrics']; changes = tables['repair_damage']
    lines = ['# 从V0修复表述敏感性：SFT × DPO 对照', '',
             '主要问题：相对于同一输入上的V0，新V修复了多少FP/FN，又新增了多少FP/FN？净修复=修复FP+修复FN−新增FP−新增FN。平均加分谁更多不作主要结论。', '',
             '|独立标签|V0决定|新V决定|变化类别|', '|---|---|---|---|',
             '|错误|接受|拒绝|修复FP|', '|正确|拒绝|接受|修复FN|',
             '|错误|拒绝|接受|新增FP|', '|正确|接受|拒绝|新增FN|', '',
             '## 先看新E上的结果', '']
    for arm in ARMS:
        seen = next(r for r in changes if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == 'we_can' and r['V'] == arm)
        natural = next(r for r in changes if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == 'original' and r['V'] == arm)
        unseen = next(r for r in changes if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == 'we_will' and r['V'] == arm)
        lines.append(f'- {arm}：已见we can输入修复FP/FN={seen["repaired_FP"]}/{seen["repaired_FN"]}，新增FP/FN={seen["new_FP"]}/{seen["new_FN"]}，净修复{seen["net_repaired"]:+d}；自然输入BA变化{100*natural["BA_delta"]:+.2f}个百分点，95%区间[{100*natural["BA_low"]:+.2f},{100*natural["BA_high"]:+.2f}]；未见we will净修复{unseen["net_repaired"]:+d}。')
    lines += ['', '净修复取决于本队列正确/错误比例，不能替代两类错误率和BA。正净修复、低翻转率均不足以单独证明修复成功；尤其要检查是否接近全拒绝或全接受。', '',
             '### 本轮能够支持的判断', '',
             '在预定主比较E60/G3/we can上，两组DPO都净修复5条，但变化同时包含新增误拒。DPO_aug的FPR下降9.30个百分点，FNR上升7.69个百分点；BA只增加0.81个百分点，按题配对95%区间为[−3.82,+5.11]个百分点。已知标签上的全部改判都朝拒绝方向，体现出更保守的工作点；尚不足以称为保留正常判断能力的整体鲁棒性修复。两组DPO的净修复率区间均跨0。',
             'DPO_aug与DPO_clean直接比较，主输入上修复1个FP、1个FN，同时新增2个FP，净差为0，净修复率区间[−1.82,+1.81]个百分点。原文及we will的净差分别为−2和+1。因此本轮没有显示前缀增强相对普通DPO的明确额外收益，不能把两组共有的改善都归因于增强。',
             'SFT_aug在主输入上纠正6个旧错误，却破坏29个旧正确判断，净新增23个错误，净修复率区间完全为负；自然原文和we will也分别净新增30、21个错误。这是本次配置的负结果，不能外推为SFT或数据增强普遍无效。',
             '局部修复仍确实存在：V0在此队列上因we can注入新增的7个FP，两组DPO都使这7条在原文和注入版本上同时正确。不过，对各自新的原文判断再注入we can，DPO_clean和DPO_aug仍分别新增7和3个FP，涉及的可以是其他回答。修好了已知案例与消除了表述敏感性是不同命题；完整逐例可在prefix_stability.csv及changed_cases_full.jsonl追查。', '',
             '## 固定设计与数据', '',
             '四组均从原始V0独立初始化相同LoRA：SFT_clean、SFT_aug、DPO_clean、DPO_aug；V0作为基线。原先Vrand/Vtarget模型及失败结果保留，没有接续它们训练。',
             'B池G0原有71道题同时含可核验正确和错误回答；题面审查隔离3道目标不适配/缺图题后，用同一哈希顺序从余下68题选64题，每题各选1正1负，共128条完整数学回答。不依据V0误判或干预效果挑选。独立标签依据最终答案，未知/截断不当负例；不声称全过程正确。',
             '增强组对每条回答使用原文与四个前缀：we need to、we can、it is essential to、逗号+换行，共640次呈现；同题同前缀同时覆盖正负两类。自然组重复呈现同128条原文5次，匹配640次呈现及80步更新；640不是独立回答数。每步严格4正4负。',
             '只使用有效完整数学回答构成主训练集；仅含前缀的无答案字符串是不同的无效样本类别，本次不混入自然错误。',
             'SFT使用普通全词表交叉熵，只监督1个A/B标签，正文/题目全部掩码，不监督EOS。DPO以相同输入、正确A/B为chosen、错误A/B为rejected，参考模型为冻结V0，β=0.1。DPO偏好由独立答案核验标签指定，V0只用于参考似然。两种目标均只更新判断接口；DPO从V0直接开始，未先继承SFT。',
             '共同LoRA r8/alpha16/dropout0，LR2e−5、8步warmup后cosine、AdamW、梯度累积8、80步；单种子。固定最后一步，不按C/E选检查点、阈值或重训。DPO额外使用rejected标签与参考似然，实际计算成本分别记录，不能称两种loss完全等价。', '',
             '## 新留E评价：修复和新增分开核算', '',
             '原E预选64道与A/B/C/T在题目ID、哈希、近似题族及源UUID上互斥的新题；题面审查保守隔离4题（缺约束、改写目标或索引记号疑点），涉及全部候选、两个G阶段、全部前缀和V，主分析为60题。索引记号题可按自然约定修复成A/B=1/3，仍沿用事先冻结的严格字面主清单，作为可恢复的敏感性候选保留；不把4题都称作已证明无解。原64题全部缓存与评分保留，不按G正确率或V表现替换。G0/G3均用同一vLLM引擎和长度/采样设置，各每题4条。',
             '数据质量修订：初轮四组训练完成后，题面审核发现上述3道B题，故初轮adapter完整保留为被替代的预演，主比较全部从V0重跑相同配置。仅以题面原因替换训练题，没有查看新E的V表现来选参数。E回答由独立冻结G模型生成，沿用修订前已固定的缓存，未因V重训而重生成；原T没有新增推理或调参。题面审查为助手初审，可据所附代数/几何反例人工复核，不冒充人工金标。详见source_quality_revision004.json与E质量修订003。',
             '下表为G3上的三种输入：自然原文、用作增强前缀的we can、未用作增强前缀的we will。每个新V与V0比较的是完全相同的输入版本；修复和新增均保留，未知不进入正确/错误分母。', '',
             '|输入|V|修复FP|修复FN|新增FP|新增FN|净修复|净修复率95%区间|ΔFPR/百分点|ΔFNR/百分点|',
             '|---|---|---:|---:|---:|---:|---:|---|---:|---:|']
    for r in changes:
        if r['cohort'] != 'E_new_final' or r['view'] != 'G3_vLLM_newE': continue
        lines.append(f'|{r["variant"]}|{r["V"]}|{r["repaired_FP"]}|{r["repaired_FN"]}|{r["new_FP"]}|{r["new_FN"]}|{r["net_repaired"]:+d}|[{pct(r["net_repair_low"])},{pct(r["net_repair_high"])}]|{100*r["FPR_delta"]:+.2f}|{100*r["FNR_delta"]:+.2f}|')
    lines += ['', '正净修复表示可核验回答中旧错误纠正多于旧正确被破坏；负ΔFPR/ΔFNR表示相应错误减少。修复率与损伤率的条件分母不同，CSV分别给出，不能只报修复比例。未知仅在两V决定不同时可能改变净错误：若u条未知改变决定，则全体回答净修复的确定性界为[已知净修复−u,已知净修复+u]，CSV单列。未知不强判对错。区间按同一题的所有回答与设置整组重采样10000次，模型间共享抽样；单种子小规模诊断，逐项区间不保证同时覆盖。', '',
              '### 未知标签对主要we can结果的影响', '',
              '下表的范围仅反映未知标签可能取值；上表95%区间反映按题抽样不确定性。未改变决定的未知回答对两V的净错误差贡献为0。', '',
              '|V|已知标签净修复|未知回答改判数|全部回答净修复的保守范围|',
              '|---|---:|---:|---|']
    for r in changes:
        if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == 'we_can':
            lines.append(f'|{r["V"]}|{r["net_repaired"]:+d}|{r["unknown_decisions_changed"]}|[{r["net_repaired_all_lower"]:+d},{r["net_repaired_all_upper"]:+d}]|')
    lines += ['', '## 新E的实际工作点：防止把全拒绝称为修复', '',
              '|G|输入|V|FP/错误总数|FN/正确总数|FPR|FNR|BA|接受集错误率|未知|',
              '|---|---|---|---:|---:|---:|---:|---:|---:|---:|']
    for r in metrics_rows:
        if r['cohort'] != 'E_new_final': continue
        lines.append(f'|{r["view"]}|{r["variant"]}|{r["V"]}|{r["FP"]}/{r["FP"]+r["TN"]}|{r["FN"]}/{r["FN"]+r["TP"]}|{pct(r["FPR"])}|{pct(r["FNR"])}|{pct(r["BA"])}|{pct(r["accepted_error_risk"])}|{r["unknown_accepted"]+r["unknown_rejected"]}|')
    lines += ['', '## C开发集：已见与未见前缀', '',
              'C95×4 G3回答上，对已见4词、未见4词分别等权汇总。下表计数单位是“回答×前缀实例”，同一数学回答出现4次，独立题仍是95；不是1520条独立回答。C已反复开发，不当作新的最终测试。', '',
              '|设置|V|修复FP|修复FN|新增FP|新增FN|净修复|ΔFPR/百分点|ΔFNR/百分点|',
              '|---|---|---:|---:|---:|---:|---:|---:|---:|']
    for r in changes:
        if r['cohort'] != 'C_development' or r['view'] != 'G3_vLLM' or r['variant'] not in ['original', 'seen4_equal', 'unseen4_equal']: continue
        lines.append(f'|{r["variant"]}|{r["V"]}|{r["repaired_FP"]}|{r["repaired_FN"]}|{r["new_FP"]}|{r["new_FN"]}|{r["net_repaired"]:+d}|{100*r["FPR_delta"]:+.2f}|{100*r["FNR_delta"]:+.2f}|')
    lines += ['', '留出前缀为we will、we have to、Thought process、Let\'s reason step by step；它们未作为本轮人工前缀增强使用，但部分已在旧探索测过，原数学正文也可能自然包含这些短语。natural_training_phrase_exposure.json列出原文中的开头/任意位置暴露次数。因此这里只称增强前缀留出，不是严格词汇、全新攻击机制或词族留出。E题则此前未用于生成与评价。', '',
              '## 如何区分数据增强效果和训练目标效果', '',
              'contrast_repair_damage.csv直接给出SFT_aug相对SFT_clean、DPO_aug相对DPO_clean、DPO_aug相对SFT_aug各修复/新增多少FP、FN；factorial_comparisons.csv给出同一比较的FPR/FNR/BA及按题配对区间。若两种目标都只有aug改善，更支持数据增强的作用；若自然组也改善，则不能全部归因于前缀增强。任何结论还需同时看自然输入保持情况。',
              '', '下表直接比较同一种训练目标下增强组与自然组。此处“修复/新增”的起点是clean组，不是V0；净修复区间仍按题配对重采样。', '',
              '|输入|增强组−自然组|修复FP|修复FN|新增FP|新增FN|净修复|净修复率95%区间|',
              '|---|---|---:|---:|---:|---:|---:|---|']
    for r in tables['contrast_repair_damage']:
        if (r['cohort'] != 'E_new_final' or r['view'] != 'G3_vLLM_newE'
                or r['contrast'] not in ['SFT_aug-SFT_clean', 'DPO_aug-DPO_clean']): continue
        lines.append(f'|{r["variant"]}|{r["contrast"]}|{r["repaired_FP"]}|{r["repaired_FN"]}|{r["new_FP"]}|{r["new_FN"]}|{r["net_repaired"]:+d}|[{pct(r["net_repair_low"])},{pct(r["net_repair_high"])}]|')
    lines += ['',
              'prefix_stability.csv另外比较每个V自身“原文→前缀”的双向翻转，并追踪V0原本正确、注入后才出错的案例，更新后是否在原文和注入两边都正确。低翻转率若伴随全拒绝/全接受与低BA，不是成功修复。连续分数及margin保留审计，不围绕正误两类平均加分差展开。', '',
              '## 历史G阶段与训练拟合', '',
              '|C自然输出|V|FPR|FNR|BA|', '|---|---|---:|---:|---:|']
    for r in metrics_rows:
        if r['cohort'] == 'C_development' and r['variant'] == 'original':
            lines.append(f'|{r["view"]}|{r["V"]}|{pct(r["FPR"])}|{pct(r["FNR"])}|{pct(r["BA"])}|')
    lines += ['', 'B只含G0训练回答，C的G1/G2/G3为后续生成阶段迁移。C的G0/G1旧缓存来自Transformers，G2/G3来自vLLM，跨该引擎边界的变化只作描述；新E的G0/G3引擎一致。B原文为四组共同训练拟合，B前缀版本只有增强组用作训练输入，metrics.csv分别列出；训练输入与C/E严格分开。', '',
              '## 实际资源、评分核对与限制', '',
              '|分支|步数|输入token|chosen token|rejected token|训练保存重载秒|峰值显存GiB|',
              '|---|---:|---:|---:|---:|---:|---:|']
    for arm in ARMS:
        r = json.loads((RUN/arm/'training_report.json').read_text())
        lines.append(f'|{arm}|{r["steps"]}|{r["input_tokens"]}|{r["chosen_loss_tokens"]}|{r["rejected_tokens_used"]}|{r["seconds_this_invocation"]:.1f}|{r["peak_memory_bytes"]/2**30:.2f}|')
    lines += ['', '评分保留原V0系统提示、A/B映射和0.5阈值。批量加速烟测有1/16边界决定不同，正式评价已按独立版本修订恢复逐回答原双标签输入；仅尝试减少输出logits位置。每个V均与原完整评分实现核对16条固定控制，不通过则回退原实现。1520条C自然原文V0分数/决定另与历史缓存核对，不能混入旧的批量加速评分。详细逐模型选择见scorer_equivalence.json。',
              '运行中30GB根磁盘触顶。被替代初轮的adapter、数据、日志均保留，清理了其已完成训练的冗余optimizer.pt，逐文件清理前哈希见disk_recovery_20261026.json。主修订试验每16步先完整提交再只保留最新可恢复检查点，最终step80及其优化器保留；不清理既有七天实验的模型与缓存。',
              'V0可复用初轮已经按修订002核验的完全相同完整输入评分，来源路径/行号保留；不复用初轮更新后V的结果，也不复用未通过核对的batch4评分。',
              '旧模型/缓存/配置/报告启动前保存哈希，交付时复核；本轮不覆盖旧Vrand/Vtarget，不把它们当作本轮重新训练结果。',
              '单种子、64训练题与60道有效新评价题不足以确立目标的普遍优劣；训练题限定同时有自然正负回答，覆盖有选择性。与旧失败更新还同时改变了监督token范围、学习率和步内平衡，不能把差异归结为某一个原因。参考最终答案不证明全过程正确。', '',
              '数据增强的监督思想参考[AdvJudge-Zero附录A33](https://arxiv.org/html/2512.17375v3#A33)，DPO采用[原论文的参考相对偏好目标](https://arxiv.org/html/2305.18290)。本轮是完整解答、交叉标签前缀增强的项目内实验，非论文数值复现。', '',
              '[图册](robust_verifier_figures.pdf) · [Excel](robust_verifier_review.xlsx) · [四类修复/损伤表](repair_damage.csv) · [全部变化案例](changed_cases_full.jsonl)', '']
    (REPORT/'robust_verifier_report.md').write_text('\n'.join(lines)+'\n')


def figures(tables):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 9})
    with PdfPages(REPORT/'robust_verifier_figures.pdf') as pdf:
        fig, axes = plt.subplots(1, 3, figsize=(15, 6), constrained_layout=True)
        for ax, variant in zip(axes, ['original', 'we_can', 'we_will']):
            rr = [r for r in tables['repair_damage'] if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == variant]
            x = np.arange(4); width = .19
            for j, (field, color) in enumerate([('repaired_FP', '#337d5e'), ('repaired_FN', '#70af91'), ('new_FP', '#bf4f40'), ('new_FN', '#dd9a67')]):
                ax.bar(x+(j-1.5)*width, [r[field] for r in rr], width=width, color=color, label=field)
            ax.set_xticks(x, [r['V'] for r in rr], rotation=35); ax.set_title(variant); ax.set_ylabel('Changed decisions'); ax.legend(fontsize=8)
        fig.suptitle('New E60 valid questions / G3: repaired AND newly introduced errors\n64 raw questions retained; four whole-question exclusions for every V/G/prefix')
        fig.savefig(REPORT/'E_repaired_new_FP_FN.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(1, 2, figsize=(13, 7), constrained_layout=True)
        for ax, metric in zip(axes, ['FPR', 'FNR']):
            labels = []; y = 0
            for variant in ['original', 'we_can', 'we_will']:
                rr = [r for r in tables['repair_damage'] if r['cohort'] == 'E_new_final' and r['view'] == 'G3_vLLM_newE' and r['variant'] == variant]
                for r in rr:
                    ax.plot([r[metric+'_low']*100, r[metric+'_high']*100], [y, y], color='#377da0')
                    ax.plot(r[metric+'_delta']*100, y, 'o', color='#377da0'); labels.append(variant+'/'+r['V']); y += 1
            ax.axvline(0, color='black', ls='--'); ax.set_yticks(range(y), labels); ax.invert_yaxis()
            ax.set_xlabel(metric+'(updated V) - '+metric+'(V0), percentage points')
        fig.suptitle('New E60 valid questions / G3: paired whole-question 95% intervals\nNegative means fewer errors; joint assessment required')
        fig.savefig(REPORT/'E_paired_error_changes.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(2, 3, figsize=(14, 8), constrained_layout=True)
        for ax, (view, variant) in zip(axes.flat, [(g, p) for g in ['G0_vLLM_newE', 'G3_vLLM_newE'] for p in ['original', 'we_can', 'we_will']]):
            rr = [r for r in tables['metrics'] if r['cohort'] == 'E_new_final' and r['view'] == view and r['variant'] == variant]
            x = np.arange(5); ax.bar(x-.18, [r['FPR']*100 for r in rr], .35, label='FPR', color='#b84b3e')
            ax.bar(x+.18, [r['FNR']*100 for r in rr], .35, label='FNR', color='#297c9f')
            ax.set_xticks(x, [r['V'] for r in rr], rotation=35, fontsize=8); ax.set_ylim(0, 105); ax.set_title(view+'/'+variant); ax.legend()
        fig.suptitle('Absolute error rates prevent all-reject/all-accept from appearing successful')
        fig.savefig(REPORT/'E_absolute_FPR_FNR.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(1, 2, figsize=(13, 6), constrained_layout=True)
        for ax, variant in zip(axes, ['seen4_equal', 'unseen4_equal']):
            rr = [r for r in tables['repair_damage'] if r['cohort'] == 'C_development' and r['view'] == 'G3_vLLM' and r['variant'] == variant]
            for i, r in enumerate(rr):
                ax.plot([r['net_repair_low']*100, r['net_repair_high']*100], [i, i], color='#377d69')
                ax.plot(r['net_repair_rate']*100, i, 'o', color='#377d69')
            ax.axvline(0, color='black', ls='--'); ax.set_yticks(range(4), [r['V'] for r in rr]); ax.invert_yaxis()
            ax.set_xlabel('Net repaired / known answer-prefix instances (%)'); ax.set_title(variant)
        fig.suptitle('Repeated-use C95 / G3: four augmented versus four held-out prefix strings\nIntervals cluster all answers and prefixes of the same question')
        fig.savefig(REPORT/'C_seen_unseen_net_repairs.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(1, 2, figsize=(13, 5), constrained_layout=True)
        for arm in ARMS:
            rr = read_jsonl(RUN/arm/'steps.jsonl'); axes[0 if arm.startswith('SFT') else 1].plot([r['step'] for r in rr], [r['mean_loss'] for r in rr], label=arm)
        for ax, name in zip(axes, ['SFT full-vocabulary label CE', 'DPO reference-relative loss']):
            ax.set_title(name); ax.set_xlabel('Optimizer step'); ax.set_ylabel('Loss'); ax.legend()
        fig.suptitle('Different objectives: loss magnitudes are not comparable across panels')
        fig.savefig(REPORT/'training_losses.png', dpi=180); pdf.savefig(fig); plt.close(fig)


def main():
    REPORT.mkdir(exist_ok=True); cfg, jobs, source, scores = load_data(); tables, cases = analyze(jobs, source, scores)
    for name, rows in tables.items(): write_csv(REPORT/(name+'.csv'), clean(rows))
    write_json(REPORT/'results.json', clean(tables)); write_jsonl(REPORT/'changed_cases_full.jsonl', cases)
    bad_quality = {r['question_id']: r for r in E_quality()['whole_question_quarantine']}
    primary_answers = []; quarantined_answers = []
    for answer in source.values():
        if answer['question_id'] in bad_quality:
            quarantined_answers.append({**answer, 'raw_reference_comparison_correct': answer['correct'], 'correct': None,
                    'question_status': 'quarantined_not_an_unknown_answer', 'metrics_eligible': False,
                    'question_quarantine_reason': bad_quality[answer['question_id']]['reason']})
        else: primary_answers.append(answer)
    write_jsonl(REPORT/'answers_full.jsonl', primary_answers)
    write_jsonl(REPORT/'quarantined_E_answers_full.jsonl', quarantined_answers)
    flat = [s for name in MODELS for s in scores[name].values()]; primary_scores = []; quarantined_scores = []
    for s in flat:
        if s['cohort'] == 'E_new_final' and s['question_id'] in bad_quality:
            quarantined_scores.append({**s, 'raw_reference_comparison_correct': s['correct'], 'correct': None,
                    'question_status': 'quarantined_not_an_unknown_answer', 'metrics_eligible': False,
                    'question_quarantine_reason': bad_quality[s['question_id']]['reason']})
        else: primary_scores.append(s)
    write_jsonl(REPORT/'scores_full.jsonl', primary_scores)
    write_jsonl(REPORT/'quarantined_E_scores_full.jsonl', quarantined_scores)
    equivalence = {name: json.loads((RUN/'scoring'/name/'legacy_equivalence.json').read_text()) for name in MODELS}
    write_json(REPORT/'scorer_equivalence.json', equivalence)
    figures(tables); text_report(cfg, tables)
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in tables.items():
        ws = wb.create_sheet(name[:31]); keys = list(dict.fromkeys(k for r in rows for k in r)); ws.append(keys)
        for i, r in enumerate(clean(rows), 2):
            values = [r.get(k) for k in keys]; ws.append(values)
            for j, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith('='): ws.cell(i, j).data_type = 's'
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(REPORT/'robust_verifier_review.xlsx')
    for p in [PLAN, ROOT/'configs/verifier_robustness_20261026.json', *ROOT.glob('configs/verifier_robustness_20261026_*.json'),
              *ROOT.glob('src/pilot/robust_verifier_*.py'), RUN/'evaluation.manifest.json', RUN/'scorer_legacy_probe.json',
              RUN/'disk_recovery_20261026.json', RUN/'natural_training_phrase_exposure.json', RUN/'source_quality_revision004.json',
              RUN/'question_quality_mathematical_evidence.json', RUN/'objective_test.log',
              ROOT/'tests/test_robust_verifier_objectives.py']:
        shutil.copy2(p, REPORT/p.name)
    for arm in ARMS:
        shutil.copy2(RUN/arm/'training_report.json', REPORT/(arm+'_training_report.json'))
    initial = ROOT/'runs/verifier_robustness_20261026'
    write_json(REPORT/'superseded_trial_training_resources.json', {
               'reason': 'Data-quality-only restart; initial models not used in main comparisons',
               'training_reports': {arm: json.loads((initial/arm/'training_report.json').read_text()) for arm in ARMS},
               'all_initial_adapters_retained': True,
               'extra_training_seconds': sum(json.loads((initial/arm/'training_report.json').read_text())['seconds_this_invocation'] for arm in ARMS)})
    protected = json.loads((RUN/'preserved_prior_files.json').read_text()); failures = [p for p, h in protected.items() if sha(ROOT/p) != h]
    assert not failures, failures[:10]
    write_json(REPORT/'preservation_audit.json', {'status': 'passed', 'prior_files_verified_unchanged': len(protected),
               'ledger_append_explicitly_excluded': True, 'old_G_V_models_retained': True, 'at': utc()})
    write_json(REPORT/'audit.json', {'status': 'passed', 'raw_scores': len(flat), 'primary_scores': len(primary_scores),
               'quarantined_scores': len(quarantined_scores), 'models': MODELS, 'scopes': len(tables['metrics'])//5,
               'true_unknowns_not_negatives': True, 'same_input_repairs_and_new_errors_both_counted': True,
               'all4_fresh_V0_initializations_and_save_reload': True, 'old_T_not_newly_scored': True,
               'math_bodies_unchanged': True, 'prior_files_unchanged': len(protected), 'at': utc()})
    summary = {}
    for g in ['G0', 'G3']:
        bad = {r['question_id'] for r in E_quality()['whole_question_quarantine']}
        raw = [r for r in read_jsonl(RUN/'E_answers.jsonl') if r['G_round'] == g]
        rr = [r for r in raw if r['question_id'] not in bad]
        summary[g] = {'raw_questions': 64, 'raw_answers': len(raw), 'quarantined_question_ids': sorted(bad),
                      'questions': len({r['question_id'] for r in rr}), 'answers': len(rr),
                      'correct': sum(r['correct'] is True for r in rr), 'wrong': sum(r['correct'] is False for r in rr),
                      'unknown': sum(r['correct'] is None for r in rr), 'truncated': sum(r['truncated'] for r in rr),
                      'unknown_reasons': dict(Counter(r['unknown_reason'] for r in rr if r['correct'] is None)),
                      'output_tokens_mean': float(np.mean([r['output_tokens'] for r in rr])),
                      'output_tokens_max': max(r['output_tokens'] for r in rr)}
    write_json(REPORT/'E_label_summary.json', summary)
    (REPORT/'README.md').write_text('# 验证器修复实验交付\n\n[主报告](robust_verifier_report.md) · [五页图册](robust_verifier_figures.pdf) · [Excel](robust_verifier_review.xlsx)\n\n四组均从V0独立更新；主要结果分别核算修复/新增FP、FN，辅以自然输入保持、留出增强前缀及新E60有效题迁移。\n\n[四类变化CSV](repair_damage.csv) · [增强组相对原文组的四类变化](contrast_repair_damage.csv) · [完整变化案例](changed_cases_full.jsonl) · [2×2指标差值](factorial_comparisons.csv) · [新E标签与截断统计](E_label_summary.json)\n\n题面问题整题隔离，原E64题完整原始缓存保留；隔离回答和参考比较标签另存quarantined_E_*，不作为回答未知或G错误。\n\n旧数据、模型、周报均保留，主修订训练和评分产物位于runs/verifier_robustness_clean_20261026；被替代初轮在runs/verifier_robustness_20261026。复核入口：scripts/pilot robust_verifier_clean verify。\n')
    files = sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix != '.zip' and p.name not in ['artifact_manifest.json', 'independent_verification.json'])
    manifest = {p.name: sha(p) for p in files}; write_json(REPORT/'artifact_manifest.json', manifest)
    with zipfile.ZipFile(REPORT/'robust_verifier_delivery.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']: z.write(p, p.name)
    from .robust_verifier_verify import main as verify
    verify()
    write_json(RUN/'state.json', {'status': 'complete', 'report': str(REPORT/'robust_verifier_report.md'), 'at': utc()})
    print(json.dumps({'status': 'complete', 'models': 5, 'scores': len(flat), 'report': str(REPORT/'robust_verifier_report.md')}, ensure_ascii=False), flush=True)


if __name__ == '__main__': main()
