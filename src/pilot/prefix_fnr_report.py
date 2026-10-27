"""Question-clustered net false-rejection changes under fixed prefix injection."""
import json
import math
import shutil
import zipfile
from pathlib import Path

import numpy as np
from scipy import stats

from .core import ROOT,read_jsonl,write_json,write_jsonl,digest
from .formal_common import messages,sha,utc
from .prefix_fnr import RUN,REPORT,PLAN,validate
from .fpr_pattern_diagnostics import bh,write_csv
from .v0_drift_analysis import clean,interval


def transitions(original_accepted,injected_accepted):
    b=np.asarray(original_accepted,bool);a=np.asarray(injected_accepted,bool)
    rescued=int((~b&a).sum());new=int((b&~a).sum())
    return {'original_FN':int((~b).sum()),'injected_FN':int((~a).sum()),
            'rescued_FN_to_TP':rescued,'new_TP_to_FN':new,'net_added_FN':new-rescued}


def cluster_contrast(y,base,qi,W):
    """Paired answer means, resample entire questions with all sibling answers."""
    y=np.asarray(y,float);base=np.asarray(base,float);G=W.shape[1]
    counts=np.bincount(qi,minlength=G);diff=np.bincount(qi,weights=y-base,minlength=G)
    sums=np.bincount(qi,weights=y,minlength=G);delta=float((y-base).mean())
    den=W@counts;bd=(W@diff)/den;br=(W@sums)/den
    influence=(diff-delta*counts)/len(y)
    se=math.sqrt(G/max(G-1,1)*float(influence@influence))
    p=float(2*stats.t.sf(abs(delta/se),G-1)) if se>1e-14 and G>1 else 1. if abs(delta)<1e-14 else None
    ci=interval(bd);ri=interval(br)
    return {'delta':delta,'low':ci['low'],'high':ci['high'],'rate_low':ri['low'],'rate_high':ri['high'],
            'p_cluster':p,'cluster_SE':se},bd


def load_scores():
    cfg=validate();rr=read_jsonl(RUN/'answers.jsonl');ss=read_jsonl(RUN/'V0_scores.jsonl')
    index={r['candidate_id']:r for r in rr};by={(s['candidate_id'],s['variant']):s for s in ss}
    vv={v['name']:v for v in cfg['variants']}
    assert len(ss)==len(by)==cfg['logical_scores'] and len(index)==986
    for s in ss:
        r=index[s['candidate_id']];v=vv[s['variant']];text=v['prefix']+r['text']
        assert r['correct'] is s['correct'] is True and not r['truncated']
        assert s['question_id']==r['question_id'] and s['split']==r['split'] and s['G_round']=='G3'
        assert s['source_accepted']==r['accepted'] and s['source_score']==r['score']
        assert s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None and s['threshold']==.5
        assert s['accepted']==(s['score']>=.5) and s['prompt_tokens']+4<=8192
        assert s['text_sha256']==digest(text) and s['prompt_sha256']==digest(messages('V',r['question'],text))
        assert s['cache_key']==digest({'plan':sha(PLAN),'candidate_id':r['candidate_id'],'variant':v['name']})
        if v['name']=='original':
            assert s['source']=='new_original_V0_score' and s['accepted']==r['accepted'] and abs(s['score']-r['score'])<2e-6
    return cfg,rr,ss,by


def historical_wrong_rates(cfg):
    """Existing original112 TN answers only, never substitute cross-round cohort."""
    paths=[ROOT/'runs/prefix_paraphrase_20261024/V0_scores.jsonl',ROOT/'runs/fpr_patterns_20261024/EFP_V0_scores.jsonl']
    from .prefix_paraphrase import validate as old_validate,prefix_variants
    old_validate();old_names={v['name']:v['prefix'] for v in prefix_variants()};index={}
    for s in read_jsonl(paths[0]):
        if s['arm']!='prefix':continue
        assert s['correct'] is False and not s['original_accepted']
        index[s['candidate_id'],old_names[s['variant']]]=s
    for s in read_jsonl(paths[1]):
        prefix=s['prefix'];key=(s['candidate_id'],prefix)
        if key not in index:
            index[key]={**s,'split':'A' if s['view']=='A_confirmation' else 'C'}
    out=[]
    for split,n in [('A',64),('C',48)]:
        for v in cfg['variants']:
            rows=[s for (cid,p),s in index.items() if p==v['prefix'] and s['split']==split]
            assert len(rows)==n and len({s['candidate_id'] for s in rows})==n
            assert all(s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None and s['threshold']==.5 for s in rows)
            out.append({'split':split,'variant':v['name'],'responses':n,'original_FP':0,
                        'injected_FP':sum(s['accepted'] for s in rows),'conditional_TN_to_FP_rate':sum(s['accepted'] for s in rows)/n,
                        'population':'Historical G3 TN cohort; not whole-pool FPR; not same questions as all correct answers'})
    write_json(REPORT/'comparison_sources.json',{str(p.relative_to(ROOT)):sha(p) for p in paths})
    return out


def analyze(cfg,rr,by):
    result=[];means=[]
    for split in ['A','C']:
        rows=[r for r in rr if r['split']==split];qs=sorted({r['question_id'] for r in rows});qix={q:i for i,q in enumerate(qs)}
        qi=np.array([qix[r['question_id']] for r in rows]);G=len(qs)
        W=np.random.default_rng(cfg['seed']).multinomial(G,np.full(G,1/G),size=cfg['bootstrap'])
        base=np.array([not r['accepted'] for r in rows],float)
        neutral=np.array([not by[r['candidate_id'],'neutral_control']['accepted'] for r in rows],float)
        all_y=[]
        for v in cfg['variants']:
            ss=[by[r['candidate_id'],v['name']] for r in rows];accepted=[s['accepted'] for s in ss]
            y=1-np.array(accepted,float);point,boot=cluster_contrast(y,base,qi,W)
            control,_=cluster_contrast(y,neutral,qi,W);tr=transitions([r['accepted'] for r in rows],accepted)
            assert tr['injected_FN']-tr['original_FN']==tr['net_added_FN']
            discordant=tr['rescued_FN_to_TP']+tr['new_TP_to_FN']
            row={'split':split,'variant':v['name'],'prefix':v['prefix'],'category':v['category'],'responses':len(rows),'questions':G,
                 **tr,'original_FNR':float(base.mean()),'injected_FNR':float(y.mean()),**point,
                 'FN_rescue_rate':tr['rescued_FN_to_TP']/tr['original_FN'],
                 'TP_damage_rate':tr['new_TP_to_FN']/(len(rows)-tr['original_FN']),
                 'delta_vs_neutral':control['delta'],'neutral_low':control['low'],'neutral_high':control['high'],
                 'p_vs_neutral':control['p_cluster'],
                 'p_McNemar_descriptive':float(stats.binomtest(tr['new_TP_to_FN'],discordant).pvalue) if discordant else 1.,
                 'mean_acceptance_log_odds_change':float(np.mean([s['log_likelihood_A']-s['log_likelihood_B']-
                    by[r['candidate_id'],'original']['log_likelihood_A']+by[r['candidate_id'],'original']['log_likelihood_B'] for r,s in zip(rows,ss)]))}
            result.append(row)
            if v['category']!='control':all_y.append(y)
        mean_y=np.mean(all_y,axis=0);point,_=cluster_contrast(mean_y,base,qi,W)
        means.append({'split':split,'triggers':len(all_y),'original_FNR':float(base.mean()),'mean_injected_FNR':float(mean_y.mean()),
                      **point,'interpretation':'Exploratory equal-weight18-prefix mean on all correct G3 answers; not a random language sample'})
    triggers=[r for r in result if r['category']!='control'];assert len(triggers)==36
    for field,pfield in [('q_BH36_vs_original','p_cluster'),('q_BH36_vs_neutral','p_vs_neutral')]:
        for r,q in zip(triggers,bh([r[pfield] if r[pfield] is not None else np.nan for r in triggers])):
            r[field]=float(q) if np.isfinite(q) else None
    return result,means


def figures(cfg,rows,wrong):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
    vv=cfg['variants'];lookup={(r['split'],r['variant']):r for r in rows}
    with PdfPages(REPORT/'prefix_FNR_figures.pdf') as pdf:
        fig,axes=plt.subplots(1,2,figsize=(13,10),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            rr=[lookup[split,v['name']] for v in vv]
            ax.barh(range(len(vv)),[100*r['injected_FNR'] for r in rr],color=['#999999' if v['category']=='control' else '#2279A4' for v in vv])
            ax.axvline(100*rr[0]['original_FNR'],color='black',ls='--',label='Original FNR')
            for i,r in enumerate(rr):ax.text(100*r['injected_FNR']+.08,i,f'{r["injected_FN"]}/{r["responses"]} ({100*r["injected_FNR"]:.2f}%)',va='center',fontsize=8)
            ax.set_yticks(range(len(vv)),[v['name'] for v in vv]);ax.invert_yaxis();ax.set_xlim(0,max(r['injected_FNR'] for r in rr)*135)
            ax.set_xlabel('False rejection among ALL correct answers (%)');ax.set_title(f'{split}: n={rr[0]["responses"]}, questions={rr[0]["questions"]}');ax.legend()
        fig.suptitle('G3 correct answers: fixed V0, unchanged full answer, same injected prefixes')
        fig.savefig(REPORT/'prefix_FNR_before_after.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,2,figsize=(13,10),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            for i,v in enumerate(vv):
                r=lookup[split,v['name']];q=r.get('q_BH36_vs_original')
                color='#B33D31' if q is not None and q<.05 else '#2279A4'
                ax.plot([r['low']*100,r['high']*100],[i,i],color=color,lw=2);ax.plot(r['delta']*100,i,'o',color=color)
            ax.axvline(0,color='black',ls='--');ax.set_yticks(range(len(vv)),[v['name'] for v in vv]);ax.invert_yaxis()
            ax.set_xlabel('Injected FNR minus original (percentage points)');ax.set_title(split)
        fig.suptitle('Negative = fewer false rejections; paired question-cluster 95% intervals\nRed = BH36 q<0.05; controls shown descriptively')
        fig.savefig(REPORT/'prefix_FNR_paired_changes.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,2,figsize=(13,10),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            rr=[lookup[split,v['name']] for v in vv]
            ax.barh(np.arange(len(vv))-.18,[r['rescued_FN_to_TP'] for r in rr],height=.35,color='#388866',label='FN rescued to TP')
            ax.barh(np.arange(len(vv))+.18,[r['new_TP_to_FN'] for r in rr],height=.35,color='#C45746',label='TP damaged to FN')
            ax.set_yticks(range(len(vv)),[v['name'] for v in vv]);ax.invert_yaxis();ax.set_xlabel('Number of correct answers');ax.set_title(split);ax.legend()
        fig.suptitle('Both directions counted: net added FN = damaged TP minus rescued FN')
        fig.savefig(REPORT/'prefix_FNR_two_way_transitions.png',dpi=180);pdf.savefig(fig);plt.close(fig)


def text_report(cfg,rows,means,wrong):
    lines=['# 前缀注入：正确回答的误拒率补测','','## 范围与指标','',
      '固定G3缓存与原V0参数、提示、Transformers评分和0.5阈值。在每条正确回答前添加与此前相同的18种触发前缀及3项控制，原正文、最终答案与独立正确标签保持。',
      '覆盖冻结清理后A598/C95缓存中全部可核验正确且未截断回答：A839条/276题（原TP774、FN65），C147条/51题（原TP126、FN21）。不按原接受或拒绝、目标词或干预后效果筛选，也不平衡或复制样本。',
      'FNR=被拒的正确回答/全部正确回答。误拒净变化=(新增TP→FN−纠正FN→TP)/正确回答总数。既保留误拒被纠正，也计入正确接受变成误拒，不能只报告纠正率。','',
      '## 1. 注入后FNR与双向翻转','']
    for split in ['A','C']:
        lines += [f'### {split}池','','|前缀|原FNR→注入FNR|原FN→TP纠正数|原TP→FN新增数|净变化/百分点|按题配对95%区间|BH36 q|',
                  '|---|---:|---:|---:|---:|---|---:|']
        for r in rows:
            if r['split']!=split:continue
            q=r.get('q_BH36_vs_original');qtext=f'{q:.4g}' if q is not None else '控制项'
            lines.append(f'|{r["variant"]}|{r["original_FNR"]:.2%}→{r["injected_FNR"]:.2%}|{r["rescued_FN_to_TP"]}|{r["new_TP_to_FN"]}|{100*r["delta"]:+.2f}|[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]|{qtext}|')
    lines += ['','18触发前缀×2题池的36项主要比较采用双侧按题聚类影响函数Student-t检验，BH校正。相对中性前缀另作独立BH36族，见CSV/Excel。原始逐回答McNemar仅作描述：同题多个回答相关，不能把其p当主要证据。区间按题整组重采样10000次；逐项区间不保证同时覆盖，稀疏翻转与全0退化不等于总体无风险。','',
      '## 2. 同前缀与此前错误回答注入并列','','左侧来自原G3错误回答TN注入子集（A64/C48，各每题1条，原误接受为0）；右侧来自本次全体正确回答。二者题目覆盖和选样不同，不能相加成整体准确率，也不能把左侧翻转率称作全池FPR。','',
      '|前缀|A旧TN子集注入误接受比例|A本次正确回答FNR（原→后）|C旧TN子集注入误接受比例|C本次正确回答FNR（原→后）|','|---|---:|---:|---:|---:|']
    lookup={(r['split'],r['variant']):r for r in rows};wl={(r['split'],r['variant']):r for r in wrong}
    for v in cfg['variants']:
        a,c=lookup['A',v['name']],lookup['C',v['name']]
        lines.append(f'|{v["name"]}|{wl["A",v["name"]]["conditional_TN_to_FP_rate"]:.2%}|{a["original_FNR"]:.2%}→{a["injected_FNR"]:.2%}|{wl["C",v["name"]]["conditional_TN_to_FP_rate"]:.2%}|{c["original_FNR"]:.2%}→{c["injected_FNR"]:.2%}|')
    lines += ['','## 3. 结果判断','']
    for split in ['A','C']:
        rr=[r for r in rows if r['split']==split and r['category']!='control']
        decreases=[r['variant'] for r in rr if r['delta']<0];increases=[r['variant'] for r in rr if r['delta']>0]
        sig=[r['variant'] for r in rr if r.get('q_BH36_vs_original') is not None and r['q_BH36_vs_original']<.05 and r['delta']<0]
        neutral=[r['variant'] for r in rr if r.get('q_BH36_vs_neutral') is not None and r['q_BH36_vs_neutral']<.05 and r['delta_vs_neutral']<0]
        lines.append(f'{split}：18前缀中{len(decreases)}个FNR点估计下降、{len(increases)}个上升、{18-len(decreases)-len(increases)}个不变；相对原文通过BH且下降者：{", ".join(sig) or "无"}；相对中性控制通过BH且下降者：{", ".join(neutral) or "无"}。')
    lines += ['','所测18前缀等权平均仅作探索性描述，不代表任意英文前缀的平均效应：','',
              '|池|原FNR|18前缀平均FNR|变化/百分点|95%区间/百分点|','|---|---:|---:|---:|---|']
    for r in means:lines.append(f'|{r["split"]}|{r["original_FNR"]:.2%}|{r["mean_injected_FNR"]:.2%}|{100*r["delta"]:+.2f}|[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]|')
    lines += ['','本次可检验这些注入是否减少G3正确回答上的误拒；即使出现同向变化，也不能直接证明它解释G训练后的自然跨轮变化。本次没有G0/G1/G2正确回答的对应注入轨迹。A为G训练域，C已反复用于开发；最终T仍封存。独立标签只保证参考最终答案一致，不证明全过程正确或所有公共参考均无误。',
      '', '全部双向变化案例保留全文、注入内容、原/新分数和候选ID。绝不只挑支持误拒下降的前缀或样本。','', '## 资源与审计','']
    compute=json.loads((RUN/'compute.json').read_text())
    lines += [f'共986条正确回答×21设置={len(rows)//2*986}条逻辑评分；新增{compute["new_scores_this_invocation"]}、完全相同完整提示复用{compute["reused_this_invocation"]}，GPU阶段{compute["seconds_this_invocation"]:.1f}秒。986条原文均新评分复现原分数与决定，所有最终答案、标签及完整上下文检查通过。',
              '无新G生成/训练，无V更新，B未使用、T未读取。旧配置/模型/报告保持原样，全部新产物位于独立目录。图、3页PDF、Excel和ZIP附完整来源哈希。']
    (REPORT/'prefix_FNR_report.md').write_text('\n'.join(lines)+'\n')


def main():
    cfg,rr,ss,by=load_scores();rows,means=analyze(cfg,rr,by);wrong=historical_wrong_rates(cfg)
    index={r['candidate_id']:r for r in rr};changes=[];per=[]
    for s in ss:
        r=index[s['candidate_id']];before=r['accepted'];after=s['accepted']
        direction='FN_to_TP' if not before and after else 'TP_to_FN' if before and not after else 'unchanged'
        per.append({'split':s['split'],'question_id':s['question_id'],'candidate_id':s['candidate_id'],'variant':s['variant'],
                    'original_accepted':before,'injected_accepted':after,'direction':direction,'original_score':r['score'],
                    'injected_score':s['score'],'correct_final_answer':r['extracted_answer']})
        if direction!='unchanged':
            changes.append({**per[-1],'question':r['question'],'prefix':s['prefix'],'original_response':r['text'],
                            'injected_response':s['prefix']+r['text'],'gold_answer':r['gold_answer']})
    tables={'FNR_summary':rows,'uniform_prefix_means':means,'historical_TN_to_FP':wrong,'per_answer_decisions':per,
            'changed_answers':[{k:v for k,v in r.items() if k not in ['original_response','injected_response','question']} for r in changes]}
    for name,table in tables.items():write_csv(REPORT/(name+'.csv'),clean(table))
    write_jsonl(REPORT/'changed_cases_full.jsonl',changes);write_json(REPORT/'results.json',clean(tables))
    figures(cfg,rows,wrong);text_report(cfg,rows,means,wrong)
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    for name,table in tables.items():
        ws=wb.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in table for k in r));ws.append(keys)
        for row_number,row in enumerate(clean(table),start=2):
            values=[row.get(k) for k in keys];ws.append(values)
            for column,value in enumerate(values,start=1):
                if isinstance(value,str) and value.startswith('='):
                    ws.cell(row_number,column).data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(REPORT/'prefix_FNR_review.xlsx')
    for p in [PLAN,RUN/'compute.json',Path(__file__),ROOT/'src/pilot/prefix_fnr.py']:shutil.copyfile(p,REPORT/p.name)
    write_json(REPORT/'audit.json',{'passed':True,'all986_originals_freshly_rescored_and_matched':True,
               'full_context_and_final_answer_preserved':True,'source_and_input_hashes_verified':True,'fixed_V0':True,
               'original_TP_and_FN_both_included':True,'logical_scores':len(ss),'changed_records':len(changes),
               'changed_unique_answers':len({r['candidate_id'] for r in changes}),'B_used':False,'T_read':False,'at':utc()})
    files=sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix!='.zip' and p.name!='artifact_manifest.json')
    manifest={p.name:sha(p) for p in files};write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'prefix_FNR_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']:z.write(p,p.name)
    import hashlib
    with zipfile.ZipFile(REPORT/'prefix_FNR_review.zip') as z:
        assert z.testzip() is None and all(hashlib.sha256(z.read(n)).hexdigest()==h for n,h in manifest.items())
    write_json(RUN/'state.json',{'status':'complete','report':str(REPORT/'prefix_FNR_report.md'),'at':utc()})
    print(json.dumps(clean({'means':means,'decreases_BH36':[r['variant']+':'+r['split'] for r in rows
          if r.get('q_BH36_vs_original') is not None and r['q_BH36_vs_original']<.05 and r['delta']<0]}),ensure_ascii=False,indent=2),flush=True)


if __name__=='__main__':main()
