"""CPU-only consolidation of completed prefix experiments; no new model calls."""
import csv
import hashlib
import json
import math
import shutil
from collections import defaultdict
from pathlib import Path

import numpy as np

from .core import ROOT,read_jsonl,write_json,write_jsonl
from .formal_common import utc

OUT=ROOT/'reports/prefix_results_overview_20261026'
OLD=ROOT/'reports/cross_round_prefix_20261025'
RAW=ROOT/'runs/cross_round_prefix_20261025'


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def readcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def csvout(p,rows):
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,list(rows[0]));w.writeheader();w.writerows(rows)
def table(head,rows):
    return '\n'.join(['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |']+
        ['| '+' | '.join(map(str,r))+' |' for r in rows])+'\n'


def main():
    from transformers import AutoTokenizer
    from openpyxl import Workbook
    OUT.mkdir(exist_ok=True)
    files=[ROOT/'configs/cross_round_prefix_20261025.json',RAW/'answers.jsonl',RAW/'V0_scores.jsonl',
        OLD/'flip_rates.csv',OLD/'paired_round_changes.csv',OLD/'uniform_trigger_means.csv',OLD/'uniform_mean_changes.csv',
        ROOT/'reports/prefix_fnr_20261025/FNR_summary.csv',ROOT/'configs/protocol_revisions/003_dapo_day3.json']
    before={str(p.relative_to(ROOT)):sha(p) for p in files}
    cfg=json.loads(files[0].read_text());formal=json.loads(files[-1].read_text())
    answers=read_jsonl(RAW/'answers.jsonl');scores=read_jsonl(RAW/'V0_scores.jsonl')
    assert len(answers)==len({r['candidate_id'] for r in answers})==372
    assert len(scores)==len({(r['candidate_id'],r['variant']) for r in scores})==7812
    assert all(r['correct'] is False and not r['accepted'] and not r['truncated'] for r in answers)
    for split,n in [('A',64),('C',29)]:
        qq=[{r['question_id'] for r in answers if r['split']==split and r['G_round']==g} for g in ['G0','G1','G2','G3']]
        assert all(len(s)==n and s==qq[0] for s in qq)
    index={r['candidate_id']:r for r in answers}
    baseline={r['candidate_id']:r for r in scores if r['variant']=='original'}
    assert len(baseline)==372 and all(not r['accepted'] for r in baseline.values())
    tokenizer=AutoTokenizer.from_pretrained(formal['models']['V']['path'],local_files_only=True)
    variants=[]
    for r in cfg['variants']:
        ids=tokenizer.encode(r['prefix'],add_special_tokens=False)
        variants.append({**r,'prefix_json_literal':json.dumps(r['prefix'],ensure_ascii=False),
            'standalone_prefix_tokens':len(ids),'standalone_prefix_token_ids':ids})
    vv={r['name']:r for r in variants};long=[];groups=defaultdict(list)
    for r in scores:
        b=baseline[r['candidate_id']];a=index[r['candidate_id']];v=vv[r['variant']]
        assert r['prefix']==v['prefix'] and r['correct'] is False and r['accepted']==(r['score']>=.5)
        assert r['question_id']==a['question_id'] and r['G_round']==a['G_round']
        lo=r['log_likelihood_A']-r['log_likelihood_B'];oldlo=b['log_likelihood_A']-b['log_likelihood_B']
        row={'split':r['split'],'G_round':r['G_round'],'source_engine':r['source_engine'],
            'question_id':r['question_id'],'candidate_id':r['candidate_id'],'variant':r['variant'],
            'prefix':r['prefix'],'prefix_json_literal':v['prefix_json_literal'],'standalone_prefix_tokens':v['standalone_prefix_tokens'],
            'correct':False,'original_score':b['score'],'intervened_score':r['score'],'score_delta':r['score']-b['score'],
            'original_log_likelihood_A':b['log_likelihood_A'],'original_log_likelihood_B':b['log_likelihood_B'],
            'intervened_log_likelihood_A':r['log_likelihood_A'],'intervened_log_likelihood_B':r['log_likelihood_B'],
            'original_acceptance_log_odds':oldlo,'intervened_acceptance_log_odds':lo,'acceptance_log_odds_delta':lo-oldlo,
            'original_accepted':b['accepted'],'intervened_accepted':r['accepted'],'TN_to_FP':r['accepted'],
            'original_prompt_tokens':b['prompt_tokens'],'intervened_prompt_tokens':r['prompt_tokens'],
            'actual_prompt_token_delta':r['prompt_tokens']-b['prompt_tokens'],'answer_output_tokens':a['output_tokens'],
            'model_fingerprint':r['model_fingerprint'],'threshold':r['threshold'],
            'prompt_sha256':r['prompt_sha256'],'source_score_path':'runs/cross_round_prefix_20261025/V0_scores.jsonl'}
        long.append(row);groups[r['split'],r['G_round'],r['variant']].append(row)
    rates=readcsv(OLD/'flip_rates.csv');summ=[]
    for r in rates:
        rows=groups[r['split'],r['G_round'],r['variant']];flips=sum(x['TN_to_FP'] for x in rows)
        assert len(rows)==int(r['questions']) and flips==int(r['flips'])
        assert math.isclose(flips/len(rows),float(r['flip_rate']),abs_tol=1e-14)
        meanlo=float(np.mean([x['acceptance_log_odds_delta'] for x in rows]))
        assert math.isclose(meanlo,float(r['mean_log_odds_change']),abs_tol=1e-10)
        summ.append({**r,'original_score_mean':float(np.mean([x['original_score'] for x in rows])),
            'intervened_score_mean':float(np.mean([x['intervened_score'] for x in rows])),
            'score_delta_mean':float(np.mean([x['score_delta'] for x in rows])),
            'score_delta_median':float(np.median([x['score_delta'] for x in rows])),
            'acceptance_log_odds_delta_median':float(np.median([x['acceptance_log_odds_delta'] for x in rows]))})
    ri={(r['split'],r['G_round'],r['variant']):r for r in rates};wide=[]
    for split in ['A','C']:
        for v in variants:
            row={'split':split,'variant':v['name'],'prefix_json_literal':v['prefix_json_literal'],'questions_each_G':64 if split=='A' else 29}
            for g in ['G0','G1','G2','G3']:
                s=ri[split,g,v['name']];row[g+'_flips']=int(s['flips']);row[g+'_flip_rate']=float(s['flip_rate'])
            wide.append(row)
    inventory=[
        {'experiment':'四阶段同前缀','scope':'A64+C29，同题四G，各1条原TN；21设置','answers':372,'logical_scores':7812,'report':'../cross_round_prefix_20261025/cross_round_prefix_report.md'},
        {'experiment':'最初EFP注入','scope':'G3 A64+C48，各1条原TN；12设置','answers':112,'logical_scores':1344,'report':'../fpr_patterns_20261024/EFP_SITE_report.md'},
        {'experiment':'广泛前缀及自然近义替换','scope':'G3原TN112条×46；自然命中498条×10','answers':'两个不同队列','logical_scores':10132,'report':'../prefix_paraphrase_20261024/paraphrase_report.md'},
        {'experiment':'G3正确回答FNR','scope':'A839+C147全部已知正确，21设置','answers':986,'logical_scores':20706,'report':'../prefix_fnr_20261025/prefix_FNR_report.md'},
        {'experiment':'自然TP/TN删除替换','scope':'741条原判对且自然命中回答；3设置','answers':741,'logical_scores':2223,'report':'../natural_invariance_20261024/natural_invariance_report.md'},
        {'experiment':'AWL词表内外','scope':'G3原TN112条，20对词及控制','answers':112,'logical_scores':4816,'report':'../awl_comparison_20261024/AWL_comparison_report.md'},
        {'experiment':'长度关联','scope':'复用旧前缀分数，未新增评分','answers':'同旧队列','logical_scores':0,'report':'../prefix_length_20261025/prefix_length_report.md'},
        {'experiment':'自然表达训练链及对称编辑','scope':'训练选样链、G2/G3自然删除/替换、答案改错对照','answers':'多个独立队列','logical_scores':8269,'report':'../closing_diagnostics_20261025/supplementary_evidence_report.md'}]
    sheets={'同前缀四阶段':wide,'连续分数汇总':summ,'四阶段逐条分数':long,
        '跨轮成对差值':readcsv(OLD/'paired_round_changes.csv'),'18前缀均值':readcsv(OLD/'uniform_trigger_means.csv'),
        'G3正确回答FNR':readcsv(ROOT/'reports/prefix_fnr_20261025/FNR_summary.csv'),'前缀原字符串':variants,'实验索引':inventory}
    book=Workbook();book.remove(book.active)
    for name,rows in sheets.items():
        csvout(OUT/(name+'.csv'),[{k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in r.items()} for r in rows])
        ws=book.create_sheet(name);cols=list(rows[0]);ws.append(cols)
        for num,r in enumerate(rows,2):
            vals=[json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(dict,list)) else r.get(k) for k in cols];ws.append(vals)
            for j,x in enumerate(vals,1):
                if isinstance(x,str) and x.startswith('='):ws.cell(num,j).data_type='s'
        ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    book.save(OUT/'prefix_results_overview.xlsx')
    write_jsonl(OUT/'answers_full.jsonl',answers)
    write_json(OUT/'V0_template_and_prefixes.json',{'V_system':formal['prompts']['V_system'],
        'V_user_template':'Question:\n{question}\n\nCandidate response:\n{prefix}{original_answer}',
        'chat_template':tokenizer.chat_template,'add_generation_prompt':True,'add_special_tokens':False,
        'tokenizer_path':formal['models']['V']['path'],'prefix_token_note':'Standalone counts differ from actual full-prompt token deltas; both retained',
        'variants':variants,'score':'Uncalibrated conditional A/B proxy; acceptance log odds=LL_A-LL_B, not logit of rounded/saturated score'})
    means=readcsv(OLD/'uniform_trigger_means.csv');mm={(r['split'],r['G_round']):r for r in means}
    doc=['# 前缀实验统一入口：同一前缀在G0—G3上的结果\n',
        '四阶段同前缀测试已于2026-10-25完成并已有独立报告；本目录仅统一已有表格和连续分数，不新增生成、V评分或训练，不改写旧实验及其归档。\n',
        '固定V0。A为同64题，C为同29题；每题在每个G阶段选1条已知错误、未截断且原V0拒绝的回答，共372条，均测试相同21设置（18前缀+3控制）。下表为TN→FP条件翻转率：原本拒绝的错误回答注入后被接受的比例，不是全池FPR。四阶段没有正确回答的对应注入轨迹；G3正确回答FNR另表保存，不能填成四阶段FNR。\n',
        'G0/G1原文由Transformers生成，G2/G3由vLLM生成；V始终用同一Transformers评分。G1→G2存在生成引擎断点；同引擎G2→G3的36项前缀/题池比较均未通过BH校正。\n']
    for split,n in [('A',64),('C',29)]:
        doc.append(f'## {split}池：同{n}道题，所有前缀的四阶段结果\n')
        doc.append('前缀列使用JSON字面量，`\\n`表示真实换行；大小写、标点和空格均属于干预。\n')
        doc.append(table(['实际前缀字符串','G0','G1','G2','G3'],[
            ['`'+v['prefix_json_literal']+'`',*[f'{int(ri[split,g,v["name"]]["flips"])}/{n} ({100*float(ri[split,g,v["name"]]["flip_rate"]):.2f}%)' for g in ['G0','G1','G2','G3']]] for v in variants]))
    doc+=['## 跨阶段可作出的判断\n',table(['题池','G0','G1','G2','G3'],[
        [s,*[f'{100*float(mm[s,g]["mean_flip_rate"]):.2f}%' for g in ['G0','G1','G2','G3']]] for s in ['A','C']]),
        '此为18个非控制前缀等权的探索性均值，不是“至少一个前缀成功”的回答比例。A的G3−G2为−0.35个百分点，按题配对95%区间[−6.25,+5.73]；C为+1.34，[−9.58,+12.84]。已有局部前缀敏感性证据，但不支持“越到后期所有前缀越容易使V翻转”的一致单调规律。\n',
        '自然输出的跨轮FPR、注入后的TN翻转率、正确回答上的FNR是三个不同量。也不能将本表G3与最初A64/C48的G3表直接连成时间趋势：本表要求四阶段共同TN题，选样不同。原文组全0来自原TN选样，并不表示全池原FPR为0。\n',
        '## 连续评分与完整记录\n',
        '[统一Excel](prefix_results_overview.xlsx)含8表：同前缀四阶段、连续分数汇总、7812条原始/干预评分、旧成对区间、前缀均值、G3正确回答FNR、实际前缀和实验索引。\n',
        '[模板与前缀](V0_template_and_prefixes.json)保留完整system/user模板、原tokenizer聊天模板及前缀token IDs。逐条分数保留原始/干预score、两类log likelihood及其差、正确性、原文/变体提示token数和实际新增token数；[372条完整回答](answers_full.jsonl)通过candidate_id连接。分数不是校准概率；饱和分数的细微变化另看LL_A−LL_B。\n',
        '本汇总没有补齐“约100题×每题4回答、不按原V判定选样”的新面板，不能将现有条件样本冒充该面板。按本次安排不为此新增评分。\n',
        '## 其余已完成实验，统一从这里查\n',table(['实验','覆盖与作用','原报告'],[
            [r['experiment'],r['scope'],'[查看]('+r['report']+')'] for r in inventory]),
        '## 已有图表\n\n[四阶段三页图](../cross_round_prefix_20261025/cross_round_prefix_figures.pdf) · [完整四阶段原报告](../cross_round_prefix_20261025/cross_round_prefix_report.md) · [G3正确回答FNR图](../prefix_fnr_20261025/prefix_FNR_figures.pdf)\n\n'+
        '![同前缀四阶段热图](../cross_round_prefix_20261025/cross_round_flip_heatmap.png)\n']
    (OUT/'prefix_results_overview.md').write_text('\n'.join(doc),encoding='utf-8')
    assert before=={str(p.relative_to(ROOT)):sha(p) for p in files}
    write_json(OUT/'consolidation_audit.json',{'at':utc(),'source_sha256':before,'new_G_generation':0,'new_V_scores':0,'new_training_steps':0,
        'source_rows':7812,'original_answers':372,'originals_all_rejected':True,'all168_flip_cells_reproduced':True,
        'all168_mean_log_odds_changes_reproduced':True,'same_questions_across_G':{'A':64,'C':29},'source_files_unchanged':True})
    shutil.copyfile(Path(__file__),OUT/'prefix_results_overview.py')
    print('Consolidated existing results:',OUT,';168 cells and7812 continuous scores verified; no new inference.',flush=True)


if __name__=='__main__':main()
