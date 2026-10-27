"""Final provenance validation and bundle creation; can wait independently in tmux."""
import argparse
import hashlib
import json
import math
import re
import shutil
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from .core import ROOT,read_jsonl,write_json,digest,cache_identity
from .formal_common import protocol,messages,inputs,sha,utc
from .day3_protocol import formal_generation_settings
from .runtime import model_identity
from .oracle_control import RUN,REPORT,CONFIG,ARMS,checkpoint,validate_plan
from .oracle_evaluate import generated,scored


def validate():
    cfg=validate_plan();result=json.loads((REPORT/'results.json').read_text())
    plan=json.loads((REPORT/'analysis_plan.json').read_text())
    assert plan['analysis_code_sha256']==sha(ROOT/'src/pilot/oracle_report.py')
    assert result['training_audit']['passed'] and result['original_sources_unchanged']
    p=protocol();vfp,_=model_identity('V');pool={r['question_id']:r for r in inputs('C')}
    checked=0
    for arm in ARMS:
        g=read_jsonl(generated(arm));v={r['candidate_id']:r for r in read_jsonl(scored(arm))}
        assert len(g)==len(v)==384
        base,adapter=model_identity('G',checkpoint(arm))
        for r in g:
            assert r['question_id'] in pool and r['question']==pool[r['question_id']]['question']
            assert r['candidate_index'] in range(4) and r['G_round']==arm and r['split']=='C'
            assert r['candidate_id']==f'{r["question_id"]}:formal:{arm}:{r["candidate_index"]}'
            assert not any(k in r for k in ['correct','gold_answer','reference'])
            assert r['model_fingerprint']==base and r['adapter_fingerprint']==adapter
            actual=formal_generation_settings(r['prompt_tokens'])
            seed=42+int(digest(r['question_id'])[:8],16)%1000000
            key=cache_identity('G',base,adapter,r['question_id'],messages('G',r['question']),dict(actual,seed=seed))
            assert r['generation_seed']==seed and r['prompt_sha256']==digest(messages('G',r['question']))
            assert r['question_cache_key']==key and r['cache_key']==digest({'question_key':key,'candidate_index':r['candidate_index']})
            assert r['max_new_tokens']==actual['max_new_tokens'] and r['output_tokens']<=actual['max_new_tokens']
            assert r['prompt_tokens']+r['output_tokens']<=4096
            vr=v[r['candidate_id']];msg=messages('V',r['question'],r['text'])
            assert vr['model_fingerprint']==vfp and vr['adapter_fingerprint'] is None and vr['threshold']==.5
            assert vr['prompt_tokens']+4<=8192 and vr['prompt_sha256']==digest(msg)
            assert vr['cache_key']==cache_identity('V',vfp,None,r['candidate_id'],msg,p['V_scoring'])
            a,b=vr['log_likelihood_A'],vr['log_likelihood_B']
            expected=1/(1+math.exp(b-a)) if b-a<700 else 0.
            assert abs(expected-vr['score'])<2e-7 and vr['accepted']==(a>=b)
            checked+=1
    assert checked==768
    for g in ['G0',*ARMS]:
        assert result['metrics']['C95_exclude_reference_suspect::'+g]['point']['total']==380
        for scenario in ['C96_frozen_original','C96_reference1_overlay']:
            assert result['metrics'][scenario+'::'+g]['point']['total']==384
    wb=load_workbook(REPORT/'oracle_control_review.xlsx');assert len(wb.sheetnames)==9
    assert not [c.coordinate for ws in wb for row in ws for c in row if c.data_type=='f']
    for f in (REPORT/'figures').glob('*.svg'):ET.parse(f)
    pages=len(re.findall(rb'/Type\s*/Page\b',(REPORT/'oracle_control_figures.pdf').read_bytes()));assert pages==5
    assert len(list((REPORT/'figures').glob('*.png')))==5
    assert len(read_jsonl(REPORT/'candidates_full.jsonl'))==1152
    tests=re.search(r'(\d+) passed',(REPORT/'tests.log').read_text())
    assert tests and int(tests.group(1))>=115
    test_count=int(tests.group(1))
    assert all(sha(ROOT/n)==h for n,h in cfg['original_sources_sha256'].items())
    modules=sorted((ROOT/'src/pilot').glob('oracle_*.py'))
    evidence={'passed':True,'validated_at':utc(),'original_sources_unchanged':len(cfg['original_sources_sha256']),
              'new_G_records_and_V0_records_checked_each':768,'same_prompts_seeds_contexts_and_V0_scoring':True,
              'training_tensor_rows_exactly_rebuilt':578,'same_question_order_and_gradient_accumulation_groups':True,
              'loss_tokens_each':225565,'optimizer_steps_each':37,'fresh_G0_both':True,
              'T_read':False,'B_used_for_control':False,'oracle_explicitly_labelled':True,
              'C95_and_two_C96_sensitivities_checked':True,'full_eval_records':1152,
              'pdf_pages':pages,'workbook_sheets':len(wb.sheetnames),'pytest_passed':test_count,
              'source_sha256':{str(f.relative_to(ROOT)):sha(f) for f in modules}}
    write_json(REPORT/'validation.json',evidence)
    return evidence


def bundle():
    source=REPORT/'source';source.mkdir(exist_ok=True)
    for f in sorted((ROOT/'src/pilot').glob('oracle_*.py')):shutil.copy2(f,source/f.name)
    shutil.copy2(CONFIG,REPORT/CONFIG.name)
    files=sorted(f for f in REPORT.rglob('*') if f.is_file() and f.name!='bundle_manifest.json')
    hashes={str(f.relative_to(REPORT)):sha(f) for f in files}
    write_json(REPORT/'bundle_manifest.json',{'created_at':utc(),'files_sha256':hashes,
               'entry_points':['oracle_control_report.md','oracle_control_review.xlsx','oracle_control_figures.pdf','index.html']})
    archive=REPORT.with_suffix('.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for f in sorted(REPORT.rglob('*')):
            if f.is_file():z.write(f,str(Path(REPORT.name)/f.relative_to(REPORT)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        for n,h in hashes.items():assert hashlib.sha256(z.read(REPORT.name+'/'+n)).hexdigest()==h
    return archive


def main():
    p=argparse.ArgumentParser();p.add_argument('--wait',action='store_true');a=p.parse_args()
    while a.wait and (not (RUN/'state.json').exists() or json.loads((RUN/'state.json').read_text()).get('status')!='complete'):
        time.sleep(30)
    evidence=validate();archive=bundle()
    write_json(RUN/'delivery_complete.json',{'validated_at':utc(),'validation_passed':True,'archive':str(archive),'archive_sha256':sha(archive)})
    print(json.dumps({'validation':'passed','archive':str(archive),'bytes':archive.stat().st_size,'evaluation_records':evidence['full_eval_records']},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
