"""CPU charts and Chinese report for the read-only cache follow-up."""
import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path
import numpy as np
from .core import ROOT, read_jsonl, write_json
from .formal_common import sha,utc
from .cache_followup import DEST,MODELS,ARMS,NOTES


def exact_checks():
 from fractions import Fraction
 from math import comb,isqrt
 import sympy as s
 assert 201-53+1==149 and comb(11,5)==462
 rest=sum(d%4==0 and (d-1)%10>=7 for d in range(1,1001));assert rest==100
 squares=[(a,b,c) for a in range(1,10) for b in range(a,10) for c in range(b,10) if a*a+b*b+c*c==81]
 assert squares==[(1,4,8),(3,6,6),(4,4,7)]
 v=Fraction(1,2)
 for _ in range(3):v=1/(2-v)
 assert v==Fraction(4,5)
 a=2**2004-1;b=2**4008-2**2004+1
 assert (a.bit_count(),b.bit_count(),(a+b).bit_count())==(2004,2005,1)
 assert (15**3-15)//3==1120 and 8*7*6*5*4-7*2*6*5*4==5040
 assert all(n%d for n in [509,1019,53,59] for d in range(2,isqrt(n)+1))
 assert 53*59==3127 and 53*61**2==197213 and 9999-1010==8989
 x,a,b,c=s.symbols('x a b c')
 expr=(x-49)/50+(x-50)/49-49/(x-50)-50/(x-49)
 roots=s.solve(s.together(expr).as_numer_denom()[0],x);assert sum(roots)==s.Rational(14702,99)
 assert s.expand((b-c)**3+(c-a)**3+(a-b)**3).subs({a:1,b:1,c:1})==0
 y=s.symbols('y',real=True)
 assert s.expand((1+s.sqrt(3))**3)==10+6*s.sqrt(3)
 write_json(DEST/'exact_checks.json',{'passed':True,'rest_days':rest,'three_squares':squares,
  'continued_fraction':str(v),'binary_popcounts':[2004,2005,1],
  'rational_equation_roots':[str(r) for r in roots],'coefficient_sum_literal_A_target':0,
  'scope':'Explicit local checks; not a full certification of all source references or all proofs.'})


def render(result):
 import matplotlib
 matplotlib.use('Agg')
 import matplotlib.pyplot as plt
 from matplotlib.backends.backend_pdf import PdfPages
 from openpyxl import Workbook,load_workbook
 from .oracle_visuals import excel_safe_value,table,pct
 exact_checks()
 fd=DEST/'figures';fd.mkdir(exist_ok=True)
 primary=[r for r in result['intervals'] if r['scenario'].startswith('C95')]
 features=list(csv.DictReader((DEST/'supervision_per_question.csv').open(encoding='utf-8-sig')))
 pair=list(csv.DictReader((DEST/'supervision_pairwise.csv').open(encoding='utf-8-sig')))
 with PdfPages(DEST/'cache_followup_figures.pdf') as pdf:
  def save(fig,name):
   fig.tight_layout()
   for ext in ['png','svg','pdf']:fig.savefig(fd/f'{name}.{ext}',bbox_inches='tight')
   pdf.savefig(fig,bbox_inches='tight');plt.close(fig)
  fig,axes=plt.subplots(1,2,figsize=(12,4.5))
  for ax,key in zip(axes,['FPR','FNR']):
   rr=[r for r in primary if r['metric']==key]
   for i,r in enumerate(rr):
    value=r['estimate']*100
    ax.errorbar(value,i,xerr=[[value-r['low']*100],[r['high']*100-value]],fmt='o',capsize=5)
   ax.axvline(0,color='black',lw=.8);ax.set_yticks(range(3),[r['contrast'] for r in rr])
   ax.set_title(key+': Random minus other');ax.set_xlabel('Percentage points');ax.grid(axis='x',alpha=.2)
  fig.suptitle('C95 paired question bootstrap: 10,000 draws; exploratory pointwise intervals')
  save(fig,'01_paired_error_rates')
  fig,axes=plt.subplots(1,2,figsize=(12,4.6))
  for arm in ARMS:
   rr=[r for r in features if r['arm']==arm]
   lengths=np.sort([int(r['loss_tokens']) for r in rr]);frac=np.sort([float(r['last_answer_loss_fraction']) for r in rr if r['last_answer_loss_fraction']])
   axes[0].plot(lengths,np.arange(1,len(lengths)+1)/len(lengths),label=arm)
   axes[1].plot(frac,np.arange(1,len(frac)+1)/len(frac),label=arm)
  axes[0].set_xlabel('Supervised tokens per question');axes[1].set_xlabel('Last explicit answer span / response loss tokens')
  for ax in axes:ax.set_ylabel('Empirical CDF');ax.legend(fontsize=8);ax.grid(alpha=.2)
  fig.suptitle('Global token matching does not match local supervision weights')
  save(fig,'02_supervision_distributions')
  fig,axes=plt.subplots(1,2,figsize=(12,4.6))
  for ax,other in zip(axes,ARMS[1:]):
   rr=[r for r in pair if r['contrast']=='Random_SFT-'+other]
   x=[int(r['random_loss_tokens']) for r in rr];y=[int(r['other_loss_tokens']) for r in rr]
   ax.scatter(x,y,s=12,alpha=.5);limit=max(x+y);ax.plot([0,limit],[0,limit],color='gray',lw=.8)
   ax.set_xlabel('Random loss tokens');ax.set_ylabel(other+' loss tokens');ax.set_title('Same 289 questions, one response each')
  save(fig,'03_same_question_supervision')
 wb=Workbook();wb.remove(wb.active)
 paths=[('Paired_error_rates','paired_FPR_FNR.csv'),('Counts_unknown_bounds','point_counts_and_unknown_bounds.csv'),
  ('Question_comparison','question_comparison_C95.csv'),('Decision_flips','same_truth_decision_flips.csv'),
  ('FPR_decomposition','FPR_question_support_decomposition.csv'),('Supervision_summary','supervision_summary.csv'),
  ('Supervision_by_question','supervision_per_question.csv'),('Supervision_pairwise','supervision_pairwise.csv'),
  ('A_feedback_metrics','A_feedback_metrics.csv'),('A_quality_sensitivity','A_training_quality_sensitivity.csv'),
  ('Case_index','reviewed_cases_index.csv')]
 for name,path in paths:
  rows=list(csv.DictReader((DEST/path).open(encoding='utf-8-sig')));ws=wb.create_sheet(name);keys=list(rows[0]);ws.append(keys)
  for row in rows:
   vals=[]
   for k in keys:
    v=row[k]
    if v=='':v=None
    elif v in ['True','False']:v=v=='True'
    elif re.fullmatch(r'-?\d+(?:\.\d*)?(?:[eE][+-]?\d+)?',v):v=float(v) if any(c in v for c in '.eE') else int(v)
    vals.append(excel_safe_value(v,row.get('candidate_id','')))
   ws.append(vals)
   for cell in ws[ws.max_row]:
    if isinstance(cell.value,str):cell.data_type='s'
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for name,rows in [('Reviewed_full',read_jsonl(DEST/'reviewed_cases_full.jsonl')),('A_problem_issues',result['issues'])]:
  ws=wb.create_sheet(name);keys=sorted(set().union(*(set(r) for r in rows)));ws.append(keys)
  for r in rows:
   ws.append([excel_safe_value(r.get(k),r.get('candidate_id','')) for k in keys])
   for cell in ws[ws.max_row]:
    if isinstance(cell.value,str):cell.data_type='s'
  ws.freeze_panes='A2'
 wb.save(DEST/'cache_followup_review.xlsx')
 check=load_workbook(DEST/'cache_followup_review.xlsx');assert len(check.sheetnames)==13
 assert not any(c.data_type=='f' for ws in check for row in ws for c in row)
 summaries=[r for r in result['supervision'] if r['arm'] in ARMS]
 pair_summary=[]
 for other in ARMS[1:]:
  rr=[r for r in pair if r['contrast']=='Random_SFT-'+other]
  dif=np.array([abs(int(r['token_difference'])) for r in rr]);rat=np.array([float(r['random_token_coefficient_over_other']) for r in rr])
  pair_summary.append({'other':other,'same_candidate':sum(r['same_candidate']=='True' for r in rr),
   'mean_abs_token_difference':float(dif.mean()),'max_abs_token_difference':int(dif.max()),
   'coefficient_ratio_min':float(rat.min()),'coefficient_ratio_max':float(rat.max()),
   'ratio_over_2_or_under_half':int(sum(rat>2)+sum(rat<.5))})
 write_json(DEST/'supervision_pairwise_summary.json',pair_summary)
 points=[r for r in result['points'] if r['scenario'].startswith('C95')]
 main=['# 四分支缓存分析及原两轮V0误判初审','',
  '**本轮发现：G正确率没有明确差异，但Random的FPR较低，不能把全部现象概括成随机波动。**独立实现先精确重现原2000次按题区间，再扩至10000次。逐题监督检查没有发现prompt泄漏或最终答案被mask，但总体token匹配没有匹配逐题监督权重。原A池另有缺图及参考目标疑点，必须与正常数学错例区分。',
  '本轮为CPU只读缓存分析：没有训练、生成、新V评分或T读取。完整结果均保留原标签，题目问题只做整题排除的影子敏感性统计。','',
  '## 1. Random与其他分支的FPR/FNR成对区间','',
  table(['Random−其他','指标','差值/百分点','成对95%区间/百分点'],[[r['contrast'],r['metric'],f'{100*r["estimate"]:+.2f}',f'[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]'] for r in primary]),
  table(['C95模型','TP/FN/FP/TN','FPR','FNR','接受比例'],[[r['model'],f'{r["TP"]}/{r["FN"]}/{r["FP"]}/{r["TN"]}',pct(r['FPR']),pct(r['FNR']),pct(r['retention'])] for r in points]),
  '按整题重采样，四候选一起进入，各分支共用题目索引；10000次、种子20261022。区间是探索性点位95%，未作多重比较校正、不含训练种子不确定性。C95排除原参考疑点整题；原C96和reference=1影子视图同样导出，不挑选有利口径。未知不当错例，未知标签识别界另列。',
  'V0参数、提示和0.5阈值完全相同；FPR分母是各G自身的已知错误集合，较低FPR反映该回答分布下更少误接受，不代表V0本身升级或G更强。接受比例下降、错误题组成以及同题错误文本变化都可能有关。',
  table(['Random−其他','共同错题组成项/pp','共同错题类内项/pp','进入/pp','退出/pp','总差/pp'],[[r['contrast'],*[f'{r[k]*100:+.2f}' for k in ['composition','within','entry','exit','total']]] for r in result['question_support_decomposition']]),
  '这项分解精确相加，但每题最多4条错例，类内值粗糙；它是描述，不是语义错误根因或独立因果效应。不同回答上V接受率改变，不等于固定文本上V遗忘。',
  '![配对区间](figures/01_paired_error_rates.png)','',
  '## 2. 同题抽查：过程、长度、解法与错误类型','',
  '95题全部逐题计数与长度导出；目的抽查6道C题、每题Random/Vmatched/Oracle同一候选索引，共18条完整回答。另保留这6题所有四分支×4候选共96条原文。相同candidate_index仅作描述，参数改变后不代表同一采样路径的因果配对。',
  '- 有理方程：Random错误通分后出现明显矛盾，被V拒绝；另外两组推导更连贯，却在约分时漏掉零分子分支，被V接受。三组都错，但错误的可识别性不同。\n- 汽车计数：同为答案11，Random主要推导成立；Vmatched有20=17矛盾，Oracle最后一步缺中间依据。另两组被拒，在当前最终答案任务中仍是FN。\n- 立方根：三组代数推导均通过此次检查，Random仍被拒；半回文也有简短正确推导被拒，不能用过程错误解释所有FN。\n- 最小合数：Random把已分解为53×59的3127说成质数，丢掉正确解；另外两组正确。\n- 三平方：Random/V漏掉(3,6,6)；Oracle最终3正确，但前文反向不等式和排序条件不成立。不能把最终答案Oracle等同于过程Oracle。',
  '每例的原文引用、数学证据、完整回答、判定和长度见[直接初审全文](case_reviews.md)。均标为助手初审，人工复核待审；定向案例不能估计总体错误类型比例。','',
  '## 3. 逐题监督token与loss mask','',
  table(['分支','样本','loss token','最短/中位/最长','prompt泄漏','assistant空洞'],[[r['arm'],r['examples'],r['loss_tokens'],f'{r["min"]}/{r["median"]:.0f}/{r["max"]}',r['prompt_leaks'],r['assistant_holes']] for r in summaries]),
  '逐条重建3个匹配分支及原2轮训练共1868条张量，核对input IDs、shift后的labels、attention、完整assistant目标、显式答案片段和结束token。未见prompt参与损失、assistant中间被静默mask，或最后显式答案未监督。多盒和无标记另计；最后显式片段不是语义正确性或过程认证。',
  '训练microbatch=1，先对每条回答的有效token平均，再按当前累积组实际样本数m平均：\n\nL_step=(1/m) Σ_q [(1/L_q) Σ_j CE(q,j)]。\n\n所以单token的loss系数为1/(mL_q)。同一题选300还是1500 token的回答，每个token系数相差5倍；总token一致并不能消除这种差别。短回答里最终答案片段占比也往往更大。这里是损失代数系数，不是AdamW实际梯度贡献或因果归因。',
  table(['Random相对分支','相同候选数','逐题token绝对差均值/最大','每token系数比范围','超2倍或低于1/2的题数'],[[r['other'],r['same_candidate'],f'{r["mean_abs_token_difference"]:.1f}/{r["max_abs_token_difference"]}',f'{r["coefficient_ratio_min"]:.3f}–{r["coefficient_ratio_max"]:.3f}',r['ratio_over_2_or_under_half']] for r in pair_summary]),
  '289条分为36组×8加末组1条。末条独占完整优化器步，样本loss系数是普通组中单条的8倍；三组题单、顺序、末题相同，这是共同设置，不能单独解释组间差异。原500/501条的末组分别4/5条，也已列出。',
  '![逐题监督分布](figures/02_supervision_distributions.png)','',
  '## 4. 原两轮反馈中FP和FN的可观测原因','',
  '两轮指G0_A经V0筛选后训练G1，以及G1_A经V0筛选后训练G2。共4800条A候选重新汇总；不把C评估误判混成A训练筛选误判。FP是最终错答被接受，FN是最终对答被拒。',
  table(['反馈轮次','视图','FP','FN','FPR','FNR'],[[r['feedback_update'],r['scenario'],r['FP'],r['FN'],pct(r['FPR']),pct(r['FNR'])] for r in result['feedback_metrics']]),
  '每轮直接初审3条FP和3条FN，共12条。FP中有边界多加1、把3天休息缩成1天、忽略排列谷底约束、把32–39日当真实日期、误用几何距离公式、丢掉连分数层级。说明V会接受形式顺畅但遗漏约束/情况的错误。',
  'FN分两类：一类最终答案正确但正文错误或冲突，例如正文推到2005却盒中写1，正文算1575却盒中写1120，或给出的质数配对构造不成立；另一类所检查主要推导成立，例如Jimmy固定座位下6720−1680=5040仍被拒。当前任务只判最终答案，前一类也必须记FN；若改判完整有效解答须另设任务，不能回头改口径。',
  '**原V0只有A/B条件似然，没有理由文本。**本次“原因”限于可见数学问题与失败模式，不能声称V内部一定因某一步错误而拒绝，也不把事后分析当作V真实推理。','',
  '## 5. 新发现的数据问题与敏感性','',
  'A缺图题4b578eb4…只说三幅天平图，却没有提供图或图示关系；两轮实际选中了凭空构造平衡方程的回答。应把整题候选与正常G数学错答分开。系数和题93eb13ae…字面多项式的系数和应为0（代a=b=c=1），冻结参考9有目标/参考冲突，标为疑点，尚未正式修订。',
  table(['反馈轮次','问题','入训条数','属于共同289题'],[[r['feedback_update'],r['issue'],len(r['actual_training_ids']),r['in_common_289']] for r in result['issues']]),
  table(['反馈轮次','影子视图','实际样本','正确/错误/未知','已核验错误率'],[[r['feedback_update'],r['scenario'],r['selected'],f'{r["correct"]}/{r["wrong"]}/{r["unknown"]}',pct(r['verified_error_rate'])] for r in result['training_quality']]),
  '整题排除仅用于本次敏感性汇总；原候选、真值、训练和模型均不重写，既有模型仍接受过这些样本。未对全部A题目作完整质量认证，不能声称只有这两题有问题，也不能用这少量记录解释整条训练轨迹。','',
  '## 交付文件','',
  '- cache_followup_review.xlsx：13张表，区间、逐题监督、两轮A反馈及30条全文初审。\n- cache_followup_figures.pdf及figures/：3张图。\n- case_reviews.md / reviewed_cases_full.jsonl：完整证据与助手初审。\n- A_misclassified_full.jsonl：原两轮所有缓存FP/FN，未冒称全部过程已审查。\n- exact_checks.json / source_integrity.json / validation.json：精确复算与来源完整性。',
  f'完成UTC时间：{utc()}。']
 supplement=DEST/'A_feedback_interval_note.md'
 if supplement.exists():main.append(supplement.read_text())
 (DEST/'cache_analysis_report.md').write_text('\n\n'.join(main))
 (DEST/'index.html').write_text('<meta charset="utf-8"><h1>四分支缓存分析</h1><p><a href="cache_analysis_report.md">报告</a> · <a href="case_reviews.md">案例初审</a> · <a href="cache_followup_review.xlsx">Excel</a> · <a href="cache_followup_figures.pdf">PDF</a></p><img width="100%" src="figures/01_paired_error_rates.png"><img width="100%" src="figures/02_supervision_distributions.png">')


def bundle():
 source=DEST/'source';source.mkdir(exist_ok=True)
 for p in [Path(__file__),ROOT/'src/pilot/cache_followup.py',NOTES]:shutil.copy2(p,source/p.name)
 files=sorted(p for p in DEST.rglob('*') if p.is_file() and p.name!='bundle_manifest.json')
 hashes={str(p.relative_to(DEST)):sha(p) for p in files}
 write_json(DEST/'bundle_manifest.json',{'created_at':utc(),'files_sha256':hashes})
 with zipfile.ZipFile(DEST.with_suffix('.zip'),'w',zipfile.ZIP_DEFLATED) as z:
  for p in sorted(DEST.rglob('*')):
   if p.is_file():z.write(p,str(Path(DEST.name)/p.relative_to(DEST)))
 with zipfile.ZipFile(DEST.with_suffix('.zip')) as z:
  assert z.testzip() is None
  for name,h in hashes.items():assert hashlib.sha256(z.read(DEST.name+'/'+name)).hexdigest()==h
