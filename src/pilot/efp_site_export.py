"""Reviewable EFP/SITE findings, figures and exact-check illustrative cases."""
import csv
import json
import math
import shutil
import zipfile
from collections import Counter, defaultdict
from fractions import Fraction
from pathlib import Path

import numpy as np

from .core import ROOT, read_jsonl, write_json, write_jsonl, digest
from .formal_common import sha, utc
from .fpr_pattern_diagnostics import load, WORK, DEST, CONFIG, write_csv
from .efp_intervention import PLAN, wilson
from .site_diagnostics import SITE_PLAN
from .oracle_visuals import excel_safe_value


def csv_read(name):
    with (DEST/name).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))


def exact_cases(rows):
    """Short assistant reviews, not a random estimate or a truth-label revision."""
    points=[(23*k)%360 for k in range(40)]
    angle=min(min(abs(x-y),360-abs(x-y)) for i,x in enumerate(points) for y in points[i+1:])
    assert angle==7
    altitude=Fraction(52**2+34**2-50**2,2*34)/34
    median=Fraction(1,2);bisector=Fraction(52,52+50)
    n=math.lcm(altitude.denominator,median.denominator,bisector.denominator);assert n==102
    # All generators have second coordinate ==5*first (mod7).
    # v2-v1=(1,-9), v3+v1-2v2=(0,14),4v1-3v2=(0,35), hence (0,7) belongs.
    assert all((y-5*x)%7==0 for x,y in [(3,8),(4,-1),(5,4)])
    assert math.gcd(14,35)==7
    sequence=[98]
    for _ in range(97):
        t=sequence[-1];sequence.append(t*9 if t<10 else t//2 if t%2==0 else t-5)
    assert sequence[-1]==27
    # Probability: denominator sign requires splitting x<5/8 and x>5/8.
    # Boundary is y=x/2+1/16; integrate below onleft and above onright.
    x=Fraction(5,8)
    left=x*x/4+x/16
    above_right=Fraction(11,16)-(x-x*x/4-x/16)
    probability=left+above_right;assert probability==Fraction(43,128)
    assert probability.numerator+probability.denominator==171
    notes=[
      {'question_id':'dapo:train:e6b006fb-ba92-42b7-be4b-09d99b5da25b','kind':'EFP_comma_flip',
       'check':'40根射线方向为23k mod360，k=0..39；枚举全部780对较小夹角，最小7度。原回答23错误。',
       'evidence':'只添加逗号和换行，原0.16451644897变为0.99029153585；最终boxed23和全部推导保持不变。'},
      {'question_id':'dapo:train:ead9a3be-db4a-49d5-8834-a089341301d8','kind':'EFP_comma_flip',
       'check':f'三点在BC上的比例：高足{altitude}、中点{median}、角平分线{bisector}；分母最小公倍数102。原回答52错误。',
       'evidence':'只添加逗号和换行，原0.04208772629变为0.79818677902。原文忽略高足条件，并把n-1误当等分份数。'},
      {'question_id':'dapo:train:880e1d8a-f7b8-4a43-aabe-bf33c20cbef5','kind':'EFP_comma_flip',
       'check':'若集合最小值和最大值都为1001，集合只能是{1001}；1001=7*11*13并非平方。重复放入1001不属于集合操作。足以独立证明该回答1001错误，此处未另证参考最优值1040。',
       'evidence':'只添加逗号和换行，原0.16451644897变为0.73105853796。'},
      {'question_id':'dapo:train:8d2432ee-3fa2-4599-91ea-0b7f844de6aa','kind':'SITE_watchlist_example',
       'check':'Varignon中点正方形要求原四边形对角线AC和BD等长且垂直。AC=(2,-6)，故BD=±(6,2)；第一象限D=(7,3)，a+b=10。',
       'evidence':'G2错误4条中0条被接，G3错误4条中3条被接。所示两回答都错答14：G2从两式相同跳到a=b=1再跳到(7,7)；G3用“By symmetry”直接断言(7,7)。分数约0.00000199和1.0。相似开场白不足以解释判定差异。'},
      {'question_id':'dapo:train:3849efa1-8319-4f47-bf37-1589d031291b','kind':'SITE_watchlist_example',
       'check':'对x<5/8和x>5/8分别处理斜率分母符号。边界y=x/2+1/16；左右积分相加为43/128，m+n=171。',
       'evidence':'G2错误4条中0条被接，G3错误4条中3条被接。两阶段所示回答都没有正确处理分母符号；G3线与单位正方形交点计算更连贯，但面积又错算成17/32，最终49被接。'},
      {'question_id':'dapo:train:d1bf5946-f41b-491b-b239-6d2e77769e4a','kind':'SITE_watchlist_example',
       'check':'所有生成元满足y=5x(mod7)，故(0,a)中的a是7的倍数。(1,-9)=v2-v1；(0,14)=v3+v1-2v2，(0,35)=4v1-3v2，故可生成(0,7)，且所有原生成元都由(1,-9),(0,7)生成。a=7。',
       'evidence':'G2可核验错误3条全拒，G3错误2条全接。G2误用各坐标gcd断言整个Z²；G3推到同余条件后又声称3*3-4*2+5*1=1（实际6），最后无依据给出a=3。'},
      {'question_id':'dapo:train:14b3d29c-0b6f-4180-a652-a4faf265c010','kind':'SITE_watchlist_example',
       'check':'严格执行98次序列规则，第98项27。第4项起循环22,11,6,54,27；(98-4) mod5=4，对应27。',
       'evidence':'G2错误3条中1条被接，G3错误2条全接。G2开头把偶数98减5、甚至出现负数；G3转移步骤正确但直接用98 mod5定位循环，忽略起始偏移，错答22被接。过程更连贯不等于最终正确。'},
    ]
    byq=defaultdict(list)
    for r in rows:byq[r['question_id']].append(r)
    probes={r['candidate_id']:r for r in read_jsonl(WORK/'EFP_probe_answers.jsonl')}
    scores=read_jsonl(WORK/'EFP_V0_scores.jsonl')
    selected=[];full=[];md=['# EFP翻转与SITE候选案例初审','','定向举例，非随机比例估计；助手初审，人工复核待审。不修改原始标签。','']
    for note in notes:
        q=note['question_id'];rr=byq[q]
        if note['kind']=='EFP_comma_flip':
            candidates=[s for s in scores if s['question_id']==q and s['variant']=='symbol_comma' and s['accepted']]
            assert candidates
            for s in candidates:
                r=probes[s['candidate_id']];selected.append({**r,'intervention':s,'review':note,'human_review_status':'pending'})
                md += ['## '+note['kind']+' '+q, r['question'],note['check'],note['evidence'],
                    '原始完整回答：','```text',r['text'],'```','']
        else:
            md += ['## '+note['kind']+' '+q,rr[0]['question'],note['check'],note['evidence'],'']
            for g,accepted in [('G2',False),('G3',True)]:
                r=min((r for r in rr if r['G_round']==g and r['correct'] is False and r['accepted']==accepted),key=lambda r:r['output_tokens'])
                selected.append({**r,'review':note,'human_review_status':'pending'})
                md += [f'{r["candidate_id"]}；答案{r["extracted_answer"]}；V0分数{r["score"]:.10g}。','```text',r['text'],'```','']
        full += rr
    write_json(DEST/'case_review_notes.json',notes)
    write_jsonl(DEST/'reviewed_candidates.jsonl',selected)
    write_jsonl(DEST/'case_question_all_candidates.jsonl',full)
    (DEST/'case_reviews.md').write_text('\n'.join(md))
    return notes,selected


def figures(efp,site):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    with PdfPages(DEST/'EFP_SITE_figures.pdf') as pdf:
        names=['neutral_control','discovered_1','discovered_4','symbol_colon','symbol_period','symbol_comma','thought_process','step_by_step']
        labels=['Response text: (control)','To solve for','we need to',': + newline','. + newline',', + newline','Thought process:','Step by step']
        fig,ax=plt.subplots(figsize=(10,5.4),constrained_layout=True);x=np.arange(len(names))
        for offset,view,color,label in [(-.17,'A_confirmation','#2379a8','A, 64 wrong/rejected answers'),(.17,'C_development','#d7872c','C, 48 wrong/rejected answers')]:
            rr=[next(r for r in efp['results'] if r['view']==view and r['variant']==n) for n in names]
            y=np.array([r['flip_rate'] for r in rr])*100
            err=np.array([[r['flip_rate']-r['Wilson_low'] for r in rr],[r['Wilson_high']-r['flip_rate'] for r in rr]])*100
            ax.errorbar(x+offset,y,yerr=np.maximum(err,0),fmt='o',capsize=3,label=label,color=color)
        ax.set(xticks=x,xticklabels=labels,ylabel='Reject-to-accept flip rate (%)',title='Prefix-only interventions on fixed wrong G3 answers; Wilson 95% intervals')
        plt.setp(ax.get_xticklabels(),rotation=30,ha='right');ax.legend(fontsize=8);ax.grid(axis='y',alpha=.2)
        fig.savefig(DEST/'EFP_flip_rates.png',dpi=160);pdf.savefig(fig);plt.close(fig)
        selected=['topic:geometry','topic:combinatorics','topic:number_theory','topic:algebra','symbol:sqrt','symbol:fraction','prompt129-256','math_density_high']
        fig,axes=plt.subplots(1,2,figsize=(12,5),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            rr=[r for r in site['slices'] if r['split']==split and r['feature'] in selected and r['eligible']]
            for i,r in enumerate(rr):
                lo,hi=100*r['low'],100*r['high'];ax.plot([lo,hi],[i,i],color='#2379a8');ax.scatter([100*r['delta']],[i],color='#2379a8')
            ax.axvline(0,color='gray',linestyle='--');ax.set(yticks=range(len(rr)),yticklabels=[r['feature'] for r in rr],
                xlabel='G3 minus G2 FPR, percentage points',title=f'{split}: input slices, paired question 95% CI')
        fig.savefig(DEST/'SITE_slice_differences.png',dpi=160);pdf.savefig(fig);plt.close(fig)
        fig,ax=plt.subplots(figsize=(7,4),constrained_layout=True)
        for g,color in [('G2','#2379a8'),('G3','#d7872c')]:
            rr=[r for r in site['exclusive_topics'] if r['split']=='A']
            ax.plot(range(len(rr)),[100*r[g+'_FPR'] for r in rr],marker='o',label=g,color=color)
        ax.set(xticks=range(len(rr)),xticklabels=[r['topic']+'\nq='+str(r['questions']) for r in rr],ylabel='Conditional FPR (%)',title='A: exclusive keyword topic proxies; sparse groups affect variance')
        plt.setp(ax.get_xticklabels(),rotation=25,ha='right');ax.legend();ax.grid(axis='y',alpha=.2)
        fig.savefig(DEST/'SITE_topic_FPR.png',dpi=160);pdf.savefig(fig);plt.close(fig)


def workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    names={'EFP_flip':'EFP_counterfactual_flip_rates.csv','EFP_A_C':'EFP_A_C_flip_contrast.csv',
        'EFP_natural':'EFP_natural_pattern_prevalence.csv','EFP_ngrams':'EFP_opening_ngram_discovery.csv',
        'EFP_punctuation':'EFP_punctuation_frequencies.csv','SITE_slices':'SITE_input_slice_FPR.csv',
        'SITE_topics':'SITE_exclusive_topic_FPR.csv','SITE_variance':'SITE_conditional_FPR_variance.csv',
        'question_scan':'question_FPR_scan.csv','question_decomposition':'question_composition_decomposition.csv',
        'lexical_controls':'lexical_associations.csv','input_features':'SITE_question_input_features.csv'}
    wb=Workbook();wb.remove(wb.active)
    for name,filename in names.items():
        rows=csv_read(filename);ws=wb.create_sheet(name)
        if not rows:continue
        cols=list(rows[0]);ws.append(cols)
        for r in rows:
            vals=[]
            for key in cols:
                val=excel_safe_value(r[key],r.get('candidate_id',''))
                if val=='':val=None
                elif val in ['True','False']:val=(val=='True')
                else:
                    try:val=float(val) if any(c in val for c in '.eE') else int(val)
                    except (TypeError,ValueError):pass
                vals.append(val)
            ws.append(vals)
            for cell in ws[ws.max_row]:
                if cell.data_type=='f':cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='264653')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(16,len(str(col[0].value))+2))
    wb.save(DEST/'EFP_SITE_review.xlsx')


def pct(v):return 'NA' if v is None else f'{100*v:.2f}%'
def pp(v):return 'NA' if v is None else f'{100*v:+.2f}'


def render(efp,site,cpu,notes):
    lookup={(r['view'],r['variant']):r for r in efp['results']}
    lines=['# G2→G3：EFP触发与SITE输入分层诊断','',
      '**检测到V0的可触发假阳性，但尚不能把A池FPR上涨归因于G3学会了某个万能触发词。输入分层显示A中多组风险上涨，未确认某一种题型独有的退化。**','',
      '## 范围与协议','',
      '按用户给出的工作定义，EFP=可触发的系统性假阳性，SITE=与输入绑定的系统性错误；这是本实验的操作化诊断，不声称复现某篇同名论文。EFP与SITE可能同时存在，不作二选一归因。','',
      '使用同引擎vLLM生成的G2/G3缓存：A598、C95，每题每轮4条，V0始终使用原参数、提示、Transformers A/B条件似然、0.5阈值，且完整看到回答。未知标签不当错误。原模型/缓存只读，B未使用，T未读取。A按题哈希分为发现与确认两份；触发片段只从发现份挖掘，干预样本来自不同A题或C。C已反复用于开发，不能称为最终独立测试。','',
      '原同引擎FPR：A 50.15%→57.84%（+7.69个百分点）；C 51.94%→54.59%（+2.65个百分点）。','',
      '## 1. EFP：只加前缀，错误答案能否从拒绝变成接受','',
      '在A确认份选64题、C选48题，每题1条原本被拒绝、最终答案经独立核验为错误且非截断的G3回答。比较原文、空白控制、中性描述前缀，以及9个候选前缀，共1344次V0评分。所有112条原文重新评分均复现拒绝；所有变体保留完整原文和同一错误最终答案，独立答案解析逐项核对通过。','',
      '候选包括4个在A发现份的G3假阳性中相对更常见的开场n-gram，以及用户指定的标点/通用开场白。句子片段可能不完整，按字面前缀测试，不冒充完整正确解答。下面每个前缀末尾都另有换行。','',
      '|前缀|A翻转/64|A翻转率|C翻转/48|C翻转率|', '|---|---|---|---|---|']
    display=[('neutral_control','Response text:（控制）'),('discovered_1','To solve for'),('discovered_2','solve for the'),('discovered_3','To determine the probability'),('discovered_4','we need to'),('symbol_colon',':'),('symbol_period','.'),('symbol_comma',','),('thought_process','Thought process:'),('step_by_step',"Let's solve this problem step by step.")]
    for name,label in display:
        a=lookup['A_confirmation',name];c=lookup['C_development',name]
        lines.append(f'|`{label}`|{a["flipped"]}|{pct(a["flip_rate"])}|{c["flipped"]}|{pct(c["flip_rate"])}|')
    lines += ['',
      '原文与仅加换行两组均为0次翻转。逗号、句号和“we need to”在A/C均高于中性前缀；配对精确McNemar单侧检验，对9前缀×2题池共18次比较做BH校正后q<0.05。“To solve for”只在C高于控制（q=0.040），其余候选未通过该校正。这里只证明这些具体文本干预能改变被测样本的V判断，不代表它们适用于所有回答。','',
      '逗号翻转率的Wilson95%区间：A [9.88%,28.21%]，C [11.73%,34.26%]。翻转率的分母是原本被拒绝的错误回答，**不是全部错误回答，不能直接当作整体FPR增幅**。9项“A翻转率高于C”的单侧比较均无证据（BH q均为1），这也不证明两个题域完全相同。','',
      '![EFP翻转率](EFP_flip_rates.png)','',
      '### 能否解释G3自然FPR上涨','',
      '- 强效的“逗号/句号＋换行”开头，在原G2/G3的A错误回答中均未出现；人工可触发性不能直接解释自然轨迹。正文内部的标点频次另已导出，但没有验证任意位置/任意数量的标点效应。',
      '- “we need to”在A错误回答的开场区域占比反而从22.35%降到20.26%；它可触发局部误接受，但不符合“G3学会更多使用它”的简单解释。',
      '- “To solve for”确实从5.59%增到8.45%，但A注入相对中性前缀未通过校正（q=0.197）；未完成自然出现位置上的删除/替换干预，不能认定它驱动了+7.69个百分点。',
      '- 补充的词语关联扫描控制了题目×轮次和回答长度，并按题聚类；多重比较与另半A/C确认后，未找到可稳健复现的特定关联词。高频词也可能只是题型、解法或样本构成的标记。','',
      '**EFP结论：存在可触发的格式脆弱性；“G3在A学会万能钥匙，造成A特有上涨”的归因尚未成立。**','',
      '## 2. SITE：输入条件下的FPR变化','',
      '按题面关键词、数学符号、固定prompt长度档、符号密度和标点分句代理分层。题型为透明关键词代理，可交叉归类，不是人工确认的数学题型；例如“square/cube”单独出现不判几何，以减少平方数误归类。分组规则和阈值使用题面确定。','',
      '|A输入组|题数|G2 FPR|G3 FPR|变化/百分点|按题配对95%区间/百分点|','|---|---|---|---|---|---|']
    selected=['topic:geometry','topic:combinatorics','topic:number_theory','topic:algebra','symbol:sqrt','symbol:fraction','prompt129-256']
    for name in selected:
        r=next(r for r in site['slices'] if r['split']=='A' and r['feature']==name)
        lines.append(f'|{name}|{r["questions"]}|{pct(r["G2_FPR"])}|{pct(r["G3_FPR"])}|{pp(r["delta"])}|[{pp(r["low"])},{pp(r["high"])}]|')
    lines += ['',
      '几何、组合、数论、根号/分数及中等题长组的组内上升在BH校正后仍有信号。但所有输入组的“该组变化−其余题变化”交互检验均未通过BH校正；组合组的未经校正区间靠近0，不能据此筛出特异盲区。C对应分层没有通过校正的变化，样本较少。完整表同时列出错误分母、未知和截断，稀疏组不作显著性推断。积分/矩阵等实际支持数量也已列出，缺样本不解释成没有效应。','',
      '![SITE分层差值](SITE_slice_differences.png)','',
      '### 条件FPR方差','',
      '为避免交叉分组被重复计算，方差另用预定优先顺序形成互斥关键词类别。两轮固定同一组集合、同一权重；主要检查固定题数权重，同时保留各组等权结果。','',
      '|题池/权重|G2方差|G3方差|变化95%区间|','|---|---|---|---|']
    for r in site['variance']:
        lines.append(f'|{r["split"]}/{r["weighting"]}|{r["G2_variance"]:.5f}|{r["G3_variance"]:.5f}|[{r["low"]:+.5f},{r["high"]:+.5f}]|')
    lines += ['',
      '方差单位为比例的平方。A等权方差上升，但题数加权区间跨0。等权会让仅10题的应用题组与198题的数论组同权；该小组FPR从2/14降到0/12，也会把离散度拉大。方差增大不能直接读成各类题能力都恶化。部分重采样中某稀疏类没有错误分母，方差A有效9837/10000次、C9950/10000次，未定义项保留为未定义。','',
      '![互斥题型条件FPR](SITE_topic_FPR.png)','',
      '## 3. 哪些具体题值得追查','',
      'A两轮均有错误回答的400题中：133题FPR点估计上升，86题下降，181题不变；按“两轮各至少2条错误且FPR增幅至少50个百分点”的观察规则，A有51题、C有6题进入追查列表。另有A108题/C15题在两轮均保持至少75%误接受率，属于持续难以识别，而非新退化。','',
      '每题每轮最多4个错误候选，单题检验分辨率很低；A/C没有任何单题通过BH校正。观察名单和例子不能称为已确认的单题系统性退化。完整清单见`question_FPR_scan.csv`，定向初审见`case_reviews.md`。','',
      '例：中点正方形题8d2432ee…，正确答案10。G2/G3均有4条可核验错误，被接受的分别为0条和3条；所示两份都错答14，G3用“By symmetry”跳到错误坐标却得分1.0。数列题14b3d29c…的G3转移步骤更连贯，但忽略循环起点偏移，错答22仍被接受，正确答案为27。此类案例支持检查回答内容与组织变化，不能只凭开场词推断机制。','',
      'A整体+7.69个百分点可按题支持分解为：共同错误题的组成项+2.65、共同题内误接受变化项+5.52、新增错误题+1.29、退出错误题−1.76个百分点，四项相加恰为总变化。这是精确描述性分解，不是语义错误类型或因果分解。事后移除贡献最大的10题，A仍上升约6.02个百分点，说明点估计并非只由一两题驱动；这种移除只是敏感性检查。','',
      '**SITE结论：A里存在输入条件下的多组风险上涨和具体追查案例；尚不能确认某一题型特异性恶化，也不能把固定V参数未变时的条件风险变化说成V参数能力下降。EFP阴性/不足本身也不能证明SITE，反之亦然。**','',
      '## 审计与交付','',
      f'EFP评分耗时{efp["compute"]["total_seconds"]:.1f}秒；无G生成、无训练、无V更新，GPU评分后已释放。原始缓存SHA256保持一致。软件检查138项通过；所有表格保留原始分母、区间、p/q和支持数量。','',
      '- `EFP_SITE_review.xlsx`：12个工作表，翻转、自然词频、输入分层、方差和逐题清单。',
      '- `EFP_SITE_figures.pdf`：3页图；PNG可单独查看。',
      '- `EFP_counterfactual_cases.jsonl`：112条完整原文及全部变体分数；实际前缀在冻结计划中。',
      '- `case_reviews.md`：7个定向案例的精确检查与完整回答；助手初审，人工待审。',
      '- `source/`、`configs/`及`artifact_manifest.json`：可复核代码、冻结计划和产物哈希。',
      '- `EFP_SITE_review.zip`：本目录打包。','']
    (DEST/'EFP_SITE_report.md').write_text('\n'.join(lines))


def main():
    rows=load();efp=json.loads((DEST/'EFP_results.json').read_text());site=json.loads((DEST/'SITE_results.json').read_text())
    cpu=json.loads((DEST/'results.json').read_text());notes,cases=exact_cases(rows)
    probes=read_jsonl(WORK/'EFP_probe_answers.jsonl');sc=read_jsonl(WORK/'EFP_V0_scores.jsonl')
    cfg=json.loads(CONFIG.read_text());disc=set(cfg['A_discovery']);confirm=set(cfg['A_confirmation'])
    assert not disc&confirm and disc|confirm==set(cfg['question_ids']['A'])
    assert len({r['question_id'] for r in probes})==112
    assert all(r['question_id'] in confirm if r['split']=='A' else r['question_id'] in set(cfg['question_ids']['C']) for r in probes)
    by=defaultdict(list)
    for r in sc:by[r['candidate_id']].append(r)
    assert len(by)==len(probes)==112 and len(sc)==1344
    for r in probes:
        ss=by[r['candidate_id']];assert len(ss)==12 and len({s['variant'] for s in ss})==12
        assert all(s['original_body_sha256']==digest(r['text']) for s in ss)
        assert all(s['final_answer_preserved'] and s['correct'] is False and s['adapter_fingerprint'] is None for s in ss)
    write_jsonl(DEST/'EFP_counterfactual_cases.jsonl',[{'original':r,'variant_scores':by[r['candidate_id']]} for r in probes])
    figures(efp,site);workbook();render(efp,site,cpu,notes)
    sources=DEST/'source';sources.mkdir(exist_ok=True);configs=DEST/'configs';configs.mkdir(exist_ok=True)
    for name in ['fpr_pattern_diagnostics.py','efp_intervention.py','site_diagnostics.py','efp_site_export.py']:
        shutil.copyfile(ROOT/'src/pilot'/name,sources/name)
    shutil.copyfile(ROOT/'tests/test_fpr_patterns.py',sources/'test_fpr_patterns.py')
    for p in [CONFIG,PLAN,SITE_PLAN,WORK/'selected_word_hypotheses.json',WORK/'EFP_runtime.manifest.json']:
        shutil.copyfile(p,configs/p.name)
    write_json(DEST/'method_definition_update.json',{'confirmed_at_analysis_stage':True,
        'base_config_term_status_is_historical':True,'definitions':'EFP=可触发的系统性假阳性；SITE=与输入绑定的系统性错误',
        'provided_by':'user detailed instructions in current conversation',
        'no_same_name_paper_reproduction_claim':True,'prefix_intervention_completed':True,'input_strata_and_variance_completed':True})
    load() # Final immutable-input hash check.
    write_json(DEST/'audit.json',{'completed_at':utc(),'same_engine_G2_G3':True,'fixed_V0':True,
        'all112_baselines_reproduced':True,'all1344_score_contexts_complete':True,'all_answers_preserved':True,
        'discovery_confirmation_question_disjoint':True,'source_hashes_unchanged':True,
        'tests_passed':138,'B_used':False,'T_read':False,'G_training_performed':False,
        'definitions_confirmed':True,'human_case_review':'pending'})
    manifest={str(p.relative_to(DEST)):sha(p) for p in sorted(DEST.rglob('*')) if p.is_file() and p.name not in ['artifact_manifest.json','EFP_SITE_review.zip']}
    write_json(DEST/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(DEST/'EFP_SITE_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(DEST.rglob('*')):
            if p.is_file() and p.name!='EFP_SITE_review.zip':z.write(p,str(p.relative_to(DEST)))
    with zipfile.ZipFile(DEST/'EFP_SITE_review.zip') as z:assert z.testzip() is None
    write_json(WORK/'state.json',{'status':'complete','completed_at':utc(),'report':str(DEST/'EFP_SITE_report.md'),
        'definitions_confirmed':True,'B_used':False,'T_read':False})
    print(json.dumps({'complete':True,'report':str(DEST/'EFP_SITE_report.md')},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
