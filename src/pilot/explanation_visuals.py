"""Figures, review workbook and Chinese report for competitive explanations."""
import json
from collections import Counter
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .core import read_jsonl
from .explanation_probes import DEST
from .formal_common import sha, utc

COLORS = ['#2379a8', '#d7872c', '#458a65']
ROUNDS = ['G0', 'G1', 'G2']
TYPE_NAMES = {'constraint_or_modeling': 'Constraint / model', 'invalid_algebra_or_equivalence': 'Invalid equivalence',
              'arithmetic_or_base_conversion': 'Arithmetic / base', 'unsupported_generalization_or_cases': 'Incomplete / unsupported'}


def pp(x): return 'NA' if x is None else f'{100*x:.2f}%'
def signed(x): return 'NA' if x is None else f'{100*x:+.2f}'
def table(head, rows):
    return '\n'.join(['| '+' | '.join(head)+' |', '| '+' | '.join(['---']*len(head))+' |']+['| '+' | '.join(map(str,r))+' |' for r in rows])+'\n'


def main():
    a = json.loads((DEST/'analysis.json').read_text())
    run = json.loads((DEST/'likelihood_run.json').read_text())
    review = read_jsonl(DEST/'semantic_review_initial18.jsonl')
    figures = DEST/'figures'; figures.mkdir(exist_ok=True)
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'axes.spines.top': False,
                         'axes.spines.right': False, 'savefig.dpi': 180})
    names = []
    pdf = PdfPages(DEST/'explanations_figures.pdf')

    def save(fig, name):
        fig.tight_layout(pad=1.4)
        for ext in ['png', 'pdf', 'svg']: fig.savefig(figures/f'{name}.{ext}', bbox_inches='tight')
        pdf.savefig(fig, bbox_inches='tight'); plt.close(fig); names.append(name)

    # C is primary and the new reference-suspect question is excluded as a whole.
    fig, axs = plt.subplots(1, 3, figsize=(13, 4))
    for ax, k, title in zip(axs, ['G_accuracy_verified', 'FPR', 'FNR'], ['G accuracy, verifiable', 'V0 false acceptance', 'V0 false rejection']):
        rs = [r for r in a['behavior_ci95'] if r['split'] == 'C' and r['metric'] == k]
        y = np.array([r['estimate'] for r in rs])*100
        ax.errorbar(range(3), y, yerr=[y-np.array([r['low'] for r in rs])*100, np.array([r['high'] for r in rs])*100-y], fmt='o-', color=COLORS[0], capsize=4)
        for i, v in enumerate(y): ax.annotate(f'{v:.1f}%', (i, v), xytext=(0, 10), textcoords='offset points', ha='center')
        ax.set_xticks(range(3), ROUNDS); ax.set_title(title); ax.set_ylabel('Percent'); ax.grid(axis='y', alpha=.2)
    fig.suptitle('C95 exploratory sensitivity: paired question-cluster 95% intervals', y=1.03)
    save(fig, '01_C95_accuracy_fpr_fnr')

    for split, idx in [('C', '02'), ('B', '03')]:
        fig, ax = plt.subplots(figsize=(10, 4.6))
        groups = ['output_length', 'answer_gap', 'question_length']
        components = ['composition', 'within_stratum', 'total']
        for j, (component, color) in enumerate(zip(components, COLORS)):
            rr = [next(r for r in a['decomposition'] if r['split'] == split and r['scheme'] == s and r['contrast'] == 'G2-G0' and r['component'] == component) for s in groups]
            x = np.arange(3)+(j-1)*.22
            y = np.array([r['estimate'] for r in rr])*100
            ax.errorbar(x, y, yerr=[y-np.array([r['low'] for r in rr])*100, np.array([r['high'] for r in rr])*100-y], fmt='o', color=color, capsize=4, label=component.replace('_',' '))
        ax.axhline(0, color='black', lw=.8); ax.set_xticks(range(3), ['Response length', 'Final-answer gap', 'Question length'])
        ax.set_ylabel('Change in FPR, percentage points'); ax.set_title(f'{split}{95 if split == "C" else 200}: G2 - G0 midpoint decomposition (exploratory)')
        ax.legend(loc='upper left', bbox_to_anchor=(1.01, 1.0)); ax.grid(axis='y', alpha=.2)
        save(fig, idx+'_fpr_decomposition_'+split)

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.6))
    groups = ['<=512', '513-1024', '>1024']
    for j, g in enumerate(ROUNDS):
        rr = [next(r for r in a['strata'] if r['split'] == 'C' and r['G_round'] == g and r['scheme'] == 'output_length' and r['stratum'] == k) for k in groups]
        x = np.arange(3)+(j-1)*.25
        axs[0].bar(x, [r['error_share']*100 for r in rr], .24, color=COLORS[j], label=g)
        bars = axs[1].bar(x, [r['FPR_within']*100 for r in rr], .24, color=COLORS[j], label=g)
        for b, r in zip(bars, rr): axs[1].text(b.get_x()+b.get_width()/2, b.get_height()+1, str(r['wrong_n']), ha='center', fontsize=8)
    for ax in axs: ax.set_xticks(range(3), groups); ax.set_xlabel('Response tokens'); ax.set_ylabel('Percent'); ax.legend(); ax.set_ylim(0, 76)
    axs[0].set_title('C95: composition of independently wrong answers'); axs[1].set_title('C95: FPR within length group (labels = wrong n)')
    save(fig, '04_error_length_composition_and_within_fpr')

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.7))
    truthnames = ['correct', 'wrong', 'unknown']
    for j, (target, contrast) in enumerate([('G1', 'G1-G0'), ('G2', 'G2-G1')]):
        rr = [next(r for r in a['likelihood_summary'] if r['cohort'] == f'{target}_selected_{k}' and r['contrast'] == contrast) for k in truthnames]
        x = np.arange(3)+(j-.5)*.28; y = np.array([r['delta_nll'] for r in rr])
        axs[0].bar(x, y, .27, color=COLORS[j], label=contrast,
                   yerr=[y-np.array([r['low'] for r in rr]), np.array([r['high'] for r in rr])-y], capsize=3)
        rr2 = [next(r for r in a['likelihood_matched_controls'] if r['target'] == target and r['correct'] is k) for k in [True,False,None]]
        yy = np.array([r['selected_minus_control_delta'] for r in rr2])
        axs[1].bar(x, yy, .27, color=COLORS[j], label=contrast,
                   yerr=[yy-np.array([r['low'] for r in rr2]), np.array([r['high'] for r in rr2])-yy], capsize=3)
    for ax in axs: ax.axhline(0, color='black', lw=.8); ax.set_xticks(range(3), truthnames); ax.legend(); ax.set_ylabel('Delta mean response NLL, nats/token')
    axs[0].set_title('Selected A texts: 12 responses per outcome/update')
    axs[1].set_title('Matched same-question alternatives: selected minus control')
    fig.suptitle('Lower NLL means fitting these exact texts; correctness is a separate outcome', y=1.02)
    save(fig, '05_fixed_text_likelihood_and_matched_controls')

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.4))
    for ax, target in zip(axs, ['G1','G2']):
        rr = [r for r in a['feedback_quality'] if r['target'] == target]
        bottom = np.zeros(3)
        for k, color in [('correct', COLORS[2]), ('wrong', '#c44b43'), ('unknown', '#b0b0b0')]:
            values = np.array([r[k]/r['n']*100 for r in rr])
            ax.bar(range(3), values, bottom=bottom, color=color, label=k)
            for i,v in enumerate(values):
                if v > 8: ax.text(i,bottom[i]+v/2,f'{v:.1f}%',ha='center',va='center',color='white')
            bottom += values
        ax.set_xticks(range(3), ['All A candidates','V0 accepted','Actual SFT']); ax.set_ylim(0,114); ax.set_ylabel('All responses, percent'); ax.set_title(target+' feedback composition'); ax.legend(loc='upper center',ncol=3,fontsize=8,frameon=False)
    save(fig, '06_feedback_quality_before_after_selection')

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.5))
    reasons = sorted({r['reason'] for r in a['unknown_reasons']})
    palette = plt.get_cmap('tab20')
    for ax, split in zip(axs, ['C','B']):
        bottom = np.zeros(3)
        for j, reason in enumerate(reasons):
            vals = np.array([sum(r['n'] for r in a['unknown_reasons'] if r['split'] == split and r['G_round'] == g and r['reason'] == reason) for g in ROUNDS])
            ax.bar(range(3), vals, bottom=bottom, color=palette(j), label=reason.replace('_',' ')); bottom+=vals
        ax.set_xticks(range(3),ROUNDS); ax.set_ylabel('Unknown answers'); ax.set_title(split+('95' if split=='C' else '200')+': unknown labels by reason')
    axs[1].legend(loc='upper left',bbox_to_anchor=(1.01,1),fontsize=8)
    save(fig, '07_unknown_reason_composition')

    fig, axs = plt.subplots(1, 2, figsize=(12, 4.5))
    modes = [('original96','Original C96'), ('exclude_suspect95','Whole-question exclusion C95'),('reference1_overlay96','Reference=1 overlay C96')]
    for j,(mode,label) in enumerate(modes):
        rr=[next(r for r in a['reference_sensitivity'] if r['scenario']==mode and r['G_round']==g) for g in ROUNDS]
        axs[0].plot(range(3),[r['FPR']*100 for r in rr], 'o-', color=COLORS[j], label=label)
    rr=[r for r in a['behavior'] if r['split']=='C']
    x=np.arange(3);lo=np.array([r['FPR_unknown_lower'] for r in rr])*100;hi=np.array([r['FPR_unknown_upper'] for r in rr])*100
    axs[1].fill_between(x,lo,hi,color=COLORS[0],alpha=.2,label='Sharp identification range')
    axs[1].plot(x,[r['FPR']*100 for r in rr],'o-',color=COLORS[0],label='Known-only FPR')
    for ax in axs: ax.set_xticks(range(3),ROUNDS); ax.set_ylabel('FPR, percent'); ax.legend(fontsize=8)
    axs[0].set_title('Reference-answer sensitivity (same frozen V0)'); axs[1].set_title('C95: unknown-label bounds, NOT confidence intervals')
    save(fig, '08_reference_and_unknown_sensitivity')

    fig, axs = plt.subplots(1, 2, figsize=(10.4,4.8))
    t=next(r for r in a['transitions'] if r['split']=='C' and r['contrast']=='G2-G0')
    c=t['wrong_to_wrong_V_transitions']; matrix=np.array([[c.get('accept_to_accept',0),c.get('accept_to_reject',0)],[c.get('reject_to_accept',0),c.get('reject_to_reject',0)]])
    axs[0].imshow(matrix,cmap='Blues',vmin=0,vmax=matrix.max())
    for i in range(2):
        for j in range(2): axs[0].text(j,i,str(matrix[i,j]),ha='center',va='center',fontsize=18,color='white' if matrix[i,j]>matrix.max()/2 else 'black')
    axs[0].set_xticks([0,1],['Accept','Reject']);axs[0].set_yticks([0,1],['Accept','Reject']);axs[0].set_xlabel('V0 on G2 answer');axs[0].set_ylabel('V0 on G0 answer');axs[0].set_title('C95: wrong in both rounds (n=166)')
    rr=[r for r in a['behavior'] if r['split']=='C']
    axs[1].plot(range(3),[r['any_correct_at4']*100 for r in rr],'o-',color=COLORS[0],label='Any observed correct in 4')
    axs[1].plot(range(3),[r['G_accuracy_lower']*100 for r in rr],'o-',color=COLORS[1],label='Correct / all responses')
    axs[1].set_xticks(range(3),ROUNDS);axs[1].set_ylabel('Percent');axs[1].set_title('C95: generator coverage');axs[1].legend(fontsize=8)
    fig.suptitle('Descriptive same-question/index pairs; different answers, not V forgetting on fixed texts',y=1.02)
    save(fig,'09_cancellation_and_generator_coverage')
    fig,axs=plt.subplots(1,2,figsize=(13.5,5))
    components=['common_composition','common_within','new_error_question_entry','old_error_question_exit','total']
    labels=['Common\ncomposition','Common\nwithin','New error\nquestions','Lost error\nquestions','Total']
    for ax,split in zip(axs,['C','B']):
        rr=[next(r for r in a['question_support_decomposition'] if r['split']==split and r['contrast']=='G2-G0' and r['component']==k) for k in components]
        y=np.array([r['estimate'] for r in rr])*100
        ax.bar(range(5),y,color=[COLORS[0],COLORS[1],'#ad6aa9','#7a888f',COLORS[2]],
               yerr=[y-np.array([r['low'] for r in rr])*100,np.array([r['high'] for r in rr])*100-y],capsize=3)
        ax.axhline(0,color='black',lw=.8);ax.set_xticks(range(5),labels,fontsize=8);ax.set_ylabel('FPR change, percentage points');ax.set_ylim(-11,9)
        for i,v in enumerate(y):ax.text(i,v+(.35 if v>=0 else -.55),f'{v:+.2f}',ha='center',fontsize=9)
        ax.set_title(f'{split}{95 if split=="C" else 200}: common {rr[0]["common_error_questions"]}, new {rr[0]["new_error_questions"]}, lost {rr[0]["lost_error_questions"]} error questions')
    fig.suptitle('G2 - G0 question-support decomposition; at most 4 wrong observations per question',y=1.02)
    save(fig,'10_question_support_decomposition')
    pdf.close()

    # Review workbook stores complete originals in JSONL when Excel would exceed cell limits.
    wb=Workbook();wb.remove(wb.active)
    sheets = {'readme': [{'item':'用途','value':'竞争性解释诊断；C95暂排除新参考疑点题，C96原结果另列；B为补充，T封存。'},
                          {'item':'标签','value':'semantic_review为助手初审，human_review_status=pending；不是人工金标准。'},
                          {'item':'语义类型','value':'仅18条抽查，不据此估计总体比例或语义类内FPR；全量分解使用长度等可观测代理。'},
                          {'item':'区间','value':'按题聚类成对bootstrap2000次，探索性95%区间；未知识别界不是置信区间。'}],
              'behavior':a['behavior'], 'behavior_ci95':a['behavior_ci95'],'paired_deltas':a['behavior_paired_deltas'],
              'decomposition':a['decomposition'],'question_support':a['question_support_decomposition'],'strata':a['strata'],'contributions':a['contributions'],
              'feedback_quality':a['feedback_quality'],'training_token_mass':a['training_token_mass'],
              'likelihood':a['likelihood_summary'],'matched_likelihood':a['likelihood_matched_controls'],
              'unknown_reasons':a['unknown_reasons'],'reference_sensitivity':a['reference_sensitivity'],
              'semantic_review':review,'suspect_all12':read_jsonl(DEST/'reference_suspect_all12.jsonl')}
    for name, rs in sheets.items():
        ws=wb.create_sheet(name)
        if not rs: continue
        keys=list(dict.fromkeys(k for r in rs for k in r))
        ws.append(keys)
        for r in rs:
            ws.append([None]*len(keys))
            for j,k in enumerate(keys,1):
                v=r.get(k)
                if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
                if isinstance(v,str) and len(v)>32000:v=f'完整原文见JSONL，candidate_id={r.get("candidate_id")}; Excel长度限制，未展示原文。'
                cell=ws.cell(ws.max_row,j,v)
                if isinstance(v,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245477')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for j,k in enumerate(keys,1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(j)].width=65 if k in ['question','text','explanation_initial'] else 25
    wb.save(DEST/'explanations_review.xlsx')

    lookup=lambda split,g:next(r for r in a['behavior'] if r['split']==split and r['G_round']==g)
    lines=['# 为什么本次没有检出明显恶化：竞争性解释诊断',
    '**更符合现有证据的表述是：G确实拟合了所选反馈，但尚未显示稳定的能力提升或更难验证的错误分布；V0误接受率的变化又小于当前估计的不确定性。不能据此推断“验证器跨轮泛化很好”。**',
    '本报告将生成器变化、反馈质量、错误组成和测量问题分开检查。C为主分析，B为单独补充；本次发现C一题参考答案疑点，因此新增C95整题排除敏感性视图。它是探索性分析修订，不覆盖冻结C96结果，也不修改正式标签、训练样本、模型、V0或阈值。T未读取。',
    '## 先报告一个真实的测量问题',
    'C题 `dapo:train:24d751b7-4429-4605-9754-d1451911fb26` 问：若整数x,y满足x³−y³=xy+61，求全部可能x的和。冻结参考为6；精确检查得到全部整数解只有(-5,-6)、(6,5)，故所求和为1。数值6恰为正解x，但不是所有整数x之和；源答案为何如此尚未查明。',
    a['reference_checks']['integer_cubic']['proof'],
    '原12条回答和原标签全部保留于`reference_suspect_all12.jsonl`。主诊断暂按题排除该题全部12条，另单列“原C96”和“仅把该题参考改为1”的影子核验结果；未知仍为未知，不借机放宽解析规则。此证明与精算属于助手初审，等待用户复核后才能正式修订参考清单。',
    table(['视图','G0 FPR','G1 FPR','G2 FPR','G2−G0/百分点'],[[label]+[pp(next(r for r in a['reference_sensitivity'] if r['scenario']==mode and r['G_round']==g)['FPR']) for g in ROUNDS]+[signed(next(r for r in a['reference_sensitivity'] if r['scenario']==mode and r['G_round']=='G2')['FPR']-next(r for r in a['reference_sensitivity'] if r['scenario']==mode and r['G_round']=='G0')['FPR'])] for mode,label in modes]),
    '原C96的+0.42个百分点变为C95的−0.52个百分点；参考=1影子视图则恰为0。**一题就能改变点估计方向**，但三种视图都没有支持明显单向恶化。不能选择最有利视图替代原结果。',
    '此外精确复算了有理方程、进制递推、三平方分解、彩色铺砖、车牌计数、红色四门车和实数幂方程等7题，所检参考一致；抽查并非全数据的参考答案认证。几何及较复杂计数只核查了明确局部错误的案例，未声称完成所有题的全解。',
    '## 解释一：G的变化是否不足',
    table(['C95','正确/错误/未知','G正确率（可核验）','正确/全部输出','4候选至少1对的题占比','输出均长'],[[g,f'{lookup("C",g)["correct"]}/{lookup("C",g)["wrong"]}/{lookup("C",g)["unknown"]}',pp(lookup('C',g)['G_accuracy_verified']),pp(lookup('C',g)['G_accuracy_lower']),pp(lookup('C',g)['any_correct_at4']),f'{lookup("C",g)["mean_output_tokens"]:.1f}'] for g in ROUNDS]),
    'C95可核验正确率42.54%→38.61%→40.51%，没有稳定提高；G2−G0为−2.03个百分点，按题成对95%区间[−7.10,+2.84]。B200相应41.81%→42.28%→43.03%，差值+1.22个百分点，区间[−2.32,+4.89]。B的正确/全部回答却39.88%→39.75%，因此不能只根据可核验分母宣布增强。至少一条正确的4候选覆盖只是观测命中率，不把未知当错，也不冒称总体真实pass@4。',
    '但“没增强”不等于“没变化”：C95的G0/G2同题同候选索引共有341个双方可核验对，只有165个最终答案相同，全文相同仅1条；36个从对转错，25个从错转对。候选索引配对只是描述，同种子不保证参数变化后每个采样路径对应同一个随机实验。单条训练轨迹也不能刻画训练随机性。',
    '固定文本似然探针进一步证实更新有可测的拟合效果。它支持“参数更新并非空跑”，却不能证明数学能力或针对V的对抗性变化足够强。当前证据更支持“能力变化不足/方向不稳定仍是可行解释”，不能断言所有回答分布都没有移动。',
    '![C95主指标](figures/01_C95_accuracy_fpr_fnr.png)',
    '## 解释二：G从高噪声反馈中学到了什么',
    table(['用于更新','原A全部候选错误率','V0接受池错误率','实际SFT错误率'],[[g]+[pp(next(r for r in a['feedback_quality'] if r['target']==g and r['group']==k)['verified_error_fraction']) for k in ['all_A','V0_accepted_A','actual_SFT']] for g in ['G1','G2']]),
    '以上错误率均只以可核验样本为分母，完整正确/错误/未知计数见Excel。V0把原A约62%–65%的错误率降至接受池约47%–49%，因此“实际SFT错误率超过50%”不能单独推出反馈比原G更差或完全无效。另一方面，每题至多随机1条让题目权重发生变化，实际SFT的错误率约53%，确实比按候选计数的接受池更差；不能用整个接受池质量代替实际训练质量。',
    '两轮实际监督token中，冻结标签为错误的占56.17%和54.96%。这不是梯度贡献率：原目标按每条回答的平均token loss再对样本平均，长回答不简单按token数增加同等倍数权重。两轮有465道共同入选题，其中444道两轮标签均已知；192道两轮都选错，84道重复了同一个错误最终答案。重复错例提供了可研究的反馈延续现象，但不单凭它宣布错误记忆是所有变化的原因。',
    '### 新增的只读似然探针',
    '固定158条文本：两轮实际A入选样本各抽正确/错误/未知12条，共72条；同题、同标签且V0接受但未入选的54条替代回答；C中G0 candidate0的正确/错误各16条，共32条。按预定哈希抽样，抽样清单在GPU评分前保存。所有158条在G0/G1/G2下都评分，共474次。诊断用真值分层不回流训练或筛选；B/T不进入此探针。',
    '沿用原BF16、SDPA和精确response-only模板，完整文本含assistant结束符，逐token计算交叉熵。没有生成、训练、梯度、V0重评分或截断；NLL是按回答先求token平均，再跨回答平均，单位nats/token。负ΔNLL表示该文本变得更容易被复现。',
    table(['实际更新','入选正确文本ΔNLL','入选错误文本ΔNLL','入选未知文本ΔNLL'],[[g]+[f'{next(r for r in a["likelihood_summary"] if r["cohort"]==g+"_selected_"+k and r["contrast"]==contrast)["delta_nll"]:+.5f}' for k in ['correct','wrong','unknown']] for g,contrast in [('G1','G1-G0'),('G2','G2-G1')]]),
    '同题同标签匹配对照避免直接比较12条入选与较少替代回答的均值。错误文本中，G1的“入选ΔNLL−替代ΔNLL”为−0.01552（10对，95%区间[−0.02083,−0.01055]）；G2为−0.02009（9对，[−0.02567,−0.01440]）。正确文本同样存在入选拟合优势。**G在拟合所选文本，对错误回答也如此；这一结果不等于错误发生率必然增加。**',
    '这些替代回答不是随机分配的无噪声训练分支，A同题本来也被训练，C仅为开发诊断。没有运行“等量等长度的更干净反馈”训练对照，因此不能据此把未增强因果归结为噪声、过拟合或训练不足其中之一。',
    '![固定文本似然](figures/05_fixed_text_likelihood_and_matched_controls.png)',
    '![反馈构成](figures/06_feedback_quality_before_after_selection.png)',
    '## 解释三：总体FPR是否掩盖组成和类内变化',
    r'令 $w_{tk}=P_t(k\mid Y=0)$，$p_{tk}=P_t(V_0接受\mid Y=0,k)$。有 $FPR_t=\sum_k w_{tk}p_{tk}$。采用对起止轮次对称的精确分解：',
    r'$$\Delta FPR=\underbrace{\sum_k(w_{tk}-w_{0k})(p_{tk}+p_{0k})/2}_{组成项}+\underbrace{\sum_k(p_{tk}-p_{0k})(w_{tk}+w_{0k})/2}_{类内项}.$$ ',
    '全量使用三套可重复的可观测分组：回答长度≤512、513–1024、>1024；错误最终答案绝对偏差=1、2–10、其他；题目提示长度≤128、129–256、>256。答案偏差用冻结标准答案定义，只作描述，不能称为推理错误根因；题目长度也不是题目难度。这三套分解是同一总变化的不同视图，不能把不同分组的项再相加。',
    table(['C95/G2−G0','组成项/百分点','类内项/百分点','总变化/百分点'],[[s]+[signed(next(r for r in a['decomposition'] if r['split']=='C' and r['scheme']==s and r['contrast']=='G2-G0' and r['component']==k)['estimate']) for k in ['composition','within_stratum','total']] for s in ['output_length','answer_gap','question_length']]),
    '最终答案偏差分层中，组成项+1.32和类内项−1.84个百分点发生数值抵消；但两者95%区间分别约[−0.30,+3.36]和[−9.48,+5.57]，都跨0，尚不足以认定可靠的抵消机制。长度分解的组成项−0.56、类内项+0.04，幅度也小。B上的三套分解另列，未与C合并。',
    '每次抽整题、保留同题4候选，跨G共用重采样索引，2000次探索性百分位95%区间；不做多重比较校正。某层只在单轮出现时，另一轮类内率不可识别，记NA而不填0；两轮都无样本的层贡献0。答案偏差C分解有效重采样1999/2000，B题长分解有效1936/2000，完整计数随表保存，避免隐藏支持不足。',
    'C95同题同索引且两轮都错的166对中，V0由拒绝转接受24对，同时由接受转拒绝也是24对。这是总体稳定下局部变化抵消的直接描述；它没有控制具体回答文本相同，不能称为V参数遗忘或归结于某一种数学错误。',
    '![FPR分解](figures/02_fpr_decomposition_C.png)',
    '![长度组成与类内率](figures/04_error_length_composition_and_within_fpr.png)',
    '### 进一步按独立题目分解，并显式处理进入/退出错题集',
    '若直接令k为题目，部分题在某一轮没有已核验错误，不能凭空填该轮类内FPR。因此只在两轮都观察到错例的共同题集计算对称组成/类内项，再加“新进入错题集的误接受质量”和“退出错题集的旧误接受质量”。四项之和精确等于总体ΔFPR；未观察到错误不代表该题真实错误概率为0。',
    table(['G2−G0/百分点','共同题组成项','共同题类内项','新错题进入项','旧错题退出项','总变化'],[[split]+[signed(next(r for r in a['question_support_decomposition'] if r['split']==split and r['contrast']=='G2-G0' and r['component']==k)['estimate']) for k in ['common_composition','common_within','new_error_question_entry','old_error_question_exit','total']] for split in ['C','B']]),
    'C95的共同错题64道，进入7道、退出6道；B200共同132道、进入12道、退出21道。这样能直接看见跨题组成的变化，而不只看全局均值。每题每轮只有最多4个错误观测，类内率很粗，区间只反映本题池聚类重采样；进入项天然非负、退出项天然非正，不能根据它们的符号或单个区间宣布因果机制。此处的题目分解也不能与上面三套分组的贡献再相加。',
    '![题目支持集分解](figures/10_question_support_decomposition.png)',
    '### 数学错误类型初审的范围',
    '按C中每轮已标错的candidate0预定哈希抽6条，共18条；先审G原文及数学证据，不使用V0分数或接受状态来定类型。错误初审按首个已见实质问题，分为约束/建模、无效等价变换、算术/进制、漏情况/无依据推广；每例保留原文引用、计算或证明、核验范围和空白人工复核列。两例涉及同一道新参考疑点题，已明确标识。',
    '抽查找到跨轮都存在的约束误读、错误代数和无依据推广，例如把白色两门车数当红白交集、删掉四次方程一次项、把格点计数近似当精确等式。这些提供实质案例，却还不足以估计全量语义错误类型比例和类内FPR：每轮仅6条且只取candidate0，部分类型在一轮缺失，不能拿小样本填完整机制分解。18条都是模型辅助初审，未冒充人工金标准。',
    '## 解释四：未知、截断及参考质量是否改变解读',
    'C95未知率4.74%→5.26%→7.11%，截断4→1→6；B未知4.62%→4.50%→7.62%，截断7→11→16。B上G2−G0未知增加3.00个百分点，探索性95%区间[+0.63,+5.38]；C相应+2.37个百分点，区间[−0.27,+5.00]。主要未知来源还有答案冲突和等价性未解，不能把未知都归为截断。',
    '冻结协议已把所有截断回答标为未知，因此“删掉截断后FPR不变”是口径所致，并不是截断无影响的证明：它影响覆盖率和我们看不到的负例分布。未补长输出，也没有把未完成文本判错。',
    '令未知中V接受ua、拒绝ur，已核验错例FP、TN。若未知最终真值完全不受约束，总体FPR的尖锐识别界为[FP/(FP+TN+ur),(FP+ua)/(FP+TN+ua)]；不是将全部未知统一判对或统一判错，也不是置信区间。C95的G0为[52.53%,56.68%]，G2为[50.44%,56.56%]，跨轮差值识别界约[−6.24,+4.03]个百分点，足以改变很小点差的方向。',
    'C95的G2−G0 FPR抽样95%区间为[−8.10,+7.17]个百分点，FNR为[−6.98,+6.52]。区间较宽，不能用“未显著”证明等价稳定；目前既有真值覆盖的不确定性，也有按题抽样的不确定性，还没有多个训练种子。',
    '![未知原因](figures/07_unknown_reason_composition.png)',
    '![参考与未知敏感性](figures/08_reference_and_unknown_sensitivity.png)',
    '## 四种解释目前能区分到哪里',
    table(['竞争性解释','现有支持','仍不能确定'],[
      ['G变化不足','无稳定正确率提升；题覆盖也未提高','文本与答案确有变化，不能说输入分布没变；未检验强G跨轮情形'],
      ['反馈未推动可靠行为','错例被实际拟合，53%左右错误仍入训；部分同一错误重复入选','V筛选相对原A有提纯；没有等量干净反馈训练对照，噪声不能独占因果解释'],
      ['组成变化抵消','偏差分层有+1.32/−1.84的数值抵消；双错配对接受翻转24/24','全量数学错误类型尚无可靠标注，组成/类内项区间跨0'],
      ['测量不完整','新参考错误、未知增加，点差方向依一题即可变化','抽查未覆盖全部题；未定继续保留，不把覆盖改善当自然错误减少']]),
    '**当前可交付的结论是“尚未构成强生成器分布变化下的验证器稳定性证据”，不是“已经找到唯一的没有恶化原因”。** 已有小幅抵消和错例拟合证据，下一步应针对竞争解释设置对照，不能直接增加轮数。',
    '## 下一步对照建议（本次未执行）',
    '先由用户复核新参考疑点及18条初审，再把明确错误类型的标注覆盖扩大到独立抽样且足够的题目；保留未定类别和抽样权重。若要区分反馈质量与更新强度，可另立协议，在A上建立等样本/尽量匹配长度和题目的V0反馈、较干净诊断反馈与无更新对照，复用C做开发、T继续封存。任何使用标准答案构建的干净分支必须明确标为oracle诊断对照，不能混入当前无真值筛选的主轨迹；也需多个训练种子。',
    '## 资源、执行与文件',
    f'主评分在tmux会话`explanation_probes_20261022`中执行，总用时{run["total_seconds"]:.1f}秒；主评分峰值已分配显存{max(v["peak_allocated_bytes"] for v in run["models_this_invocation"].values())/1024**3:.2f} GiB。此外在GPU上用两条固定文本对照模型原生response-only loss，与新增分块评分的差值均低于1e-5。两项均为只读推理，工作记入原预算流水，未重启正式训练；统计和绘图在CPU完成。',
    '本次没有需要重新采样的缺失回答，不为测试引擎而制造新数据。vLLM尚未安装；本次固定完整文本交叉熵复用原Transformers即可，未测vLLM吞吐，因此不声称它更快。以后批量自回归补生成时可在独立环境做vLLM小样本速度/格式/LoRA兼容检查，并把引擎变更的输出单列。',
    '- `explanations_review.xlsx`：指标、分解、反馈质量、似然与匹配对照、未知原因、参考敏感性、18条初审和疑点题全部12条。',
    '- `explanations_figures.pdf`和`figures/`：10张图，PNG/PDF/SVG；`index.html`为浏览入口。',
    '- `analysis.json`及CSV：完整数值；`reference_checks.json`保存精算证明；`semantic_review_initial18.jsonl`保存完整审查文本。',
    '- `likelihood_inputs.jsonl`、`likelihood_plan.json`、`likelihood_scores.jsonl`、`likelihood_run.json`：评分前抽样清单、哈希、474条评分及资源记录。',
    '- `analysis_plan.json`、`source_integrity.json`、`validation.json`：探索性分组、来源哈希和验证。原始32个文件哈希保持一致；冻结源码链保持通过。',
    '复现CPU分析：`scripts/pilot explanation_analysis`、`scripts/pilot explanation_review`、`scripts/pilot explanation_visuals`。已完成的似然探针不会自动重复GPU工作；`scripts/pilot explanation_probes score`读取完成缓存。',
    '时间采用北京时间2026-10-22；原记录中的UTC时间不改写。']
    (DEST/'explanations_report.md').write_text('\n\n'.join(lines)+'\n')
    case_lines=['# 18条数学错误模型辅助初审','按预定哈希选每轮6条C中candidate0已标错回答；未按V分数选择。人工复核列仍为空。不能把本样本的类型比例作为总体比例。']
    for r in review:
        case_lines += [f'## {r["candidate_id"]}',f'初审类型：{r["primary_error_type_initial"]}；核查范围：{r["verification_scope"]}；人工复核：待审。',
                       '题目：'+r['question'],f'冻结参考：{r["gold_answer"]}；G最终答案：{r["extracted_answer"]}；新参考疑点：{r["reference_suspect"]}。',
                       '原文证据：\n\n> '+r['literal_evidence'],'初审依据：'+r['explanation_initial'], '<details><summary>G完整回答</summary>\n\n'+r['text']+'\n\n</details>']
    (DEST/'semantic_review_cases.md').write_text('\n\n'.join(case_lines)+'\n')
    html='<!doctype html><html lang="zh"><meta charset="utf-8"><title>没有明显恶化：解释诊断</title><style>body{max-width:1200px;margin:40px auto;font:17px system-ui}img{width:100%;border:1px solid #ddd}section{margin:50px 0}a{color:#176a9d}</style><h1>为什么未检出明显恶化</h1><p>C95为新参考疑点整题排除的探索视图；原C96保留，B单列补充，T封存。</p><p><a href="explanations_report.md">中文报告</a> · <a href="explanations_review.xlsx">审查表</a> · <a href="explanations_figures.pdf">图册</a> · <a href="semantic_review_cases.md">18条初审</a></p>'
    html+=''.join(f'<section><h2>{name}</h2><a href="figures/{name}.pdf"><img src="figures/{name}.png" alt="{name}"></a></section>' for name in names)+'</html>'
    (DEST/'index.html').write_text(html)
    from .core import write_json
    write_json(DEST/'visuals_manifest.json', {'created_at':utc(),'figures':names,'figure_count':len(names),'source_analysis_sha256':sha(DEST/'analysis.json'),
                'review_sha256':sha(DEST/'semantic_review_initial18.jsonl'),'code_sha256':sha(Path(__file__))})
    print(json.dumps({'figures':len(names),'workbook_sheets':len(wb.sheetnames),'report':str(DEST/'explanations_report.md')},ensure_ascii=False))


if __name__ == '__main__': main()
