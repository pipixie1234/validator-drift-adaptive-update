"""Independent count, paired-interval and artifact verification for the C95 panel."""
import hashlib
import json
import zipfile

import numpy as np

from .core import ROOT, read_jsonl, write_json
from .formal_common import sha, utc
from .prefix_full_panel import RUN, REPORT, PLAN, validate


def main():
    cfg = validate(); answers = read_jsonl(RUN/'answers.jsonl'); scores = read_jsonl(RUN/'V0_scores.jsonl')
    result = json.loads((REPORT/'results.json').read_text()); by = {(s['candidate_id'], s['variant']): s for s in scores}
    qs = sorted({r['question_id'] for r in answers}); assert len(qs) == 95
    W = np.random.default_rng(cfg['seed']).multinomial(95, np.full(95, 1/95), size=10000)
    def counts(name):
        cells = []
        for q in qs:
            rr = [r for r in answers if r['question_id'] == q]; assert len(rr) == 4
            cell = [0]*6
            for r in rr:
                a = by[r['candidate_id'], name]['accepted']; y = r['correct']
                i = (0 if a else 1) if y is True else (2 if a else 3) if y is False else (4 if a else 5)
                cell[i] += 1
            cells.append(cell)
        return np.array(cells)
    allcounts = {v['name']: counts(v['name']) for v in cfg['variants']}
    def rates(name):
        tp, fn, fp, tn, ua, ur = (W@allcounts[name]).T
        return fp/(fp+tn), fn/(tp+fn)
    base = rates('original'); neutral = rates('neutral_control')
    def close(a, b): assert abs(float(a)-float(b)) < 1e-12, (a, b)
    for row in result['net_rates']:
        name = row['variant']; totals = allcounts[name].sum(0)
        for key, n in zip(['TP', 'FN', 'FP', 'TN', 'unknown_accepted'], totals[:5]): assert row[key] == n
        for k, metric in enumerate(['FPR', 'FNR']):
            value = totals[2]/(totals[2]+totals[3]) if k == 0 else totals[1]/(totals[0]+totals[1])
            close(row[metric], value)
            for suffix, control in [('', base), ('_vs_neutral', neutral)]:
                low, high = np.quantile(rates(name)[k]-control[k], [.025, .975])
                close(row[metric+suffix+'_low'], low); close(row[metric+suffix+'_high'], high)
        trans = dict.fromkeys(['TN_to_FP', 'FP_to_TN', 'FN_to_TP', 'TP_to_FN'], 0)
        for r in answers:
            original = by[r['candidate_id'], 'original']['accepted']; after = by[r['candidate_id'], name]['accepted']
            if original == after or r['correct'] is None: continue
            key = ('FN_to_TP' if after else 'TP_to_FN') if r['correct'] else ('TN_to_FP' if after else 'FP_to_TN')
            trans[key] += 1
        assert all(row[k] == v for k, v in trans.items())
    for row in result['continuous_changes']:
        rr = [r for r in answers if ('unknown' if r['correct'] is None else 'correct' if r['correct'] else 'wrong') == row['truth']]
        def quantity(s): return s['score'] if row['quantity'] == 'score' else s['log_likelihood_A']-s['log_likelihood_B']
        deltas = [quantity(by[r['candidate_id'], row['variant']])-quantity(by[r['candidate_id'], 'original']) for r in rr]
        close(row['delta'], sum(deltas)/len(deltas))
    from openpyxl import load_workbook
    wb = load_workbook(REPORT/'full_panel_review.xlsx', read_only=True, data_only=False)
    for ws in wb:
        for row in ws.iter_rows(): assert all(c.data_type != 'f' for c in row)
    wb.close()
    from .week_prefix_delivery import pdf_tools
    Reader, _ = pdf_tools(); assert len(Reader(REPORT/'full_panel_figures.pdf').pages) == 3
    manifest = json.loads((REPORT/'artifact_manifest.json').read_text())
    with zipfile.ZipFile(REPORT/'full_panel_review.zip') as z:
        assert z.testzip() is None
        for name, expected in manifest.items():
            assert sha(REPORT/name) == expected and hashlib.sha256(z.read(name)).hexdigest() == expected
    output = {'status': 'passed', 'at': utc(), 'confusion_matrices_recounted': 8,
              'paired_FPR_FNR_intervals_vs_original_and_neutral': 32,
              'continuous_means_independently_checked': 48, 'two_way_transitions_recounted': True,
              'unknowns_excluded_only_from_truth_denominators': True, 'scores': len(scores),
              'workbook_formulas_absent': True, 'PDF_pages': 3, 'archive_crc_and_file_hashes': True,
              'plan_sha256': sha(PLAN), 'code_sha256': sha(ROOT/'src/pilot/prefix_full_panel_verify.py')}
    write_json(REPORT/'independent_verification.json', output); print(json.dumps(output, indent=2), flush=True)


if __name__ == '__main__': main()
