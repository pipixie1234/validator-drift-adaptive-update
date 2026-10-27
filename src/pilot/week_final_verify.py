"""Independent raw-count/rank comparison and final artifact integrity audit."""
import json
import math
import re
import zipfile

import numpy as np

from .core import ROOT,read_jsonl,write_json
from .formal_common import sha,utc
from .week_final_common import RUN,REPORT,GS,VS,validate


def close(a,b):
    assert (a is None and b is None) or (a is not None and b is not None and abs(a-b)<1e-12),(a,b)


def main():
    from openpyxl import load_workbook
    cfg=validate();quality=json.loads((REPORT/'006_T_source_quality_20261026.json').read_text())
    bad={r['question_id'] for r in quality['whole_question_quarantine']}
    result=json.loads((REPORT/'week_final_results.json').read_text());tables=result['tables']
    points={(r['G'],r['V']):r for r in tables['metrics']};raw={};labels={};scores={}
    for g in GS:
        raw[g]=read_jsonl(RUN/f'{g}_T.jsonl');labels[g]={r['candidate_id']:r for r in read_jsonl(RUN/f'labels_{g}_T.jsonl')}
        assert len(raw[g])==1200 and len({r['question_id'] for r in raw[g]})==300
        assert not any(any(k in r for k in ['correct','gold_answer','reference']) for r in raw[g])
    question_sets=[{r['question_id'] for r in raw[g]} for g in GS]
    assert all(s==question_sets[0] for s in question_sets)
    splits=[read_jsonl(ROOT/p) for p in ['data/dapo/inputs/dapo100.jsonl','data/dapo/formal/inputs/A_G_update.jsonl',
        'data/dapo/formal/inputs/B_V_update.jsonl','data/dapo/formal/inputs/T_final.jsonl']]
    for key in ['question_id','question_hash','family_id']:
        ss=[{r[key] for r in rows} for rows in splits]
        assert all(len(s)==len(rows) for s,rows in zip(ss,splits)),key
        assert len(set.union(*ss))==sum(map(len,ss)),key
    uuids=[{u for r in rows for u in r['family_source_uuids']} for rows in splits]
    assert len(set.union(*uuids))==sum(map(len,uuids))
    ids=sorted(question_sets[0]-bad);assert len(ids)==299
    for v in VS:scores[v]={r['candidate_id']:r for r in read_jsonl(RUN/'scoring'/v/'scores.jsonl')}
    boot=np.random.default_rng(cfg['analysis']['seed']).integers(299,size=(10000,299));reps={}
    for g in GS:
        for v in VS:
            cells=[];known=[]
            for q in ids:
                qq=[r for r in raw[g] if r['question_id']==q];assert len(qq)==4
                count=[0]*6
                for r in qq:
                    y=labels[g][r['candidate_id']]['correct'];s=scores[v][r['candidate_id']]
                    accept=s['score']>=.5;assert accept==s['accepted']
                    if y is None:count[4 if accept else 5]+=1
                    else:
                        count[(0 if accept else 1) if y else (2 if accept else 3)]+=1
                        known.append((bool(y),s['score']))
                cells.append(count)
            arr=np.asarray(cells);tp,fn,fp,tn,ua,ur=arr.sum(0);p=points[g,v]
            for k,n in zip(['TP','FN','FP','TN','unknown','total'],[tp,fn,fp,tn,ua+ur,1196]):close(p[k],int(n))
            close(p['FPR'],float(fp/(fp+tn)));close(p['FNR'],float(fn/(tp+fn)))
            close(p['V_accuracy'],float((tp+tn)/(tp+fn+fp+tn)))
            close(p['balanced_accuracy'],float((tp/(tp+fn)+tn/(fp+tn))/2))
            close(p['accepted_error_FDR'],float(fp/(tp+fp)))
            pos=np.array([s for y,s in known if y]);neg=np.array([s for y,s in known if not y])
            auc=np.mean((pos[:,None]>neg).astype(float)+.5*(pos[:,None]==neg))
            close(p['AUROC'],float(auc))
            bc=arr[boot].sum(1).T;b_tp,b_fn,b_fp,b_tn=bc[:4]
            reps[g,v]={'FPR':b_fp/(b_fp+b_tn),'FNR':b_fn/(b_tp+b_fn)}
            reps[g,v]['balanced_accuracy']=1-(reps[g,v]['FPR']+reps[g,v]['FNR'])/2
    for category in ['paired_V_differences','paired_G_differences']:
        for row in tables[category]:
            key=row['metric']
            if key not in ['FPR','FNR','balanced_accuracy']:continue
            after,before=row['contrast'].split('-')
            if category=='paired_V_differences':delta=reps[row['G'],after][key]-reps[row['G'],before][key]
            else:delta=reps[after,row['V']][key]-reps[before,row['V']][key]
            low,high=np.quantile(delta,[.025,.975]);close(row['low'],float(low));close(row['high'],float(high))
    reasons=read_jsonl(REPORT/'reason_reviews_full.jsonl');assert len(reasons)==60
    assert all(r['rationale_validity_initial'] in ['valid','invalid','unverifiable','not_provided'] for r in reasons)
    for v in VS:
        rr=[r for r in reasons if r['V']==v];assert len(rr)==20
        for r in rr:
            if r.get('evidence_quote'):assert r['evidence_quote'] in r['V_rationale_full_text']
            assert r['human_review_status']=='pending' and r['human_label'] is None
    sources={r['candidate_id']:r for r in read_jsonl(RUN/'reason_inputs_C20.jsonl')}
    steps=read_jsonl(REPORT/'G_steps_initial.jsonl');assert len(steps)==114
    for r in steps:assert sources[r['candidate_id']]['text'][r['character_start']:r['character_end']]==r['step_text']
    for name,expected in [('week_review.xlsx',['回答与V理由审查','G步骤审查','标签规范与来源']),('week_final_metrics.xlsx',list(tables))]:
        book=load_workbook(REPORT/name,read_only=True,data_only=False);assert book.sheetnames==expected
        for ws in book:
            for row in ws.iter_rows():assert all(cell.data_type!='f' for cell in row)
        book.close()
    pdf=(REPORT/'week_final_figures.pdf').read_bytes()
    assert pdf.startswith(b'%PDF-') and b'%%EOF' in pdf[-64:]
    prefix_figures=None
    if (REPORT/'figure_sections.json').exists():
        from .week_prefix_delivery import verify_figures
        prefix_figures=verify_figures(REPORT);pdf_pages=prefix_figures['total_pages']
        integration=json.loads((REPORT/'prefix_integration_audit.json').read_text())
        for entry in integration['source_files']:
            assert sha(ROOT/entry['source'])==entry['sha256']
            assert sha(REPORT/entry['copy'])==entry['sha256']
    else:pdf_pages=6
    assert len(re.findall(rb'/Type\s*/Page\b',pdf))==pdf_pages
    from PIL import Image
    pngs=list(REPORT.glob('T_*.png'));assert len(pngs)==6
    for path in pngs:
        with Image.open(path) as im:im.verify()
    sums=json.loads((REPORT/'delivery_SHA256.json').read_text());archive=REPORT/'week_final_delivery.zip'
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None and set(z.namelist())==set(sums)|{'delivery_SHA256.json'}
        import hashlib
        for name,expected in sums.items():
            assert sha(REPORT/name)==expected and hashlib.sha256(z.read(name)).hexdigest()==expected
    output={'at':utc(),'status':'passed','T_raw_questions':300,'T_valid_questions':299,'whole_question_exclusion_same_for_all':True,
        'confusion_matrices_verified':12,'AUROC_independent_pair_comparisons':12,'paired_FPR_FNR_BA_intervals_independently_recomputed':81,
        'reason_reviews':60,'original_steps_offsets_verified':114,'human_gold':False,
        'workbook_formulas_absent':True,'archive_crc_and_all_file_hashes':'pass','archive_sha256':sha(archive),
        'ABC_T_question_hash_family_UUID_disjointness':True,'semantic_dedup_guarantee':False,
        'PDF_pages_verified':pdf_pages,'PNG_images_decoded':6,'prefix_figure_integration':prefix_figures,
        'code_sha256':sha(ROOT/'src/pilot/week_final_verify.py')}
    write_json(REPORT/'independent_verification.json',output);print(json.dumps(output,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
