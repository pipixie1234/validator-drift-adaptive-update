"""Frozen-answer V0/Vrand/Vtarget comparison and preservation audit."""
import json
import shutil
import zipfile
from pathlib import Path

import numpy as np

from .core import ROOT,read_jsonl,recover_jsonl,write_json,write_jsonl,append_jsonl,ensure_run_manifest,digest
from .formal_common import messages,sha,utc,gpu_lock,WorkBudget
from .verifier_update import RUN,REPORT,PLAN,ARMS,validate,targeted_selection,matched_random
from .v0_drift_analysis import analyze,clean,interval,divide
from .fpr_pattern_diagnostics import write_csv

VIEWS=['G0_TF','G1_TF','G2_TF','G2_vLLM','G3_vLLM']
VERSIONS=['V0',*ARMS]
METRICS=['FPR','FNR','balanced_accuracy','V_accuracy','accepted_error_FDR','retention_all','AUROC','Brier_proxy']


def training_fit_diagnostic(cfg):
    """Post-hoc read-only check after near-universal acceptance on C; no updates."""
    import gc
    import torch
    from .runtime import load,conditional_label_scores,model_identity
    summaries=[]
    for arm in ARMS:
        d=RUN/arm;rows=read_jsonl(d/'selected.jsonl');fp,ad=model_identity('V',d/'adapter')
        ensure_run_manifest(d/'postfit.manifest.json',{'plan_sha256':sha(PLAN),'selected_sha256':sha(d/'selected.jsonl'),
            'adapter_fingerprint':ad,'purpose':'Post-hoc B training-fit diagnosis after C acceptance shift; no optimizer or threshold changes',
            'runtime_sha256':sha(ROOT/'src/pilot/runtime.py')})
        path=d/'B_postfit_scores.jsonl';saved=recover_jsonl(path);done={r['candidate_id'] for r in saved}
        if len(done)<len(rows):
            with gpu_lock():
                budget=WorkBudget('verifier_update/postfit/'+arm);model,tok=load('V',d/'adapter')
                for r in rows:
                    if r['candidate_id'] in done:continue
                    budget.check();msgs=messages('V',r['question'],r['text']);point=conditional_label_scores(model,tok,msgs)
                    assert point['prompt_tokens']+4<=8192
                    append_jsonl(path,{'candidate_id':r['candidate_id'],'question_id':r['question_id'],'split':'B','correct':r['correct'],
                        **point,'model_fingerprint':fp,'adapter_fingerprint':ad,'prompt_sha256':digest(msgs),'threshold':.5})
                    done.add(r['candidate_id']);budget.charge(r['candidate_id'])
                    if len(done)%32==0:
                        write_json(RUN/'state.json',{'status':'postfit_diagnostic','arm':arm,'done':len(done),'total':256,'at':utc()})
                del model;gc.collect();torch.cuda.empty_cache()
        scores=read_jsonl(path);assert len(scores)==len(done)==256
        lookup={r['candidate_id']:r for r in rows};positive=sum(r['correct'] for r in scores)
        assert positive==128
        for s in scores:
            r=lookup[s['candidate_id']]
            assert s['correct'] is r['correct'] and s['adapter_fingerprint']==ad and s['model_fingerprint']==fp
            assert s['prompt_sha256']==digest(messages('V',r['question'],r['text'])) and s['accepted']==(s['score']>=.5)
        fpcount=sum(not r['correct'] and r['accepted'] for r in scores);fncount=sum(r['correct'] and not r['accepted'] for r in scores)
        ce=float(np.mean([np.logaddexp(r['log_likelihood_A'],r['log_likelihood_B'])-
                         (r['log_likelihood_A'] if r['correct'] else r['log_likelihood_B']) for r in scores]))
        summaries.append({'V':arm,'B_train_examples':256,'FP':fpcount,'FN':fncount,'FPR':fpcount/128,'FNR':fncount/128,
            'accuracy':(256-fpcount-fncount)/256,'binary_A_B_cross_entropy':ce,
            'exact_score_0p5':sum(r['score']==.5 for r in scores),'score_min':min(r['score'] for r in scores),
            'score_max':max(r['score'] for r in scores),'posthoc':True})
    return summaries


def load_all():
    cfg=validate();rows=read_jsonl(RUN/'C_evaluation.jsonl');index={r['candidate_id']:r for r in rows};by={}
    for q in {r['question_id'] for r in rows}:
        qr=[r for r in rows if r['question_id']==q]
        assert len(qr)==20 and len({r['question'] for r in qr})==1
        assert all(sum(r['view']==view for r in qr)==4 for view in VIEWS)
    for arm in ARMS:
        d=RUN/arm;report=json.loads((d/'training_report.json').read_text());scores=read_jsonl(d/'C_scores.jsonl')
        assert len(scores)==len({r['candidate_id'] for r in scores})==1900
        assert report['save_reload_passed'] and report['base_parameters_unchanged'] and report['supervised_tokens']==768
        by[arm]={}
        for s in scores:
            r=index[s['candidate_id']]
            assert s['question_id']==r['question_id'] and s['view']==r['view'] and s['correct'] is r['correct']
            assert s['model_fingerprint']==cfg['models']['V']['fingerprint'] and s['adapter_fingerprint']==report['adapter_fingerprint']
            assert s['threshold']==.5 and s['accepted']==(s['score']>=.5) and s['prompt_tokens']+4<=8192
            assert s['text_sha256']==digest(r['text']) and s['prompt_sha256']==digest(messages('V',r['question'],r['text']))
            assert s['cache_key']==digest({'plan':sha(PLAN),'arm':arm,'candidate_id':r['candidate_id']})
            by[arm][r['candidate_id']]=s
    init=[json.loads((RUN/a/'initial_adapter.json').read_text())['hash'] for a in ARMS];assert init[0]==init[1]
    # Same G,label,length-bin order at every microbatch, same step budgets.
    selected={a:read_jsonl(RUN/a/'selected.jsonl') for a in ARMS}
    eligible=read_jsonl(RUN/'B_eligible.jsonl');reproduced_target=targeted_selection(eligible)
    reproduced_random,matching=matched_random(eligible,reproduced_target)
    assert matching==cfg['selection']['random_matching']
    assert [r['candidate_id'] for r in reproduced_target]==[r['candidate_id'] for r in selected['Vtarget']]
    assert [r['candidate_id'] for r in reproduced_random]==[r['candidate_id'] for r in selected['Vrand']]
    assert [(r['G_round'],r['correct'],r['length_bin']) for r in selected['Vrand']]==[(r['G_round'],r['correct'],r['length_bin']) for r in selected['Vtarget']]
    logs={a:read_jsonl(RUN/a/'steps.jsonl') for a in ARMS}
    assert len(logs['Vrand'])==len(logs['Vtarget'])==32
    for a in ARMS:
        lookup={r['candidate_id']:r for r in selected[a]}
        assert all(r['split']=='B' and r['correct'] is not None for r in selected[a])
        from transformers import AutoTokenizer
        tok=AutoTokenizer.from_pretrained(cfg['models']['V']['path'],local_files_only=True)
        for ex in read_jsonl(RUN/a/'tokenized.jsonl'):
            raw=lookup[ex['candidate_id']]
            prompt=tok.apply_chat_template(messages('V',raw['question'],raw['text']),tokenize=True,add_generation_prompt=True)
            assert ex['prefix_length']==len(prompt) and ex['input_ids'][:len(prompt)]==prompt
            assert all(x==-100 for x in ex['labels'][:ex['prefix_length']])
            assert sum(x!=-100 for x in ex['labels'])==3 and len(ex['input_ids'])<=8192
            assert ex['input_ids'][len(prompt):]==[x for x in ex['labels'] if x!=-100]
            assert tok.decode(ex['input_ids'][len(prompt):])==('A' if raw['correct'] else 'B')+'<|im_end|>\n'
        for step in logs[a]:
            assert step['examples']==8 and step['supervised_tokens']==24
            step['strata']=[(lookup[cid]['G_round'],lookup[cid]['correct'],lookup[cid]['length_bin']) for cid in step['candidate_ids']]
    assert [s['strata'] for s in logs['Vrand']]==[s['strata'] for s in logs['Vtarget']]
    return cfg,rows,by


def judgment_transition(rows,old,new,view,arm):
    known=[r for r in rows if r['correct'] is not None];cells=np.zeros((2,2),int)
    for r in known:
        cid=r['candidate_id'];cells[int(old[cid]!=r['correct']),int(new[cid]!=r['correct'])]+=1
    a,b,c,d=map(int,cells.ravel())
    return clean({'view':view,'new_V':arm,'old_V':'V0','a_old_right_new_right':a,'b_old_right_new_wrong':b,
                  'c_old_wrong_new_right':c,'d_old_wrong_new_wrong':d,'unknown_excluded':len(rows)-len(known),
                  'retained_judgment_correctness':divide(a,a+b),'degraded_judgment_correctness':divide(b,a+b),
                  'repaired_judgment_correctness':divide(c,c+d),'net_repaired':c-b})


def compute(cfg,rows,by):
    points=[];cis=[];contrasts=[];transitions=[];decisions=[];curves=[];score_summary=[];equal_retention=[]
    qs=sorted({r['question_id'] for r in rows});assert len(qs)==95
    indices=np.random.default_rng(cfg['seed']).integers(95,size=(cfg['bootstrap'],95))
    for view in VIEWS:
        write_json(RUN/'state.json',{'status':'analyzing','view':view,'bootstrap':cfg['bootstrap'],'at':utc()})
        rr=[r for r in rows if r['view']==view];assert len(rr)==380
        replicas={};pointraw={}
        for version in VERSIONS:
            data=[{**r,'accepted':r['V0_accepted'] if version=='V0' else by[version][r['candidate_id']]['accepted'],
                    'score':r['V0_score'] if version=='V0' else by[version][r['candidate_id']]['score']} for r in rr]
            point,ci,rep,curve=analyze(data,qs,indices);replicas[version]=rep;pointraw[version]=point
            points.append({'view':view,'V':version,**point})
            for key in METRICS:
                cis.append({'view':view,'V':version,'metric':key,'estimate':point[key],**ci[key]} if key in ci else
                           {'view':view,'V':version,'metric':key,'estimate':point[key],**interval(rep[key])})
            curves += [{'view':view,'V':version,**dict(zip(curve.keys(),items))} for items in zip(*curve.values())]
            decisions += [{'view':view,'V':version,'question_id':r['question_id'],'candidate_id':r['candidate_id'],
                'correct':r['correct'],'accepted':r['accepted'],'score':r['score']} for r in data]
            score_summary.append({'view':view,'V':version,'score_mean':float(np.mean([r['score'] for r in data])),
                'score_p05':float(np.quantile([r['score'] for r in data],.05)),
                'score_median':float(np.median([r['score'] for r in data])),
                'score_p95':float(np.quantile([r['score'] for r in data],.95)),
                'exact_score_0p5':sum(r['score']==.5 for r in data),
                'fraction_0p4_to_0p6':sum(.4<=r['score']<=.6 for r in data)/len(data)})
            k=sum(r['V0_accepted'] for r in rr)
            ranked=sorted(data,key=lambda r:(-r['score'],digest(r['candidate_id'])))
            selected=ranked[:k];correct=sum(r['correct'] is True for r in selected)
            wrong=sum(r['correct'] is False for r in selected);unknown=k-correct-wrong
            equal_retention.append(clean({'view':view,'V':version,'all_candidates':len(data),'selected_top_k':k,
                'retention':k/len(data),'accepted_correct':correct,'accepted_wrong':wrong,'accepted_unknown':unknown,
                'accepted_error_FDR_known':divide(wrong,correct+wrong),
                'unknown_fraction_in_selected':unknown/k if k else None,
                'ranking_boundary_proxy':ranked[k-1]['score'] if k else None,
                'boundary_tie_spans_selection':bool(k and k<len(data) and ranked[k-1]['score']==ranked[k]['score']),
                'method':'Descriptive topK at original V0 acceptance count on all380 C answers; deterministic hash ties. Not deployed threshold tuning, no held-out claim.'}))
        old={r['candidate_id']:r['V0_accepted'] for r in rr}
        for arm in ARMS:
            new={r['candidate_id']:by[arm][r['candidate_id']]['accepted'] for r in rr}
            transitions.append(judgment_transition(rr,old,new,view,arm))
        for after,before in [('Vtarget','Vrand'),('Vrand','V0'),('Vtarget','V0')]:
            for key in METRICS:
                primary=view=='G2_TF' and after=='Vtarget' and before=='Vrand' and key in ['FPR','FNR']
                contrasts.append({'view':view,'contrast':after+'-'+before,'metric':key,'role':'primary' if primary else 'secondary_exploratory',
                    'delta':pointraw[after][key]-pointraw[before][key],**interval(replicas[after][key]-replicas[before][key])})
    return {'metrics':points,'metric_intervals':cis,'paired_contrasts':contrasts,'judgment_transitions':transitions,
            'all_decisions':decisions,'ROC_PR_curves':curves,'score_distribution':score_summary,'equal_retention_C':equal_retention}


def render(tables):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
    lookup={(r['view'],r['V']):r for r in tables['metrics']}
    with PdfPages(REPORT/'verifier_update_figures.pdf') as pdf:
        fig,axes=plt.subplots(2,2,figsize=(12,7),constrained_layout=True)
        for ax,key in zip(axes.flat,['FPR','FNR','balanced_accuracy','retention_all']):
            matrix=np.array([[lookup[view,v][key]*100 for view in VIEWS] for v in VERSIONS])
            im=ax.imshow(matrix,aspect='auto',cmap='YlOrRd',vmin=0,vmax=100)
            for i in range(3):
                for j in range(5):ax.text(j,i,f'{matrix[i,j]:.2f}',ha='center',va='center',color='white' if matrix[i,j]>70 else 'black')
            ax.set_xticks(range(5),VIEWS,rotation=20);ax.set_yticks(range(3),VERSIONS);ax.set_title(key+' (%)');fig.colorbar(im,ax=ax,shrink=.7)
        fig.suptitle('C95 fixed cached answers: V updates trained on B only; threshold 0.5')
        fig.savefig(REPORT/'V_by_G_metrics.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(5,3,figsize=(10,15),constrained_layout=True)
        vmax=max(max(r[k] for k in ['TP','FN','FP','TN']) for r in tables['metrics'])
        for i,view in enumerate(VIEWS):
            for j,v in enumerate(VERSIONS):
                r=lookup[view,v];m=np.array([[r['TP'],r['FN']],[r['FP'],r['TN']]])
                ax=axes[i,j];ax.imshow(m,cmap='Blues',vmin=0,vmax=vmax)
                for a in range(2):
                    for b in range(2):ax.text(b,a,str(int(m[a,b])),ha='center',va='center',fontsize=14,color='white' if m[a,b]>vmax*.55 else 'black')
                ax.set_xticks([0,1],['Accept','Reject']);ax.set_yticks([0,1],['Correct','Wrong']);ax.set_title(f'{view} / {v}; unknown={int(r["unknown"])}')
        fig.suptitle('All15 confusion matrices; unknowns excluded, identical answers within each row')
        fig.savefig(REPORT/'V_by_G_confusion_matrices.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,3,figsize=(13,4),constrained_layout=True)
        for ax,key in zip(axes,['FPR','FNR','balanced_accuracy']):
            for i,view in enumerate(VIEWS):
                r=next(r for r in tables['paired_contrasts'] if r['view']==view and r['contrast']=='Vtarget-Vrand' and r['metric']==key)
                ax.plot([r['low']*100,r['high']*100],[i,i],color='#2279A4',lw=2);ax.plot(r['delta']*100,i,'o',color='#2279A4')
            ax.axvline(0,color='black',ls='--');ax.set_yticks(range(5),VIEWS);ax.invert_yaxis();ax.set_title(key);ax.set_xlabel('Target minus random (percentage points)')
        fig.suptitle('Paired question-cluster95% intervals; lower FPR/FNR and higher BA preferable')
        fig.savefig(REPORT/'Vtarget_minus_Vrand.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(5,2,figsize=(9,14),constrained_layout=True)
        for i,view in enumerate(VIEWS):
            for j,arm in enumerate(ARMS):
                r=next(r for r in tables['judgment_transitions'] if r['view']==view and r['new_V']==arm)
                m=np.array([[r['a_old_right_new_right'],r['b_old_right_new_wrong']],[r['c_old_wrong_new_right'],r['d_old_wrong_new_wrong']]])
                ax=axes[i,j];ax.imshow(m,cmap='Purples',vmin=0,vmax=300)
                for a in range(2):
                    for b in range(2):ax.text(b,a,str(m[a,b]),ha='center',va='center',fontsize=13,color='white' if m[a,b]>160 else 'black')
                ax.set_xticks([0,1],['New right','New wrong']);ax.set_yticks([0,1],['Old right','Old wrong']);ax.set_title(view+' / V0 to '+arm)
        fig.suptitle('Judgment correctness transitions on identical answers; unknowns excluded')
        fig.savefig(REPORT/'V0_to_updated_judgment_transitions.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,2,figsize=(11,4),constrained_layout=True)
        cfg=json.loads(PLAN.read_text())
        for arm in ARMS:
            logs=read_jsonl(RUN/arm/'steps.jsonl');axes[0].plot([l['step'] for l in logs],[l['mean_loss'] for l in logs],label=arm)
        axes[0].set_xlabel('Optimizer step');axes[0].set_ylabel('Mean A/B + EOS suffix training loss');axes[0].legend()
        summary=cfg['selection']['summaries'];errors=[summary[a]['V0_mistakes'] for a in ARMS]
        axes[1].bar(ARMS,errors,color='#C65B4B',label='Original V0 mistakes');axes[1].bar(ARMS,[256-n for n in errors],bottom=errors,color='#4C9674',label='Original V0 right')
        axes[1].set_ylabel('Selected B examples');axes[1].legend();fig.suptitle('Equal256 examples /768 supervised tokens /32 steps; selection strategies differ')
        fig.savefig(REPORT/'training_and_selection.png',dpi=180);pdf.savefig(fig);plt.close(fig)


def report_text(cfg,t):
    mainrows={r['V']:r for r in t['metrics'] if r['view']=='G2_TF'}
    lines=['# 验证器更新：V0、随机更新与针对性更新','',
      f'本轮在固定0.5工作点没有得到优于V0的更新验证器。主视图G2_TF上，V0的FPR/FNR为{mainrows["V0"]["FPR"]:.2%}/{mainrows["V0"]["FNR"]:.2%}，随机更新为{mainrows["Vrand"]["FPR"]:.2%}/{mainrows["Vrand"]["FNR"]:.2%}，针对更新为{mainrows["Vtarget"]["FPR"]:.2%}/{mainrows["Vtarget"]["FNR"]:.2%}。平衡准确率分别{mainrows["V0"]["balanced_accuracy"]:.2%}、{mainrows["Vrand"]["balanced_accuracy"]:.2%}、{mainrows["Vtarget"]["balanced_accuracy"]:.2%}。FNR下降不能单独解释为改善：更新后接受偏移很强，针对组几乎全接受。',
      f'随机组主视图AUROC为{mainrows["Vrand"]["AUROC"]:.4f}，原V0为{mainrows["V0"]["AUROC"]:.4f}；针对组为{mainrows["Vtarget"]["AUROC"]:.4f}。这提示随机组的工作点变化与排序能力需要分开看，不能仅靠换阈值的猜测解释针对组结果。所有失败结果、原数据和新adapter均保留，不据C结果追加训练或更换主阈值。',
      '', '## 设计与数据保留','',
      '两组均从原始V0独立新建同种子LoRA，未继承开发短训练或彼此adapter。训练仅用B200题的G0/G1/G2自然回答与独立最终答案标签；未知不进入监督。G不更新，不生成新回答，T继续封存。',
      '每组256条，正确/错误各128；G0/G1/G2分别64/64/128条，G0/G1历史回放占50%。匹配G来源、真假与V输入长度档，并将随机组输入token总量控制在针对组±1%内；每步类别和长度档顺序匹配。每组1epoch、32步、学习率1e-4、LoRA r16/alpha32/dropout0.05。每回答监督A/B加模板EOS及换行，共3token，两组各768个loss token，其中决定A/B的标签token各256。',
      'Vtarget在每个G×真假层请求75%原V0误判及25%合法对照，误拒独立题不足时补合法对照；Vrand按匹配层随机抽取，不参考V分数或原判定。每个G×真假×题至多一条，每组不重复完整提示。B共2400候选已获取独立标签，成本按全部2400计，不能只报选中256。','',
      '|分支|回答|独立题|输入token|监督token|原V0误判|','|---|---:|---:|---:|---:|---:|']
    for arm in ARMS:
        r=cfg['selection']['summaries'][arm];lines.append(f'|{arm}|256|{r["questions"]}|{r["input_tokens"]}|768|{r["V0_mistakes"]}|')
    lines += ['','整体题目覆盖未完全匹配，是选样策略与学习差异的解释限制。随机组是标签、轮次和长度条件下的随机对照，不是对整个B无条件随机抽样。',
      '', '## 同一批C回答上的主要指标','',
      '每个视图均为同样95道开发题×4回答，V0/Vrand/Vtarget看到完全相同的原文。主适应视图为G2_TF，历史保持看G0_TF；G2_vLLM/G3_vLLM另作运行引擎和未训练生成阶段迁移。B均来自原Transformers，不能把跨生成引擎变化解释为纯G训练效应。C已反复开发，结果不是最终测试成绩。','',
      '|G缓存|V|FPR|FNR|平衡准确率|接受集错误率|全部保留率|','|---|---|---:|---:|---:|---:|---:|']
    for r in t['metrics']:
        lines.append('|'+r['view']+'|'+r['V']+'|'+'|'.join(f'{r[k]:.2%}' for k in ['FPR','FNR','balanced_accuracy','accepted_error_FDR','retention_all'])+'|')
    lines += ['','## 针对更新相对随机更新的配对差值','','负FPR/FNR差表示针对组更低，正平衡准确率差表示针对组更高。全部区间按同一95道题及各题所有候选共同重采样10000次，各模型/视图共享抽样。区间是逐项95%区间，不保证所有比较同时覆盖；除G2_TF的FPR/FNR两个预定主要指标外，其余比较为探索性，不能挑一项有利结果当总体胜出。','',
      '|视图|指标|Vtarget−Vrand/百分点|配对95%区间/百分点|','|---|---|---:|---|']
    for r in t['paired_contrasts']:
        if r['contrast']!='Vtarget-Vrand' or r['metric'] not in ['FPR','FNR','balanced_accuracy']:continue
        lines.append(f'|{r["view"]}|{r["metric"]}|{100*r["delta"]:+.2f}|[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]|')
    lines += ['','## 固定回答上的判断修复与损伤','','a=旧对新对，b=旧对新错，c=旧错新对，d=旧错新错；保持率a/(a+b)，损伤率b/(a+b)，修复率c/(c+d)。这里旧新判断同一个回答，区别于不同G生成不同回答的逐题转移。未知不当错，单列排除。','',
      '|视图|新V|a|b|c|d|净修复c−b|','|---|---|---:|---:|---:|---:|---:|']
    for r in t['judgment_transitions']:
        lines.append(f'|{r["view"]}|{r["new_V"]}|{r["a_old_right_new_right"]}|{r["b_old_right_new_wrong"]}|{r["c_old_wrong_new_right"]}|{r["d_old_wrong_new_wrong"]}|{r["net_repaired"]:+d}|')
    lines += ['','## 实际训练与解释限制','','|分支|实际步数|实际输入token|loss token|训练/保存/重载秒|最大实际上下文|保存重载|','|---|---:|---:|---:|---:|---:|---|']
    for arm in ARMS:
        r=json.loads((RUN/arm/'training_report.json').read_text());lines.append(f'|{arm}|{r["steps"]}|{r["input_tokens"]}|{r["supervised_tokens"]}|{r["seconds_this_invocation"]:.1f}|{r["max_actual_context"]}|通过|')
    lines += ['','## 相同保留量下的描述性检查','','为区分0.5工作点的接受倾向与排序表现，另按每个视图原V0接受数K，从各V分数中取topK；包含未知并单列，边界同分按候选ID哈希打破。该检查在随机组C评分显示接受偏移后补充，属于开发诊断，不改变主比较的0.5阈值，不作为预注册确认性终点。完整分数分布见score_distribution.csv，边界同分情况见equal_retention_C.csv。','',
      '|视图|V|相同保留数K|接受的正确/错误/未知|已核验接受集错误率|AUROC|','|---|---|---:|---|---:|---:|']
    for r in t['equal_retention_C']:
        point=next(p for p in t['metrics'] if p['view']==r['view'] and p['V']==r['V'])
        lines.append(f'|{r["view"]}|{r["V"]}|{r["selected_top_k"]}|{r["accepted_correct"]}/{r["accepted_wrong"]}/{r["accepted_unknown"]}|{r["accepted_error_FDR_known"]:.2%}|{point["AUROC"]:.4f}|')
    lines += ['','## 训练集拟合的补充检查','','在观察到C上接受偏移后，另对每组原256条B训练回答进行只读A/B评分，无新增梯度、训练或阈值修改。二分类交叉熵按A/B归一化概率计算；不是包含EOS/换行的训练平均loss。0.5按原冻结规则接受，不能将并列分数当明确偏好A。','',
      '|V|训练集FPR|训练集FNR|训练集判对率|A/B二分类交叉熵|恰0.5数量|','|---|---:|---:|---:|---:|---:|']
    for r in t['B_training_fit']:
        lines.append(f'|{r["V"]}|{r["FPR"]:.2%}|{r["FNR"]:.2%}|{r["accuracy"]:.2%}|{r["binary_A_B_cross_entropy"]:.4f}|{r["exact_score_0p5"]}/256|')
    lines += ['','分数集中与大量并列是本轮工作点的实际现象，不直接证明具体数值精度或优化机制是唯一原因。训练loss下降本身不能说明A/B判别已学好。',
      '单种子、单训练预算，未按C调学习率或选训练检查点。A/B似然分数未校准，Brier/AUROC仅在可核验子集描述。超长不截断、未知不当负例；独立参考标签不保证所有解题步骤或公共参考均无误。新V没有放回G反馈循环，因此不能声称已验证完整共同适应。',
      '已有runs/reports/configs文件在启动前保存哈希清单，完成后逐一核验；共享GPU计时账本只追加并明确排除，旧模型、数据、图表及评分未覆盖或清理。所有新adapter、训练样本、张量、优化器检查点与评分位于独立runs/verifier_update_20261025。',
      '交付15个混淆矩阵、V×G指标图、配对差值图、判断转移矩阵、训练曲线、5页PDF、Excel、完整决策CSV与审计ZIP。T未解封，后续研究可继续使用原始缓存。']
    (REPORT/'verifier_update_report.md').write_text('\n'.join(lines)+'\n')


def main():
    write_json(RUN/'state.json',{'status':'auditing_and_analyzing','at':utc()})
    cfg,rows,by=load_all();fit=training_fit_diagnostic(cfg);t=compute(cfg,rows,by);t['B_training_fit']=fit
    changes=[]
    for arm in ARMS:
        for r in rows:
            s=by[arm][r['candidate_id']]
            if r['correct'] is None or s['accepted']==r['V0_accepted']:continue
            changes.append({'view':r['view'],'new_V':arm,'question_id':r['question_id'],'candidate_id':r['candidate_id'],
                'correct':r['correct'],'old_accepted':r['V0_accepted'],'new_accepted':s['accepted'],
                'old_score':r['V0_score'],'new_score':s['score'],
                'direction':'repaired' if s['accepted']==r['correct'] else 'damaged',
                'question':r['question'],'response':r['text'],'gold_answer':r['gold_answer'],
                'extracted_answer':r['extracted_answer']})
    write_jsonl(REPORT/'changed_judgments_full.jsonl',changes)
    for name,table in t.items():write_csv(REPORT/(name+'.csv'),clean(table))
    write_json(REPORT/'results.json',clean(t));render(t);report_text(cfg,t)
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    for name,table in t.items():
        ws=wb.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in table for k in r));ws.append(keys)
        for row_number,row in enumerate(clean(table),2):
            values=[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in keys];ws.append(values)
            for column,value in enumerate(values,1):
                if isinstance(value,str) and value.startswith('='):ws.cell(row_number,column).data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(REPORT/'verifier_update_review.xlsx')
    protected=json.loads((RUN/'preserved_existing_files.json').read_text())
    changed=[name for name,r in protected.items() if not (ROOT/name).exists() or sha(ROOT/name)!=r['sha256']]
    assert not changed,changed
    for p in [PLAN,Path(__file__),ROOT/'src/pilot/verifier_update.py',RUN/'preserved_existing_files.json']:
        shutil.copyfile(p,REPORT/p.name)
    for arm in ARMS:
        for name in ['training_report.json','evaluation_compute.json','initial_adapter.json','longest_backward_probe.json']:
            shutil.copyfile(RUN/arm/name,REPORT/(arm+'_'+name))
    write_json(REPORT/'audit.json',{'passed':True,'preserved_old_files':len(protected),'old_files_changed':changed,
        'source_hashes_and_all_score_inputs_checked':True,'independent_same_initialization':True,'B_only_supervision':True,
        'B_selection_reproduced':True,'training_tensors_reconstructed_from_raw_prompts':True,'posthoc_B_fit_scores':512,
        'matched_label_round_length_strata_and_per_step_order':True,'supervised_tokens_each':768,'all_inputs_full_context':True,
        'training_examples_never_truncated':True,'G_unchanged':True,'T_decoded':False,'at':utc()})
    files=sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix!='.zip' and p.name!='artifact_manifest.json')
    manifest={p.name:sha(p) for p in files};write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'verifier_update_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']:z.write(p,p.name)
    import hashlib
    with zipfile.ZipFile(REPORT/'verifier_update_review.zip') as z:
        assert z.testzip() is None and all(hashlib.sha256(z.read(n)).hexdigest()==h for n,h in manifest.items())
    write_json(RUN/'state.json',{'status':'complete','report':str(REPORT/'verifier_update_report.md'),'at':utc()})
    print(json.dumps(clean({'G2_TF':[r for r in t['metrics'] if r['view']=='G2_TF'],
          'primary_differences':[r for r in t['paired_contrasts'] if r['role']=='primary']}),ensure_ascii=False,indent=2),flush=True)


if __name__=='__main__':main()
