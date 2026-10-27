"""CPU audit and paired-question analysis of the isolated G2 -> G3 extension."""
import csv
import json
import math
import random
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np

from .core import ROOT, read_jsonl, write_json, write_jsonl, digest, cache_identity, training_example
from .formal_common import protocol, messages, sha, utc
from .g3_common import RUN, REPORT, REVISION, ENGINE, pool, adapter, raw, scores, validate, status
from .g3_engine import request_seed
from .v0_drift_analysis import count_array, count_metrics, clean, interval, divide

METRICS = ['FPR', 'FNR', 'balanced_accuracy', 'accepted_error_FDR',
           'G_accuracy_verified', 'G_accuracy_lower_all', 'unknown_fraction',
           'retention_all', 'truncated_fraction', 'mean_output_tokens']


def metrics(a):
    result = count_metrics(a)
    tp, fn, fp, tn, ua, ur, _, _ = np.moveaxis(np.asarray(a), -1, 0)
    result.update({'FPR_unknown_lower': divide(fp, fp+tn+ur),
                   'FPR_unknown_upper': divide(fp+ua, fp+tn+ua),
                   'FNR_unknown_lower': divide(fn, tp+fn+ua),
                   'FNR_unknown_upper': divide(fn+ur, tp+fn+ur)})
    return result


def paired(data, questions, draws=10000, seed=20261024):
    """Resample the SAME question indices for every branch, not candidate slots."""
    assert questions and len(set(questions)) == len(questions)
    expected = set(questions)
    assert all({r['question_id'] for r in rr} == expected for rr in data.values())
    samples = np.random.default_rng(seed).integers(len(questions), size=(draws, len(questions)))
    point, boot = {}, {}
    for name, rows in data.items():
        a = count_array(rows, questions)
        point[name] = clean(metrics(a.sum(0)))
        # Chunking avoids a draws x questions x 8 allocation for the A pool.
        sums = np.concatenate([a[ix].sum(1) for ix in np.array_split(samples, 20)])
        boot[name] = metrics(sums)
    return point, boot


def contrast(point, boot, after, before, split, role):
    result = []
    for metric in METRICS:
        a, b = point[after][metric], point[before][metric]
        ci = interval(boot[after][metric] - boot[before][metric])
        result.append({'split': split, 'role': role, 'contrast': after + '-' + before,
                       'metric': metric, 'estimate': a - b if a is not None and b is not None else None,
                       **ci, 'draws': 10000, 'seed': 20261024, 'unit': 'whole_question',
                       'unit_of_value': 'tokens' if metric == 'mean_output_tokens' else 'fraction'})
    return result


def load_new(name, split, tokenizer):
    from .g3_pipeline import label
    from .runtime import model_identity
    rows = read_jsonl(raw(name, split)); labs = label(name, split)
    vv = read_jsonl(scores(name, split)); v = {r['candidate_id']: r for r in vv}
    lab = {r['candidate_id']: r for r in labs}
    assert len(rows) == len(vv) == len(v) == len(lab) == 4 * len(pool(split))
    assert len({r['cache_key'] for r in rows}) == len(rows)
    base, ah = model_identity('G', adapter(name)); vf, va = model_identity('V')
    assert va is None and vf == protocol()['models']['V']['fingerprint']
    byq = {r['question_id']: r for r in pool(split)}; seen = defaultdict(set); out = []
    for r in rows:
        q, cid = r['question_id'], r['candidate_id']; vr = v[cid]
        assert r['question'] == byq[q]['question'] and r['G_round'] == name and r['split'] == split
        assert r['model_fingerprint'] == base and r['adapter_fingerprint'] == ah
        assert not any(k in r for k in ['gold_answer', 'correct', 'reference'])
        assert r['prompt_sha256'] == digest(messages('G', r['question']))
        assert r['engine_settings_sha256'] == digest(ENGINE)
        assert r['generation_seed'] == request_seed(q)
        assert r['finish_reason'] in ['stop', 'length']
        assert r['truncated'] == (r['finish_reason'] == 'length')
        assert r['output_tokens'] == len(r['generated_token_ids'])
        assert 0 < r['output_tokens'] <= r['max_new_tokens'] <= 3840
        assert r['prompt_tokens'] + r['output_tokens'] <= 4096
        assert tokenizer.decode(r['generated_token_ids'], skip_special_tokens=True) == r['text']
        assert vr['model_fingerprint'] == vf and vr['adapter_fingerprint'] is None
        assert vr['V_version'] == 'V0' and vr['threshold'] == .5 and vr['scoring_status'] == 'ok'
        assert vr['prompt_tokens'] + 4 <= 8192
        assert vr['accepted'] == (vr['score'] >= .5)
        x = vr['log_likelihood_A'] - vr['log_likelihood_B']
        recomputed = 1 / (1 + math.exp(-x)) if x >= 0 else math.exp(x) / (1 + math.exp(x))
        assert abs(vr['score'] - recomputed) < 2e-7
        msgs = messages('V', r['question'], r['text'])
        assert vr['prompt_sha256'] == digest(msgs)
        assert vr['cache_key'] == cache_identity('V', vf, None, cid, msgs, protocol()['V_scoring'])
        assert r['cache_key'] == digest({'protocol': sha(REVISION), 'adapter': ah, 'question': q,
            'index': r['candidate_index'], 'seed': request_seed(q), 'engine': digest(ENGINE)})
        seen[q].add(r['candidate_index'])
        out.append({**r, **lab[cid], 'accepted': vr['accepted'], 'score': vr['score'],
                    'V_prompt_tokens': vr['prompt_tokens']})
    assert set(seen) == set(byq) and all(ix == {0, 1, 2, 3} for ix in seen.values())
    return out


def old_C():
    allowed = {r['question_id'] for r in pool('C')}
    rr = read_jsonl(ROOT/'runs/formal/G2_C.jsonl')
    v = {r['candidate_id']: r for r in read_jsonl(ROOT/'runs/formal/V0_G2_C.jsonl')}
    labs = {r['candidate_id']: r for r in read_jsonl(ROOT/'runs/formal/labels_G2_C.jsonl')}
    out = [{**r, **labs[r['candidate_id']], 'accepted': v[r['candidate_id']]['accepted'],
            'score': v[r['candidate_id']]['score']} for r in rr if r['question_id'] in allowed]
    assert len(out) == 380
    return out


def supervision_audit(tokenizer):
    from .day3_protocol import select_feedback
    p = RUN/'stage_train/train/G2'
    selected = read_jsonl(p/'selected.jsonl'); tensors = read_jsonl(p/'tokenized.jsonl')
    steps = read_jsonl(p/'steps.jsonl'); report = json.loads((p/'report.json').read_text())
    actual = json.loads((RUN/'training_complete.json').read_text())
    cfg = validate()
    assert actual['actual_source'] == 'G2' and actual['actual_target'] == 'G3'
    assert actual['parent_adapter_fingerprint'] == cfg['parent_adapter_fingerprint']
    assert report['parent_adapter_fingerprint'] == cfg['parent_adapter_fingerprint']
    assert report['parent_probe_max_logit_difference_before_update'] <= .02
    assert report['save_reload_passed'] and report['base_frozen'] and report['only_LoRA_updated']
    assert report['consumed_examples_once'] and not report['truth_used_by_trainer']
    assert report['training_examples_truncated'] == 0
    assert actual['G3_adapter_fingerprint'] == report['final_adapter_fingerprint']
    assert actual['G3_adapter_fingerprint'] != cfg['parent_adapter_fingerprint']
    assert len(selected) == len(tensors) == report['selected_examples'] <= 512
    manifest = json.loads((p/'selection.manifest.json').read_text())
    vv = {r['candidate_id']: r for r in read_jsonl(scores('G2', 'A'))}
    replay = select_feedback(read_jsonl(raw('G2', 'A')), vv, .5, 42, 512, max_length=4096)
    expected = [r['candidate_id'] for r in replay if training_example(tokenizer, messages('G', r['question']), r['text'], 4096) is not None]
    assert expected == manifest['selected_ids'] == [r['candidate_id'] for r in selected]
    assert manifest['truth_used'] is False
    order = list(range(len(selected))); random.Random(42).shuffle(order)
    assert [[selected[i]['candidate_id'] for i in order[j:j+8]] for j in range(0, len(order), 8)] == [s['candidate_ids'] for s in steps]
    batch = {cid: (s['step'], len(s['candidate_ids'])) for s in steps for cid in s['candidate_ids']}
    assert len(batch) == len(selected) and len(steps) == report['optimizer_steps']
    rows = []
    for r, ex in zip(selected, tensors):
        assert not any(k in r for k in ['correct', 'reference', 'gold_answer'])
        assert r['selection_used_gold'] is False and r['V0_score'] >= .5
        rebuilt = training_example(tokenizer, messages('G', r['question']), r['text'], 4096)
        assert rebuilt is not None and all(rebuilt[k] == ex[k] for k in ['input_ids', 'labels', 'attention_mask', 'prefix_length'])
        prefix = ex['prefix_length']; L = len(ex['labels']) - prefix
        assert prefix > 0 and all(t == -100 for t in ex['labels'][:prefix])
        assert ex['labels'][prefix:] == ex['input_ids'][prefix:] and ex['input_ids'][-1] == 151645
        assert r['supervised_tokens'] == L and len(ex['input_ids']) <= 4096
        step, m = batch[r['candidate_id']]
        rows.append({'question_id': r['question_id'], 'candidate_id': r['candidate_id'], 'step': step,
                     'input_tokens': len(ex['input_ids']), 'supervised_tokens': L,
                     'prefix_tokens': prefix, 'accumulation_size': m,
                     'token_loss_coefficient': 1/(m*L), 'example_loss_coefficient': 1/m})
    assert sum(r['supervised_tokens'] for r in rows) == report['supervised_tokens'] == sum(s['supervised_tokens'] for s in steps)
    assert sum(r['input_tokens'] for r in rows) == report['input_tokens']
    csv_out(REPORT/'supervision_audit.csv', rows)
    write_json(REPORT/'actual_G3_training_report.json', {**report, 'checkpoint': 'G3_SFT',
        'source_generator': 'G2', 'adapter_path': str(adapter('G3')),
        'private_frozen_trainer_alias': 'Private G1 parent=actual G2; private G2 target=actual G3',
        'raw_private_report': str(p/'report.json')})
    return report, rows


def csv_out(path, rows):
    if not rows:
        path.write_text(''); return
    columns = list(dict.fromkeys(k for r in rows for k in r))
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, columns); w.writeheader()
        w.writerows([{k: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in clean(r).items()} for r in rows])


def summarize(data):
    points = []; differences = []; intervals = []
    for split in ['A', 'C']:
        qs = sorted(r['question_id'] for r in pool(split))
        view = {n: data[n, split] for n in ['G2', 'G3']}
        p, b = paired(view, qs)
        for n in ['G2', 'G3']:
            points.append({'split': split, 'model': n, 'engine': 'vllm', 'questions': len(qs), **p[n]})
            for m in METRICS:
                intervals.append({'split': split, 'model': n, 'metric': m, 'estimate': p[n][m], **interval(b[n][m])})
        differences += contrast(p, b, 'G3', 'G2', split, 'training_domain' if split == 'A' else 'development_held_out_from_training')
    bridge = {'G2_new_vllm': data['G2', 'C'], 'G2_old_TF': old_C()}
    p, b = paired(bridge, sorted(r['question_id'] for r in pool('C')))
    for name in bridge:
        points.append({'split': 'C_engine_bridge', 'model': name,
                       'engine': 'vllm' if name.endswith('vllm') else 'Transformers', 'questions': 95, **p[name]})
    differences += contrast(p, b, 'G2_new_vllm', 'G2_old_TF', 'C_engine_bridge', 'engine_and_resampling_bridge_not_training_effect')
    return points, differences, intervals


def diagnostics(data):
    perq = []; lengths = []; unknown = []; support = []
    for (name, split), rows in data.items():
        groups = defaultdict(list)
        for r in rows: groups[r['question_id']].append(r)
        for q, rr in groups.items():
            p = clean(count_metrics(count_array(rr, [q]).sum(0)))
            perq.append({'split': split, 'model': name, 'question_id': q,
                         **{k: p[k] for k in ['correct', 'wrong', 'unknown', 'TP', 'FN', 'FP', 'TN', 'mean_output_tokens']},
                         'answers': [r['extracted_answer'] for r in sorted(rr, key=lambda r: r['candidate_index'])]})
        for label, lo, hi in [('<=512', 0, 512), ('513-1024', 513, 1024), ('1025-2048', 1025, 2048), ('>2048', 2049, 10000)]:
            rr = [r for r in rows if lo <= r['output_tokens'] <= hi]
            if rr:
                p = clean(count_metrics(count_array(rr, sorted({r['question_id'] for r in rr})).sum(0)))
                lengths.append({'split': split, 'model': name, 'length_group': label, **p})
        for reason, n in Counter(r['unknown_reason'] for r in rows if r['correct'] is None).items():
            unknown.append({'split': split, 'model': name, 'reason': reason, 'count': n})
        for truth in [True, False, None]:
            rr = [r for r in rows if r['correct'] is truth]
            support.append({'split': split, 'model': name, 'truth': str(truth), 'answers': len(rr),
                            'questions': len({r['question_id'] for r in rr}),
                            'accepted': sum(r['accepted'] for r in rr),
                            'mean_tokens': float(np.mean([r['output_tokens'] for r in rr])) if rr else None})
    return perq, lengths, unknown, support


def workbook(tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in tables.items():
        ws = wb.create_sheet(name[:31]); columns = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(columns)
        for r in clean(rows):
            ws.append([json.dumps(r.get(k), ensure_ascii=False) if isinstance(r.get(k), (list, dict)) else r.get(k) for k in columns])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='264653')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(42, max(15, len(str(col[0].value or ''))+2))
    wb.save(REPORT/'g3_results.xlsx')


def figures(points, deltas):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    lookup = {(r['split'], r['model']): r for r in points}
    with PdfPages(REPORT/'g3_figures.pdf') as pdf:
        fig, axes = plt.subplots(2, 2, figsize=(10, 8), constrained_layout=True)
        for i, split in enumerate(['A', 'C']):
            for j, name in enumerate(['G2', 'G3']):
                ax = axes[i, j]; p = lookup[split, name]
                a = np.array([[p['TP'], p['FN']], [p['FP'], p['TN']]])
                normalized = a / a.sum(1, keepdims=True)
                ax.imshow(normalized, vmin=0, vmax=1, cmap='Blues')
                for y in range(2):
                    for x in range(2):
                        ax.text(x, y, f'{int(a[y,x])}\n{normalized[y,x]:.1%}', ha='center', va='center',
                                color='white' if normalized[y,x]>.6 else 'black')
                ax.set(xticks=[0,1], xticklabels=['V0 accepts', 'V0 rejects'], yticks=[0,1],
                       yticklabels=['Answer correct', 'Answer wrong'],
                       title=f'{split}: {name}, {p["questions"]} questions; unknown={int(p["unknown"])}')
        fig.suptitle('Fixed V0 confusion matrices; row percentages, unknown excluded')
        fig.savefig(REPORT/'confusion_matrices.png', dpi=160); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(1, 2, figsize=(11, 4), constrained_layout=True)
        for ax, split in zip(axes, ['A','C']):
            rr = [r for r in deltas if r['split']==split and r['metric'] in ['FPR','FNR','balanced_accuracy','unknown_fraction']]
            for i, r in enumerate(rr):
                v, lo, hi = [100*r[k] for k in ['estimate','low','high']]
                ax.plot([lo,hi], [i,i], color='#1976A3', linewidth=2); ax.scatter([v],[i], color='#1976A3')
            ax.axvline(0, color='grey', linestyle='--'); ax.set(yticks=range(len(rr)), yticklabels=[r['metric'] for r in rr],
                title=f'{split}: G3 minus G2 (same vLLM)', xlabel='Difference, percentage points; paired 95% CI')
        fig.savefig(REPORT/'paired_differences.png', dpi=160); pdf.savefig(fig); plt.close(fig)


def pct(value): return '未定义' if value is None else f'{100*value:.2f}%'
def pp(value): return '未定义' if value is None else f'{100*value:+.2f}'


def render(points, deltas, train, quality, audit):
    lookup = {(r['split'], r['model']): r for r in points}
    cfg = validate(); bench = cfg['engine_decision']
    lines = ['# 固定V0：新增G2→G3一轮SFT结果', '',
      '本轮检查生成器再更新一次后，固定V0是否更容易误接受或误拒绝。这里只检验输出分布上的验证风险变化，不以G能力提高为前提。', '',
      '## 协议与数据', '',
      '- 原始正式G2 adapter作为起点，只新增一轮response-only LoRA SFT；V0参数、提示、原Transformers A/B似然评分及0.5阈值保持固定。',
      '- G2/G3均用同一vLLM版本、BF16 LoRA推理、提示、采样与长度协议重新生成；每题4条。旧G2回答保留，仅作引擎桥接比较，不与新G3直接拼接。',
      '- A为598题（隔离缺图题与目标/参考冲突题），C为95题（原4道隔离题及另1道参考疑点题全部隔离）；两模型分别生成2392+380条。A为训练反馈题域；C未用于本轮G训练，但已反复用于开发分析，不能视为最终测试。B未使用，T继续封存。',
      '- 本轮隔离不追溯清除原G1/G2已经学过的两道问题题影响。原模型、缓存及003协议均保留。修订004单独记录本轮设置。',
      '- 按题配对bootstrap 10000次，seed=20261024；同题的4条回答整体重采样。候选序号不是相同随机轨迹或逐条反事实配对。区间为探索性点位95%区间，未校正多重比较，不包含训练种子变异。', '',
      '## 同引擎主结果', '',
      '|题池|模型|正确/错误/未知|FP/FN|FPR|FNR|平衡准确率|G已核验正确率|保留率|截断|均长token|',
      '|---|---|---|---|---|---|---|---|---|---|---|']
    for split in ['A','C']:
        for n in ['G2','G3']:
            r = lookup[split,n]
            lines.append(f'|{split}|{n}|{int(r["correct"])}/{int(r["wrong"])}/{int(r["unknown"])}|{int(r["FP"])}/{int(r["FN"])}|{pct(r["FPR"])}|{pct(r["FNR"])}|{pct(r["balanced_accuracy"])}|{pct(r["G_accuracy_verified"])}|{pct(r["retention_all"])}|{int(r["truncated"])}|{r["mean_output_tokens"]:.1f}|')
    lines += ['', 'FPR=错误回答被接受/可核验错误回答；FNR=正确回答被拒绝/可核验正确回答。未知不并入正确或错误；接受池已核验错误率FDR与FPR分母不同。完整CSV另含未知上下界、正确回答占全部候选比例和各类独立题数。', '',
      '![混淆矩阵](confusion_matrices.png)', '',
      '|题池|G3−G2指标|差值/百分点|按题配对95%区间/百分点|', '|---|---|---|---|']
    for r in deltas:
        if r['split'] in ['A','C'] and r['metric'] in ['FPR','FNR','balanced_accuracy','G_accuracy_verified','unknown_fraction','retention_all']:
            lines.append(f'|{r["split"]}|{r["metric"]}|{pp(r["estimate"])}|[{pp(r["low"])}, {pp(r["high"])}]|')
    lines += ['', '![配对差值](paired_differences.png)', '', '## 如何解释这次新增一轮', '']
    for split in ['A','C']:
        r = next(r for r in deltas if r['split']==split and r['metric']=='FPR')
        if r['low']>0: statement='误接受风险上升，按题区间未跨0，支持本轮该题域上出现退化信号。'
        elif r['high']<0: statement='误接受风险下降，按题区间未跨0，本轮该题域未支持更容易误接受。'
        else: statement='按题区间跨0；方向和量级如上，尚不能区分稳定变化与当前抽样波动。'
        lines.append(f'- {split}：FPR变化{pp(r["estimate"])}个百分点，区间[{pp(r["low"])}, {pp(r["high"])}]；{statement}')
    lines += ['', 'FPR与FNR有不同分母，不能用“FPR升3个百分点、FNR降1个百分点”相减当作总体准确率变化；平衡准确率同时反映两者。固定V不等于测到V参数退化；变化描述V面对更新后G所产生回答的适配程度。即使有同方向信号，也不能凭单次训练、单个开发集证明普遍规律或具体内部机制。', '',
      '## 引擎桥接检查：同一个G2、C95', '',
      '|引擎|FPR|FNR|G已核验正确率|未知比例|', '|---|---|---|---|---|']
    for n in ['G2_old_TF','G2_new_vllm']:
        r=lookup['C_engine_bridge',n]
        lines.append(f'|{n}|{pct(r["FPR"])}|{pct(r["FNR"])}|{pct(r["G_accuracy_verified"])}|{pct(r["unknown_fraction"])}|')
    lines += ['', '这组变化同时包含引擎、LoRA运行精度及新一批采样差异；不能单独识别某一实现的因果影响。它用于说明为什么保留同引擎训练前基线。', '',
      '## 实际训练数据与审计', '',
      f'实际选择{train["selected_examples"]}题，每题至多1条；正确/错误/未知={quality["correct"]}/{quality["wrong"]}/{quality["unknown"]}。可核验部分错误率{pct(quality["verified_error_fraction"])}。选择先固定，再连接独立答案记录质量；未按真答案筛选。', '',
      f'从原始G2继续同一个adapter，训练1 epoch，共{train["optimizer_steps"]}个优化步、{train["supervised_tokens"]:,}个有效监督token；学习率1e-4，LoRA r16/alpha32/dropout0.05，micro batch1、梯度累积8，seed42。训练与检查点耗时{train["training_and_checkpoint_seconds_this_invocation"]:.1f}秒，训练阶段总耗时{train["total_seconds_this_invocation"]:.1f}秒。', '',
      '训练仅更新LoRA，底座冻结；逐条重建input/label，prompt全部mask、assistant（含结束符）完整受监督，无静默截断；固定种子重放筛选及训练顺序通过。独立工程探针为“Compute 17 + 23.”，不读B/C/T；父adapter初始化、保存重载均核对通过。', '',
      '冻结训练器的私有目录沿用G1/G2编号，分别映射实际父G2/新G3；`actual_G3_training_report.json`、`actual_lineage.json`记录真实身份，不将私有别名误作新的原G2。', '',
      '## 速度与产物', '',
      f'匹配8题×4完整生成测速：vLLM {bench["vllm"]["total_seconds"]:.1f}秒，原4进程Transformers {bench["transformers_4_processes"]["total_seconds"]:.1f}秒；耗时下降{100*(1-bench["vllm"]["total_seconds"]/bench["transformers_4_processes"]["total_seconds"]):.1f}%。输出token数不同，该比例是此匹配题组的端到端耗时比，不能视为等token吞吐比。正式生成采用连续请求调度，实测成本见`resource_costs.json`。', '',
      '可复核文件：`point_metrics.csv`、`paired_deltas.csv`、`point_intervals.csv`、`question_metrics.csv`、`length_strata.csv`、`unknown_reasons.csv`、`label_question_support.csv`、`supervision_audit.csv`、`g3_results.xlsx`、`g3_figures.pdf`。完整回答、V分数与独立答案标签分别保留在runs/g3_extension_20261024，合并审查文本为`answers_joined.jsonl`。', '',
      f'审计完成：{audit["completed_at"]}。协议004 SHA256：`{sha(REVISION)}`。', '']
    (REPORT/'g3_report.md').write_text('\n'.join(lines))


def resource_costs(train):
    stages = []
    for name in ['G2','G3']:
        path = RUN/f'engine/{name}_continuous/complete.json'
        if path.exists(): stages.append({'stage': name+'_generation', **json.loads(path.read_text())})
        for split in ['A','C']:
            for r in read_jsonl(RUN/f'V0_scoring/{name}_{split}/inference_invocations.jsonl'):
                stages.append({'stage': f'V0_{name}_{split}', **r})
    ledger = [r for r in read_jsonl(ROOT/'runs/formal/work_ledger.jsonl') if 'g3_extension' in r.get('stage','')]
    write_json(REPORT/'resource_costs.json', {'stages': stages, 'training': train,
        'extension_GPU_ledger_entries': ledger, 'training_seconds': train['total_seconds_this_invocation'],
        'recorded_extension_GPU_hours': sum(r['seconds'] for r in ledger)/3600,
        'not_equal_to_all_historical_GPU_cost': True})


def main():
    from transformers import AutoTokenizer
    cfg=validate(); REPORT.mkdir(exist_ok=True)
    status('G3_report_CPU_audit')
    tok=AutoTokenizer.from_pretrained(protocol()['models']['G']['path'], local_files_only=True)
    data={(name, split): load_new(name, split, tok) for name in ['G2','G3'] for split in ['A','C']}
    train, supervision = supervision_audit(tok)
    points, deltas, intervals = summarize(data)
    perq, lengths, unknown, support = diagnostics(data)
    quality=json.loads((REPORT/'G3_training_quality.json').read_text())['quality']
    tables={'point_metrics': points, 'paired_deltas': deltas, 'point_intervals': intervals,
            'question_metrics': perq, 'length_strata': lengths, 'unknown_reasons': unknown,
            'label_question_support': support, 'supervision_audit': supervision}
    for n, rr in tables.items(): csv_out(REPORT/(n+'.csv'), rr)
    write_jsonl(REPORT/'answers_joined.jsonl', [r for rr in data.values() for r in rr])
    audit={'completed_at': utc(), 'rows': sum(map(len,data.values())), 'V0_fixed_and_full_context': True,
           'selection_replay_passed': True, 'supervision_rebuild_passed': True,
           'parent_actual_original_G2_verified': True, 'original_artifact_hashes_unchanged': True,
           'B_used': False, 'T_read': False, 'revision_sha256': sha(REVISION),
           'analysis_code_sha256': sha(Path(__file__)), 'new_engine_both_rounds': True}
    write_json(REPORT/'audit.json', audit); resource_costs(train)
    write_json(REPORT/'results.json', {'point': points, 'paired_deltas': deltas, 'intervals': intervals,
                                     'training_quality': quality, 'audit': audit})
    workbook(tables); figures(points,deltas); render(points,deltas,train,quality,audit)
    validate()
    manifest={str(p.relative_to(REPORT)): sha(p) for p in sorted(REPORT.rglob('*')) if p.is_file() and p.name not in ['bundle_manifest.json','g3_extension_review.zip']}
    write_json(REPORT/'bundle_manifest.json', manifest)
    with zipfile.ZipFile(REPORT/'g3_extension_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(REPORT.rglob('*')):
            if p.is_file() and p.name!='g3_extension_review.zip':z.write(p,str(p.relative_to(REPORT)))
    with zipfile.ZipFile(REPORT/'g3_extension_review.zip') as z: assert z.testzip() is None
    write_json(RUN/'state.json', {'status':'complete','completed_at':utc(), 'T_read':False,
                                'report':str(REPORT/'g3_report.md'),'revision_sha256':sha(REVISION)})
    status('complete', report=str(REPORT/'g3_report.md'))
    print(json.dumps({'complete':True,'report':str(REPORT/'g3_report.md')}, ensure_ascii=False),flush=True)


if __name__=='__main__':main()
