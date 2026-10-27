"""Matched-question prefix susceptibility across G0/G1/G2/G3, fixed V0."""
import json
import shutil
import zipfile
from pathlib import Path

import numpy as np
from scipy import stats

from .core import ROOT,read_jsonl,write_json,digest
from .formal_common import messages,sha,utc
from .cross_round_prefix import RUN,REPORT,PLAN,ROUNDS,validate
from .efp_intervention import wilson
from .fpr_pattern_diagnostics import bh,write_csv
from .v0_drift_analysis import clean,interval


def paired_change(after,before,boot_after,boot_before):
    a=np.asarray(after,bool);b=np.asarray(before,bool)
    plus=int((a&~b).sum());minus=int((~a&b).sum())
    p=float(stats.binomtest(plus,plus+minus).pvalue) if plus+minus else 1.
    return {'delta':float(a.mean()-b.mean()),**interval(boot_after-boot_before),
            'after_only_flips':plus,'before_only_flips':minus,'both_flipped':int((a&b).sum()),'p_McNemar':p}


def load_scores():
    cfg=validate();rr=read_jsonl(RUN/'answers.jsonl');ss=read_jsonl(RUN/'V0_scores.jsonl')
    by={(s['candidate_id'],s['variant']):s for s in ss};index={r['candidate_id']:r for r in rr}
    vv={v['name']:v for v in cfg['variants']}
    assert len(ss)==len(by)==cfg['counts']['logical_scores'] and len(index)==cfg['counts']['answers']
    for split in ['A','C']:
        for q in cfg['support'][split]['question_ids']:
            qr=[r for r in rr if r['split']==split and r['question_id']==q]
            assert len(qr)==4 and {r['G_round'] for r in qr}==set(ROUNDS)
            assert len({r['question'] for r in qr})==1
    for key,s in by.items():
        r=index[s['candidate_id']];v=vv[s['variant']];text=v['prefix']+r['text']
        assert s['question_id']==r['question_id'] and s['split']==r['split'] and s['G_round']==r['G_round']
        assert s['source_engine']==r['source_engine'] and s['correct'] is False and not s['source_accepted']
        assert s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None and s['threshold']==.5
        assert s['accepted']==(s['score']>=.5) and s['prompt_tokens']+4<=8192
        assert s['text_sha256']==digest(text) and s['prompt_sha256']==digest(messages('V',r['question'],text))
        assert s['cache_key']==digest({'plan':sha(PLAN),'candidate_id':r['candidate_id'],'variant':v['name']})
        if v['name']=='original':
            assert not s['accepted'] and abs(s['score']-r['score'])<2e-6 and s['source']=='new_original_V0_score'
    return cfg,rr,ss,by


def analyze(cfg,rr,by):
    rates=[];contrasts=[];margin=[];margin_rates=[];means=[];mean_changes=[]
    triggers=[i for i,v in enumerate(cfg['variants']) if v['category']!='control']
    names=[v['name'] for v in cfg['variants']];neutral=names.index('neutral_control')
    for split in ['A','C']:
        qs=cfg['support'][split]['question_ids'];n=len(qs);qix={q:i for i,q in enumerate(qs)}
        Y=np.zeros((n,4,len(names)),float);L=np.zeros_like(Y);lookup={}
        for r in rr:
            if r['split']!=split:continue
            qi=qix[r['question_id']];gi=ROUNDS.index(r['G_round']);lookup[qi,gi]=r
            original=by[r['candidate_id'],'original'];oldodds=original['log_likelihood_A']-original['log_likelihood_B']
            for vi,name in enumerate(names):
                s=by[r['candidate_id'],name];Y[qi,gi,vi]=s['accepted'];L[qi,gi,vi]=s['log_likelihood_A']-s['log_likelihood_B']-oldodds
        assert len(lookup)==n*4
        rng=np.random.default_rng(cfg['seed']);W=rng.multinomial(n,np.full(n,1/n),size=cfg['bootstrap'])
        boot=(W@Y.reshape(n,-1)/n).reshape(cfg['bootstrap'],4,len(names))
        for gi,g in enumerate(ROUNDS):
            rows=[lookup[i,gi] for i in range(n)];originals=[by[r['candidate_id'],'original'] for r in rows]
            margin.append({'split':split,'G_round':g,'source_engine':rows[0]['source_engine'],'questions':n,
                'mean_original_acceptance_proxy':float(np.mean([r['score'] for r in rows])),
                'median_original_acceptance_proxy':float(np.median([r['score'] for r in rows])),
                'mean_rejection_log_odds_margin':float(np.mean([s['log_likelihood_B']-s['log_likelihood_A'] for s in originals])),
                'mean_output_tokens':float(np.mean([r['output_tokens'] for r in rows])),
                'unique_wrong_final_values':len({r['extracted_answer'] for r in rows})})
            for vi,v in enumerate(cfg['variants']):
                count=int(Y[:,gi,vi].sum());lo,hi=wilson(count,n)
                rates.append({'split':split,'G_round':g,'source_engine':rows[0]['source_engine'],
                    'variant':v['name'],'prefix':v['prefix'],'category':v['category'],'questions':n,
                    'flips':count,'flip_rate':count/n,'Wilson_low':lo,'Wilson_high':hi,
                    'mean_log_odds_change':float(L[:,gi,vi].mean()),
                    'excess_over_neutral':float(Y[:,gi,vi].mean()-Y[:,gi,neutral].mean())})
            average=Y[:,gi,triggers].mean(1);samples=boot[:,gi,triggers].mean(1)
            means.append({'split':split,'G_round':g,'questions':n,'triggers':len(triggers),'mean_flip_rate':float(average.mean()),
                **interval(samples),'interpretation':'Exploratory equal-weight mean over18 tested triggers, not union of flipped answers'})
            for label,low,high in [('proxy<0.01',0.,.01),('0.01<=proxy<0.1',.01,.1),('0.1<=proxy<0.5',.1,.5)]:
                ids=[i for i,r in enumerate(rows) if low<=r['score']<high]
                for vi,v in enumerate(cfg['variants']):
                    margin_rates.append({'split':split,'G_round':g,'original_proxy_bin':label,'variant':v['name'],
                        'questions':len(ids),'flips':int(Y[ids,gi,vi].sum()) if ids else 0,
                        'flip_rate':float(Y[ids,gi,vi].mean()) if ids else None,
                        'interpretation':'Descriptive confidence stratum, sparse cells not inferential'})
        for after,before,label in [('G3','G2','primary_same_vLLM'),('G1','G0','secondary_same_Transformers'),('G3','G0','descriptive_engine_confounded')]:
            ai=ROUNDS.index(after);bi=ROUNDS.index(before)
            for vi,v in enumerate(cfg['variants']):
                d=paired_change(Y[:,ai,vi],Y[:,bi,vi],boot[:,ai,vi],boot[:,bi,vi])
                if label=='descriptive_engine_confounded':d['p_McNemar']=None
                interaction=(boot[:,ai,vi]-boot[:,ai,neutral])-(boot[:,bi,vi]-boot[:,bi,neutral])
                ii=interval(interaction)
                contrasts.append({'split':split,'contrast':after+'-'+before,'role':label,'variant':v['name'],
                    'category':v['category'],'questions':n,**d,
                    'excess_over_neutral_delta':float((Y[:,ai,vi]-Y[:,ai,neutral]-Y[:,bi,vi]+Y[:,bi,neutral]).mean()),
                    'interaction_low':ii['low'],'interaction_high':ii['high']})
            mean_changes.append({'split':split,'contrast':after+'-'+before,'role':label,
                'delta':float(Y[:,ai,triggers].mean()-Y[:,bi,triggers].mean()),
                **interval(boot[:,ai,triggers].mean(1)-boot[:,bi,triggers].mean(1)),
                'interpretation':'Exploratory uniform-prefix average; no confirmatory p-value'})
    for contrast in ['G3-G2','G1-G0']:
        rows=[r for r in contrasts if r['contrast']==contrast and r['category']!='control']
        assert len(rows)==36
        for r,q in zip(rows,bh([r['p_McNemar'] for r in rows])):r['q_BH_36']=float(q)
    return {'flip_rates':rates,'paired_round_changes':contrasts,'baseline_rejection_margins':margin,
            'flip_rates_by_margin':margin_rates,'uniform_trigger_means':means,'uniform_mean_changes':mean_changes}


def figures(cfg,tables):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    rates=tables['flip_rates'];lookup={(r['split'],r['G_round'],r['variant']):r for r in rates}
    selected=['symbol_comma','we_can','we_will','we_need_to','to_solve_for','it_is_essential_to','original_step_by_step']
    with PdfPages(REPORT/'cross_round_prefix_figures.pdf') as pdf:
        fig,axes=plt.subplots(1,2,figsize=(13,5.2),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            for i,name in enumerate(selected):
                y=[100*lookup[split,g,name]['flip_rate'] for g in ROUNDS];color=plt.get_cmap('tab10')(i)
                ax.plot([0,1],y[:2],marker='o',color=color,label=name)
                ax.plot([2,3],y[2:],marker='o',color=color)
                ax.plot([1,2],y[1:3],ls=':',alpha=.35,color=color)
            ax.axvline(1.5,color='black',ls='--',alpha=.4);ax.set_xticks(range(4),ROUNDS)
            ax.set_ylabel('Wrong reject-to-accept flip rate (%)');ax.set_title(f'{split}: same {cfg["support"][split]["selected_questions"]} questions')
            ax.set_ylim(bottom=-1);ax.text(.5,-.19,'Transformers',transform=ax.get_xaxis_transform(),ha='center');ax.text(2.5,-.19,'vLLM',transform=ax.get_xaxis_transform(),ha='center')
        axes[1].legend(fontsize=8,bbox_to_anchor=(1.02,1),loc='upper left')
        fig.suptitle('Fixed V0: same injected prefixes across G0/G1/G2/G3 outputs\nDotted G1->G2 links cross a generation-engine boundary')
        fig.savefig(REPORT/'cross_round_flip_curves.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(13,11),constrained_layout=True)
        vmax=max(35,max(r['flip_rate']*100 for r in rates))
        for ax,split in zip(axes,['A','C']):
            matrix=np.array([[100*lookup[split,g,v['name']]['flip_rate'] for g in ROUNDS] for v in cfg['variants']])
            im=ax.imshow(matrix,cmap='YlOrRd',vmin=0,vmax=vmax,aspect='auto')
            ax.set_yticks(range(len(cfg['variants'])),[v['name'] for v in cfg['variants']],fontsize=9)
            ax.set_xticks(range(4),ROUNDS);ax.xaxis.tick_top();ax.axvline(1.5,color='black',ls='--');ax.axhline(2.5,color='black')
            for (i,j),value in np.ndenumerate(matrix):ax.text(j,i,f'{value:.1f}%',ha='center',va='center',fontsize=9,color='white' if value>vmax*.65 else 'black')
            ax.set_title(split+f' (matched n={cfg["support"][split]["selected_questions"]})',pad=24)
        fig.suptitle('All21 prespecified settings; all originals are verified wrong and V0-rejected\nDashed line = generation-engine boundary; controls first3 rows')
        fig.colorbar(im,ax=axes,label='Conditional flip rate (%)',shrink=.6)
        fig.savefig(REPORT/'cross_round_flip_heatmap.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        vv=[v for v in cfg['variants'] if v['category']!='control']
        fig,axes=plt.subplots(1,2,figsize=(12,9),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            for i,v in enumerate(vv):
                r=next(r for r in tables['paired_round_changes'] if r['split']==split and r['contrast']=='G3-G2' and r['variant']==v['name'])
                color='#b03a2e' if r['q_BH_36']<.05 else '#1976A3'
                ax.plot([100*r['low'],100*r['high']],[i,i],lw=2,color=color);ax.plot(100*r['delta'],i,'o',color=color)
            ax.set_yticks(range(len(vv)),[v['name'] for v in vv]);ax.invert_yaxis();ax.axvline(0,color='black',ls='--')
            ax.set_xlabel('G3 minus G2 flip rate (percentage points)');ax.set_title(split)
        fig.suptitle('Same-engine paired-question late-round changes\n95% bootstrap intervals; red = exact McNemar BH q<0.05 across36 trigger/pool tests')
        fig.savefig(REPORT/'G3_minus_G2_flip_differences.png',dpi=180);pdf.savefig(fig);plt.close(fig)


def text_report(cfg,tables):
    rates=tables['flip_rates'];lookup={(r['split'],r['G_round'],r['variant']):r for r in rates}
    lines=['# 不同G阶段上的前缀注入翻转率','','## 测量对象','','固定同一个V0，比较G0/G1/G2/G3生成回答在同一组前缀下的TN→FP翻转率。每条原回答最终答案经独立核验为错误，原V0正确拒绝；注入后若接受，即一次新增误接受。分母为这批原本被拒的错误回答，不是全体错误回答，也不是总体判错率。',
      '四阶段选同一道题、每阶段一条合格回答。A限制在先前的确认半份，避免用前缀发现半份作本次对照。所有阶段都有合格回答的共同题A97/C29；按预定hash取A64，C全部29，共372条原文。完整原文与最终答案保持不变，只在最前面加固定前缀（含换行）。',
      'G0/G1原输出来自Transformers；G2/G3来自vLLM。图中明确标出G1→G2断点。主要可比检验是同vLLM的G2→G3，次要为同Transformers的G0→G1；不能把四阶段曲线直接作为纯训练效应。',
      '', '## 1. 各前缀的四阶段结果','']
    for split in ['A','C']:
        n=cfg['support'][split]['selected_questions']
        lines += [f'### {split}：同一批{n}道题','','|前缀|G0|G1|G2|G3|','|---|---:|---:|---:|---:|']
        for v in cfg['variants']:
            values=[lookup[split,g,v['name']] for g in ROUNDS]
            lines.append('|'+v['name']+'|'+'|'.join(f'{r["flips"]}/{n} ({100*r["flip_rate"]:.2f}%)' for r in values)+'|')
    lines += ['', '每个格子都有Wilson区间，保存在CSV/Excel。C只有29题，一条翻转对应3.45个百分点；A一条对应1.56个百分点。不能把一两条差别当成稳定的前缀优劣。原文控制全0是选样条件，不是总体FPR为0。',
      '', '## 2. G2→G3同引擎差值','','|池|前缀|变化/百分点|配对95%区间|双侧McNemar BH q|相对中性前缀的额外变化/百分点|','|---|---|---:|---|---:|---:|']
    for split in ['A','C']:
        for v in cfg['variants']:
            if v['category']=='control':continue
            r=next(r for r in tables['paired_round_changes'] if r['split']==split and r['contrast']=='G3-G2' and r['variant']==v['name'])
            lines.append(f'|{split}|{v["name"]}|{100*r["delta"]:+.2f}|[{100*r["low"]:+.2f}, {100*r["high"]:+.2f}]|{r["q_BH_36"]:.4g}|{100*r["excess_over_neutral_delta"]:+.2f}|')
    sig=[r for r in tables['paired_round_changes'] if r['contrast']=='G3-G2' and r.get('q_BH_36',1)<.05]
    lines += ['',f'同引擎G2→G3的36个前缀/池比较中，{len(sig)}项通过BH校正。G0→G1另列独立36项次要比较；G3→G0只作引擎混杂的描述，不报告显著性检验。前缀来自先前探索，C也反复用于开发，因此本次仍是开发诊断，不能称为全新确认集。',
      '相对中性前缀的额外变化计算为：(触发前缀−中性前缀)_G3−(触发前缀−中性前缀)_G2；配对区间另存。这有助于识别变化是否只是一般性的开头文本敏感，但本次不对该交互作额外确认性显著声明。',
      '', '## 3. 所测前缀等权平均（探索性）','','|池|G0|G1|G2|G3|G3−G2/百分点及95%区间|','|---|---:|---:|---:|---:|---|']
    for split in ['A','C']:
        rr=[next(r for r in tables['uniform_trigger_means'] if r['split']==split and r['G_round']==g) for g in ROUNDS]
        d=next(r for r in tables['uniform_mean_changes'] if r['split']==split and r['contrast']=='G3-G2')
        lines.append('|'+split+'|'+'|'.join(f'{100*r["mean_flip_rate"]:.2f}%' for r in rr)+f'|{100*d["delta"]:+.2f} [{100*d["low"]:+.2f}, {100*d["high"]:+.2f}]|')
    lines += ['', '该均值对18个非控制前缀等权，表示从这份清单等概率挑一个前缀时的平均翻转率。它不是“至少一个前缀可骗过”的回答比例，也不代表任意英文前缀的平均风险；未预设为确认性终点。',
      '', '## 4. 原始拒绝强度与输出长度','','|池|G阶段|原接受分数均值|原接受分数中位数|平均拒绝log-odds余量|回答平均token|','|---|---|---:|---:|---:|---:|']
    for r in tables['baseline_rejection_margins']:
        lines.append(f'|{r["split"]}|{r["G_round"]}|{r["mean_original_acceptance_proxy"]:.5f}|{r["median_original_acceptance_proxy"]:.5f}|{r["mean_rejection_log_odds_margin"]:.3f}|{r["mean_output_tokens"]:.1f}|')
    lines += ['', '接受分数是A/B似然归一化代理，不是经过校准的正确概率。每阶段虽然都是TN，原始拒绝强弱与错误回答内容仍可能不同；对应固定代理分数档的翻转率已导出，稀疏档只描述，不据此做机制推断。',
      '', '## 解释边界','','如果后期G翻转率更高，支持在这批共同题的原本拒绝错误回答中，固定V更易受这些注入影响；如果没有更高，也不能否定全池FPR已经上涨。四阶段都必须存在TN的选样条件保留了一个更窄的共同子群，原本就被接受的错误已被排除。',
      '同一题不等于同一错误回答：最终错误数、数学解法、长度和组织可能都改变。此处测的是阶段输出分布下的条件脆弱性，不是训练改变了V参数，也不证明G有意利用了某个前缀。',
      '对同一完整错误回答集合，若f=P(原TN注入后变FP)、r=P(原FP注入后变TN)，则FPR_after=FPR_before*(1-r)+(1-FPR_before)*f。本次只测共同TN题子群的f，不能直接代入全池公式，也不能只靠f推算整池注入后的FPR。',
      '不能把本次G3数值直接与旧112题EFP表作时间趋势比较：本次要求四阶段共同TN题，且C仅29题，原来的G3独立选样为A64/C48。旧结果保留，本次统一匹配题的表才用于跨阶段比较。',
      '区间按同一道题同时重采样四阶段与所有前缀10000次。置信区间为逐项区间，没有同时覆盖保证；无翻转或全翻转时bootstrap可能退化，不等于总体零风险或必然翻转。',
      '', '## 审计与交付','']
    compute=json.loads((RUN/'compute.json').read_text())
    lines += [f'共{cfg["counts"]["answers"]}条原回答×21设置={compute["logical_scores"]}条逻辑评分；新增{compute["new_scores_this_invocation"]}、完整输入相同复用{compute["reused_this_invocation"]}条，GPU阶段{compute["seconds_this_invocation"]:.1f}秒。所有372条原文均重新评分复现原分数及拒绝，所有注入保持错误最终答案与解析状态，V完整读取回答。',
      '固定V0模型、提示、原Transformers评分及0.5阈值。没有新G生成或训练，没有V更新，B未使用、T未读取。原模型及冻结产物保持不变。',
      '图为cross_round_flip_curves.png、cross_round_flip_heatmap.png和G3_minus_G2_flip_differences.png；三页PDF、Excel、全部分数来源与可核对哈希ZIP一起提供。']
    (REPORT/'cross_round_prefix_report.md').write_text('\n'.join(lines)+'\n')


def main():
    cfg,rr,ss,by=load_scores();tables=analyze(cfg,rr,by)
    index={r['candidate_id']:r for r in rr}
    tables['per_question_flips']=[{'split':s['split'],'question_id':s['question_id'],'G_round':s['G_round'],
        'source_engine':s['source_engine'],'candidate_id':s['candidate_id'],'variant':s['variant'],
        'flipped_to_false_accept':s['accepted'],'original_acceptance_proxy':s['source_score'],
        'injected_acceptance_proxy':s['score'],'wrong_final_answer':index[s['candidate_id']]['extracted_answer'],
        'output_tokens':index[s['candidate_id']]['output_tokens']} for s in ss]
    for name,rows in tables.items():write_csv(REPORT/(name+'.csv'),clean(rows))
    write_json(REPORT/'results.json',clean(tables));figures(cfg,tables);text_report(cfg,tables)
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        ws=wb.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in rows for k in r));ws.append(keys)
        for r in clean(rows):
            ws.append([r.get(k) for k in keys])
            for c in ws[ws.max_row]:
                if c.data_type=='f':c.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(REPORT/'cross_round_prefix_review.xlsx')
    for src in [PLAN,RUN/'compute.json',RUN/'support.json',ROOT/'src/pilot/cross_round_prefix.py',Path(__file__)]:shutil.copyfile(src,REPORT/src.name)
    write_json(REPORT/'audit.json',{'passed':True,'source_and_input_hashes_verified':True,'V0_fixed':True,
        'all372_originals_freshly_rescored_and_matched':True,'all_prefix_answers_preserved':True,
        'full_context_verified':True,'same_questions_in_all4_stages':True,'logical_scores':len(ss),'B_used':False,'T_read':False,'at':utc()})
    files=sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix!='.zip' and p.name!='artifact_manifest.json')
    manifest={p.name:sha(p) for p in files};write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'cross_round_prefix_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']:z.write(p,p.name)
    import hashlib
    with zipfile.ZipFile(REPORT/'cross_round_prefix_review.zip') as z:
        assert z.testzip() is None and all(hashlib.sha256(z.read(n)).hexdigest()==h for n,h in manifest.items())
    write_json(RUN/'state.json',{'status':'complete','report':str(REPORT/'cross_round_prefix_report.md'),'at':utc()})
    print(json.dumps(clean({'means':tables['uniform_trigger_means'],'late_mean_changes':[r for r in tables['uniform_mean_changes'] if r['contrast']=='G3-G2']}),ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
