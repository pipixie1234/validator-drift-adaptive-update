"""Audit and export all AWL contrasts, including negative results."""
import json
import shutil
import zipfile

import numpy as np
from scipy import stats

from .core import ROOT, read_jsonl, write_json, digest
from .formal_common import messages, sha, utc
from .awl_comparison import RUN, REPORT, PLAN, validate, membership
from .prefix_paraphrase import RUN as PRIOR, PLAN as PRIOR_PLAN
from .fpr_pattern_diagnostics import bh, write_csv
from .efp_intervention import wilson
from .v0_drift_analysis import clean


def matrix_contrast(awl, non, draws=10000):
    """Keep all lexical pairs together in primary resampling of answer/question."""
    a=np.asarray(awl,float);b=np.asarray(non,float)
    assert a.shape==b.shape and a.ndim==2 and min(a.shape)>1
    d=a-b;n,k=d.shape;per_answer=d.mean(1);delta=float(d.mean())
    rng=np.random.default_rng(20261024)
    qi=rng.integers(n,size=(draws,n));boot=per_answer[qi].mean(1)
    null=(rng.choice([-1,1],size=(draws,n))*per_answer).mean(1)
    p=(1+int((abs(null)>=abs(delta)-1e-12).sum()))/(draws+1)
    crossed=[]
    for _ in range(draws):
        ix=rng.integers(n,size=n);jx=rng.integers(k,size=k)
        crossed.append(d[np.ix_(ix,jx)].mean())
    lo,hi=np.quantile(boot,[.025,.975]);cl,ch=np.quantile(crossed,[.025,.975])
    return {'questions':n,'lexical_pairs':k,'AWL_mean_flip_rate':float(a.mean()),
        'nonAWL_mean_flip_rate':float(b.mean()),'delta_AWL_minus_nonAWL':delta,
        'question_CI_low':float(lo),'question_CI_high':float(hi),
        'crossed_CI_low':float(cl),'crossed_CI_high':float(ch),'p_question_signflip':p,
        'increased_questions':int((per_answer>0).sum()),'decreased_questions':int((per_answer<0).sum()),
        'equal_questions':int((per_answer==0).sum()),'AWL_higher_lexical_pairs':int((d.mean(0)>0).sum()),
        'AWL_lower_lexical_pairs':int((d.mean(0)<0).sum()),'equal_lexical_pairs':int((d.mean(0)==0).sum())}


def load_scores():
    cfg=validate();rr=read_jsonl(RUN/'answers.jsonl');ss=read_jsonl(RUN/'V0_scores.jsonl')
    assert len(ss)==cfg['scores'] and len({s['cache_key'] for s in ss})==len(ss)
    index={r['candidate_id']:r for r in rr};vv={v['name']:v for v in cfg['variants']};by={}
    for s in ss:
        r=index[s['candidate_id']];v=vv[s['variant']];text=v['prefix']+r['text']
        assert s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None
        assert s['threshold']==.5 and s['accepted']==(s['score']>=.5) and s['correct'] is False
        assert s['question_id']==r['question_id'] and s['split']==r['split'] and s['group']==v['group']
        assert s['text_sha256']==digest(text) and s['prompt_sha256']==digest(messages('V',r['question'],text))
        assert s['cache_key']==digest({'plan':sha(PLAN),'candidate_id':r['candidate_id'],'variant':v['name']})
        assert s['prompt_tokens']+4<=8192
        if v['name']=='original':assert not s['accepted'] and abs(s['score']-r['score'])<2e-6
        by[s['candidate_id'],s['variant']]=s
    assert len(by)==len(rr)*len(vv)
    return cfg,rr,ss,by


def analyze(cfg,rr,by):
    rows=[];pairs=[];groups=[]
    for split in ['A','C']:
        answers=[r for r in rr if r['split']==split];n=len(answers)
        for v in cfg['variants']:
            scores=[by[r['candidate_id'],v['name']] for r in answers]
            y=np.array([s['accepted'] for s in scores],float);lo,hi=wilson(int(y.sum()),n)
            rows.append({'split':split,**v,'responses':n,'false_accepts':int(y.sum()),'flip_rate':float(y.mean()),
                'Wilson_low':lo,'Wilson_high':hi,'mean_score':float(np.mean([s['score'] for s in scores])),
                'median_added_prompt_tokens':float(np.median([s['prompt_tokens']-by[r['candidate_id'],'original']['prompt_tokens'] for r,s in zip(answers,scores)]))})
        a=np.array([[by[r['candidate_id'],p['AWL']]['accepted'] for p in cfg['pairs']] for r in answers],float)
        b=np.array([[by[r['candidate_id'],p['nonAWL']]['accepted'] for p in cfg['pairs']] for r in answers],float)
        groups.append({'split':split,**matrix_contrast(a,b)})
        for j,p in enumerate(cfg['pairs']):
            d=a[:,j]-b[:,j];gained=int((d>0).sum());lost=int((d<0).sum())
            rng=np.random.default_rng(20261024);boot=d[rng.integers(n,size=(10000,n))].mean(1)
            lo,hi=np.quantile(boot,[.025,.975])
            pv=float(stats.binomtest(gained,gained+lost).pvalue) if gained+lost else 1.
            pairs.append({'split':split,'pair':j,**p,'responses':n,'AWL_flip_rate':float(a[:,j].mean()),
                'nonAWL_flip_rate':float(b[:,j].mean()),'delta':float(d.mean()),'low':float(lo),'high':float(hi),
                'nonAWL_reject_AWL_accept':gained,'nonAWL_accept_AWL_reject':lost,'p_McNemar':pv})
    for r,q in zip(groups,bh([r['p_question_signflip'] for r in groups])):r['q_BH_2_pools']=float(q)
    for r,q in zip(pairs,bh([r['p_McNemar'] for r in pairs])):r['q_BH_40_pairs']=float(q)
    return rows,pairs,groups


def existing_classification():
    members=membership();cfg=json.loads(PRIOR_PLAN.read_text());ss=read_jsonl(PRIOR/'V0_scores.jsonl');rows=[]
    for v in cfg['variants']['prefix']:
        words=[w.lower() for w in __import__('re').findall(r"[A-Za-z]+(?:['’][A-Za-z]+)?",v['prefix'])]
        matched=sorted(set(w for w in words if w in members))
        group='control' if v['category']=='control' else 'AWL_present' if matched else 'AWL_absent'
        for split in ['A','C']:
            selected=[s for s in ss if s['arm']=='prefix' and s['variant']==v['name'] and s['split']==split]
            assert len(selected)==({'A':64,'C':48}[split])
            rows.append({'split':split,'variant':v['name'],'prefix':v['prefix'],'group':group,
                'AWL_members':matched,'AWL_families':sorted({f['headword'] for w in matched for f in members[w]}),
                'responses':len(selected),'flip_rate':sum(s['accepted'] for s in selected)/len(selected)})
    return rows


def audit_transcription():
    from bs4 import BeautifulSoup
    from collections import Counter
    soup=BeautifulSoup((RUN/'awl_source.html').read_text(),'html.parser')
    table=next(t for t in soup.find_all('table') if len(t.find_all('tr'))==571)
    families=[]
    for tr in table.find_all('tr')[1:]:
        cells=tr.find_all('td',recursive=False)
        head=cells[0].get_text(strip=True);sub=int(cells[1].get_text(strip=True))
        forms={head}|{s.strip() for s in cells[2].get_text().split(',') if s.strip()}
        families.append({'headword':head,'sublist':sub,'forms':sorted(forms)})
    assert families==json.loads((RUN/'awl_570_transcription.json').read_text())
    assert len(families)==len({f['headword'] for f in families})==570
    out={'passed':True,'family_count':570,'sublist_counts':dict(Counter(f['sublist'] for f in families)),
         'unique_surface_forms':len(membership()),'raw_HTML_sha256':sha(RUN/'awl_source.html'),
         'raw_HTML_to_transcription_exact_match':True,'transcription_source':'EAPFoundation; university direct downloads HTTP410'}
    write_json(REPORT/'source_audit.json',out)
    return out


def figures(cfg,rows,groups,existing):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    by={(r['split'],r['name']):r for r in rows}
    with PdfPages(REPORT/'AWL_comparison_figures.pdf') as pdf:
        data=[];labels=[]
        for p in cfg['pairs']:
            data.append([100*by[s,p[g]]['flip_rate'] for s in ['A','C'] for g in ['AWL','nonAWL']])
            labels.append(p['AWL']+'  /  '+p['nonAWL'])
        for v in cfg['variants'][:3]:
            data.append([100*by[s,v['name']]['flip_rate'] for s in ['A','C'] for _ in range(2)])
            labels.append('Control: '+v['name'])
        data.append([100*next(r for r in groups if r['split']==s)[g+'_mean_flip_rate'] for s in ['A','C'] for g in ['AWL','nonAWL']])
        labels.append('Mean across 20 pairs')
        m=np.array(data);fig,ax=plt.subplots(figsize=(11.8,13.5),constrained_layout=True)
        im=ax.imshow(m,cmap='YlOrRd',vmin=0,vmax=max(30,float(m.max())),aspect='auto')
        ax.set_yticks(range(len(labels)),labels);ax.set_xticks(range(4),['A: AWL','A: non-AWL','C: AWL','C: non-AWL'])
        ax.xaxis.tick_top();ax.axvline(1.5,color='black',lw=2)
        for k in [19.5,22.5]:ax.axhline(k,color='black',lw=1.5)
        for (i,j),v in np.ndenumerate(m):ax.text(j,i,f'{v:.1f}%',ha='center',va='center',fontsize=10,color='white' if v>m.max()*.65 else 'black')
        ax.set_title('AWL vs non-AWL matched verb headers\nSame wrong, originally rejected answers: A n=64; C n=48\nPrefix = Capitalized word + colon + newline',pad=48)
        fig.colorbar(im,ax=ax,label='Wrong reject-to-accept flip rate (%)',shrink=.6)
        fig.savefig(REPORT/'AWL_matched_pairs_heatmap.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        vv=cfg['variants'][:3]+[v for g in ['AWL','nonAWL'] for v in cfg['variants'] if v['group']==g]
        m=np.array([[100*by[s,v['name']]['flip_rate'] for s in ['A','C']] for v in vv])
        fig,ax=plt.subplots(figsize=(7,19),constrained_layout=True)
        im=ax.imshow(m,cmap='YlOrRd',vmin=0,vmax=max(30,float(m.max())),aspect='auto')
        ax.set_yticks(range(len(vv)),[f'[{v["group"]}] {v["prefix"].strip() or "(unchanged)"}' for v in vv],fontsize=10)
        ax.set_xticks([0,1],['A (64 answers)','C (48 answers)']);ax.xaxis.tick_top()
        for k in [2.5,22.5]:ax.axhline(k,color='black',lw=1.5)
        for (i,j),v in np.ndenumerate(m):ax.text(j,i,f'{v:.1f}%',ha='center',va='center',color='white' if v>m.max()*.65 else 'black')
        ax.set_title('All prespecified AWL / non-AWL headers\nWrong reject-to-accept flips; unchanged answer bodies',pad=40)
        fig.colorbar(im,ax=ax,label='Flip rate (%)',shrink=.5)
        fig.savefig(REPORT/'AWL_membership_heatmap.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        fig,ax=plt.subplots(figsize=(9,3.5),constrained_layout=True)
        for i,r in enumerate(groups):
            x=100*r['delta_AWL_minus_nonAWL']
            ax.plot([100*r['crossed_CI_low'],100*r['crossed_CI_high']],[i+.07,i+.07],lw=3,color='#90a4ae',label='Questions + lexical pairs' if i==0 else None)
            ax.plot([100*r['question_CI_low'],100*r['question_CI_high']],[i-.07,i-.07],lw=3,color='#b03a2e',label='Questions, fixed vocabulary' if i==0 else None)
            ax.plot(x,i,'ko');ax.text(x,i+.17,f'{x:+.2f} pp',ha='center')
        ax.axvline(0,color='black',ls='--');ax.set_yticks([0,1],['A','C']);ax.set_ylim(-.4,1.5)
        ax.set_xlabel('Mean flip-rate difference: AWL minus non-AWL (percentage points)')
        ax.set_title('Group contrast: paired 95% bootstrap intervals\n20 purposively selected lexical pairs; not all AWL vocabulary')
        ax.legend(loc='lower right');fig.savefig(REPORT/'AWL_group_difference.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        names=[r['variant'] for g in ['control','AWL_present','AWL_absent'] for r in existing if r['split']=='A' and r['group']==g]
        lookup={(r['split'],r['variant']):r for r in existing}
        m=np.array([[100*lookup[s,v]['flip_rate'] for s in ['A','C']] for v in names])
        labels=[f'[{lookup["A",v]["group"]}] '+(lookup['A',v]['prefix'].strip() or '(unchanged)') for v in names]
        fig,ax=plt.subplots(figsize=(10,21),constrained_layout=True)
        im=ax.imshow(m,cmap='YlOrRd',vmin=0,vmax=30,aspect='auto')
        ax.set_yticks(range(len(names)),labels,fontsize=9);ax.set_xticks([0,1],['A (64 answers)','C (48 answers)']);ax.xaxis.tick_top()
        for (i,j),v in np.ndenumerate(m):ax.text(j,i,f'{v:.1f}%',ha='center',va='center',color='white' if v>=20 else 'black')
        ax.set_title('Earlier 46 prefixes annotated by AWL membership\nDescriptive only: phrase length, syntax and punctuation differ',pad=40)
        fig.colorbar(im,ax=ax,label='Wrong reject-to-accept flip rate (%)',shrink=.5)
        fig.savefig(REPORT/'existing_prefixes_AWL_heatmap.png',dpi=150);pdf.savefig(fig);plt.close(fig)


def workbook(tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        ws=wb.create_sheet(name);cols=list(dict.fromkeys(k for r in rows for k in r));ws.append(cols)
        for r in clean(rows):
            ws.append([json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(list,dict)) else r.get(k) for k in cols])
            for c in ws[ws.max_row]:
                if c.data_type=='f':c.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True)
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
    wb.save(REPORT/'AWL_review.xlsx')


def main():
    cfg,rr,ss,by=load_scores();audit_transcription();rows,pairs,groups=analyze(cfg,rr,by);old=existing_classification()
    tables={'group_contrast':groups,'word_pairs':pairs,'all_headers':rows,'earlier_prefix_annotation':old}
    for name,data in tables.items():write_csv(REPORT/(name+'.csv'),clean(data))
    write_json(REPORT/'results.json',clean(tables));workbook(tables);figures(cfg,rows,groups,old)
    lines=['# AWL 内外词汇对固定 V0 的影响','','## 结论与组间对比','',
        '同一批 G3 错误且原本被拒绝的回答，分别注入 AWL 内外各20个动词。正文、最终答案、冒号、换行、V0参数/提示/评分/0.5阈值保持相同。',
        '表中是原本拒绝→误接受的翻转率，不是全部错误回答上的总体 FPR。两组均值各对20个词等权，再对同一批回答等权。',
        '', '| 池 | 独立题/回答 | AWL均值 | 非AWL均值 | 差值/百分点 | 题目配对95%区间/百分点 | 题目+词对重采样95%区间 | BH q（两池） |',
        '|---|---:|---:|---:|---:|---|---|---:|']
    for r in groups:
        lines.append(f'| {r["split"]} | {r["questions"]} | {100*r["AWL_mean_flip_rate"]:.2f}% | {100*r["nonAWL_mean_flip_rate"]:.2f}% | {100*r["delta_AWL_minus_nonAWL"]:+.2f} | [{100*r["question_CI_low"]:+.2f}, {100*r["question_CI_high"]:+.2f}] | [{100*r["crossed_CI_low"]:+.2f}, {100*r["crossed_CI_high"]:+.2f}] | {r["q_BH_2_pools"]:.4g} |')
    for r in groups:
        sign='高于' if r['delta_AWL_minus_nonAWL']>0 else '低于' if r['delta_AWL_minus_nonAWL']<0 else '等于'
        evidence='题目条件检验通过两池BH校正' if r['q_BH_2_pools']<.05 else '题目条件检验未通过两池BH校正'
        lines += ['',f'{r["split"]}池：AWL均值{sign}非AWL，{evidence}；20词对中AWL较高{r["AWL_higher_lexical_pairs"]}对、较低{r["AWL_lower_lexical_pairs"]}对、相同{r["equal_lexical_pairs"]}对。']
    lines += ['', '## 图的读法','','- `AWL_matched_pairs_heatmap.png`：每行一组近义动词，左半为A池（AWL/非AWL两格），右半为C池。所有格用同一色标；数字为翻转率。最后一行为20词对的均值。',
      '- `AWL_membership_heatmap.png`：与此前图一致，A/C两列；按控制、AWL、非AWL分块，全部43个前缀展示。',
      '- `AWL_group_difference.png`：AWL减非AWL的差值；红线只重采样题目，灰线同时重采样题目和词对，竖线0表示无差异。',
      '- `existing_prefixes_AWL_heatmap.png`：给上一次46个前缀事后标注AWL成员。句式、长度不同，只作描述，不当成干净的词表效应对照。',
      '', '## AWL来源与分类','','AWL是Coxhead（2000）的570词族，按照跨学科范围、频率及均匀度选取，并排除最常见2000词。因此“非AWL”不表示“不学术”，AWL子表序号也是频率分组，不是学术性强弱等级。',
      '定义、范围、词族规则见[大学官方说明](https://www.wgtn.ac.nz/lals/resources/academicwordlist/information)、[选择原则](https://www.wgtn.ac.nz/lals/resources/academicwordlist/information/principles)和[官方词族文件](https://www.wgtn.ac.nz/lals/resources/academicwordlist/publications/awlsublists1.pdf)。',
      '本次大学源可由搜索索引查证，但直接下载返回HTTP410；机器可读成员表取自[EAPFoundation转录](https://www.eapfoundation.com/vocab/academic/awllists/)，原始HTML保留，已校验570个不同词族以及1–9子表各60、10子表30。分类按词族已列成员精确匹配，包含屈折/派生，不用猜词干。此转录来源限制明确保留。',
      '`require`及其变形、`analysis`在表内；`essential`、`necessary`、`determine`、`reasoning`不在表内。上次`essential`效果较高本身不能支持AWL效应。',
      '', '## 完整词对','','| AWL词 | 词族/子表 | 非AWL近义动词 |','|---|---|---|']
    members=membership()
    for p in cfg['pairs']:lines.append(f'| {p["AWL"]} | '+', '.join(f'{f["headword"]}/{f["sublist"]}' for f in members[p['AWL']])+f' | {p["nonAWL"]} |')
    significant=[r for r in pairs if r['q_BH_40_pairs']<.05]
    compute=json.loads((RUN/'compute.json').read_text())
    lines += ['', '## 协议、核验和解释限制','',
      f'共112道独立题，A确认集64、C开发集48；每题43种设置，共{len(ss)}条评分。原回答全部错误、未截断且被V0拒绝。未更换题目或重生成回答。',
      '20个词对在新评分之前冻结，均为单个动词；各词统一首字母大写、冒号、单换行。所有改写独立答案提取、标签和状态保持一致；V读取完整回答，逐条检查8192上下文。所有112条原文重新评分，与历史值/决定一致。',
      '两池均值差为预设主比较：固定词表、以题目为单位做10000次配对bootstrap及双侧题目符号翻转检验，两池BH校正。每个词对用双侧精确McNemar，跨40个词对比较BH校正。所有95%区间为逐项区间，不是同时置信区间。',
      f'40个词对比较中{len(significant)}个通过BH校正。题目+词对的交叉bootstrap仅作为词汇选择敏感性检查；由于20组词为人工挑选，这不能替代对570词族的随机抽样。',
      '近义词语气/含义只能尽量接近，词长、分词长度、词频未严格匹配。统一标点排除了不同标点数量的混杂，但没有排除词汇与冒号格式的交互。结果只适用于这些人工标题，不解释自然G3 FPR上涨，更不能证明G有意利用学术词。',
      f'新评分{compute["new_scores_this_invocation"]}条、复用字节/提示一致旧评分{compute["reused_scores_this_invocation"]}条，本次GPU阶段{compute["seconds_this_invocation"]:.1f}秒。G未生成/训练、V未更新，B未使用、T未读取。',
      '', '原冻结产物保持不变。完整分词增量、所有词对差值、原始评分与来源哈希见Excel/CSV及runs目录。']
    (REPORT/'AWL_comparison_report.md').write_text('\n'.join(lines)+'\n')
    for src in [PLAN,RUN/'awl_570_transcription.json',RUN/'awl_source.html',RUN/'sources.json',RUN/'compute.json']:
        shutil.copyfile(src,REPORT/src.name)
    for name in ['awl_comparison.py','awl_comparison_report.py']:shutil.copyfile(ROOT/'src/pilot'/name,REPORT/name)
    write_json(REPORT/'audit.json',{'passed':True,'scores':len(ss),'unique_questions':len(rr),'fresh_originals_matched':112,
        'labels_preserved':True,'full_V_context':True,'source_hashes_verified':True,'V0_fixed':True,'at':utc()})
    paths=sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix!='.zip' and p.name!='artifact_manifest.json')
    manifest={p.name:sha(p) for p in paths};write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'AWL_comparison_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in paths+[REPORT/'artifact_manifest.json']:z.write(p,p.name)
    with zipfile.ZipFile(REPORT/'AWL_comparison_review.zip') as z:
        assert z.testzip() is None
        import hashlib
        assert all(hashlib.sha256(z.read(name)).hexdigest()==h for name,h in manifest.items())
    write_json(RUN/'state.json',{'status':'complete','report':str(REPORT/'AWL_comparison_report.md'),'at':utc()})
    print(json.dumps(clean(groups),ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
