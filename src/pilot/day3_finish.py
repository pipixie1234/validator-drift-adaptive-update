"""Audit Day3 artifacts, export complete review materials and report measured limits."""
import json
from collections import Counter
from .core import ROOT,read_jsonl,write_json,write_jsonl,digest
from .dapo_common import BASE,RUN,sha,messages
from .day3_protocol import OUT,DOC,quarantined_ids,quality,formal_generation_settings
from .runtime import model_identity

def excel_literal(text):
    # XML cannot represent several C0 controls present in upstream math strings.
    # Show explicit escapes in Excel; the separate JSONL retains the exact original.
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub(lambda m:f'\\u{ord(m[0]):04x}',text)

def audit():
    protected=json.loads((RUN/'protected_day2.json').read_text())
    assert all(sha(ROOT/p)==h for p,h in protected.items()),'Changed frozen day2 artifact'
    protocol=json.loads((ROOT/'configs/protocol_revisions/003_dapo_day3.json').read_text())
    assert protocol['preserved_original_G400_sha256']==sha(RUN/'G0_dapo100.jsonl')
    assert protocol['preserved_original_V400_sha256']==sha(RUN/'V0_dapo100.jsonl')
    assert protocol['preserved_day2_config_sha256']==sha(ROOT/'configs/formal_protocol.json')
    for key,path in [('answer_protocol_sha256','configs/dapo_answer_v3.json'),
                     ('working_point_sha256','reports/day3/working_point.json'),
                     ('length_decision_sha256','reports/day3/length_decision.json'),
                     ('quarantine_sha256','data/dapo/manifests/question_quarantine.jsonl'),
                     ('formal_split_sha256','data/dapo/formal/manifest.json')]:assert protocol[key]==sha(ROOT/path)
    g=read_jsonl(RUN/'G0_dapo100.jsonl');v=read_jsonl(RUN/'V0_dapo100.jsonl');labels=read_jsonl(OUT/'labels_C.jsonl')
    assert len(g)==len(v)==len(labels)==400 and len({r['candidate_id'] for r in g})==400
    bad=quarantined_ids();quarantine=read_jsonl(OUT/'quarantined_C_candidates.jsonl')
    assert len(quarantine)==16 and Counter(r['question_id'] for r in quarantine)==Counter({q:4 for q in bad})
    assert all(r['correct'] is None for r in labels if r['question_id'] in bad)
    byid={r['candidate_id']:r for r in g};bylabel={r['candidate_id']:r for r in labels}
    assert all(not r['truncated'] or bylabel[r['candidate_id']]['correct'] is None for r in g)
    for role,rows in [('G',g),('V',v)]:
        fp,adapter=model_identity(role);assert adapter is None and fp==protocol['models'][role]['fingerprint']
        assert all(r['model_fingerprint']==fp and r['adapter_fingerprint'] is None for r in rows)
        assert all('gold_answer' not in r and 'correct' not in r for r in rows)
    for r in v:
        source=byid[r['candidate_id']]
        assert r['prompt_sha256']==digest(messages('V',source['question'],source['text']))
        assert r['scoring_status']=='ok' and r['prompt_tokens']+4<=8192
    replay=read_jsonl(OUT/'length_replay/G0_native_context.jsonl')
    assert len(replay)==32 and all(r['same_original_decoded_prefix'] for r in replay)
    assert all(r['prompt_tokens']+r['output_tokens']<=4096 for r in replay)
    pools=json.loads((BASE/'formal/manifest.json').read_text())
    members=[read_jsonl(BASE/'inputs/dapo100.jsonl')]
    for name,record in pools['splits'].items():
        path=BASE/f'formal/inputs/{name}.jsonl';rows=read_jsonl(path);members.append(rows)
        assert sha(path)==record['input_sha256'] and len(rows)==record['questions']
        assert sha(BASE/f'formal/labels/{name}.jsonl')==record['labels_sha256']
        assert all('gold_answer' not in r for r in rows)
    for field in ['question_id','question_hash','family_id']:
        groups=[{r[field] for r in rs} for rs in members];assert len(set.union(*groups))==sum(map(len,groups))
    uuids=[{u for r in rs for u in r['family_source_uuids']} for rs in members];assert len(set.union(*uuids))==sum(map(len,uuids))
    dev={r['question_id'] for r in members[0]};smokes={}
    for role in ['G','V']:
        report=json.loads((OUT/f'smoke_{role}_report.json').read_text());smokes[role]=report
        assert all(report[k] for k in ['finite_loss','nonzero_gradient','only_lora_trainable','base_gradients_absent','base_parameter_versions_unchanged','save_reload_passed','prompt_and_padding_masks_checked'])
        assert report['config_sha256']==sha(ROOT/'configs/day3_smoke.json')
        assert sha(OUT/f'smoke_{role}_adapter/adapter_model.safetensors')==report['adapter_sha256']
        selected=json.loads((OUT/f'smoke_{role}_selection.json').read_text())['selected']
        assert len(selected)==16 and all(r['question_id'] in dev-bad for r in selected)
        manifest=json.loads((OUT/f'smoke_{role}.manifest.json').read_text())
        candidates=[ROOT/'src/pilot/day3_smoke.py',DOC/'code_snapshots/day3_smoke_before_V_probe_fix.py']
        assert any(p.exists() and sha(p)==manifest['runner_sha256'] for p in candidates)
        if role=='V':
            assert report['label_counts']=={'A':8,'B':8}
            assert all(row['consumed_target_counts']=={'A':4,'B':4} for row in report['logs'])
        assert report['length_backward_only_probe']['optimizer_steps']==0
    # Every inference recorded so far belongs to C; no A/B/T outputs have been generated.
    assert {r['question_id'] for r in g+v+replay}<=dev
    implementations={str(p.relative_to(ROOT)):sha(p) for p in [ROOT/'src/pilot/core.py',ROOT/'src/pilot/runtime.py',ROOT/'src/pilot/day2_protocol.py',ROOT/'src/pilot/day3_protocol.py',ROOT/'src/pilot/day3_smoke.py',ROOT/'src/pilot/day3_length_replay.py',ROOT/'src/pilot/day3_finish.py']}
    result={'passed':True,'day2_files_preserved':len(protected),'original400_and_scores_unchanged':True,
        'C_original_questions':100,'C_eligible_questions':96,'C_quarantined_questions':4,
        'eligible_candidates':384,'quarantined_candidates':16,'supplementary_candidates_separate':32,
        'question_hash_family_UUID_split_disjointness':True,'semantic_dedup_guarantee':False,
        'SFT_selection_has_no_gold_input':True,'A_B_T_inference_started':False,'T_inferred':False,
        'G1_G2_formal_training_started':False,'G_and_V_natural_label_smokes_passed':True,
        'formal_protocol_sha256':sha(ROOT/'configs/protocol_revisions/003_dapo_day3.json'),
        'frozen_implementation_sha256':implementations}
    write_json(DOC/'artifact_audit.json',result)
    return protocol,smokes,result

def export_review():
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    g=read_jsonl(RUN/'G0_dapo100.jsonl');v={r['candidate_id']:r for r in read_jsonl(RUN/'V0_dapo100.jsonl')}
    labels={r['candidate_id']:r for r in read_jsonl(OUT/'labels_C.jsonl')}
    audits={r['candidate_id']:r.get('assistant_review') for r in read_jsonl(DOC/'answer_audit_sample.jsonl')}
    records=[]
    for r in g:
        lab=labels[r['candidate_id']];sv=v[r['candidate_id']]
        records.append({'question_id':r['question_id'],'candidate_id':r['candidate_id'],'split':'C','G_round':'G0','V_version':'V0',
            'protocol_version':'A_B_only_boxed_v3','question':r['question'],'source_reference_answer':lab['gold_answer'],
            'question_status':lab['question_status'],'question_issue':lab['question_issue'],'G_full_text':r['text'],
            'independent_result_label':'correct' if lab['correct'] is True else 'wrong' if lab['correct'] is False else 'question_quarantined' if lab['question_status']=='quarantined' else 'unknown',
            'label_source':lab['label_source'],'parse_status':lab['parse_status'],'comparison':lab['comparison'],
            'grader_version':lab['grader_version'],'V_main_verdict':'A' if sv['accepted'] else 'B','V_score':sv['score'],
            'V_rationale_text':None,'V_rationale_status':'not_provided_AB_only','rationale_initial_review':None,
            'rationale_evidence':None,'rationale_review_status':'pending_not_started','human_review_label':None,'human_review_comment':None,
            'assistant_final_answer_spotcheck':audits.get(r['candidate_id']),'G_step_review_status':'pending_not_started',
            'truncated':r['truncated'],'output_tokens':r['output_tokens'],'label_revision_history':lab['revision_history']})
    jsonpath=DOC/'review_records.jsonl';write_jsonl(jsonpath,records)
    wb=Workbook();ws=wb.active;ws.title='回答与V理由审查';headers=list(records[0]);ws.append(headers)
    overflow=[];control_escapes=[]
    for rownum,r in enumerate(records,2):
        values=[]
        for key in headers:
            value=r[key]
            if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
            if isinstance(value,str):
                escaped=excel_literal(value)
                if escaped!=value:control_escapes.append({'candidate_id':r['candidate_id'],'field':key,'operation':'XML-illegal controls displayed as literal Unicode escapes; exact original in JSONL'})
                value=escaped
            if isinstance(value,str) and len(value)>32767:
                overflow.append({'candidate_id':r['candidate_id'],'field':key,'characters':len(value)})
                value=f"[超出Excel单元格长度；完整原文见reports/day3/review_records.jsonl，candidate_id={r['candidate_id']}，字段={key}]"
            values.append(value)
        ws.append(values)
    steps=wb.create_sheet('G步骤审查')
    steps.append(['question_id','candidate_id','step_id','raw_start','raw_end','step_text','previous_context','step_label','first_error_position','error_type','evidence','reviewer','review_status'])
    for r in records:steps.append([r['question_id'],r['candidate_id'],None,None,None,None,None,None,None,None,None,None,'未切分/待审'])
    norms=wb.create_sheet('标签规范与来源');norms.append(['field','definition_or_source'])
    definitions={'G_result':'correct/wrong/unknown independently from source reference; invalid question quarantines all4 candidates.',
        'V_rationale':'有效/无效/无法核验/未提供; AB-only has no rationale; all initial/human review labels remain blank.',
        'G_steps':'未切分/待审. No process correctness or first-error claim.',
        'automatic_labels':'Exact rational comparison with frozen boxed extraction; unknown is never a negative.',
        'assistant_audit':'6 targeted final-answer spotchecks are model-assisted, not human gold; distinct from V-rationale review.',
        'human_review':'No human verification performed; blank labels do not mean correct.',
        'question_quarantine':'4 invalid questions16 answers kept in originals and quarantine file; excluded from primary metrics.',
        'source_revision':'open-r1/DAPO-Math-17k-Processed@31dd309567e3da778038cc87d868b6097a3ccf68',
        'protocol':'configs/protocol_revisions/003_dapo_day3.json','full_text':'reports/day3/review_records.jsonl UTF-8; never silently truncated forExcel.',
        'sampling':'All original400 development candidates; automatic labels and separate6 targeted final-answer spotchecks; no V-rationale review yet.'}
    for k,val in definitions.items():norms.append([k,val])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='225566')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=min(55,max(18,len(str(column[0].value))+2))
    path=DOC/'day3_review.xlsx';wb.save(path)
    check=load_workbook(path,read_only=True);assert len(check.sheetnames)==3 and check.worksheets[0].max_row==401;check.close()
    write_json(DOC/'review_export_audit.json',{'records':400,'worksheets':3,'full_text_jsonl_sha256':sha(jsonpath),
        'xlsx_sha256':sha(path),'overlength_cells_redirected':overflow,'control_character_cells_escaped':control_escapes,'rationale_reviews_performed':0,'human_reviews_performed':0,
        'original_jsonl_text_complete':True,'string_cells_written_as_literal_text':True})

def budget(smokes):
    g=read_jsonl(RUN/'G0_dapo100.jsonl');v=read_jsonl(RUN/'V0_dapo100.jsonl')
    gs=sum({r['question_id']:r['generation_seconds_for_question'] for r in g}.values());vs=sum(r['seconds'] for r in v)
    length=json.loads((DOC/'length_decision.json').read_text());supp=length['supplementary_G_seconds']
    raw_projection=gs/400*12000/3600
    projection={'assumed_G_candidates':12000,'assumed_V_scores':19200,'G_at_observed3072_distribution_hours':raw_projection,
        'G_at_1_25_length_growth_scenario_hours':raw_projection*1.25,
        'V_AB_at_observed_distribution_hours':vs/400*19200/3600,
        'V_AB_at_1_25_length_growth_scenario_hours':vs/400*19200/3600*1.25,
        'G_two_rounds_512_examples_proxy_hours':smokes['G']['training_seconds']/16*512*2/3600,
        'V_two_branches_128_examples_proxy_hours':smokes['V']['training_seconds']/16*128*2/3600,
        'rationale_generation_hours':None,'rationale_note':'Not measured and not estimated from AB-only scoring; deferred until needed.',
        'load_and_save_seconds_measured_smokes':sum(r['total_seconds_including_load_save_reload']-r['training_seconds'] for r in smokes.values()),
        'warning':'Linear planning scenarios, not guarantees. Generator updates may lengthen responses; smoke includes a separate stress backward and has only2 optimizer steps. CPU/assistant/human review times not GPU time.'}
    result={'global_gpu_work_cap_hours':40,'recorded_DAPO_primary_G_seconds':gs,'recorded_DAPO_primary_V_seconds':vs,
        'recorded_supplementary_length_G_seconds':supp,'recorded_smoke_training_seconds':{k:r['training_seconds'] for k,r in smokes.items()},
        'recorded_smoke_peak_allocated_GiB':{k:r['peak_allocated_bytes']/2**30 for k,r in smokes.items()},
        'recorded_DAPO_and_smoke_work_hours':(gs+vs+supp+sum(r['total_seconds_including_load_save_reload'] for r in smokes.values()))/3600,
        'cost_per_verified_error_primary_seconds':(gs+vs)/212,
        'cost_per_verified_error_including_length_supplement_seconds':(gs+vs+supp)/212,
        'projection':projection,'formal_training_started':False,
        'budget_enforcement':'Before next GPU stage, include measured previous work and re-estimate remaining work against40h. Reduce noncore rationale/V extensions first; never add anotherGPU or silently reduce round-specific generation settings.',
        'accounting_limit':'Saved successful question batches and recorded smoke times only. Interrupted partial computations and idle/approval/review time are not fully measured; do not interpret as exact rental billing.'}
    write_json(DOC/'budget_reestimate.json',result);return result

def report(protocol,smokes,aud,bud):
    from .day2_evaluate import pct
    m=json.loads((DOC/'C_valid_metrics.json').read_text());length=json.loads((DOC/'length_decision.json').read_text())
    lab={r['candidate_id']:r for r in read_jsonl(OUT/'labels_C.jsonl')}
    gs=json.loads((OUT/'smoke_G_selection.json').read_text())['selected'];gquality=quality([r['candidate_id'] for r in gs],lab)
    write_json(DOC/'smoke_G_actual_selected_quality.json',{'purpose':'Quality of actual16 engineering-smoke SFT responses, measured after selection; not A/G1 training data.','quality':gquality,'selection_used_gold':False,'formal_A_training_data_not_yet_generated':True})
    rows=['# Day 3：DAPO 正式协议与训练前检查','',
        '执行依据：一周前置实验方案_20261020.md及用户最新五项协议。DAPO已确定为主任务。本次完成题目级隔离、独立答案规则、单独长度复查、V工作点、正式题池和自然正负短检查；尚未启动正式G0→G1→G2或V更新分支。', '',
        '## 题目级隔离后的主指标','',
        '原100题×4候选全部保留为开发C。4道缺目标、缺图或改写冲突题的全部16个候选整体隔离，主分析仅使用其余96题384个候选。没有替换坏题或只删除其中某条回答。', '',
        '| 指标 | 有效题目上的结果 |','|---|---:|',
        f"| 可核验覆盖 | {m['verified']}/384（{pct(m['verifiable_coverage'])}） |",
        f"| 正确／自然错误／回答级未知 | {m['correct']}／{m['natural_errors']}／{m['unknown']} |",
        f"| G正确率（可核验子集） | {pct(m['G_accuracy_verified'])} |",
        f"| 出现错误的独立题目 | {m['error_questions']}/96 |",
        f"| V误接受 FP/(FP+TN) | {m['FP']}/{m['FP']+m['TN']}（{pct(m['false_accept_rate'])}） |",
        f"| V误拒绝 FN/(TP+FN) | {m['FN']}/{m['TP']+m['FN']}（{pct(m['false_reject_rate'])}） |",
        f"| 接受集错误率（有标签部分） | {m['FP']}/{m['TP']+m['FP']}（{pct(m['accepted_error_fraction'])}） |",
        f"| 全候选保留率 | {m['accepted_total']}/384（{pct(m['acceptance_rate_all_scored'])}） |", '',
        f"全部有效输出的正确率识别区间为[{pct(m['G_correctness_lower_bound_all_outputs'])}, {pct(m['G_correctness_upper_bound_all_outputs'])}]，不是抽样置信区间。按问题聚类1000次bootstrap另存于 C_valid_metrics.json；一条G轨迹不能代表训练随机性。",'',
        f"原34条未定现在明确拆为16条题目级隔离、18条有效题目回答级未知；其中坏题内2条也曾截断，不能把原因简单相加。有效回答级未知：{m['answer_unknown_reasons']}。原始9条截断保留，其中有效题7条、坏题2条。",'',
        '## 长度检查与统一预算','',
        f"对所有出现原3072截断的8道题，按同种子4候选重放到G原生4096总上下文，共32条补充输出。32/32原前缀一致，原已结束对照均未变化；9条原截断中4条出现EOS、5条仍触及上下文上限。4条EOS输出依然缺显式最终答案，未因此新增任何正确或错误标签。补充数据独立保存在 runs/day3/length_replay，没有混入原400条。",'',
        '正式G0/G1/G2统一：max_new_tokens=min(3840,4096-实际渲染题目提示token数)，题目提示≤1024；温度0.7、top_p0.9、每题4候选不变。不扩展RoPE。V完整输入窗口8192，不截去候选尾部。G训练总长4096、V短检查总长上限8192，超长样本记录后跳过，绝不静默切掉最终答案、EOS或A/B标签。', '',
        '原100题诊断使用3072上限。正式生成预算已改，下一阶段必须用正式预算重新缓存同一轮次比较所需的G0/C基线；不能将旧3072诊断与新G1/G2曲线直接拼接。长度补充只说明增加预算的行为及原生边界，不证明延长能解决循环或所有难题。', '',
        '## 独立核验与抽查','',
        'configs/dapo_answer_v3.json固定精确整数/有理数比较、显式答案定位、多盒冲突、等价表示、单位与尾缀、截断和未知规则，实际使用保留的day2解析实现及其SHA256；不执行生成代码，参考答案不进入G/V提示。',
        '判对、判错、未知各抽查2例，证据在 answer_audit_sample.jsonl。半回文数差8989与200π以下整数计数628支持原正确标签；53被误称非质数、三题成绩组合被误算512，两例独立推导支持错误标签；进制例1102_6数学上正确但超出当前纯数值解析范围，仍保留自动未知并单存模型辅助复核建议。未定没有为了增加负例而改错。抽查选择较短案例，不代表全体核验准确率；未开展人工金标准复核。', '',
        '## V工作点及实际SFT质量','',
        '开发集比较0.5/0.6/0.7/0.8/0.9/0.95/0.99，保留0.5。提高到0.99时入选独立题目86→77，而每题随机选一条后的已核验错误比例53.57%→54.05%；本次没有证据表明值得以保留率换取这个调整。开发诊断只用于本次冻结，T未参与。',
        f"0.5下257条被接受回答中，247条有独立标签、10条未知，标签覆盖{pct(m['accepted_label_coverage'])}；全接受集错误率只可识别为[{pct(m['accepted_error_lower_bound_all'])}, {pct(m['accepted_error_upper_bound_all'])}]，不能把46.96%解释为所有接受回答的精确风险。", '',
        '冻结的反馈选样函数不接收独立标签：按题目随机选一条被V接受、未截断且满足完整训练长度的回答，每题最多一条；超过512题时固定种子抽取，绝不用标准答案暗中过滤。C开发预演86条=39正确、45错误、2未知，已核验错误45/84=53.57%，完整选样风险界52.33%—54.65%。这与整个接受池46.96%不同。',
        f"实际G工程短检查使用的16条SFT回答质量：{gquality}。这是独立测试adapter的数据，绝非正式A池或G1训练集。正式A尚未生成；每轮A实际选样完成后必须另外统计该轮正确／错误／未知与风险界。后续称为生成器更新，不预设能力提高。",'',
        '## 正式划分与配置','',
        'A_G_update600、B_V_update200、C原100（有效96、隔离4）、T_final300。按固定题族库存顺序分配，额外排除文本显式依赖缺失图的题；问题ID、规范化哈希、来源UUID题族均跨池互斥。全部开发100题及其可识别近似题族排除出A/B/T；旧GSM8K T封存。DAPO T是公开训练部分的自定义留出，不是官方测试成绩。词面近似去重不能保证消除语义改写或预训练暴露。',
        '正式协议：configs/protocol_revisions/003_dapo_day3.json；当前指针：configs/current_protocol.json。第二天 configs/formal_protocol.json 保持原字节，未覆盖。冻结模型路径与指纹、提示、采样、长度、阈值、A/B/T清单、G每轮至多512条/1epoch的response-only LoRA配置，G2继续G1的adapter。短检查adapter永远不作为正式G1/G2或V分支的初始化。', '',
        '## 短训练检查与实测资源','']
    for role,r in smokes.items():
        rows += [f"{role}：16条独立候选、2个优化器step、micro-batch1、累积8，不复制循环样本；V自然标签8A+8B，每step均4A+4B。训练loss/梯度有限且非零、基座无梯度且版本不变、仅LoRA参数变化、prompt与padding遮罩、保存重载均通过。",
            f"{role}耗时含一次额外无优化器step的长度反传探针={r['training_seconds']:.2f}秒，峰值allocated={r['peak_allocated_bytes']/2**30:.2f}GiB；实际16条训练样本最大长度={r['actual_max_length']}；独立长度探针={r['length_backward_only_probe']}。",'']
    rows += ['V补充长回答没有可用独立答案标签，首次检查在优化器step之前停止，修正记录与原代码/manifest均保留。重试用已有自然标签真实样本，向右填充attention0、labels-100测试容量，未给未知回答贴A/B标签。V的真实输入长度与填充后容量分别记录，不能将填充测试宣传为已训练真实长样本。', '',
        f"主诊断G+V约{(bud['recorded_DAPO_primary_G_seconds']+bud['recorded_DAPO_primary_V_seconds'])/60:.2f}分钟；独立长度重放约{bud['recorded_supplementary_length_G_seconds']/60:.2f}分钟。每条自然错误主诊断成本约{bud['cost_per_verified_error_primary_seconds']:.2f}秒，计入长度补充约{bud['cost_per_verified_error_including_length_supplement_seconds']:.2f}秒。",'',
        f"按原12000个G候选、19200次A/B评分预算，当前3072分布线性生成约{bud['projection']['G_at_observed3072_distribution_hours']:.2f}小时；长度增长25%情景约{bud['projection']['G_at_1_25_length_growth_scenario_hours']:.2f}小时，A/B评分另约{bud['projection']['V_AB_at_1_25_length_growth_scenario_hours']:.2f}小时。另加训练/加载及已用时间；这不是硬上界，生成器更新可能进一步增长长度。理由生成未测速，不能用A/B评分速度替代；助手/人工审查时间另计。全局40GPU小时上限不变，开下一阶段前复估，优先缩减非核心扩展。详见 budget_reestimate.json。",'',
        '## 审查材料与下一阶段','',
        '已导出 day3_review.xlsx 三个工作表及完整UTF-8 review_records.jsonl。保留全部原400条，包括题目隔离标记和六例最终答案抽查；V理由未提供、初审/人工意见为空并明确待审，G步骤未切分/待审。没有冒充本地助手完成V理由审查或人工确认，超长文本只在Excel重定向原JSONL，不静默截断。', '',
        'Day3核心检查完成，可以按冻结配置进入下一阶段：生成正式A候选、V0筛选、记录实际SFT质量，再沿同一G谱系作两次更新并重新缓存同设置B/C输出。当前没有开始这两次正式更新，也没有对T推理；不需要先完成全部理由或过程研究。', '',
        f"最终产物一致性审计={aud['passed']}，历史57个第二天文件保持不变；回归结果见 runs/day3/pytest.log。"]
    (DOC/'day3_report.md').write_text('\n'.join(rows)+'\n')

def main():
    protocol,smokes,aud=audit();export_review();bud=budget(smokes);report(protocol,smokes,aud,bud)
    write_json(OUT/'state.json',{'status':'day3_complete','primary_task':'DAPO English','formal_protocol':'configs/protocol_revisions/003_dapo_day3.json',
        'question_quarantine_complete':True,'length_check_complete':True,'G_V_smokes_complete':True,'formal_G1_G2_started':False,'T_inferred':False})
    write_json(ROOT/'configs/stage_status.json',{'phase':'day3_complete_ready_for_generator_updates','formal_training_started':False,'T_inferred':False})
    print(json.dumps({'day3':'complete','audit':aud['passed'],'G_smoke':True,'V_smoke_label_counts':smokes['V']['label_counts'],'formal_training_started':False},indent=2))

if __name__=='__main__':main()
