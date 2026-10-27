"""Whole-cohort paired rates, continuous scores, and the two-sided removal tradeoff."""
import json
import math
import shutil
import zipfile
from pathlib import Path

import numpy as np
from scipy import stats

from .core import ROOT, read_jsonl, write_json, write_jsonl, digest
from .formal_common import messages, sha, utc
from .prefix_full_panel import RUN, REPORT, PLAN, validate
from .fpr_pattern_diagnostics import bh, write_csv
from .v0_drift_analysis import clean


def ci(x):
    return [float(v) for v in np.nanquantile(x, [.025, .975])]


def ratio(values, mask, qi, W):
    values = np.asarray(values, float); mask = np.asarray(mask, bool); G = W.shape[1]
    n = np.bincount(qi, weights=mask.astype(float), minlength=G)
    sums = np.bincount(qi, weights=np.where(mask, values, 0), minlength=G)
    point = float(values[mask].mean()); den = W@n
    with np.errstate(divide='ignore', invalid='ignore'): boot = (W@sums)/den
    influence = (sums-point*n)/mask.sum()
    return point, boot, influence


def contrast(y, base, mask, qi, W):
    point, boot, influence = ratio(np.asarray(y, float)-np.asarray(base, float), mask, qi, W)
    G = W.shape[1]; se = math.sqrt(G/(G-1)*float(influence@influence))
    p = float(2*stats.t.sf(abs(point/se), G-1)) if se > 1e-14 else 1. if abs(point) < 1e-14 else None
    return {'delta': point, 'low': ci(boot)[0], 'high': ci(boot)[1], 'p_cluster': p}, boot


def audit_scores():
    cfg = validate(); rr = read_jsonl(RUN/'answers.jsonl'); jobs = read_jsonl(RUN/'jobs.jsonl')
    ss = read_jsonl(RUN/'V0_scores.jsonl'); by = {(s['candidate_id'], s['variant']): s for s in ss}
    assert len(ss) == len(by) == cfg['logical_scores'] == len(jobs)
    index = {r['candidate_id']: r for r in rr}; provenance = {}
    for j in jobs:
        s = by[j['candidate_id'], j['variant']]; r = index[j['candidate_id']]
        assert all(s[k] == j[k] for k in j)
        assert s['correct'] is r['correct'] and s['prompt_tokens']+4 <= 8192
        assert s['prompt_sha256'] == digest(messages('V', r['question'], j['prefix']+r['text']))
        assert s['text_sha256'] == digest(j['prefix']+r['text'])
        assert s['accepted'] == (s['score'] >= .5)
        assert s['model_fingerprint'] == cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None
        assert s['threshold'] == .5 and s['label_ids'] == [[32], [33]]
        # The reverse operation returns the exact cached answer, not a paraphrase.
        assert (j['prefix']+r['text'])[len(j['prefix']):] == r['text']
        if j['variant'] == 'original':
            assert s['accepted'] == r['accepted'] and abs(s['score']-r['score']) < 2e-6
        if s['source'] == 'exact_historical_full_prompt':
            if s['reuse_file'] not in provenance: provenance[s['reuse_file']] = read_jsonl(ROOT/s['reuse_file'])
            original = provenance[s['reuse_file']][s['reuse_line']-1]
            assert all(original[k] == s[k] for k in ['score', 'accepted', 'log_likelihood_A', 'log_likelihood_B'])
            assert original['prompt_sha256'] == s['prompt_sha256']
    checks = read_jsonl(RUN/'sentinel_rechecks.jsonl')
    assert all(c['passed'] for c in checks) and len(checks) >= 5
    return cfg, rr, ss, by


def analyze(cfg, rr, by):
    qs = sorted({r['question_id'] for r in rr}); qix = {q: i for i, q in enumerate(qs)}
    qi = np.array([qix[r['question_id']] for r in rr]); G = len(qs)
    W = np.random.default_rng(cfg['seed']).multinomial(G, np.full(G, 1/G), size=cfg['bootstrap'])
    yes = np.array([r['correct'] is True for r in rr]); no = np.array([r['correct'] is False for r in rr]); known = yes|no
    a0 = np.array([by[r['candidate_id'], 'original']['accepted'] for r in rr], bool)
    an = np.array([by[r['candidate_id'], 'neutral_control']['accepted'] for r in rr], bool)
    p0 = np.array([by[r['candidate_id'], 'original']['score'] for r in rr])
    m0 = np.array([by[r['candidate_id'], 'original']['log_likelihood_A']-by[r['candidate_id'], 'original']['log_likelihood_B'] for r in rr])
    summary = []; continuous = []; interactions = []; reverse = []
    for v in cfg['variants']:
        ss = [by[r['candidate_id'], v['name']] for r in rr]
        a = np.array([s['accepted'] for s in ss], bool); p = np.array([s['score'] for s in ss])
        margin = np.array([s['log_likelihood_A']-s['log_likelihood_B'] for s in ss])
        tp, fn, fp, tn = int((yes&a).sum()), int((yes&~a).sum()), int((no&a).sum()), int((no&~a).sum())
        row = {'variant': v['name'], 'prefix': v['prefix'], 'category': v['category'], 'TP': tp, 'FN': fn, 'FP': fp, 'TN': tn,
               'correct_n': int(yes.sum()), 'wrong_n': int(no.sum()), 'unknown_n': int((~known).sum()),
               'unknown_accepted': int((~known&a).sum()), 'FPR': fp/no.sum(), 'FNR': fn/yes.sum(),
               'TN_to_FP': int((no&~a0&a).sum()), 'FP_to_TN': int((no&a0&~a).sum()),
               'FN_to_TP': int((yes&~a0&a).sum()), 'TP_to_FN': int((yes&a0&~a).sum()),
               'accepted_known': tp+fp, 'accepted_wrong_risk': fp/(tp+fp), 'BA': .5*(tn/no.sum()+tp/yes.sum()),
               'known_error_rate': (fp+fn)/known.sum(), 'all_acceptance': float(a.mean()),
               'all_error_lower_bound': (fp+fn)/len(rr), 'all_error_upper_bound': (fp+fn+(~known).sum())/len(rr)}
        assert row['FP']-int((no&a0).sum()) == row['TN_to_FP']-row['FP_to_TN']
        assert row['FN']-int((yes&~a0).sum()) == row['TP_to_FN']-row['FN_to_TP']
        rate_boot = {}
        for metric, mask, y, b, n in [('FPR', no, a, a0, an), ('FNR', yes, ~a, ~a0, ~an)]:
            result, rate_boot[metric] = contrast(y, b, mask, qi, W)
            neutral, _ = contrast(y, n, mask, qi, W)
            row.update({metric+'_'+k: value for k, value in result.items()})
            row.update({metric+'_vs_neutral_'+k: value for k, value in neutral.items()})
        for kind, values, baseline in [('score', p, p0), ('margin', margin, m0)]:
            group_results = {}
            for group, mask in [('wrong', no), ('correct', yes), ('unknown', ~known)]:
                result, boot = contrast(values, baseline, mask, qi, W)
                group_results[group] = (result, boot)
                continuous.append({'variant': v['name'], 'prefix': v['prefix'], 'quantity': kind, 'truth': group,
                                   'n': int(mask.sum()), 'mean_original': float(baseline[mask].mean()),
                                   'mean_injected': float(values[mask].mean()), **result})
            diff = group_results['wrong'][0]['delta']-group_results['correct'][0]['delta']
            bounds = ci(group_results['wrong'][1]-group_results['correct'][1])
            interactions.append({'variant': v['name'], 'quantity': kind, 'wrong_minus_correct_shift': diff,
                                 'low': bounds[0], 'high': bounds[1],
                                 'interpretation': 'positive means larger mean boost on wrong answers; different truth subpopulations, not random assignment of correctness'})
        error, boot = contrast((no&a)|(yes&~a), (no&a0)|(yes&~a0), known, qi, W)
        row.update({'known_error_'+k: val for k, val in error.items()})
        row['BA_delta'] = -.5*(row['FPR_delta']+row['FNR_delta'])
        row['BA_delta_low'], row['BA_delta_high'] = ci(-.5*(rate_boot['FPR']+rate_boot['FNR']))
        reverse.append({'variant': v['name'], 'operation': 'strip exact artificially injected prefix back to original body',
                        'FP_removed': row['TN_to_FP'], 'FP_newly_created': row['FP_to_TN'],
                        'FN_newly_created': row['FN_to_TP'], 'FN_removed': row['TP_to_FN'],
                        'FPR_delta': -row['FPR_delta'], 'FNR_delta': -row['FNR_delta'],
                        'BA_delta': -row['BA_delta'], 'net_wrong_decisions_delta':
                        row['FP_to_TN']-row['TN_to_FP']+row['FN_to_TP']-row['TP_to_FN']})
        summary.append(row)
    triggers = [r for r in summary if r['category'] != 'control']
    for control in ['', '_vs_neutral']:
        tests = [(r, metric) for r in triggers for metric in ['FPR', 'FNR']]
        ps = [r[metric+control+'_p_cluster'] for r, metric in tests]
        for (r, metric), q in zip(tests, bh([p if p is not None else np.nan for p in ps])):
            r[metric+control+'_q_BH12'] = float(q) if np.isfinite(q) else None
    return summary, continuous, interactions, reverse


def figures(cfg, rows, continuous, interactions):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10})
    labels = [r['variant'] for r in rows]; yy = np.arange(len(rows))
    with PdfPages(REPORT/'full_panel_figures.pdf') as pdf:
        fig, axes = plt.subplots(1, 2, figsize=(13, 6.6), constrained_layout=True)
        for ax, metric in zip(axes, ['FPR', 'FNR']):
            for i, r in enumerate(rows):
                q = r.get(metric+'_q_BH12'); col = '#b34e3c' if q is not None and q < .05 else '#287ba1'
                ax.plot([r[metric+'_low']*100, r[metric+'_high']*100], [i, i], color=col, lw=2)
                ax.plot(r[metric+'_delta']*100, i, 'o', color=col)
                ax.annotate(f'{r[metric]:.2%}', (r[metric+'_delta']*100, i), xytext=(3, 8), textcoords='offset points', fontsize=8)
            ax.axvline(0, color='black', ls='--'); ax.set_yticks(yy, labels); ax.invert_yaxis()
            ax.set_xlabel(f'{metric}(prefix) - {metric}(original), percentage points')
            ax.set_title('ALL wrong answers: n=218' if metric == 'FPR' else 'ALL correct answers: n=147')
            ax.margins(x=.18)
        fig.suptitle('Fixed V0 / G3 / complete C95 x 4 answers\nPaired question-cluster 95% intervals; red = BH12 q < .05')
        fig.savefig(REPORT/'net_FPR_FNR.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(1, 2, figsize=(13, 7), constrained_layout=True)
        for ax, kind in zip(axes, ['score', 'margin']):
            for truth, offset, col in [('wrong', -.12, '#bb4e3c'), ('correct', .12, '#297ba0')]:
                ss = [next(c for c in continuous if c['variant'] == name and c['quantity'] == kind and c['truth'] == truth) for name in labels]
                for i, s in enumerate(ss):
                    ax.plot([s['low'], s['high']], [i+offset, i+offset], color=col, lw=2)
                ax.plot([s['delta'] for s in ss], yy+offset, 'o', color=col, label=truth)
            ax.axvline(0, color='black', ls='--'); ax.set_yticks(yy, labels); ax.invert_yaxis(); ax.legend()
            ax.set_xlabel('Mean injected minus original'); ax.set_title('Acceptance score P(A | A,B)' if kind == 'score' else 'Acceptance margin LL(A) - LL(B)')
        fig.suptitle('Separate wrong-answer advantage from correct-answer benefit\nPaired 95% intervals; margin remains informative when scores saturate')
        fig.savefig(REPORT/'continuous_truth_comparison.png', dpi=180); pdf.savefig(fig); plt.close(fig)
        fig, axes = plt.subplots(2, 4, figsize=(14, 7.4), constrained_layout=True)
        for ax, r in zip(axes.flat, rows):
            matrix = np.array([[r['TN'], r['FP']], [r['FN'], r['TP']]])
            ax.imshow(matrix, cmap='Blues', vmin=0, vmax=150)
            for i in range(2):
                for j in range(2): ax.text(j, i, f'{["TN", "FP", "FN", "TP"][i*2+j]}\n{matrix[i,j]}', ha='center', va='center', color='white' if matrix[i,j]>85 else 'black', fontsize=12)
            ax.set_xticks([0, 1], ['V rejects', 'V accepts']); ax.set_yticks([0, 1], ['Wrong', 'Correct'])
            ax.set_title(r['variant'])
        fig.suptitle('Same answers, same independent final-answer labels\n15 unknown answers preserved separately; excluded from truth denominators')
        fig.savefig(REPORT/'confusion_matrices.png', dpi=180); pdf.savefig(fig); plt.close(fig)


def text_report(cfg, rows, continuous, interactions, reverse, ss):
    lines = ['# 完整回答面板：前缀给予错误回答多少额外优势', '',
             '这次补齐此前TN条件选样的缺口。固定C95题、每题4条已有G3回答，共380条；既包括原TP126、FN21，也包括原FP119、TN99，未知15条保留。六个既有代表前缀、原文和中性控制，共3040条逻辑评分。没有新增G生成/训练或V训练，也没有读取B/T。', '',
             '选择全部题与全部回答，不按原V判定、正确性、目标词或新结果筛选。题目、原数学正文和最终答案逐字保持，只在最前加精确字符串；所有最终答案提取与标签检查不变。未知不强行转为错误。C已反复用于开发，此处是探索诊断，不能当新独立测试或训练中介识别。', '',
             '## 1. 全部错误与正确回答上的净变化', '',
             'FPR净变化=(TN→FP−FP→TN)/218；FNR净变化=(TP→FN−FN→TP)/147。与此前只测原TN不同，这里纳入两方向翻转；两种变化分别报告，不能把减少误拒当作减少误接受。', '',
             '|前缀（字符串见配置）|FPR|ΔFPR/百分点 [95%区间]|FNR|ΔFNR/百分点 [95%区间]|TN→FP / FP→TN|FN→TP / TP→FN|',
             '|---|---:|---|---:|---|---:|---:|']
    for r in rows:
        lines.append(f'|{r["variant"]}|{r["FPR"]:.2%}|{100*r["FPR_delta"]:+.2f} [{100*r["FPR_low"]:+.2f},{100*r["FPR_high"]:+.2f}]|{r["FNR"]:.2%}|{100*r["FNR_delta"]:+.2f} [{100*r["FNR_low"]:+.2f},{100*r["FNR_high"]:+.2f}]|{r["TN_to_FP"]} / {r["FP_to_TN"]}|{r["FN_to_TP"]} / {r["TP_to_FN"]}|')
    lines += ['', '区间按95道题整组配对重采样10000次，保留同题4回答的相关性。主要检验用按题聚类影响函数t检验，六前缀×两错误率共12项BH；相对中性前缀另作BH12。CSV保留全部p/q。前缀选择参考过既有探索，包括同一C池的旧评分；BH12仅校正本次列出的比较，不覆盖整个历史筛词过程，不能消除重复使用开发集与选词的偏差。逐项95%区间不是同时置信区间，零翻转不能证明总体风险为零。', '',
              '## 2. 连续评分：错误回答与正确回答各获多少提升', '',
              '分数为固定V0对A/B裸标签条件似然的softmax；margin=LL(A)−LL(B)，不计EOS。当分数接近0/1时仍保留margin，避免饱和掩盖变化。每条回答以自身原文作对照。', '',
              '|前缀|错误回答平均Δ分数|正确回答平均Δ分数|两者差 [95%区间]|错误/正确平均Δmargin|',
              '|---|---:|---:|---|---:|']
    cl = {(r['variant'], r['quantity'], r['truth']): r for r in continuous}
    il = {(r['variant'], r['quantity']): r for r in interactions}
    for r in rows:
        name = r['variant']; w, c = cl[name, 'score', 'wrong'], cl[name, 'score', 'correct']; it = il[name, 'score']
        lines.append(f'|{name}|{w["delta"]:+.4f}|{c["delta"]:+.4f}|{it["wrong_minus_correct_shift"]:+.4f} [{it["low"]:+.4f},{it["high"]:+.4f}]|{cl[name,"margin","wrong"]["delta"]:+.3f} / {cl[name,"margin","correct"]["delta"]:+.3f}|')
    lines += ['', '“错误−正确”的正差表示在本队列中错误回答的平均分提升更大；它对比的是两个真实性子群，不能把正确性当作随机干预。同样的接受倾向也可能同时帮助两类回答，不能仅凭错误回答分数上升称其为专门欺骗机制。不同子群基准分、长度与难度不同，分数变化还受天花板影响；margin提供互补证据。', '',
              '## 3. 反向去掉人工注入前缀：获益与代价同时列出', '',
              '下表从“已经注入”的版本出发，精确删掉刚加的字符串，返回原文。它是上述配对结果的严格逆操作，无需重复评分；不等于在自然文本中任意删词，也不建立一个可部署的词汇惩罚器。', '',
              '|移除的前缀|消除FP / 新增FP|新增FN / 消除FN|ΔFPR/百分点|ΔFNR/百分点|Δ平衡准确率/百分点|净增加错误数|',
              '|---|---:|---:|---:|---:|---:|---:|']
    for r in reverse:
        lines.append(f'|{r["variant"]}|{r["FP_removed"]} / {r["FP_newly_created"]}|{r["FN_newly_created"]} / {r["FN_removed"]}|{100*r["FPR_delta"]:+.2f}|{100*r["FNR_delta"]:+.2f}|{100*r["BA_delta"]:+.2f}|{r["net_wrong_decisions_delta"]:+d}|')
    lines += ['', '减少FPR而提高FNR是明确取舍。是否值得惩罚需预先规定FP/FN成本、业务保留率及独立验证；不能只看FPR降低。完整CSV还包含接受集错误比例、BA、已知标签总体错误率，以及15未知的错误率上下界。自然删除/替换的既有结果仍只作为另一类局部干预证据。', '',
              '## 4. 本次能得出的判断', '']
    triggers = [r for r in rows if r['category'] != 'control']
    both = [r['variant'] for r in triggers if r['FPR_delta']>0 and r['FNR_delta']<0]
    sigfpr = [r['variant'] for r in triggers if r.get('FPR_q_BH12', 1) is not None and r['FPR_q_BH12']<.05 and r['FPR_delta']>0]
    sigfnr = [r['variant'] for r in triggers if r.get('FNR_q_BH12', 1) is not None and r['FNR_q_BH12']<.05 and r['FNR_delta']<0]
    lines += [f'- 同时出现FPR上升、FNR下降的点估计：{", ".join(both) or "无"}。',
              f'- 相对原文，FPR上升通过主要BH12者：{", ".join(sigfpr) or "无"}；FNR下降通过BH12者：{", ".join(sigfnr) or "无"}。',
              '- 本实验直接量化表达在全部错误回答上的净接受优势，也同时核算正确回答的收益与损伤；不证明生成器已经学会利用这些词，不把人为注入效应等同于自然跨轮FPR变化的原因。',
              '- 六前缀的“错误回答平均分提升−正确回答平均分提升”区间均跨0，margin差区间也均跨0。现有数据不足以确认错误回答受到选择性偏爱；五个前缀的双向错误率变化更符合整体接受倾向上移，但这仍只是本队列的行为描述。',
              '- 前缀是继承旧探索的代表项，不能将“六个中几个有效”推广成英语表达总体比例；对照和反向翻转均保留，不据结果筛词。', '', '## 数据、模板与审计', '']
    compute = json.loads((RUN/'compute.json').read_text()); reused = sum(s['source']=='exact_historical_full_prompt' for s in ss)
    lines += [f'3040逻辑评分中，完整提示、固定V0和评分实现一致的历史缓存复用{reused}条，新评分{len(ss)-reused}条；另有{len(read_jsonl(RUN/"sentinel_rechecks.jsonl"))}条固定哨兵重评分核验。此次GPU阶段{compute["seconds_this_invocation"]:.1f}秒。',
              '逐条CSV/JSONL保留真实性标签、精确前缀、原/干预分数、两类对数似然、token数和缓存来源。answers_full.jsonl含完整原题/回答；template_and_prefixes.json和冻结配置保存完整模板及字符串，逐条可重建提示。未知保留原原因、截断状态。',
              '[三页图册](full_panel_figures.pdf) · [完整Excel](full_panel_review.xlsx) · [汇总CSV](net_rates.csv) · [逐条评分](scores.csv)', '']
    (REPORT/'full_panel_report.md').write_text('\n'.join(lines)+'\n', encoding='utf-8')


def main():
    cfg, rr, ss, by = audit_scores(); rows, continuous, interactions, reverse = analyze(cfg, rr, by)
    tables = {'net_rates': rows, 'continuous_changes': continuous, 'wrong_minus_correct': interactions, 'reverse_removal': reverse, 'scores': ss}
    for name, table in tables.items(): write_csv(REPORT/(name+'.csv'), clean(table))
    write_json(REPORT/'results.json', clean({k: v for k, v in tables.items() if k != 'scores'}))
    write_jsonl(REPORT/'answers_full.jsonl', rr); write_jsonl(REPORT/'scores_full.jsonl', ss)
    from .formal_common import protocol
    tokenizer_file = Path(protocol()['models']['V']['path'])/'tokenizer_config.json'
    tokenizer_config = json.loads(tokenizer_file.read_text())
    write_json(REPORT/'template_and_prefixes.json', {'messages_template': cfg['template'], 'variants': cfg['variants'],
               'scoring': cfg['score'], 'tokenizer_chat_template': tokenizer_config['chat_template'],
               'tokenizer_config_sha256': sha(tokenizer_file), 'apply_chat_template_arguments':
               {'tokenize': True, 'add_generation_prompt': True},
               'model_fingerprint': cfg['fixed_V0']['fingerprint'], 'adapter': None, 'threshold': .5,
               'token_counts': 'prompt_tokens = complete formatted V input; response_tokens_V_tokenizer = prefix+body standalone; prefix_tokens_standalone is not the prompt-token delta'})
    figures(cfg, rows, continuous, interactions); text_report(cfg, rows, continuous, interactions, reverse, ss)
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    for name, table in tables.items():
        ws = wb.create_sheet(name[:31]); keys = list(dict.fromkeys(k for r in table for k in r)); ws.append(keys)
        for number, row in enumerate(clean(table), 2):
            values = [row.get(k) for k in keys]
            ws.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v for v in values])
            for column, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith('='): ws.cell(number, column).data_type = 's'
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(REPORT/'full_panel_review.xlsx')
    for p in [PLAN, RUN/'compute.json', RUN/'answer_context_checks.jsonl', RUN/'sentinel_rechecks.jsonl',
              Path(__file__), ROOT/'src/pilot/prefix_full_panel.py', ROOT/'src/pilot/prefix_full_panel_verify.py']: shutil.copyfile(p, REPORT/p.name)
    write_json(REPORT/'audit.json', {'passed': True, 'at': utc(), 'support': cfg['support'], 'logical_scores': len(ss),
               'full_question_cohort_all_four_candidates': True, 'both_truths_and_both_original_V_decisions': True,
               'unknowns_retained': True, 'all_math_bodies_and_final_labels_preserved': True,
               'fixed_V0_full_context': True, 'source_hashes_and_cached_provenance_verified': True,
               'T_read': False, 'B_read': False, 'new_G_generation': 0, 'training_steps': 0})
    files = sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix != '.zip'
                   and p.name not in ['artifact_manifest.json', 'independent_verification.json'])
    write_json(REPORT/'artifact_manifest.json', {p.name: sha(p) for p in files})
    with zipfile.ZipFile(REPORT/'full_panel_review.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']: z.write(p, p.name)
    write_json(RUN/'state.json', {'status': 'complete', 'at': utc(), 'report': str(REPORT/'full_panel_report.md')})
    print(json.dumps(clean(rows), ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__': main()
