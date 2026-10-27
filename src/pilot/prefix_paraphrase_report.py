"""Paired, question-clustered analysis of prespecified prefix substitutions."""
import json
import math
import shutil
import zipfile
from collections import defaultdict
from pathlib import Path

import numpy as np
from scipy import stats

from .core import ROOT, read_jsonl, write_json, write_jsonl, digest
from .formal_common import sha, utc, messages
from .prefix_paraphrase import RUN, REPORT, PLAN, validate, edited
from .fpr_pattern_diagnostics import load as source_rows, bh, write_csv
from .v0_drift_analysis import clean, interval
from .oracle_visuals import excel_safe_value
from .efp_intervention import wilson


def paired_stats(values, original, questions, draws=10000):
    """Every row is a response pair. Bootstrap whole questions, keep siblings."""
    y=np.asarray(values,float);base=np.asarray(original,float)
    assert len(y)==len(base)==len(questions) and len(y)>0
    qs=sorted(set(questions));ix={q:i for i,q in enumerate(qs)};qi=np.array([ix[q] for q in questions]);G=len(qs)
    counts=np.bincount(qi,minlength=G).astype(float)
    observed=y.mean();before=base.mean();d=y-base;delta=float(d.mean())
    sums=np.bincount(qi,weights=y,minlength=G);diff=np.bincount(qi,weights=d,minlength=G)
    rng=np.random.default_rng(20261024);idx=rng.integers(G,size=(draws,G));den=counts[idx].sum(1)
    boot_delta=diff[idx].sum(1)/den;boot_rate=sums[idx].sum(1)/den
    infl=(diff-delta*counts)/len(y)
    se=math.sqrt(G/max(G-1,1)*float(infl@infl))
    p=float(2*stats.t.sf(abs(delta/se),G-1)) if G>1 and se>1e-14 else 1. if abs(delta)<1e-14 else None
    gained=int(((y==1)&(base==0)).sum());lost=int(((y==0)&(base==1)).sum())
    exact=float(stats.binomtest(gained,gained+lost,p=.5).pvalue) if gained+lost else 1.
    ci=interval(boot_delta);ri=interval(boot_rate)
    return {'responses':len(y),'questions':G,'original_rate':float(before),'rate':float(observed),
            'delta':delta,'low':ci['low'],'high':ci['high'],'rate_low':ri['low'],'rate_high':ri['high'],
            'gained':gained,'lost':lost,'cluster_SE':se,'p_cluster':p,'p_McNemar':exact,
            'low_question_support':G<10,'all_pairs_unchanged':bool(np.all(d==0))}


def load_scores():
    cfg=validate();scores=read_jsonl(RUN/'V0_scores.jsonl')
    assert len(scores)==cfg['planned_scores'] and len({r['cache_key'] for r in scores})==len(scores)
    data={arm:read_jsonl(RUN/f'{arm}_answers.jsonl') for arm in ['prefix','natural']}
    index={arm:{r['candidate_id']:r for r in rr} for arm,rr in data.items()}
    by=defaultdict(dict)
    for s in scores:
        arm=s['arm'];r=index[arm][s['candidate_id']];v=next(v for v in cfg['variants'][arm] if v['name']==s['variant'])
        assert s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None and s['threshold']==.5
        assert s['correct'] is r['correct'] and s['question_id']==r['question_id'] and s['split']==r['split']
        assert s['accepted']==(s['score']>=.5) and s['prompt_tokens']+4<=8192 and s['answer_and_status_preserved']
        assert s['cache_key']==digest({'plan_sha256':sha(PLAN),'arm':arm,'candidate_id':r['candidate_id'],'variant':v['name']})
        text=edited(r,v,arm)
        assert s['original_text_sha256']==digest(r['text']) and s['edited_text_sha256']==digest(text)
        assert s['prompt_sha256']==digest(messages('V',r['question'],text))
        if v['name']=='original':assert s['accepted']==r['accepted'] and abs(s['score']-r['score'])<2e-6
        by[arm,s['candidate_id']][v['name']]=s
    for arm,rr in data.items():
        assert all(len(by[arm,r['candidate_id']])==len(cfg['variants'][arm]) for r in rr)
    return cfg,data,by,scores


def analyze_prefix(cfg,data,by):
    out=[]
    for split in ['A','C']:
        rr=[r for r in data['prefix'] if r['split']==split];qs=[r['question_id'] for r in rr]
        assert len(qs)==len(set(qs)) and all(r['correct'] is False and not r['accepted'] for r in rr)
        base=[by['prefix',r['candidate_id']]['we_need_to']['accepted'] for r in rr]
        neutral=[by['prefix',r['candidate_id']]['neutral_control']['accepted'] for r in rr]
        for v in cfg['variants']['prefix']:
            ss=[by['prefix',r['candidate_id']][v['name']] for r in rr]
            row={'arm':'prefix','split':split,'variant':v['name'],'prefix':v['prefix'],'category':v['category'],
                 'metric':'wrong_reject_to_accept_rate','reference':'we_need_to',
                 **paired_stats([s['accepted'] for s in ss],base,qs)}
            row['reference_rate']=row.pop('original_rate')
            row['unchanged_body_accept_rate']=0.
            row['flipped_original_wrong_rejections']=sum(s['accepted'] for s in ss)
            reference_flips=sum(base)
            row['reference_flips']=reference_flips
            row['retained_reference_flips']=reference_flips-row['lost']
            row['retention_of_reference_flips']=(reference_flips-row['lost'])/reference_flips if reference_flips else None
            lo,hi=wilson(sum(s['accepted'] for s in ss),len(ss));row.update({'Wilson_rate_low':lo,'Wilson_rate_high':hi})
            n=paired_stats([s['accepted'] for s in ss],neutral,qs)
            row.update({'delta_vs_neutral':n['delta'],'neutral_low':n['low'],'neutral_high':n['high'],
                        'p_vs_neutral':n['p_McNemar'],
                        'mean_log_odds_change_vs_we_need_to':float(np.mean([s['log_likelihood_A']-s['log_likelihood_B']-
                            by['prefix',r['candidate_id']]['we_need_to']['log_likelihood_A']+
                            by['prefix',r['candidate_id']]['we_need_to']['log_likelihood_B'] for r,s in zip(rr,ss)]))})
            out.append(row)
    primary=[r for r in out if r['variant']!='we_need_to']
    for r,q in zip(primary,bh([r['p_McNemar'] for r in primary])):r['q_BH_vs_we_need_to']=float(q)
    secondary=[r for r in out if r['variant']!='neutral_control']
    for r,q in zip(secondary,bh([r['p_vs_neutral'] for r in secondary])):r['q_BH_vs_neutral']=float(q)
    return out


def analyze_natural(cfg,data,by):
    out=[]
    for split in ['A','C']:
        for truth,label in [(False,'FPR'),(True,'FNR'),(None,'unknown_acceptance')]:
            rr=[r for r in data['natural'] if r['split']==split and r['correct'] is truth]
            if not rr:continue
            qs=[r['question_id'] for r in rr]
            base=[not r['accepted'] if truth is True else r['accepted'] for r in rr]
            for v in cfg['variants']['natural']:
                ss=[by['natural',r['candidate_id']][v['name']] for r in rr]
                values=[not s['accepted'] if truth is True else s['accepted'] for s in ss]
                row={'arm':'natural','split':split,'variant':v['name'],'replacement':v['replacement'],
                     'category':v['category'],'metric':label,'reference':'original',**paired_stats(values,base,qs)}
                if label=='FPR':
                    row.update({'reject_to_false_accept':row['gained'],'false_accept_to_reject':row['lost']})
                elif label=='FNR':
                    row.update({'accept_correct_to_false_reject':row['gained'],'false_reject_to_accept':row['lost']})
                row['mean_log_odds_change']=float(np.mean([s['log_likelihood_A']-s['log_likelihood_B']-
                    by['natural',r['candidate_id']]['original']['log_likelihood_A']+
                    by['natural',r['candidate_id']]['original']['log_likelihood_B'] for r,s in zip(rr,ss)]))
                out.append(row)
    for metric in ['FPR','FNR']:
        rr=[r for r in out if r['metric']==metric and r['variant']!='original']
        qvalues=bh([r['p_cluster'] if r['p_cluster'] is not None else np.nan for r in rr])
        for r,q in zip(rr,qvalues):r['q_BH_vs_original']=float(q) if np.isfinite(q) else None
    return out


def projected_G3(cfg,data,by):
    full=[r for r in source_rows() if r['G_round']=='G3'];out=[]
    for split in ['A','C']:
        rr=[r for r in full if r['split']==split]
        natural={r['candidate_id']:r for r in data['natural'] if r['split']==split}
        wrong=sum(r['correct'] is False for r in rr);correct=sum(r['correct'] is True for r in rr)
        baseline_FP=sum(r['correct'] is False and r['accepted'] for r in rr)
        baseline_FN=sum(r['correct'] is True and not r['accepted'] for r in rr)
        for v in cfg['variants']['natural']:
            def accepted(r):return by['natural',r['candidate_id']][v['name']]['accepted'] if r['candidate_id'] in natural else r['accepted']
            fp=sum(r['correct'] is False and accepted(r) for r in rr);fn=sum(r['correct'] is True and not accepted(r) for r in rr)
            out.append({'split':split,'variant':v['name'],'wrong_total':wrong,'correct_total':correct,
                'original_FP':baseline_FP,'edited_FP':fp,'whole_pool_original_FPR':baseline_FP/wrong,'whole_pool_edited_FPR':fp/wrong,
                'whole_pool_FPR_delta':(fp-baseline_FP)/wrong,'original_FN':baseline_FN,'edited_FN':fn,
                'whole_pool_original_FNR':baseline_FN/correct,'whole_pool_edited_FNR':fn/correct,
                'whole_pool_FNR_delta':(fn-baseline_FN)/correct,
                'only_matching_G3_openings_edited':True,'not_G2_G3_training_mediation':True})
    return out


def cases(cfg,data,by):
    output=[];lines=['# 自然短语改写的判定变化案例','','定向展示变化较大的案例，不是随机效果估计。正文仅替换记录中的短语，数学内容和最终答案保持不变。','']
    for split in ['A','C']:
        rr=[r for r in data['natural'] if r['split']==split and r['correct'] is False]
        # Three prespecified expressions, prefer a new true-negative and a new false-positive.
        for name in ['we_should','we_have_to','require']:
            for direction in ['accept_to_reject','reject_to_accept']:
                if direction=='accept_to_reject':
                    selected=[r for r in rr if r['accepted'] and not by['natural',r['candidate_id']][name]['accepted']]
                else:
                    selected=[r for r in rr if not r['accepted'] and by['natural',r['candidate_id']][name]['accepted']]
                if not selected:continue
                def difference(r):
                    a=by['natural',r['candidate_id']]['original'];b=by['natural',r['candidate_id']][name]
                    return abs((b['log_likelihood_A']-b['log_likelihood_B'])-(a['log_likelihood_A']-a['log_likelihood_B']))
                r=max(selected,key=difference);v=next(v for v in cfg['variants']['natural'] if v['name']==name)
                output.append({'original':r,'edited_text':edited(r,v,'natural'),'variant':v,
                    'original_V0':by['natural',r['candidate_id']]['original'],'edited_V0':by['natural',r['candidate_id']][name],
                    'direction':direction,'human_review_status':'pending'})
                lines += ['## '+split+' '+name+' '+direction,r['question'],
                    f'ID={r["candidate_id"]}；最终答案={r["extracted_answer"]}；参考={r["gold_answer"]}；原分数={r["score"]:.9g}；改写分数={by["natural",r["candidate_id"]][name]["score"]:.9g}。',
                    '原开场：','```text',r['text'][:500],'```','改写开场：','```text',edited(r,v,'natural')[:500],'```','']
    write_jsonl(REPORT/'changed_case_texts.jsonl',output);(REPORT/'changed_cases.md').write_text('\n'.join(lines))
    return output


def workbook(tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();wb.remove(wb.active)
    for name,rr in tables.items():
        ws=wb.create_sheet(name[:31]);cols=list(dict.fromkeys(k for r in rr for k in r));ws.append(cols)
        for r in clean(rr):
            ws.append([excel_safe_value(r.get(k)) for k in cols])
            for cell in ws[ws.max_row]:
                if cell.data_type=='f':cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='264653')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(44,max(16,len(str(col[0].value))+2))
    wb.save(REPORT/'paraphrase_review.xlsx')


def figures(prefix,natural):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    variants=['we_should','we_have_to','we_must','it_is_necessary_to','the_task_requires_us_to','we_require_to','require','we_can']
    with PdfPages(REPORT/'paraphrase_figures.pdf') as pdf:
        for arm,rr in [('prefix',prefix),('natural',natural)]:
            fig,axes=plt.subplots(1,2,figsize=(13,5),constrained_layout=True)
            for ax,split in zip(axes,['A','C']):
                ss=[next(r for r in rr if r['split']==split and r['variant']==v and (arm=='prefix' or r['metric']=='FPR')) for v in variants]
                for i,r in enumerate(ss):
                    ax.plot([100*r['low'],100*r['high']],[i,i],color='#2379a8');ax.scatter([100*r['delta']],[i],color='#2379a8')
                ax.axvline(0,color='gray',linestyle='--');ax.set(yticks=range(len(ss)),yticklabels=variants,
                    xlabel='Paired difference, percentage points',title=f'{split}: '+('injected prefix vs "we need to"' if arm=='prefix' else 'natural phrase replacement vs original FPR'))
            fig.savefig(REPORT/(arm+'_paired_effects.png'),dpi=160);pdf.savefig(fig);plt.close(fig)
        names=list(dict.fromkeys(r['variant'] for r in prefix))
        heat=np.array([[next(r['rate'] for r in prefix if r['split']==split and r['variant']==name)*100 for split in ['A','C']] for name in names])
        fig,ax=plt.subplots(figsize=(7,13),constrained_layout=True)
        im=ax.imshow(heat,vmin=0,vmax=max(30,float(heat.max())),cmap='YlOrRd',aspect='auto')
        for i in range(len(names)):
            for j in range(2):ax.text(j,i,f'{heat[i,j]:.1f}',ha='center',va='center',fontsize=7,
                                    color='white' if heat[i,j]>.65*max(30,float(heat.max())) else 'black')
        ax.set(xticks=[0,1],xticklabels=['A (64 originally rejected wrong)','C (48 originally rejected wrong)'],
               yticks=range(len(names)),yticklabels=names,title='All 46 prespecified prefixes: reject-to-accept rate (%)')
        ax.tick_params(axis='y',labelsize=7);fig.colorbar(im,ax=ax,shrink=.6)
        fig.savefig(REPORT/'all_prefix_flip_rates.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,2,figsize=(13,5),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            ss=[next(r for r in natural if r['split']==split and r['variant']==v and r['metric']=='FNR') for v in variants]
            for i,r in enumerate(ss):
                ax.plot([100*r['low'],100*r['high']],[i,i],color='#d7872c');ax.scatter([100*r['delta']],[i],color='#d7872c')
            ax.axvline(0,color='gray',linestyle='--');ax.set(yticks=range(len(ss)),yticklabels=variants,
                xlabel='FNR change, percentage points',title=f'{split}: natural correct answers; paired question 95% CI')
        fig.savefig(REPORT/'natural_FNR_changes.png',dpi=160);pdf.savefig(fig);plt.close(fig)


def pct(x):return 'NA' if x is None else f'{100*x:.2f}%'
def pp(x):return 'NA' if x is None else f'{100*x:+.2f}'


def render(cfg,prefix,natural,projection):
    p={(r['split'],r['variant']):r for r in prefix};n={(r['split'],r['variant'],r['metric']):r for r in natural}
    labels={v['name']:v['replacement'] for v in cfg['variants']['natural']}
    labels['original']='we need to（原文）';labels['delete_target']='删除该短语'
    names=['we_need_to','we_should','we_have_to','we_must','it_is_necessary_to','the_task_requires_us_to','we_require_to','require','we_can']
    intro=[]
    for split in ['A','C']:
        base=n[split,'original','FPR']
        items='，'.join(labels[v]+' '+pct(n[split,v,'FPR']['rate']) for v in ['we_should','we_have_to','require'])
        intro.append(f'{split}自然错误组：原文FPR为{pct(base["rate"])}；改写为{items}。')
    lines=['# 前缀近义改写与句式替换实验','',*intro,'',
      '以下把人工注入和自然原位替换分别比较。数值差异、配对区间和多重校正结果一起看；某些句式改变语气或语法，不只改变字面词形。','',
      '固定V0参数、提示、原Transformers A/B似然评分及0.5阈值。无G生成、无训练、B未使用、T未读取。全部变体在新评分前冻结，保留全部结果。','',
      '## 两个不同问题','','注入组：沿用此前112道题各1条原本被拒绝的错误G3回答（A确认份64、C48），把人工注入的独立前缀替换为46种词语、语气、句式、格式或控制条件。回答全文不变。这里的接受率是“原本被拒错误回答的翻转率”，不是整池FPR。','',
      '自然组：G3全部自然命中we need to开场区域的498条回答（A410、C88）；只在原位置替换第一处短语，10种处理，含原文、删除、近义和语法压力表达。A含87正确/283错误/40未知，C含23正确/59错误/6未知；未知不并入错误。多候选同题依赖按题聚类处理。C为已反复使用的开发集，不是封存的最终T。','',
      '`we should`、`we can`改变模态强度；`require`、`we require to`允许语法不自然。本实验称“词形/模态/句式改写”，不把所有变体都当作完全语义等价的近义词。所有变体独立核对最终答案、真假标签、解析状态不变，完整数学后缀保持字节一致。','',
      '## 1. 同一批错误回答：换注入前缀后的翻转率','','|前缀|A接受率|相对we need to/百分点|C接受率|相对we need to/百分点|','|---|---|---|---|---|']
    for name in names:
        a=p['A',name];c=p['C',name];label=a['prefix'].strip();lines.append(f'|{label}|{pct(a["rate"])}|{pp(a["delta"])}|{pct(c["rate"])}|{pp(c["delta"])}|')
    lines += ['','46种完整结果见`prefix_paired_results.csv`，含大小写、标点、空格、To solve for同族和推理标题。每个前缀末尾的换行也属于干预；不得把不同位置/格式当作同一个词的普遍效果。完整图为`all_prefix_flip_rates.png`。','',
      '![注入组](prefix_paired_effects.png)','','主比较为相对we need to的双侧精确McNemar检验，跨所有非参照变体×A/C作BH校正；相对中性前缀另列一族。按题配对bootstrap10000次、seed20261024。探索性区间未作同时覆盖校正，少数翻转时也不等同于精确检验的反演区间；显著性以预定p/q判断。比较族比上次9前缀实验更大且此次为双侧检验，q值不能与上次直接对照。筛出的最大效果也不代表预先知道最优表达。','',
      '## 2. 自然出现的短语：原位替换后FPR','','|改写|A错误回答FPR|A变化/百分点|A配对95%区间/百分点|C错误回答FPR|C变化/百分点|','|---|---|---|---|---|---|']
    for name in ['original',*names[1:],'delete_target']:
        a=n['A',name,'FPR'];c=n['C',name,'FPR'];lines.append(f'|{labels[name]}|{pct(a["rate"])}|{pp(a["delta"])}|[{pp(a["low"])},{pp(a["high"])}]|{pct(c["rate"])}|{pp(c["delta"])}|')
    lines += ['','![自然组](natural_paired_effects.png)','','自然组对每种处理以同一回答原文为对照，错误回答看FPR、正确回答看FNR，未知只描述接受率。标准误和bootstrap按题聚类；FPR的9处理×2题池共18项作BH校正，FNR另作一族。C某些标签组独立题数较少，完整表标记支持量。','']
    for split in ['A','C']:
        sig=[r for r in natural if r['split']==split and r['metric']=='FPR' and r['variant']!='original' and r.get('q_BH_vs_original') is not None and r['q_BH_vs_original']<.05]
        lines.append(f'- {split}自然错误组：通过BH校正的变体为'+('、'.join(r['variant']+'('+pp(r['delta'])+'个百分点)' for r in sig) if sig else '无')+'。')
    lines += ['','### 同一条错误回答的双向变化','','|题池/处理|原误接受→拒绝|原拒绝→误接受|净新增误接受|','|---|---|---|---|']
    for split in ['A','C']:
        for name in ['we_should','we_have_to','require']:
            r=n[split,name,'FPR'];lines.append(f'|{split}/{labels[name]}|{r["lost"]}|{r["gained"]}|{r["gained"]-r["lost"]:+d}|')
    lines += ['','总体变化小也可能包含双向翻转；不能仅凭净FPR稳定，就认为同一回答对词形改写不敏感。','']
    lines += ['','## 3. 改写全部自然命中后，G3整池指标','','这不是抽样外推：已覆盖全部498条自然命中回答，未命中者复用原判定。这里只计算当前G3缓存经该编辑后的完整题池指标；尚未对G2做相同编辑，不能解释为G2→G3训练差值的因果贡献。','','|题池/处理|整池FPR|整池FPR变化/百分点|整池FNR|整池FNR变化/百分点|','|---|---|---|---|---|']
    for r in projection:
        if r['variant'] in ['original','we_should','we_have_to','require','delete_target']:
            lines.append(f'|{r["split"]}/{labels[r["variant"]]}|{pct(r["whole_pool_edited_FPR"])}|{pp(r["whole_pool_FPR_delta"])}|{pct(r["whole_pool_edited_FNR"])}|{pp(r["whole_pool_FNR_delta"])}|')
    lines += ['','正确回答的FNR结果必须同时看：若某种改写使V整体更倾向拒绝，FPR下降不自动说明验证能力改善。单独前缀注入组与自然原位组的样本、位置和上下文不同，不能只凭两组效果差异认定位置是唯一原因。','',
      '## 审计与复核','','每条原文均重新评分并核对历史分数与判定；只有完整输入完全相同的既有注入变体才复用旧分数。评分均未截断。`scores`总调用条目不等于独立题数；自然组同题多个候选始终一起重采样。','',
      '`paraphrase_review.xlsx`包含完整比较、总体指标和变体清单；`paraphrase_figures.pdf`为4页图（含完整46种前缀及正确回答FNR）；`changed_cases.md`与JSONL保留定向变化案例；原文、跨度及每次评分分别保留在runs/prefix_paraphrase_20261024。新发现的强变体属于本次探索，不自动泛化为万能钥匙或G已学会利用它。','']
    (REPORT/'paraphrase_report.md').write_text('\n'.join(lines))


def main():
    cfg,data,by,scores=load_scores();REPORT.mkdir(exist_ok=True)
    prefix=analyze_prefix(cfg,data,by);natural=analyze_natural(cfg,data,by);projection=projected_G3(cfg,data,by)
    inventory=[{'arm':arm,**v} for arm,vs in cfg['variants'].items() for v in vs]
    tables={'prefix_paired_results':prefix,'natural_paired_results':natural,'whole_G3_pool_results':projection,'variant_inventory':inventory}
    for name,rows in tables.items():write_csv(REPORT/(name+'.csv'),rows)
    for arm,rr in data.items():
        write_jsonl(REPORT/(arm+'_counterfactuals.jsonl'),
            [{'original':r,'variant_scores':list(by[arm,r['candidate_id']].values())} for r in rr])
    cases(cfg,data,by);workbook(tables);figures(prefix,natural);render(cfg,prefix,natural,projection)
    write_json(REPORT/'results.json',clean({'completed_at':utc(),'prefix':prefix,'natural':natural,'whole_G3_pool':projection,
              'compute':json.loads((RUN/'compute.json').read_text()),'plan_sha256':sha(PLAN)}))
    copied=REPORT/'source';copied.mkdir(exist_ok=True)
    for name in ['prefix_paraphrase.py','prefix_paraphrase_report.py']:shutil.copyfile(ROOT/'src/pilot'/name,copied/name)
    shutil.copyfile(PLAN,REPORT/'frozen_plan.json')
    shutil.copyfile(ROOT/'tests/test_prefix_paraphrase.py',copied/'test_prefix_paraphrase.py')
    audit={'completed_at':utc(),'scores':len(scores),'source_hashes_unchanged':True,'V0_fixed':True,
        'original_rescores_matched':True,'all_answers_preserved':True,'all_scoring_contexts_complete':True,
        'B_used':False,'T_read':False,'model_training':False,'plan_sha256':sha(PLAN),
        'analysis_code_sha256':sha(Path(__file__))}
    validate(cfg);write_json(REPORT/'audit.json',audit)
    manifest={str(p.relative_to(REPORT)):sha(p) for p in sorted(REPORT.rglob('*')) if p.is_file() and p.name not in ['artifact_manifest.json','paraphrase_review.zip']}
    write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'paraphrase_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(REPORT.rglob('*')):
            if p.is_file() and p.name!='paraphrase_review.zip':z.write(p,str(p.relative_to(REPORT)))
    with zipfile.ZipFile(REPORT/'paraphrase_review.zip') as z:assert z.testzip() is None
    write_json(RUN/'state.json',{'status':'complete','completed_at':utc(),'report':str(REPORT/'paraphrase_report.md'),'T_read':False})
    print(json.dumps({'complete':True,'report':str(REPORT/'paraphrase_report.md')},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
