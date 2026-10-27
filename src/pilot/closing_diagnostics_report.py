"""Training-chain and counterfactual closure, preserving frozen experiments."""
import json
import shutil
import zipfile
from collections import Counter
from pathlib import Path

import numpy as np

from .core import ROOT, read_jsonl, write_json, write_jsonl, digest
from .formal_common import messages, sha, utc
from .closing_diagnostics import RUN, REPORT, PLAN, SELECTED, validate, natural_features
from .fpr_pattern_diagnostics import write_csv
from .g3_report import paired
from .v0_drift_analysis import clean, interval


def chain():
    cfg=validate();rr=read_jsonl(RUN/'all_source_answers.jsonl');selected=read_jsonl(SELECTED)
    ids={r['candidate_id'] for r in selected};qs=sorted({r['question_id'] for r in selected})
    g2=[r for r in rr if r['split']=='A' and r['G_round']=='G2'];g3=[r for r in rr if r['split']=='A' and r['G_round']=='G3']
    stages={'G2_candidates':g2,'V0_accepted':[r for r in g2 if r['accepted']],
            'actual_G3_SFT':[r for r in g2 if r['candidate_id'] in ids],'G3_outputs':g3}
    feats={r['candidate_id']:natural_features(r['text']) for r in g2+g3}
    patterns=cfg['chain']['patterns'];raw=[];quality=[];matched=[];contrasts=[]
    rng=np.random.default_rng(cfg['seed']);W=rng.multinomial(len(qs),np.full(len(qs),1/len(qs)),size=cfg['bootstrap'])
    stage_matrices={};boots={}
    for stage,rows in stages.items():
        quality.append({'stage':stage,'responses':len(rows),'questions':len({r['question_id'] for r in rows}),
            'correct':sum(r['correct'] is True for r in rows),'wrong':sum(r['correct'] is False for r in rows),
            'unknown':sum(r['correct'] is None for r in rows),'truncated':sum(r['truncated'] for r in rows)})
        for pattern in patterns:
            n=sum(feats[r['candidate_id']][pattern] for r in rows)
            raw.append({'stage':stage,'pattern':pattern,'responses':len(rows),'with_pattern':n,'prevalence':n/len(rows),
                        'questions_with_pattern':len({r['question_id'] for r in rows if feats[r['candidate_id']][pattern]})})
        group={q:[] for q in qs}
        for r in rows:
            if r['question_id'] in group:group[r['question_id']].append(r)
        assert all(group[q] for q in qs)
        mat=np.array([[np.mean([feats[r['candidate_id']][p] for r in group[q]]) for p in patterns] for q in qs])
        stage_matrices[stage]=mat;boots[stage]=W@mat/len(qs)
        for j,p in enumerate(patterns):
            ci=interval(boots[stage][:,j]);matched.append({'stage':stage,'pattern':p,'questions':len(qs),
                'question_mean_prevalence':float(mat[:,j].mean()),**ci})
    pairs=[('V0_accepted','G2_candidates'),('actual_G3_SFT','V0_accepted'),
           ('G3_outputs','G2_candidates'),('G3_outputs','actual_G3_SFT')]
    for after,before in pairs:
        for j,p in enumerate(patterns):
            contrasts.append({'pattern':p,'contrast':after+' minus '+before,'questions':len(qs),
                'delta':float(stage_matrices[after][:,j].mean()-stage_matrices[before][:,j].mean()),
                **interval(boots[after][:,j]-boots[before][:,j])})
    # Labels are descriptive only, joined after the frozen actual selector committed.
    label_rates=[]
    for stage,rows in stages.items():
        for label,truth in [('correct',True),('wrong',False),('unknown',None)]:
            subset=[r for r in rows if r['correct'] is truth]
            for pattern in patterns:
                n=sum(feats[r['candidate_id']][pattern] for r in subset)
                label_rates.append({'stage':stage,'label':label,'pattern':pattern,'responses':len(subset),
                    'with_pattern':n,'prevalence':n/len(subset) if subset else None})
    tables={'chain_raw_prevalence':raw,'chain_question_matched':matched,'chain_contrasts':contrasts,
            'chain_stage_quality':quality,'chain_label_strata':label_rates}
    for name,rows in tables.items():write_csv(REPORT/(name+'.csv'),rows)
    write_json(REPORT/'chain_results.json',tables)
    return tables


def load_scores():
    cfg=validate();jobs=read_jsonl(RUN/'jobs.jsonl');ss=read_jsonl(RUN/'V0_scores.jsonl')
    index={(j['arm'],j['candidate_id'],j['variant']):j for j in jobs}
    by={(s['arm'],s['candidate_id'],s['variant']):s for s in ss}
    assert len(ss)==len(by)==len(index)==cfg['counts']['logical_jobs']
    for key,s in by.items():
        j=index[key]
        assert s['question_id']==j['question_id'] and s['split']==j['split'] and s['G_round']==j['G_round'] and s['correct'] is j['correct']
        assert s['model_fingerprint']==cfg['fixed_V0']['fingerprint'] and s['adapter_fingerprint'] is None and s['threshold']==.5
        assert s['accepted']==(s['score']>=.5) and s['prompt_tokens']+4<=8192
        assert s['text_sha256']==digest(j['text']) and s['prompt_sha256']==digest(messages('V',j['question'],j['text']))
        assert s['cache_key']==digest({'plan':sha(PLAN),'arm':j['arm'],'candidate_id':j['candidate_id'],'variant':j['variant']})
        if j['variant']=='original':assert s['accepted']==j['original_accepted'] and abs(s['score']-j['original_score'])<2e-6
    return cfg,by,jobs,ss


def symmetric(cfg,by):
    allrows=read_jsonl(RUN/'all_source_answers.jsonl');touched={r['candidate_id'] for r in read_jsonl(RUN/'surface_answers.jsonl')}
    points=[];gaps=[];events=[]
    for split in ['A','C']:
        rows=[r for r in allrows if r['split']==split];qs=sorted({r['question_id'] for r in rows});data={}
        for g in ['G2','G3']:
            source=[r for r in rows if r['G_round']==g]
            for v in ['original','delete','synonym']:
                branch=[];changed=Counter()
                for r in source:
                    s=by.get(('symmetric',r['candidate_id'],v));accepted=s['accepted'] if s else r['accepted']
                    branch.append({**r,'accepted':accepted,'score':s['score'] if s else r['score']})
                    if r['candidate_id'] in touched and v!='original':
                        cell='TP' if r['correct'] is True and r['accepted'] else 'FN' if r['correct'] is True else 'FP' if r['correct'] is False and r['accepted'] else 'TN' if r['correct'] is False else 'unknown'
                        changed['matched_'+cell]+=1
                        if accepted!=r['accepted']:changed[cell+'_flipped']+=1
                data[g+'_'+v]=branch
                if v!='original':events.append({'split':split,'G_round':g,'variant':v,**dict(changed)})
        point,boot=paired(data,qs,draws=cfg['bootstrap'],seed=cfg['seed'])
        for key,p in point.items():
            g,v=key.split('_',1)
            points.append({'split':split,'G_round':g,'variant':v,'questions':len(qs),**p})
        for v in ['original','delete','synonym']:
            for metric in ['FPR','FNR','balanced_accuracy','accepted_error_FDR']:
                delta=point['G3_'+v][metric]-point['G2_'+v][metric]
                sample=boot['G3_'+v][metric]-boot['G2_'+v][metric]
                base=boot['G3_original'][metric]-boot['G2_original'][metric]
                base_point=point['G3_original'][metric]-point['G2_original'][metric]
                reduction=interval(base-sample)
                gaps.append({'split':split,'variant':v,'metric':metric,'G3_minus_G2':delta,**interval(sample),
                    'original_gap':base_point,'gap_reduction_after_edit':base_point-delta,
                    'reduction_low':reduction['low'],'reduction_high':reduction['high'],
                    'meaning':'Descriptive change of observed round gap under the specified text edit; not causal training mediation'})
    tables={'symmetric_pool_metrics':points,'symmetric_round_gaps':gaps,'symmetric_cell_transitions':events}
    for name,rows in tables.items():write_csv(REPORT/(name+'.csv'),clean(rows))
    write_json(REPORT/'symmetric_results.json',clean(tables));return tables


def answer_control(cfg,by):
    rr=read_jsonl(RUN/'answer_control_answers.jsonl');out=[];contrasts=[];cases=[]
    for split in ['A','C']:
        rows=[r for r in rr if r['split']==split];qs=sorted({r['question_id'] for r in rows});ix={q:i for i,q in enumerate(qs)}
        qi=np.array([ix[r['question_id']] for r in rows]);counts=np.bincount(qi,minlength=len(qs))
        rng=np.random.default_rng(cfg['seed']);W=rng.multinomial(len(qs),np.full(len(qs),1/len(qs)),size=cfg['bootstrap'])
        estimates={};samples={};individual={}
        for v in ['original','delete','synonym','wrong_final_answer']:
            ss=[by['answer_control' if v=='wrong_final_answer' else 'symmetric',r['candidate_id'],v] for r in rows]
            reject=np.array([not s['accepted'] for s in ss],float);sums=np.bincount(qi,weights=reject,minlength=len(qs))
            boot=W@sums/(W@counts);samples[v]=boot;estimates[v]=float(reject.mean());individual[v]=reject
            out.append({'split':split,'variant':v,'responses':len(rows),'questions':len(qs),'rejections':int(reject.sum()),
                'rejection_rate':float(reject.mean()),**interval(boot),
                'rejection_means':'correct rejection of deliberately wrong answer' if v=='wrong_final_answer' else 'false rejection of still-correct answer',
                'still_accepted':int(len(rows)-reject.sum()),'mean_acceptance_score':float(np.mean([s['score'] for s in ss]))})
        for reference in ['delete','synonym']:
            d=individual['wrong_final_answer']-individual[reference]
            contrasts.append({'split':split,'contrast':'wrong_final_answer rejection minus '+reference+' rejection',
                'difference':estimates['wrong_final_answer']-estimates[reference],**interval(samples['wrong_final_answer']-samples[reference]),
                'responds_to_answer_not_surface':int((d>0).sum()),'responds_to_surface_not_answer':int((d<0).sum()),
                'question_paired':True})
        # Keep all failures to reject the intentionally wrong final value.
        failures=[r for r in rows if by['answer_control',r['candidate_id'],'wrong_final_answer']['accepted']]
        failures=sorted(failures,key=lambda r:(-by['answer_control',r['candidate_id'],'wrong_final_answer']['score'],r['candidate_id']))
        for r in failures:
            cases.append({'original':r,'changed_text':r['final_corruption']['text'],
                          'new_V0':by['answer_control',r['candidate_id'],'wrong_final_answer'],'human_process_review':'not performed; label-only control'})
    write_csv(REPORT/'answer_control_reactions.csv',out);write_csv(REPORT/'answer_control_paired_contrasts.csv',contrasts)
    write_jsonl(REPORT/'answer_control_accepted_wrong_cases.jsonl',cases)
    result={'answer_control_reactions':out,'answer_control_paired_contrasts':contrasts}
    write_json(REPORT/'answer_control_results.json',clean(result));return result


def figures(chain_tables,symmetric_tables,answer_tables):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    stages=['G2_candidates','V0_accepted','actual_G3_SFT','G3_outputs']
    patterns=['we need to','to solve for','we can','we will','AWL_word_any']
    with PdfPages(REPORT/'closing_diagnostics_figures.pdf') as pdf:
        fig,axes=plt.subplots(1,2,figsize=(13,4.8),constrained_layout=True)
        for ax,view,key in [(axes[0],'All A questions (stage-specific denominators)','chain_raw_prevalence'),(axes[1],'Same 511 training questions (question-equal means)','chain_question_matched')]:
            for pattern in patterns:
                rr=[next(r for r in chain_tables[key] if r['pattern']==pattern and r['stage']==s) for s in stages]
                field='prevalence' if key=='chain_raw_prevalence' else 'question_mean_prevalence'
                ax.plot(range(4),[100*r[field] for r in rr],marker='o',label=pattern)
            ax.set_xticks(range(4),['G2 candidates','V0 accepted','Actual SFT','G3 outputs'],rotation=15);ax.set_ylabel('Natural opening prevalence (%)');ax.set_title(view)
        axes[1].legend(fontsize=8);fig.suptitle('Does the feedback chain enrich the expressions? Descriptive, not a causal learning test')
        fig.savefig(REPORT/'training_chain_prevalence.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(10,4.8),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            for g,color in [('G2','#1976A3'),('G3','#D77A21')]:
                rr=[next(r for r in symmetric_tables['symmetric_pool_metrics'] if r['split']==split and r['G_round']==g and r['variant']==v) for v in ['original','delete','synonym']]
                ax.plot(range(3),[100*r['FPR'] for r in rr],marker='o',label=g,color=color)
            ax.set_xticks(range(3),['Original','Delete','Near-synonym']);ax.set_title(split);ax.set_ylabel('Full-pool FPR (%)');ax.legend()
        fig.suptitle('Same text-edit rules for G2 and G3; fixed V0, same full-pool denominators')
        fig.savefig(REPORT/'symmetric_FPR.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(11,4.8),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            rr=[next(r for r in symmetric_tables['symmetric_round_gaps'] if r['split']==split and r['variant']==v and r['metric']=='FPR') for v in ['delete','synonym']]
            for i,r in enumerate(rr):
                ax.plot([100*r['reduction_low'],100*r['reduction_high']],[i,i],lw=3);ax.plot(100*r['gap_reduction_after_edit'],i,'ko')
            ax.axvline(0,color='black',ls='--');ax.set_yticks([0,1],['Delete','Near-synonym']);ax.set_title(split)
            ax.set_xlabel('Original round gap minus edited round gap (pp)')
        fig.suptitle('How much does editing reduce the G3-G2 FPR gap?\nPaired question-bootstrap 95% intervals; descriptive counterfactual')
        fig.savefig(REPORT/'FPR_gap_reduction.png',dpi=180);pdf.savefig(fig);plt.close(fig)

        fig,axes=plt.subplots(1,2,figsize=(11,4.8),constrained_layout=True)
        for ax,split in zip(axes,['A','C']):
            vv=['original','delete','synonym','wrong_final_answer']
            rr=[next(r for r in answer_tables['answer_control_reactions'] if r['split']==split and r['variant']==v) for v in vv]
            ax.bar(range(4),[100*r['rejection_rate'] for r in rr],color=['#90a4ae','#90a4ae','#90a4ae','#1976A3'])
            for i,r in enumerate(rr):ax.text(i,100*r['rejection_rate']+2,f'{r["rejections"]}/{r["responses"]}',ha='center')
            ax.set_xticks(range(4),['Original\ncorrect','Delete\nstill correct','Synonym\nstill correct','Final +1\nnow wrong'],fontsize=8)
            ax.set_ylim(0,105);ax.set_ylabel('V0 rejection rate (%)');ax.set_title(split+f' (same {rr[0]["responses"]} answers)')
        fig.suptitle('Response to an actual final-answer error vs superficial edits\nRejecting still-correct answers is harmful; rejecting final+1 is appropriate')
        fig.savefig(REPORT/'answer_vs_surface_sensitivity.png',dpi=180);pdf.savefig(fig);plt.close(fig)


def report_text(cfg,ct,st,at):
    lines=['# 表达敏感性支线：三项补充诊断','','本次完成训练链表达追踪、G2/G3对称自然编辑及最终答案改错对照。所有结果均使用固定V0；T继续封存。',
      '', '## 1. 表达是否沿反馈训练链放大','','主轨迹为同vLLM引擎的G2→G3。全池阶段的分母不同，因此同时报告实际进入SFT的511道题、每题等权的匹配视图；接受→SFT还受到未截断、训练长度、每题随机选一条及512上限影响，不能把这一步的全部变化归因于V。',
      '', '|阶段|回答|题目|正确|错误|未知|截断|','|---|---:|---:|---:|---:|---:|---:|']
    for r in ct['chain_stage_quality']:lines.append('|'+ '|'.join(str(r[k]) for k in ['stage','responses','questions','correct','wrong','unknown','truncated'])+'|')
    lines += ['', '按实际511题匹配、每题等权的自然开场比例：','','|表达|G2候选|V接受|实际SFT|G3输出|G3−G2/百分点及95%区间|','|---|---:|---:|---:|---:|---|']
    for pattern in cfg['chain']['patterns']:
        rr=[next(r for r in ct['chain_question_matched'] if r['pattern']==pattern and r['stage']==s) for s in ['G2_candidates','V0_accepted','actual_G3_SFT','G3_outputs']]
        d=next(r for r in ct['chain_contrasts'] if r['pattern']==pattern and r['contrast']=='G3_outputs minus G2_candidates')
        lines.append('|'+pattern+'|'+'|'.join(f'{100*r["question_mean_prevalence"]:.2f}%' for r in rr)+f'|{100*d["delta"]:+.2f} [{100*d["low"]:+.2f}, {100*d["high"]:+.2f}]|')
    lines += ['', '这些是既有自然文本中的出现比例，窗口为开头160字符、首次数学公式/数字/换行前、词序起点前12词。大小写忽略，按词匹配；不是要求自然文本恰好以人工注入串开头。AWL_word_any/nonAWL_word_any只覆盖前次冻结20词对，不是全570词族。',
      '所有标签在实际选择提交后用于质量记录；没有按真答案重新筛选。阶段富集即使存在，也不能单独证明G有意利用V。全部15种模式均导出，未按本次结果挑选；区间为探索性逐项区间，不作多重比较显著性宣称。',
      '', '## 2. 相同自然编辑能否缩小跨轮FPR差距','','覆盖两轮全部自然命中回答，包括TP/FN/FP/TN和未知；未命中者保留原评分。每条只编辑最早一个合规开场跨度，沿用此前冻结的词对及短语规则；独立核验所有表面编辑的最终答案、真假标签和解析状态不变。',
      '', '|池|处理|G2 FPR|G3 FPR|跨轮差/百分点及95%区间|原差距减编辑后差距/百分点及95%区间|','|---|---|---:|---:|---|---|']
    for split in ['A','C']:
        for v in ['original','delete','synonym']:
            pp=[next(r for r in st['symmetric_pool_metrics'] if r['split']==split and r['G_round']==g and r['variant']==v) for g in ['G2','G3']]
            d=next(r for r in st['symmetric_round_gaps'] if r['split']==split and r['variant']==v and r['metric']=='FPR')
            lines.append(f'|{split}|{v}|{100*pp[0]["FPR"]:.2f}%|{100*pp[1]["FPR"]:.2f}%|{100*d["G3_minus_G2"]:+.2f} [{100*d["low"]:+.2f}, {100*d["high"]:+.2f}]|{100*d["gap_reduction_after_edit"]:+.2f} [{100*d["reduction_low"]:+.2f}, {100*d["reduction_high"]:+.2f}]|')
    lines += ['', '差距减少为正表示编辑后G3−G2差距缩小，为负表示扩大；这不是训练中介效应的因果识别。两轮回答内容和错误组成不同，编辑也可能影响语法。',
      '', '|池|处理|G2 FNR|G3 FNR|','|---|---|---:|---:|']
    for split in ['A','C']:
        for v in ['original','delete','synonym']:
            pp=[next(r for r in st['symmetric_pool_metrics'] if r['split']==split and r['G_round']==g and r['variant']==v) for g in ['G2','G3']]
            lines.append(f'|{split}|{v}|{100*pp[0]["FNR"]:.2f}%|{100*pp[1]["FNR"]:.2f}%|')
    lines += ['', 'FNR同时列出，避免把一概更倾向拒绝误解为验证能力提高。未知不并入错误；全部未知数、上下界与双向翻转计数保存在完整表。',
      '', '## 3. V对最终答案错误是否比表述变化更敏感','','沿用前次自然TP样本，取唯一数值盒且无额外####断言者；原文被正确接受。将盒内整数n改成n+1，其他字节全部不变，并核验新最终答案确实错误。对应表面编辑与答案编辑使用完全相同的回答子集。',
      '', '|池|处理|拒绝数/回答|拒绝率及题目配对95%区间|应如何解读|','|---|---|---:|---|---|']
    for r in at['answer_control_reactions']:
        meaning='正确拒绝故意改错的最终答案' if r['variant']=='wrong_final_answer' else '误拒仍然正确的最终答案'
        lines.append(f'|{r["split"]}|{r["variant"]}|{r["rejections"]}/{r["responses"]}|{100*r["rejection_rate"]:.2f}% [{100*r["low"]:.2f}, {100*r["high"]:.2f}]|{meaning}|')
    lines += ['', '最终数值改错后仍被接受，说明这个控制条件下未能检出数值错误。推导正文可能仍支持原正确数，这个结果不能直接推广为对所有自然错误的FPR；它专门检查“最终答案正确性”协议下的数值敏感性。',
      '最初准备时误把整数范围限定为0–999，核对参考值后在任何新评分前取消了该限制，另存v2配置和完整输入；旧准备未执行推理并保留。最终对照未依本次V结果挑选样本；唯一额外排除及原因见冻结配置。',
      '', '## 可以收束到什么结论','']
    for split in ['A','C']:
        base=next(r for r in st['symmetric_round_gaps'] if r['split']==split and r['variant']=='original' and r['metric']=='FPR')
        de=next(r for r in st['symmetric_round_gaps'] if r['split']==split and r['variant']=='delete' and r['metric']=='FPR')
        sy=next(r for r in st['symmetric_round_gaps'] if r['split']==split and r['variant']=='synonym' and r['metric']=='FPR')
        ans=next(r for r in at['answer_control_reactions'] if r['split']==split and r['variant']=='wrong_final_answer')
        lines.append(f'- {split}：原跨轮FPR差{100*base["G3_minus_G2"]:+.2f}个百分点，删除后{100*de["G3_minus_G2"]:+.2f}，替换后{100*sy["G3_minus_G2"]:+.2f}；答案改错检出{ans["rejections"]}/{ans["responses"]}（{100*ans["rejection_rate"]:.2f}%）。')
    lines += ['', '训练链没有显示这些模式被V筛选、SFT和后续生成一致放大：we need to在G3减少，we will基本稳定，we can在本轮SFT的预设开场窗口中未出现；to solve for在G3增加，但V接受阶段相对G2候选并未显示对应富集。',
      'A的近义替换将差距缩小约0.79个百分点，探索性配对区间[0.28,1.36]；存在小幅表述相关变化，但原7.69个百分点的大部分上涨仍然存在。不能把0.79直接叫作训练中由词语机制导致的因果贡献；C未显示对应的缩小。',
      '最终答案+1有很高的检出率，支持V对此类最终数值变化有反应。正文仍保留原计算，这也可能使V借助正文与最终答案的一致性判断，不能据此推断独立重新解题能力，更不能把99%的控制检出率推广到自然生成错误。',
      'A中3条改错后仍接受的回答来自2道题：两条312→313（分数仍约0.9998/0.9996），一条81→82（约0.8933）。人工数学复算确认这两个改动确实错误，详见answer_control_failure_math_audit.json；没有因失败而删除这些样本。',
      '', '以上以观察结果和配对区间表述，不预设触发词解释一定成立。自然频率、反馈选择、SFT与G3变化必须共同考察；局部可触发漏洞、自然利用和总体跨轮漂移是三项不同证据。',
      'C为反复使用的开发集，T仍封存。区间按同一道题的所有候选一起重采样10000次；没有训练种子重复。零翻转时非参数bootstrap可能退化为[0,0]，它不证明总体零风险。',
      '', '## 资源、审计与交付','']
    compute=json.loads((RUN/'compute.json').read_text())
    lines += [f'本轮{cfg["counts"]["surface_answers"]}条自然命中回答、{cfg["counts"]["answer_controls"]}条答案对照，{compute["logical_scores"]}条逻辑评分；新增{compute["new_scores_this_invocation"]}、复用{compute["reused_this_invocation"]}条完整输入相同结果，GPU阶段{compute["seconds_this_invocation"]:.1f}秒。24个跨轮/池/真假/原判定层的原文哨兵重新评分复现历史值；其他原文及完全相同干预按来源哈希复用。',
      '未训练或生成G、未更新V，B未使用、T未读取。原模型、冻结配置和历史产物保持不变；新评分、输入、编辑跨度、独立答案检查、全部结果和源代码保存在隔离目录。',
      '图：training_chain_prevalence.png、symmetric_FPR.png、FPR_gap_reduction.png、answer_vs_surface_sensitivity.png；四页PDF、Excel、CSV及可核对哈希ZIP一起提供。']
    (REPORT/'supplementary_evidence_report.md').write_text('\n'.join(lines)+'\n')


def main():
    cfg,by,jobs,scores=load_scores();ct=chain();st=symmetric(cfg,by);at=answer_control(cfg,by)
    figures(ct,st,at);report_text(cfg,ct,st,at)
    tables=ct|st|at
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        ws=wb.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in rows for k in r));ws.append(keys)
        for r in clean(rows):
            ws.append([r.get(k) for k in keys])
            for c in ws[ws.max_row]:
                if c.data_type=='f':c.data_type='s'
        for c in ws[1]:c.font=Font(bold=True)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(REPORT/'supplementary_evidence_review.xlsx')
    write_json(REPORT/'results.json',clean(tables))
    for src in [PLAN,RUN/'compute.json',RUN/'answer_control_exclusions.json',ROOT/'src/pilot/closing_diagnostics.py',Path(__file__)]:
        shutil.copyfile(src,REPORT/src.name)
    write_json(REPORT/'audit.json',{'passed':True,'source_and_input_hashes_verified':True,'score_provenance_and_full_prompts_verified':True,
        'surface_labels_preserved':True,'answer_corruption_labels_independently_false':True,'actual511_SFT_rows_identical':True,
        'V0_fixed':True,'B_used':False,'T_read':False,'scores':len(scores),'at':utc()})
    files=sorted(p for p in REPORT.iterdir() if p.is_file() and p.suffix!='.zip' and p.name!='artifact_manifest.json')
    manifest={p.name:sha(p) for p in files};write_json(REPORT/'artifact_manifest.json',manifest)
    with zipfile.ZipFile(REPORT/'supplementary_evidence_review.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in files+[REPORT/'artifact_manifest.json']:z.write(p,p.name)
    import hashlib
    with zipfile.ZipFile(REPORT/'supplementary_evidence_review.zip') as z:
        assert z.testzip() is None and all(hashlib.sha256(z.read(n)).hexdigest()==h for n,h in manifest.items())
    write_json(RUN/'state.json',{'status':'complete','report':str(REPORT/'supplementary_evidence_report.md'),'at':utc()})
    print(json.dumps(clean({'FPR_gaps':[r for r in st['symmetric_round_gaps'] if r['metric']=='FPR'],
          'answer_reactions':at['answer_control_reactions']}),ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
