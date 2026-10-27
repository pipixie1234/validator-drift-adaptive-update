"""One-shot final T matrices and seven-day delivery; original data immutable."""
import json
import shutil
import zipfile
from collections import Counter
from pathlib import Path

import numpy as np

from .core import ROOT,read_jsonl,write_json,write_jsonl,digest
from .formal_common import messages,sha,utc
from .week_final_common import RUN,REPORT,PLAN,GS,VS,validate,state
from .week_final_reasons import reason_messages
from .v0_drift_analysis import analyze,count_array,count_metrics,score_metrics,clean,interval,divide
from .fpr_pattern_diagnostics import write_csv

QUALITY=ROOT/'configs/protocol_revisions/006_T_source_quality_20261026.json'
KEYS=['FPR','FNR','balanced_accuracy','V_accuracy','accepted_error_FDR','retention_all','AUROC',
      'G_accuracy_verified','G_accuracy_lower_all','unknown_fraction']


def load_final():
    cfg=validate();quality=json.loads(QUALITY.read_text());assert quality['base_final_protocol_sha256']==sha(PLAN)
    assert quality['T_input_sha256']==sha(RUN/'T_inputs.jsonl')
    excluded={r['question_id'] for r in quality['whole_question_quarantine']}
    allrows=[]
    for g in GS:
        raw=read_jsonl(RUN/f'{g}_T.jsonl');labels=read_jsonl(RUN/f'labels_{g}_T.jsonl');ll={r['candidate_id']:r for r in labels}
        assert len(raw)==len(ll)==1200 and {r['candidate_id'] for r in raw}==set(ll)
        for r in raw:
            l=ll[r['candidate_id']]
            assert r['question_id']==l['question_id'] and r['G_round']==g and r['split']=='T'
            assert r['model_fingerprint']==cfg['checkpoints'][g]['base_fingerprint'] and r['adapter_fingerprint']==cfg['checkpoints'][g]['adapter_fingerprint']
            assert r['plan_sha256']==sha(PLAN) and r['prompt_tokens']+r['output_tokens']<=4096
            assert r['max_new_tokens']==min(3840,4096-r['prompt_tokens']) and r['prompt_tokens']<=1024
            assert r['prompt_sha256']==digest(messages('G',r['question']))
            assert not r['truncated'] or l['correct'] is None
            allrows.append({**r,**l,'question_status':'quarantined_missing_essential_diagram' if r['question_id'] in excluded else 'eligible'})
    index={r['candidate_id']:r for r in allrows};assert len(index)==4800
    for q in {r['question_id'] for r in allrows}:
        qr=[r for r in allrows if r['question_id']==q]
        assert len(qr)==16 and len({r['question'] for r in qr})==1
        assert all(sum(r['G_round']==g for r in qr)==4 for g in GS)
    predictions={}
    for v in VS:
        pp=read_jsonl(RUN/'scoring'/v/'scores.jsonl');assert len(pp)==len({r['candidate_id'] for r in pp})==4800
        predictions[v]={}
        for s in pp:
            r=index[s['candidate_id']]
            assert s['question_id']==r['question_id'] and s['G_round']==r['G_round'] and s['V']==v
            assert s['model_fingerprint']==cfg['checkpoints'][v]['base_fingerprint'] and s['adapter_fingerprint']==cfg['checkpoints'][v]['adapter_fingerprint']
            assert s['plan_sha256']==sha(PLAN) and s['threshold']==.5 and s['accepted']==(s['score']>=.5)
            assert s['prompt_tokens']+4<=8192 and s['text_sha256']==digest(r['text'])
            assert s['prompt_sha256']==digest(messages('V',r['question'],r['text']))
            predictions[v][s['candidate_id']]=s
    return cfg,quality,allrows,predictions


def compute(cfg,quality,raw,predictions):
    valid=[r for r in raw if r['question_status']=='eligible'];qs=sorted({r['question_id'] for r in valid});assert len(qs)==299
    rng=np.random.default_rng(cfg['analysis']['seed']);indices=rng.integers(len(qs),size=(10000,len(qs)))
    tables={k:[] for k in ['metrics','metric_intervals','paired_V_differences','paired_G_differences','judgment_transitions',
                           'generator_quality','archival300_metrics','score_distribution','all_decisions']}
    points={};replicas={}
    for g in GS:
        rr=[r for r in valid if r['G_round']==g];assert len(rr)==1196
        allg=[r for r in raw if r['G_round']==g]
        for v in VS:
            state('T_final_analysis',G=g,V=v)
            data=[{**r,**{k:predictions[v][r['candidate_id']][k] for k in ['score','accepted']}} for r in rr]
            point,ci,rep,curve=analyze(data,qs,indices);points[g,v]=point;replicas[g,v]=rep
            tables['metrics'].append({'G':g,'V':v,'scope':'T299_valid',**point})
            for key in KEYS:tables['metric_intervals'].append({'G':g,'V':v,'metric':key,'estimate':point[key],**(ci[key] if key in ci else interval(rep[key]))})
            ss=[r['score'] for r in data]
            tables['score_distribution'].append({'G':g,'V':v,'score_mean':float(np.mean(ss)),
                'score_p05':float(np.quantile(ss,.05)),'score_median':float(np.median(ss)),'score_p95':float(np.quantile(ss,.95)),
                'exact_score_0p5':sum(s==.5 for s in ss),'score_between0p4_and0p6':sum(.4<=s<=.6 for s in ss),
                'interpretation':'Uncalibrated conditional A/B proxy; exact0.5 accepted under fixed rule'})
            archival=[{**r,**{k:predictions[v][r['candidate_id']][k] for k in ['score','accepted']}} for r in allg]
            ap=count_metrics(count_array(archival,sorted({r['question_id'] for r in archival})).sum(0))
            tables['archival300_metrics'].append(clean({'G':g,'V':v,'scope':'raw300_includes_invalid_question_not_capability_estimate',**ap}))
            if v=='V0':
                tables['generator_quality'].append({'G':g,**{k:point[k] for k in ['questions','correct','wrong','unknown','total',
                    'G_accuracy_verified','G_accuracy_lower_all','G_accuracy_upper_all','unknown_fraction','mean_output_tokens','p95_output_tokens','truncated','truncated_fraction']}})
            else:
                cells=np.zeros((2,2),int)
                for r in rr:
                    if r['correct'] is None:continue
                    old=predictions['V0'][r['candidate_id']]['accepted'];new=predictions[v][r['candidate_id']]['accepted']
                    cells[int(old!=r['correct']),int(new!=r['correct'])]+=1
                a,b,c,d=map(int,cells.ravel())
                tables['judgment_transitions'].append(clean({'G':g,'old_V':'V0','new_V':v,'a_old_right_new_right':a,
                    'b_old_right_new_wrong':b,'c_old_wrong_new_right':c,'d_old_wrong_new_wrong':d,
                    'retained_correctness':divide(a,a+b),'degraded_correctness':divide(b,a+b),'repair_rate':divide(c,c+d),'net_repaired':c-b}))
        for after,before in [('Vtarget','Vrand'),('Vrand','V0'),('Vtarget','V0')]:
            for key in KEYS:
                tables['paired_V_differences'].append({'G':g,'contrast':after+'-'+before,'metric':key,
                    'delta':points[g,after][key]-points[g,before][key],**interval(replicas[g,after][key]-replicas[g,before][key])})
    for v in VS:
        for after,before in [('G1','G0'),('G2','G1'),('G2','G0'),('G3','G2'),('G3','G0')]:
            for key in KEYS:
                tables['paired_G_differences'].append({'V':v,'contrast':after+'-'+before,'metric':key,
                    'delta':points[after,v][key]-points[before,v][key],**interval(replicas[after,v][key]-replicas[before,v][key])})
    for r in raw:
        for v in VS:
            s=predictions[v][r['candidate_id']]
            tables['all_decisions'].append({'G':r['G_round'],'V':v,'question_id':r['question_id'],'candidate_id':r['candidate_id'],
                'question_status':r['question_status'],'correct_by_source_reference':r['correct'],'accepted':s['accepted'],'score':s['score'],
                'unknown_reason':r['unknown_reason'],'truncated':r['truncated'],'output_tokens':r['output_tokens']})
    return clean(tables)


def rationale_tables():
    inputs=read_jsonl(RUN/'reason_inputs_C20.jsonl');index={r['candidate_id']:r for r in inputs}
    G_reviews={r['candidate_id']:r for r in read_jsonl(RUN/'initial_review/G_response_initial20.jsonl')}
    review_path=RUN/'initial_review/rationale_reviews_initial60.jsonl'
    reviews=read_jsonl(review_path) if review_path.exists() else []
    reviewed={(r['candidate_id'],r['V']):r for r in reviews};rows=[];metrics=[]
    assert len(reviewed)==len(reviews)
    for v in VS:
        generated=read_jsonl(RUN/'rationales'/(v+'.jsonl'))
        assert len(generated)==len({r['candidate_id'] for r in generated})==20
        for s in generated:
            r=index[s['candidate_id']]
            assert s['prompt_sha256']==digest(reason_messages(r))
            assert s['protocol_sha256']==sha(RUN/'rationales/protocol.json')
            review=reviewed.get((r['candidate_id'],v))
            if review:assert review['rationale_source_sha256']==sha(RUN/'rationales'/(v+'.jsonl'))
            decision=None if s['verdict'] is None else s['verdict']=='A'
            rows.append({**r,'V':v,'V_rationale_full_text':s['rationale_full_text'],'V_rationale_body':s['rationale_body'],
                'reason_protocol_sha256':s['protocol_sha256'],'reason_prompt_sha256':s['prompt_sha256'],
                'reason_model_fingerprint':s['model_fingerprint'],'reason_adapter_fingerprint':s['adapter_fingerprint'],
                'reason_output_tokens':s['output_tokens'],
                'G_process_status_initial':G_reviews[r['candidate_id']]['process_status_initial'],
                'G_process_evidence_initial':G_reviews[r['candidate_id']]['evidence_initial'],
                'reason_protocol_verdict':s['verdict'],'reason_protocol_parse_status':s['parse_status'],
                'reason_protocol_decision_correct':None if decision is None else decision==r['correct'],
                'reason_truncated':s['truncated'],**(review or {'rationale_validity_initial':'pending_assistant_review'}),
                'human_review_status':'pending','human_label':None,'human_comment':None})
        rr=[r for r in rows if r['V']==v];counts=Counter(r['rationale_validity_initial'] for r in rr)
        judged=[r for r in rr if r['reason_protocol_decision_correct'] is not None]
        evaluable=[r for r in rr if r['rationale_validity_initial'] in ['valid','invalid']]
        metrics.append({'V':v,'answers':20,'decision_available':len(judged),'decision_correct':sum(r['reason_protocol_decision_correct'] for r in judged),
            'decision_accuracy_available':sum(r['reason_protocol_decision_correct'] for r in judged)/len(judged) if judged else None,
            **{k:counts[k] for k in ['valid','invalid','unverifiable','not_provided','pending_assistant_review']},
            'rationale_evaluable':len(evaluable),'rationale_validity_rate_evaluable':counts['valid']/len(evaluable) if evaluable else None,
            'decision_right_rationale_invalid':sum(r['reason_protocol_decision_correct'] is True and r['rationale_validity_initial']=='invalid' for r in rr),
            'decision_and_rationale_right':sum(r['reason_protocol_decision_correct'] is True and r['rationale_validity_initial']=='valid' for r in rr),
            'decision_right_rationale_invalid_fraction_all':sum(r['reason_protocol_decision_correct'] is True and r['rationale_validity_initial']=='invalid' for r in rr)/20,
            'decision_and_rationale_right_fraction_all':sum(r['reason_protocol_decision_correct'] is True and r['rationale_validity_initial']=='valid' for r in rr)/20,
            'human_review':'pending; assistant labels not gold','selection':'5 hand-auditable C questions x4 G, not population prevalence'})
    return rows,metrics,len(reviews)==60


def figures(t):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
    points={(r['G'],r['V']):r for r in t['metrics']}
    with PdfPages(REPORT/'week_final_figures.pdf') as pdf:
        fig,axes=plt.subplots(2,2,figsize=(12,7),constrained_layout=True)
        for ax,key in zip(axes.flat,['FPR','FNR','balanced_accuracy','retention_all']):
            matrix=np.array([[points[g,v][key]*100 for g in GS] for v in VS]);im=ax.imshow(matrix,aspect='auto',cmap='YlOrRd',vmin=0,vmax=100)
            for i in range(3):
                for j in range(4):ax.text(j,i,f'{matrix[i,j]:.2f}',ha='center',va='center',color='white' if matrix[i,j]>70 else 'black')
            ax.set_xticks(range(4),GS);ax.set_yticks(range(3),VS);ax.set_title(key+' (%)');fig.colorbar(im,ax=ax,shrink=.7)
        fig.suptitle('Final T299: identical cached answers per G; all G use vLLM, V threshold0.5')
        fig.savefig(REPORT/'T_V_by_G_metrics.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(4,3,figsize=(10,12),constrained_layout=True)
        maxcount=max(max(r[k] for k in ['TP','FN','FP','TN']) for r in t['metrics'])
        for i,g in enumerate(GS):
            for j,v in enumerate(VS):
                r=points[g,v];m=np.array([[r['TP'],r['FN']],[r['FP'],r['TN']]])
                ax=axes[i,j];ax.imshow(m,cmap='Blues',vmin=0,vmax=maxcount)
                for a in range(2):
                    for b in range(2):ax.text(b,a,str(int(m[a,b])),ha='center',va='center',fontsize=13,color='white' if m[a,b]>maxcount*.55 else 'black')
                ax.set_xticks([0,1],['Accept','Reject']);ax.set_yticks([0,1],['Correct','Wrong']);ax.set_title(g+' / '+v+f'; unknown={int(r["unknown"])}')
        fig.suptitle('T299 all12 confusion matrices; unknowns excluded, whole invalid question isolated')
        fig.savefig(REPORT/'T_confusion_matrices.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,2,figsize=(11,4),constrained_layout=True)
        for ax,key in zip(axes,['FPR','FNR']):
            rr=[next(r for r in t['metric_intervals'] if r['G']==g and r['V']=='V0' and r['metric']==key) for g in GS]
            ax.plot(range(4),[100*r['estimate'] for r in rr],'o-',color='#2279A4')
            for i,r in enumerate(rr):ax.plot([i,i],[100*r['low'],100*r['high']],color='#2279A4')
            ax.set_xticks(range(4),GS);ax.set_title('Fixed V0 / '+key);ax.set_ylabel('Percent');ax.grid(alpha=.2)
        fig.suptitle('Uniform-engine G trajectory on held-out T; paired question uncertainty, single training trajectory')
        fig.savefig(REPORT/'T_V0_cross_round_errors.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,3,figsize=(13,4),constrained_layout=True)
        for ax,key in zip(axes,['FPR','FNR','balanced_accuracy']):
            for i,g in enumerate(GS):
                r=next(r for r in t['paired_V_differences'] if r['G']==g and r['contrast']=='Vtarget-Vrand' and r['metric']==key)
                ax.plot([100*r['low'],100*r['high']],[i,i],color='#2279A4',lw=2);ax.plot(100*r['delta'],i,'o',color='#2279A4')
            ax.axvline(0,color='black',ls='--');ax.set_yticks(range(4),GS);ax.invert_yaxis();ax.set_title(key);ax.set_xlabel('Target minus random (percentage points)')
        fig.suptitle('T299 same-answer paired95% intervals; G3 is secondary extension')
        fig.savefig(REPORT/'T_updated_verifier_differences.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(1,3,figsize=(13,4),constrained_layout=True)
        for ax,key,title in zip(axes,['G_accuracy_verified','G_accuracy_lower_all','unknown_fraction'],['Correct among verifiable','Correct / all (lower bound)','Unknown / all']):
            rr=[next(r for r in t['metric_intervals'] if r['G']==g and r['V']=='V0' and r['metric']==key) for g in GS]
            ax.bar(GS,[r['estimate']*100 for r in rr],color='#498870')
            for i,r in enumerate(rr):ax.plot([i,i],[100*r['low'],100*r['high']],color='black')
            ax.set_title(title);ax.set_ylabel('Percent')
        fig.suptitle('Final generator outcome quality; final answer correct does not certify process')
        fig.savefig(REPORT/'T_generator_quality.png',dpi=180);pdf.savefig(fig);plt.close(fig)
        fig,axes=plt.subplots(4,2,figsize=(9,12),constrained_layout=True)
        for i,g in enumerate(GS):
            for j,v in enumerate(['Vrand','Vtarget']):
                r=next(r for r in t['judgment_transitions'] if r['G']==g and r['new_V']==v)
                m=np.array([[r['a_old_right_new_right'],r['b_old_right_new_wrong']],
                            [r['c_old_wrong_new_right'],r['d_old_wrong_new_wrong']]])
                ax=axes[i,j];ax.imshow(m,cmap='Blues',vmin=0,vmax=1196)
                for a in range(2):
                    for b in range(2):ax.text(b,a,str(m[a,b]),ha='center',va='center',fontsize=12)
                ax.set_xticks([0,1],['New right','New wrong']);ax.set_yticks([0,1],['V0 right','V0 wrong'])
                ax.set_title(g+' / V0 to '+v)
        fig.suptitle('Same-answer judgment correctness transitions; reference-unknown excluded')
        fig.savefig(REPORT/'T_judgment_transitions.png',dpi=180);pdf.savefig(fig);plt.close(fig)


def table(head,rows):
    return '\n'.join(['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |']+
                     ['| '+' | '.join(str(x).replace('|',' / ').replace('\n',' ') for x in r)+' |' for r in rows])+'\n'


def pct(x):return 'NA' if x is None else f'{100*x:.2f}%'
def pp(x):return 'NA' if x is None else f'{100*x:+.2f}'


def workbook(path,sheets,raw_filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from .oracle_visuals import EXCEL_CONTROL
    book=Workbook();book.remove(book.active)
    for name,rows in sheets.items():
        ws=book.create_sheet(name[:31]);keys=list(dict.fromkeys(k for r in rows for k in r))
        ws.append(keys or ['empty'])
        for cell in ws[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='DCEAF7')
        for rowno,row in enumerate(rows,2):
            values=[]
            for key in keys:
                value=row.get(key)
                if isinstance(value,(dict,list)):value=json.dumps(value,ensure_ascii=False)
                if isinstance(value,str):
                    value=EXCEL_CONTROL.sub(lambda m:'\\u'+f'{ord(m.group()):04x}',value)
                    if len(value)>32000:
                        value=f'完整原文见{raw_filename}; candidate_id={row.get("candidate_id", "")}; field={key}'
                values.append(value)
            ws.append(values)
            for column,value in enumerate(values,1):
                if isinstance(value,str) and value.startswith('='):ws.cell(rowno,column).data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(45,max(14,len(str(col[0].value))+2))
    book.save(path)


def old_file_audit():
    original=json.loads((RUN/'preserved_existing_files.json').read_text());missing=[];changed=[]
    for relative,expected in original.items():
        p=ROOT/relative
        if not p.exists():missing.append(relative)
        elif sha(p)!=expected:changed.append(relative)
    result={'at':utc(),'original_files':len(original),'missing':missing,'changed':changed,
        'all_original_files_unchanged':not missing and not changed,
        'scope':'Existing runs/reports/configs at start; excludes new run/report, shared append-only ledger and GPU lock',
        'new_protocol_files':['005_week_final_20261026.json','006_T_source_quality_20261026.json'],
        'prior_research_plan_preserved':'runs/week_final_20261026/prior_research_plan.md'}
    write_json(REPORT/'preservation_audit.json',result)
    assert result['all_original_files_unchanged'],result
    return result


def execution_records(cfg):
    rows=[]
    for g in ['G1','G2','G3']:
        p=ROOT/f'runs/formal/train/{g}/report.json' if g!='G3' else ROOT/'runs/g3_extension_20261024/stage_train/train/G2/report.json'
        r=json.loads(p.read_text());parent=GS[GS.index(g)-1]
        assert r['final_adapter_fingerprint']==cfg['checkpoints'][g]['adapter_fingerprint']
        assert r['parent_adapter_fingerprint']==cfg['checkpoints'][parent]['adapter_fingerprint']
        qp=ROOT/f'reports/formal/{g}_SFT_data_quality.json' if g!='G3' else ROOT/'reports/g3_extension_20261024/G3_training_quality.json'
        qq=json.loads(qp.read_text());q=qq['actual_SFT_quality'] if g!='G3' else qq['quality']
        assert q['selected']==r['selected_examples']
        rows.append({'model':g,'parent':parent,'training_pool':'A','examples':r['selected_examples'],
            'steps':r['optimizer_steps'],'epochs':r['epochs'],'input_tokens':r['input_tokens'],'loss_tokens':r['supervised_tokens'],
            'selected_correct_by_original_reference':q['correct'],'selected_wrong_by_original_reference':q['wrong'],
            'selected_unknown':q['unknown'],'selected_verified_error_fraction':q['verified_error_fraction'],
            'seconds_train_save_reload':r['total_seconds_this_invocation'],'truncated_examples':r['training_examples_truncated'],
            'base_frozen':r['base_frozen'],'adapter_fingerprint':r['final_adapter_fingerprint'],'source':str(p.relative_to(ROOT)),
            'note':'actualG3 has private-stage G2 alias; actual parent/target verified by fingerprints' if g=='G3' else ''})
    for v in ['Vrand','Vtarget']:
        p=ROOT/f'runs/verifier_update_20261025/{v}/training_report.json';r=json.loads(p.read_text())
        assert r['adapter_fingerprint']==cfg['checkpoints'][v]['adapter_fingerprint']
        rows.append({'model':v,'parent':'V0','training_pool':'B','examples':r['examples'],'steps':r['steps'],'epochs':r['epochs'],
            'input_tokens':r['input_tokens'],'loss_tokens':r['supervised_tokens'],'decision_tokens':r['decision_label_tokens'],
            'seconds_train_save_reload':r['seconds_this_invocation'],'truncated_examples':r['truncated_training_examples'],
            'base_frozen':r['base_parameters_unchanged'],'adapter_fingerprint':r['adapter_fingerprint'],'source':str(p.relative_to(ROOT))})
    compute=[]
    for group,names,filename in [('generation',GS,'complete.json'),('scoring',VS,'compute.json'),('rationales',VS,None)]:
        for name in names:
            p=RUN/group/name/filename if filename else RUN/group/(name+'_compute.json')
            compute.append({'stage':group,'model':name,**json.loads(p.read_text()),'source':str(p.relative_to(ROOT))})
    ledger=read_jsonl(ROOT/'runs/formal/work_ledger.jsonl');current=[r for r in ledger if r['stage'].startswith('week_final/')]
    write_json(REPORT/'compute_accounting.json',{'new_week_final_charged_seconds':sum(r['seconds'] for r in current),
        'new_week_final_charged_hours':sum(r['seconds'] for r in current)/3600,'ledger_sha256_at_export':sha(ROOT/'runs/formal/work_ledger.jsonl'),
        'interpretation':'Wall time charged while holding exclusive GPU lock, includes loading; excludes CPU analysis and assistant/human review',
        'stage_records':compute,'training_records_are_prior_runs_no_retraining':True})
    return rows,compute


def select_T_cases(raw,predictions):
    cases=[]
    for g in GS:
        used=set()
        for target in ['TP','FN','FP','TN','unknown']:
            candidates=[]
            for r in raw:
                if r['G_round']!=g or r['question_status']!='eligible' or r['question_id'] in used:continue
                y=r['correct'];a=predictions['V0'][r['candidate_id']]['accepted']
                group='unknown' if y is None else ('TP' if a else 'FN') if y else ('FP' if a else 'TN')
                if group==target:candidates.append(r)
            if not candidates:continue
            r=min(candidates,key=lambda r:digest('week-final-descriptive-cases:'+r['candidate_id']));used.add(r['question_id'])
            cases.append({k:value for k,value in {**r,'case_stratum':target,
                'V_scores':{v:predictions[v][r['candidate_id']] for v in VS},
                'case_selection':'post-evaluation descriptive panel: hash-first per G xV0 cell, no repeated question within G; not prevalence sample',
                'mathematical_human_review':'pending; program final-answer label not process gold'}.items() if k!='generated_token_ids'})
    return cases


def narrative(cfg,quality,t,reason_metrics,review_done,preservation):
    point={(r['G'],r['V']):r for r in t['metrics']}
    delta=lambda g,v,key:next(r for r in t['paired_V_differences'] if r['G']==g and r['contrast']==v and r['metric']==key)
    gd=lambda contrast,key:next(r for r in t['paired_G_differences'] if r['V']=='V0' and r['contrast']==contrast and r['metric']==key)
    ci=lambda r:f'[{pp(r["low"])}, {pp(r["high"])}]'
    change=gd('G2-G0','FPR');change_fn=gd('G2-G0','FNR');change_ba=gd('G2-G0','balanced_accuracy');extension=gd('G3-G2','FPR');adapt=delta('G2','Vtarget-Vrand','balanced_accuracy')
    signal=lambda r:'区间未跨0' if r['low'] is not None and r['high'] is not None and (r['low']>0 or r['high']<0) else '区间跨0或触及0'
    parts=['# 七天实验收尾与首次最终测试\n',
      '生成时间：'+utc()+'。主任务为DAPO，原七天计划的核心计算与导出'+('已完成' if review_done else '已完成，V理由助手初审待补')+'；人工复核仍待用户进行。正负结果均保留，没有为获得改善再搜索训练设置。\n',
      f'首次T299的主要结果：固定V0的FPR为G0 {pct(point["G0","V0"]["FPR"])} → G1 {pct(point["G1","V0"]["FPR"])} → G2 {pct(point["G2","V0"]["FPR"])}；G2−G0为{pp(change["delta"])}个百分点，95%区间{ci(change)}（{signal(change)}）。次要G3−G2为{pp(extension["delta"])}个百分点，{ci(extension)}（{signal(extension)}）。Vtarget−Vrand在G2上的平衡准确率差{pp(adapt["delta"])}个百分点，{ci(adapt)}。方向、分母及其他误差须与下表一起解释，不预设改善或退化。\n',
      f'固定V0的FNR为G0 {pct(point["G0","V0"]["FNR"])} → G1 {pct(point["G1","V0"]["FNR"])} → G2 {pct(point["G2","V0"]["FNR"])} → G3 {pct(point["G3","V0"]["FNR"])}；主比较G2−G0为{pp(change_fn["delta"])}个百分点，{ci(change_fn)}（{signal(change_fn)}）。T上FPR/FNR点估计同升，与此前A上FNR下降方向不同；不能把前缀注入的接受偏移直接等同于全部自然跨轮变化。\n',
      f'主轨迹V0平衡准确率从G0的{pct(point["G0","V0"]["balanced_accuracy"])}降至G2的{pct(point["G2","V0"]["balanced_accuracy"])}，变化{pp(change_ba["delta"])}个百分点，95%区间{ci(change_ba)}。这是该轨迹下的可靠性下降信号，不以G能力提高作为前提，也不表明已找到下降机制。\n',
      '## 原计划完成到哪里\n',
      table(['阶段','状态','可复查交付'],[
        ['第1天：环境、模型、初测','完成','[Day1](../day1_report.md)'],
        ['第2天：解析与GSM8K/MATH诊断','完成','[Day2](../day2/day2_report.md)'],
        ['第3阶段：DAPO协议与正式划分','完成','[Day3](../day3/day3_report.md)；冻结003'],
        ['第4—5天：固定V0，G0→G1→G2','完成','原formal模型和B/C缓存；G3为另授权扩展'],
        ['第4—5天附属：V理由20条初审','完成' if review_done else '待补助手初审','同20回答×3V，60份协议输出，未提供理由单列；G原文20条/114段初审'],
        ['第6天：随机/定向V更新','完成，当前配方未成功','[独立V更新对照](../verifier_update_20261025/verifier_update_report.md)'],
        ['第7天：冻结后T统一评价、图表、案例','完成','本报告、12格矩阵、按题成对区间、审查表和归档'],
        ['后续：人工金标准、大规模过程训练、GRPO、多模型网格','未开展，属于后续','不将这些列为本周漏做的核心计算']]),
      '## 最终测试协议与题面质量\n',
      '005在第一次解码T前冻结所有G/V检查点、提示、采样和评分代码。T与A/B/C题目互斥；T的参考答案仅由独立核验器读取，不进入G/V输入。没有在T上训练、选模型、改阈值或挑前缀。G0/G1/G2是原主轨迹，G3单列为后续扩展。所有T生成均为相同vLLM BF16运行时，每题4候选，temperature0.7/top_p0.9/top_k50，总上下文4096、输出min(3840,4096−提示长度)。V评分仍为原Transformers完整A/B条件似然实现，完整窗口8192、阈值0.5、恰0.5接受。\n',
      'T原300题/每G1200回答全部缓存。仅根据题面进行的盲质量复核发现1道缺必需分割图的矩形题：16×20分成10个全等矩形可以是16×2（周长36）或8×4（周长24），仅题文不能唯一确定。依据原坏题规则以006记录，对所有G/V整题排除，主统计为299题/每G1196回答；原300题结果另存archival300_metrics，仅作存档。其他4道文本风险标记经题面检查仍可确定，未按G/V结果排除。'+
      '排除ID：`'+quality['whole_question_quarantine'][0]['question_id']+'`。\n',
      '标签只衡量最终答案与冻结参考的精确等价；正确最终答案不保证过程正确。截断、冲突、提取失败等继续为未知，不纳入FPR/FNR分母，也不冒充负例。源数据其他参考错误仍可能存在，本次未宣称逐题人工金标准。\n',
      'T是DAPO公开训练部分的自定义留出，不是官方隐藏测试；A/B/C/T按题目、规范化哈希、近似词面题族及来源UUID分离。近似去重不保证语义完全互斥，也无法证明基座预训练未见这些公开题。\n',
      '## G的结果质量与固定V0跨轮变化\n',
      table(['G','正确/错误/未知','可核验G正确率','正确/全部下界','未知比例','平均/p95输出token','截断'],[
        [r['G'],f'{int(r["correct"])}/{int(r["wrong"])}/{int(r["unknown"])}',pct(r['G_accuracy_verified']),pct(r['G_accuracy_lower_all']),pct(r['unknown_fraction']),
         f'{r["mean_output_tokens"]:.1f}/{r["p95_output_tokens"]:.1f}',int(r['truncated'])] for r in t['generator_quality']]),
      table(['G','V0 FPR','V0 FNR','平衡准确率','全部接受基线准确率','保留率','AUROC'],[
        [g,pct(point[g,'V0']['FPR']),pct(point[g,'V0']['FNR']),pct(point[g,'V0']['balanced_accuracy']),
         pct(point[g,'V0']['G_accuracy_verified']),pct(point[g,'V0']['retention_all']),f'{point[g,"V0"]["AUROC"]:.4f}'] for g in GS]),
      table(['固定V0比较','指标','差值/百分点','按题成对95%区间/百分点'],[
        [contrast,key,pp(gd(contrast,key)['delta']),ci(gd(contrast,key))]
        for contrast in ['G1-G0','G2-G1','G2-G0','G3-G2','G3-G0'] for key in ['FPR','FNR','balanced_accuracy']]),
      '区间使用同299道题的配对聚类bootstrap10000次，seed20261026；同题4候选共同重采样。区间只反映该训练轨迹下的题目抽样不确定性，不包含训练种子波动；多项区间为点位描述、未控制整张表的多重比较。跨G的错误分母和回答内容本身会变，因此差值描述固定V在不同G输出分布上的表现，不等于V参数退化或已识别的训练因果机制。旧C的Transformers曲线与本次vLLM最终测试不能直接拼接。\n',
      '若某V在本批样本全接受，FPR=100%、FNR=0%的bootstrap区间可能退化为点；这不表示总体风险或未来差值已无不确定性。零分母一律NA，不填0。\n',
      '原G训练沿A池选样：G1/G2分别500/501条、63/63步、414830/419895回答监督token；G3为511条、64步、414590监督token，继续G2 adapter。训练/保存/重载约162.6/166.7/121.8秒；这些是既有更新，本次没有重训。G3私有训练目录里的G2名称为复用脚本别名，本次按父adapter与目标指纹核实为实际G2→G3。旧任务曾因EOS尾缀检查报错，但已有完整权重、保存重载和来源记录不因此重跑。原A训练后续发现两道问题题分别进入G1/G2，保留为轨迹解释限制，未追溯修改模型；详见原[来源与缓存审计](../cache_followup_20261024/cache_analysis_report.md)。\n',
      table(['更新','原实际SFT条数','原参考正确/错误/未知','有标签错误率','回答loss token','优化步'],[
          [r['model'],r['examples'],f'{r["selected_correct_by_original_reference"]}/{r["selected_wrong_by_original_reference"]}/{r["selected_unknown"]}',
           pct(r['selected_verified_error_fraction']),r['loss_tokens'],r['steps']] for r in t['training_lineage'] if r['training_pool']=='A']),
      '实际入训数据与全部V接受池是不同分母；真标签仅在选样后记录质量，不参与G反馈路线回筛。这里展示原训练发生时的来源参考口径，不冒充剔除疑问题后的纯净训练。这条轨迹研究含噪验证反馈下的更新，不能默认G越来越强。\n',
      '统一T引擎消除了本次评价时的运行时断点，但没有回溯改变训练候选来源：原G1/G2使用Transformers候选，后续G3反馈候选改用vLLM。G3因此继续单列为扩展，不将四点当作只改变优化次数的完全受控因果试验。\n',
      '## V更新：对同一份最终测试回答的比较\n',
      table(['G','V','TP/FN/FP/TN','FPR','FNR','BA','总体准确率(已知)','有标签接受集错误率','保留率','AUROC'],[
        [r['G'],r['V'],'/'.join(str(int(r[k])) for k in ['TP','FN','FP','TN']),pct(r['FPR']),pct(r['FNR']),pct(r['balanced_accuracy']),
         pct(r['V_accuracy']),pct(r['accepted_error_FDR']),pct(r['retention_all']),f'{r["AUROC"]:.4f}'] for r in t['metrics']]),
      table(['G','Vtarget−Vrand','差值/百分点','成对95%区间/百分点'],[
        [g,key,pp(delta(g,'Vtarget-Vrand',key)['delta']),ci(delta(g,'Vtarget-Vrand',key))]
        for g in GS for key in ['FPR','FNR','balanced_accuracy']]),
      'Vrand与Vtarget从同一V0独立新建LoRA，只用B自然已知标签训练，各256条、正负各128、1epoch/32步、768监督token（其中决定token256）。来源轮次与长度档匹配、输入token总量差0.527%，但覆盖题数165/149不同。当前固定0.5工作点的结果必须连同FPR、FNR、BA及同分集中一起判断：降低FNR可能来自接近全部接受，不能单独当作改进。分数为未校准的A/B代理，未在T上另选阈值。C上已发现拟合和工作点问题，本次照原冻结配方如实最终评价；失败属于实验结果，不能改写成未做第6天。\n',
      table(['G','新V','旧对→新对 a','旧对→新错 b','旧错→新对 c','旧错→新错 d','保持a/(a+b)','退化b/(a+b)','修复c/(c+d)'],[
        [r['G'],r['new_V'],r['a_old_right_new_right'],r['b_old_right_new_wrong'],r['c_old_wrong_new_right'],r['d_old_wrong_new_wrong'],
         pct(r['retained_correctness']),pct(r['degraded_correctness']),pct(r['repair_rate'])] for r in t['judgment_transitions']]),
      '以上转移严格固定每条G回答及独立结果标签，排除未知；保持/修复与数据保留率是不同分母。所有原始分数、0.5同分数与分位数在Excel及CSV中保留。\n',
      '## V理由与G步骤：小样本助手初审\n',
      '预先从已有C选择5道可精确复算的题×4G共20条原回答，覆盖历史V0的TP/FN/FP/TN。相同20条分别给V0/Vrand/Vtarget，真实生成60份“理由→判定”输出；提示不提供参考答案。该独立理由协议与正式A/B评分不同，公开理由不是已证实的内部决策原因。样本定向选择，比例不可外推为总体发生率。\n',
      table(['V','可解析判定/20','判定正确/可解析','理由有效','无效','无法核验','未提供','有效率(有效+无效分母)','判对但理由无效','判对且理由有效'],[
        [r['V'],r['decision_available'],f'{r["decision_correct"]}/{r["decision_available"]}',r['valid'],r['invalid'],r['unverifiable'],r['not_provided'],
         pct(r['rationale_validity_rate_evaluable']),r['decision_right_rationale_invalid'],r['decision_and_rationale_right']] for r in reason_metrics]),
      '理由是否有效要看事实/计算、是否对应G原文及能否支持“最终答案正确/错误”；只找到过程瑕疵不足以拒绝正确最终答案。空泛肯定不算有效证明，只有A/B没有理由记“未提供”。V理由初审'+('已逐条完成' if review_done else '尚未全部完成')+'；人工作为独立列全部待复核，未用这些初审标签训练。G先于新V理由独立初审20条原文/114个原段落，记录首个已见且未撤回的实质错误、原字符位置和证据；首错之后记未审，不能自动标全部后续错误。11条见实质错误、8条已查未见错误、1条最小性完整性待定，均非人工金标准。\n',
      '具体观察：V0在该定向小样本的新理由协议上20/20判定正确，初审理由17有效、3无效（其中1条对G原过程过度概括、2条仅核对根式转换而缺关键原方程检验）。Vrand与Vtarget各20条都只生成VERDICT: B，没有实质理由，判定各9/20正确。后两者在理由提示下全部拒绝，与正式A/B条件似然评分的高接受率不是同一输入协议，不能相互替换或混成一条指标；也不能把V0在选定20条上的表现推广成总体100%正确。\n',
      '## 可解释的结论边界与下一步\n',
      '本周已建立可复查的固定V跨轮轨迹、随机/定向V更新对照，以及首次独立T评价。模型更新不等于能力提高；V更新也不保证误接受改善。T表可以判断此前A/C现象在未参与开发的题目上是否重现，但一条轨迹和有限题数不足以证明通用退化规律。原Oracle/随机G对照、前缀注入/删除与错误类型分析继续原样保留；本次没有再搜索触发词或修改训练，以免扩展支线拖延七天交付。\n',
      '争议理由和G步骤已留作用户复核；后续可继续做自然错误组成、分组内风险与反馈质量的缓存分析。初审标签用于训练或正式效果声明前仍须相应复核。若要修复V训练或增加种子，建立新的开发协议并使用新的确认数据；本次T已打开，不能继续称封存测试或用于反复调参后报告无偏成绩。GRPO、大规模PRM和多模型网格仍留后续。\n',
      '## 文件与保存\n',
      '- [最终指标工作簿](week_final_metrics.xlsx)：T主表、区间、转移、G质量、原300存档、理由汇总及逐条判定。\n'+
      '- [三表审查工作簿](week_review.xlsx)：回答与V理由、G原段落步骤、标签规范与来源；完整原文另存JSONL，超长Excel单元格仅给ID和原文路径。\n'+
      '- [六页图表](week_final_figures.pdf)：V×G热图、12格混淆矩阵、V0曲线、更新差值、G质量和判断正确性转移。\n'+
      '- [全部T原文](T_candidates_full.jsonl)、[理由原文与初审](reason_reviews_full.jsonl)、[逐段审查](G_steps_initial.jsonl)。\n'+
      '- [T分层待审案例](T_review_examples.jsonl)：按G与V0混淆类型/未知作确定性hash抽样，原文及三个V分数齐全，人工数学复核待做；不将案例集比例当总体结果。\n'+
      '- [原产物保存审计](preservation_audit.json)：原'+str(preservation['original_files'])+'个runs/reports/configs文件逐一哈希一致，原数据/模型未覆盖；共享算力账本仅追加。\n'+
      '- [原始缓存与检查点清单](raw_cache_manifest.json)、005/006协议、计算耗时记录和SHA256归档清单随报告保存。\n']
    return '\n'.join(parts)


def main():
    REPORT.mkdir(exist_ok=True);cfg,quality,raw,predictions=load_final()
    t=compute(cfg,quality,raw,predictions);rr,rm,done=rationale_tables();t['rationale_metrics']=rm
    t['training_lineage'],t['week_final_compute']=execution_records(cfg)
    t['T_review_examples']=select_T_cases(raw,predictions)
    write_jsonl(REPORT/'T_review_examples.jsonl',t['T_review_examples'])
    write_json(REPORT/'week_final_results.json',{'at':utc(),'protocol':cfg,'source_quality':quality,'tables':t})
    for key,rows in t.items():write_csv(REPORT/(key+'.csv'),rows)
    full=[]
    for r in raw:
        full.append({k:value for k,value in {**r,'V_scores':{v:predictions[v][r['candidate_id']] for v in VS}}.items() if k!='generated_token_ids'})
    write_jsonl(REPORT/'T_candidates_full.jsonl',full);write_jsonl(REPORT/'reason_reviews_full.jsonl',rr)
    steps=read_jsonl(RUN/'initial_review/G_steps_initial.jsonl')
    enriched=[];source={r['candidate_id']:r for r in read_jsonl(RUN/'reason_inputs_C20.jsonl')}
    G_review={r['candidate_id']:r for r in read_jsonl(RUN/'initial_review/G_response_initial20.jsonl')}
    for r in steps:
        s=source[r['candidate_id']];enriched.append({**r,'question':s['question'],'G_full_text':s['text'],
            'preceding_original_text':s['text'][:r['character_start']],
            'response_process_status_initial':G_review[r['candidate_id']]['process_status_initial'],
            'response_evidence_initial':G_review[r['candidate_id']]['evidence_initial']})
    write_jsonl(REPORT/'G_steps_initial.jsonl',enriched)
    norm=[{'item':'rationale_protocol','definition':json.loads((RUN/'rationales/protocol.json').read_text()),'source':'independent public reason prompt; not A/B hidden rationale'},
          {'item':'step_review','definition':json.loads((RUN/'initial_review/annotation_protocol.json').read_text()),'source':'assistant initial only; human pending'},
          {'item':'sample_selection','definition':'5 fixed C questions xG0/G1/G2/G3 candidate0; historical TP4/FN7/FP8/TN1; deliberately auditable, not prevalence sample'},
          {'item':'rationale_valid','definition':'factually and mathematically sound, corresponds to G text and sufficiently supports final-answer judgment'},
          {'item':'rationale_invalid','definition':'false math/facts, text mismatch, unsupported final-answer verdict, or process-only rejection of correct final answer'},
          {'item':'rationale_unverifiable','definition':'incomplete/ambiguous text prevents evaluation; not forced into invalid'},
          {'item':'rationale_not_provided','definition':'bare label or no substantive reason; separate from invalid'},
          {'item':'G_step_first_error','definition':'first witnessed unretracted substantive flaw; earlier steps not certified gold; subsequent steps left not reviewed'},
          {'item':'human_review','definition':'blank human fields explicitly pending; assistant initial annotations not independent human gold'},
          {'item':'time_and_sources','definition':{'at':utc(),'source_SHA256':sha(RUN/'reason_inputs_C20.jsonl'),'review_code_SHA256':sha(ROOT/'src/pilot/week_initial_review.py')}}]
    write_jsonl(REPORT/'annotation_definitions.jsonl',norm)
    workbook(REPORT/'week_final_metrics.xlsx',t,'T_candidates_full.jsonl')
    workbook(REPORT/'week_review.xlsx',{'回答与V理由审查':rr,'G步骤审查':enriched,'标签规范与来源':norm},'reason_reviews_full.jsonl / G_steps_initial.jsonl')
    figures(t);preservation=old_file_audit()
    (REPORT/'week_final_report.md').write_text(narrative(cfg,quality,t,rm,done,preservation),encoding='utf-8')
    write_json(REPORT/'week_completion_status.json',{'at':utc(),
        'status':'week_core_complete_human_review_pending' if done else 'week_T_complete_assistant_review_pending',
        'original_seven_day_core_computations_complete':done,'T_is_sealed':False,'T_opened_for_final_only':True,
        'remaining_core_GPU_tasks':[],'new_G_generation':4800,'new_V_T_scores':14400,
        'V_reason_generations':60,'assistant_reason_reviews_complete':done,'assistant_G_answers_reviewed':20,
        'human_gold_review_complete':False,'remaining_human_work':['Review disputed V reasons and initial G process labels; independently spot-check accepted initial labels'],
        'V_training_outcome':'completed experiment; no successful improved verifier established by current recipe',
        'original_files_unchanged':preservation['original_files'],'T_primary_questions':299,'T_archival_questions':300,
        'future_not_week_core':['Large process annotation/PRM','GRPO','Multiple training seeds and G/V model grids'],
        'future_data_rule':'T cannot become a tuning/training set; use a fresh confirmation set after further changes',
        'report':'week_final_report.md','results':'week_final_results.json'})
    copies=[PLAN,QUALITY,RUN/'T_first_unseal.json',RUN/'T_blind_source_flags.jsonl',RUN/'initial_review/exact_math_checks.json',
            RUN/'initial_review/annotation_protocol.json',RUN/'initial_review/rationale_review_protocol.json',RUN/'initial_review/G_response_initial20.jsonl',ROOT/'configs/verifier_update_20261025.json',
            ROOT/'tests/test_week_final.py']
    copies += [ROOT/'src/pilot'/f for f in ['week_final_common.py','week_final_generate.py','week_final_run.py','week_final_reasons.py','week_final_report.py','week_initial_review.py','week_reason_initial_review.py','week_final_verify.py']]
    for p in copies:shutil.copyfile(p,REPORT/p.name)
    manifest=[]
    for p in sorted(RUN.rglob('*')):
        if p.is_file() and not p.is_symlink() and p.name!='state.json' and '/generation/' not in str(p):
            manifest.append({'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':sha(p)})
    write_json(REPORT/'raw_cache_manifest.json',{'files':manifest,'checkpoint_identities':cfg['checkpoints'],
        'original_raw_generation_contains_all_output_token_ids':True,'preserved_originals':preservation})
    package_delivery(done)
    state('week_core_complete_human_review_pending' if done else 'week_T_complete_assistant_review_pending',
          final_report=str(REPORT/'week_final_report.md'),T_primary_questions=299,old_data_preserved=True,human_review='pending')
    print('Week final delivery:',REPORT,'; assistant rationale review complete:',done,flush=True)


def package_delivery(done=True):
    """Repackage a completed delivery after dated research-plan status is appended."""
    archive=REPORT/'week_final_delivery.zip';checks=REPORT/'delivery_SHA256.json'
    files=[p for p in sorted(REPORT.rglob('*')) if p.is_file() and p not in [archive,checks] and p.name not in ['delivery_verification.json','independent_verification.json']]
    write_json(checks,{str(p.relative_to(REPORT)):sha(p) for p in files})
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in [*files,checks]:z.write(p,p.relative_to(REPORT))
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    write_json(REPORT/'delivery_verification.json',{'at':utc(),'archive_sha256':sha(archive),'archive_crc':'pass','listed_files':len(files),
        'original_files_unchanged':True,'assistant_reason_initial_review_complete':done,'human_review_complete':False,
        'T_primary_questions':299,'T_archival_questions':300,'G_candidates_each_archival':1200,'V_scores_each':4800})


if __name__=='__main__':
    import argparse
    argparse.ArgumentParser(description=__doc__).parse_args()
    main()
