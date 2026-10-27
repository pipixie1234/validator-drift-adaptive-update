"""Export review tables, a Chinese report and vector/raster scientific figures."""
import html
import json
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from .v0_drift_analysis import ROUNDS, SPLITS, COLORS, LENGTHS, PRIMARY, LABELS, clean


def pct(x, digits=2): return 'NA' if x is None else f'{100*x:.{digits}f}%'
def num(x, digits=3): return 'NA' if x is None else f'{x:.{digits}f}'
def ci_text(estimate, ci, percent=True):
    f = pct if percent else num
    return f'{f(estimate)} [{f(ci["low"])}, {f(ci["high"])}]'
def markdown_table(headers, rows):
    return '\n'.join(['| '+' | '.join(headers)+' |', '| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str, row))+' |' for row in rows])


def workbook(path, tables, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in tables.items():
        ws = wb.create_sheet(name[:31])
        if not rows: continue
        headers = list(rows[0]); ws.append(headers)
        for r in rows: ws.append([clean(r.get(k)) for k in headers])
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='225577')
        for column in ws.columns:
            name_col = str(column[0].value)
            width = 60 if name_col in ['question', 'text', 'response_preview'] else 50 if 'question_id' == name_col else 24
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # model text must never become an Excel formula
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                elif isinstance(cell.value, float):
                    header = str(ws.cell(1, cell.column).value)
                    counts = {'TP', 'FN', 'FP', 'TN', 'known', 'unknown', 'correct', 'wrong', 'total', 'accepted_total', 'accepted_unknown', 'truncated', 'questions', 'error_questions', 'false_accept_questions', 'false_reject_questions', 'valid_replicates', 'common_fully_verifiable_questions'}
                    cell.number_format = '0' if header in counts else '0.0000'
    ws = wb.create_sheet('readme_中文')
    for row in [
        ['用途', '固定V0面对G0/G1/G2输出的开发诊断；B200题、C96题，每题4候选；T封存。'],
        ['正类', '独立核验最终答案正确；V接受为预测正。'],
        ['FNR', 'FN/(TP+FN)：正确回答中被拒绝的比例。'],
        ['FPR', 'FP/(FP+TN)：错误回答中被接受的比例。'],
        ['FDR', 'FP/(TP+FP)：有标签接受集内的错误比例，不能与FPR混淆。'],
        ['置信区间', '2000次按题聚类bootstrap，同题4候选一起重采样，跨轮使用相同题目索引；点位95%区间，未校正多重比较。'],
        ['未知', '未知不计入TP/FN/FP/TN，但保留在未知比例、全候选保留率和风险上下界中。'],
        ['delta单位', 'paired_deltas中的delta为比例差；乘100得到百分点。CI跨0不意味着证明等价或无漂移。'],
        ['分数', 'A/B条件似然归一化代理值；Brier/ECE为辅助分数诊断，不将其视为已校准置信度。'],
        ['候选文本', 'candidates_brief只给前1200字符，完整原文见candidates_full.jsonl；cases给完整抽查案例。'],
        ['case选择', '每个B/C和G轮次按ID排序各抽TP/FN/FP/TN/未知最多2条；描述性审查，非随机总体错误率估计。'],
        ['过程标签', '仅核验最终答案；未对V理由或G过程给出新标签。']]: ws.append(row)
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 110
    for row in ws:
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)


def build_outputs(destination, result, data, curves, risks, summaries, long, paired, lengths, per_question, sensitivity):
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'axes.spines.top': False,
                         'axes.spines.right': False, 'savefig.dpi': 220, 'pdf.fonttype': 42, 'svg.fonttype': 'none'})
    metrics = result['metrics']; generated = []; pdf = PdfPages(destination/'v0_drift_figures.pdf')
    def m(r, s): return metrics[f'{r}_{s}']['point']
    def emit(fig, name, title, note=''):
        if note: fig.text(.5, .014, note, ha='center', va='bottom', fontsize=8, color='#444444')
        colorbar = getattr(fig, '_drift_colorbar', None)
        fig.tight_layout(rect=[0, .045, .88 if colorbar else 1, .95]); fig.suptitle(title, fontsize=14, y=.99)
        if colorbar:
            cax = fig.add_axes([.91, .17, .015, .64])
            fig.colorbar(colorbar[0], cax=cax, label=colorbar[1])
        for extension in ['png', 'pdf', 'svg']: fig.savefig(destination/'figures'/f'{name}.{extension}', bbox_inches='tight')
        pdf.savefig(fig, bbox_inches='tight'); generated.append((name, title)); plt.close(fig)

    fig, ax = plt.subplots(figsize=(11, 7)); ax.axis('off')
    ax.text(.02, .95, 'Fixed V0 across generator updates', fontsize=22, fontweight='bold', transform=ax.transAxes)
    ax.text(.02, .86, 'DAPO internal splits: B = 200 questions; C = 96 eligible questions\n4 candidates/question/round. G0 -> G1_SFT -> G2_SFT. Threshold = 0.5.', fontsize=12, transform=ax.transAxes)
    table = [['Split / G', 'FNR', 'FPR', 'Balanced acc.', 'ROC AUC', 'Unknown']]
    for s in SPLITS:
        for r in ROUNDS:
            p = m(r, s); table.append([f'{s} / {r}', pct(p['FNR']), pct(p['FPR']), pct(p['balanced_accuracy']), num(p['AUROC']), f'{int(p["unknown"])}/{int(p["total"])}'])
    ax.table(cellText=table[1:], colLabels=table[0], bbox=[.01, .30, .98, .44], cellLoc='center')
    ax.text(.02, .20, 'Primary drift: paired changes in FNR/FPR relative to G0.\n95% question-cluster percentile bootstrap; 2,000 replicates.\nB/C diagnostics only. T remains sealed. No new inference or training.', fontsize=11, transform=ax.transAxes)
    pdf.savefig(fig); plt.close(fig)

    for normalized in [False, True]:
        fig, axes = plt.subplots(2, 3, figsize=(12, 7))
        vmax = 1 if normalized else max(max(m(r, s)[k] for k in ['TP', 'FN', 'FP', 'TN']) for s in SPLITS for r in ROUNDS)
        for i, s in enumerate(SPLITS):
            for j, r in enumerate(ROUNDS):
                p = m(r, s); counts = np.array([[p['TP'], p['FN']], [p['FP'], p['TN']]])
                values = counts/counts.sum(1, keepdims=True) if normalized else counts
                ax = axes[i, j]; image = ax.imshow(values, cmap='Blues', vmin=0, vmax=vmax)
                for (y, x), v in np.ndenumerate(values):
                    label = [['TP', 'FN'], ['FP', 'TN']][y][x]
                    text = f'{label} {100*v:.1f}%\n{int(counts[y,x])}/{int(counts[y].sum())}' if normalized else f'{label}\n{int(v)}'
                    ax.text(x, y, text, ha='center', va='center', fontsize=11, color='white' if v > .55*vmax else '#123044')
                ax.set_xticks([0, 1], ['Accept', 'Reject']); ax.set_yticks([0, 1], ['Correct (Y=1)', 'Wrong (Y=0)'] if j == 0 else ['', ''])
                ax.set_title(f'{s} / {r}\nlabelled n={int(p["known"])}, unknown={int(p["unknown"])}')
        fig._drift_colorbar = (image, 'Conditional proportion' if normalized else 'Candidates')
        emit(fig, '02_confusion_row_normalized' if normalized else '01_confusion_counts',
             'V0 confusion matrices: row-normalized rates' if normalized else 'V0 confusion matrices: counts',
             'Positive = independently correct final answer. Unknown labels excluded. FN/top-right = false rejection; FP/bottom-left = false acceptance.')

    fig, axes = plt.subplots(2, 2, figsize=(10.5, 7))
    for i, s in enumerate(SPLITS):
        for j, field in enumerate(['FNR', 'FPR']):
            ax = axes[i, j]; values = [m(r, s)[field] for r in ROUNDS]
            lo = [metrics[f'{r}_{s}']['ci95'][field]['low'] for r in ROUNDS]; hi = [metrics[f'{r}_{s}']['ci95'][field]['high'] for r in ROUNDS]
            ax.errorbar(range(3), values, yerr=[np.array(values)-lo, np.array(hi)-values], marker='o', color='#244B65', capsize=5)
            for x, y in enumerate(values): ax.annotate(pct(y), (x, y), xytext=(0, 12), textcoords='offset points', ha='center', fontsize=9)
            ax.set_xticks(range(3), ROUNDS); ax.set_ylim(0, 1); ax.set_xlim(-.3, 2.3); ax.set_title(f'{s}: {LABELS[field]}'); ax.set_ylabel('Rate'); ax.grid(alpha=.18)
    emit(fig, '03_fnr_fpr_cluster_ci', 'Fixed V0 error rates across G rounds', 'Bars: 95% question-cluster bootstrap intervals; 2,000 replicates. Same 0-1 y scale; V0 and threshold are fixed.')

    fig, axes = plt.subplots(2, 2, figsize=(11, 7))
    primary_deltas = [d for d in paired if d['metric'] in ['FNR', 'FPR']]
    span = max(abs(d[k]) for d in primary_deltas for k in ['low', 'high'])*100+2
    for i, s in enumerate(SPLITS):
        for j, metric in enumerate(['FNR', 'FPR']):
            ax = axes[i, j]; ds = [d for d in paired if d['split'] == s and d['metric'] == metric]
            for y, d in enumerate(ds):
                x, lo, hi = [d[k]*100 for k in ['delta', 'low', 'high']]
                ax.errorbar(x, y, xerr=[[x-lo], [hi-x]], color=[COLORS[1], COLORS[2], '#666666'][y], marker='o', capsize=4)
                ax.text(-span+.2, y+.25, f'{x:+.2f} [{lo:+.2f}, {hi:+.2f}] pp', fontsize=8)
            ax.axvline(0, color='black', lw=.8, ls='--'); ax.set_xlim(-span, span); ax.set_ylim(2.7, -.5)
            ax.set_yticks(range(3), [d['comparison'] for d in ds]); ax.set_xlabel('Paired change (percentage points)'); ax.set_title(f'{s}: delta {metric}'); ax.grid(axis='x', alpha=.2)
    emit(fig, '04_paired_drift_forest', 'Paired changes in V0 error rates', 'Same question clusters resampled across rounds. Pointwise 95% intervals; exploratory, unadjusted. Crossing zero does not establish equivalence.')

    for kind in ['ROC', 'PR']:
        fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))
        for ax, s in zip(axes, SPLITS):
            for r, color in zip(ROUNDS, COLORS):
                c = curves[f'{r}_{s}']; p = m(r, s)
                if kind == 'ROC':
                    x, y = np.r_[0, c['FPR']], np.r_[0, c['TPR']]
                    ci = metrics[f'{r}_{s}']['ci95']['AUROC']
                    ax.plot(x, y, color=color, label=f'{r}: AUC={p["AUROC"]:.3f} [{ci["low"]:.3f},{ci["high"]:.3f}]')
                    ax.scatter([p['FPR']], [p['TPR']], color=color, s=30, zorder=4)
                else:
                    ax.step(np.r_[0, c['recall']], np.r_[1, c['precision']], where='post', color=color, label=f'{r}: AP={p["average_precision"]:.3f}')
                    ax.axhline(p['G_accuracy_verified'], color=color, lw=.8, ls=':', alpha=.6)
                    ax.scatter([p['TPR']], [p['precision_PPV']], color=color, s=30, zorder=4)
            if kind == 'ROC': ax.plot([0, 1], [0, 1], '--', color='grey', lw=.8)
            ax.set_xlabel('False acceptance rate' if kind == 'ROC' else 'Recall of correct answers')
            ax.set_ylabel('Recall of correct answers' if kind == 'ROC' else 'Accepted precision')
            ax.set_title(s); ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.legend(fontsize=8, loc='lower right' if kind == 'ROC' else 'lower left'); ax.grid(alpha=.18)
        emit(fig, '05_roc_curves' if kind == 'ROC' else '06_precision_recall_curves',
             'V0 ranking: ROC curves' if kind == 'ROC' else 'V0 ranking: precision-recall curves',
             'Verifiable subset only. Dots: frozen threshold 0.5. PR dotted lines: per-round correct-answer prevalence. AP is noninterpolated average precision.')

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.6))
    for ax, s in zip(axes, SPLITS):
        for r, color in zip(ROUNDS, COLORS):
            rows = risks[f'{r}_{s}']; x = [d['retention_all'] for d in rows]
            y = [np.nan if d['risk_verified'] is None else d['risk_verified'] for d in rows]
            low = [np.nan if d['risk_lower_all'] is None else d['risk_lower_all'] for d in rows]; high = [np.nan if d['risk_upper_all'] is None else d['risk_upper_all'] for d in rows]
            ax.plot(x, y, color=color, label=r); ax.fill_between(x, low, high, color=color, alpha=.09)
            ax.scatter(m(r, s)['retention_all'], m(r, s)['accepted_error_FDR'], color=color, s=40, zorder=4)
        ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.set_title(s); ax.set_xlabel('Retained / all candidates'); ax.set_ylabel('Error fraction among accepted'); ax.legend(); ax.grid(alpha=.18)
    emit(fig, '07_risk_coverage', 'Risk versus retention: descriptive threshold sweep', 'Lines: error rate in labelled accepted subset. Bands: unknown-label identification bounds over ALL accepted. Dots: unchanged threshold 0.5.')

    fig, axes = plt.subplots(2, 3, figsize=(12, 6.5), sharex=True, sharey=True)
    for i, s in enumerate(SPLITS):
        for j, r in enumerate(ROUNDS):
            ax = axes[i, j]; rows = data[f'{r}_{s}']
            for truth, label, color in [(True, 'Correct', '#1976A3'), (False, 'Wrong', '#BB4949')]:
                scores = [x['score'] for x in rows if x['correct'] is truth]
                ax.hist(scores, bins=np.linspace(0, 1, 21), weights=np.ones(len(scores))/len(scores), histtype='step', linewidth=1.6, color=color, label=f'{label} n={len(scores)}')
            ax.axvline(.5, color='black', ls='--', lw=.8); ax.set_title(f'{s} / {r}'); ax.legend(fontsize=8); ax.set_xlabel('V0 A/B score'); ax.set_ylabel('Within-class probability mass')
    emit(fig, '08_score_distributions', 'V0 score distributions by independent answer correctness', 'Each class is normalized separately. Score is a conditional-label proxy, not an established calibrated probability.')

    calibration = []; fig, axes = plt.subplots(1, 2, figsize=(11, 4.6))
    for ax, s in zip(axes, SPLITS):
        for r, color in zip(ROUNDS, COLORS):
            rows = [x for x in data[f'{r}_{s}'] if x['correct'] is not None]; xvals = []; yvals = []; counts = []
            for b in range(10):
                subset = [x for x in rows if min(int(x['score']*10), 9) == b]
                if not subset: continue
                x = float(np.mean([z['score'] for z in subset])); y = float(np.mean([z['correct'] for z in subset]))
                xvals.append(x); yvals.append(y); counts.append(len(subset))
                calibration.append({'split': s, 'G_round': r, 'bin_lower': b/10, 'bin_upper': (b+1)/10, 'n': len(subset), 'mean_score': x, 'observed_correct_fraction': y})
            p = m(r, s); dense = np.array(counts) >= 10
            ax.plot(np.where(dense, xvals, np.nan), np.where(dense, yvals, np.nan), marker='o', color=color, label=f'{r}: ECE={p["ECE10_proxy"]:.3f}, Brier={p["Brier_proxy"]:.3f}')
            if (~dense).any(): ax.scatter(np.array(xvals)[~dense], np.array(yvals)[~dense], s=16, facecolors='none', edgecolors=color, alpha=.5)
        ax.plot([0, 1], [0, 1], '--', color='grey'); ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.set_title(s)
        ax.set_xlabel('Mean positive-class proxy score'); ax.set_ylabel('Observed correct fraction'); ax.legend(fontsize=8); ax.grid(alpha=.18)
    emit(fig, '09_proxy_score_reliability', 'Auxiliary reliability diagram of the A/B proxy score', 'Filled points: bin n>=10; small hollow points: n<10. All bins enter ECE. No fitting/recalibration. Bin counts in workbook.')

    fig, axes = plt.subplots(2, 3, figsize=(12, 7))
    for i, s in enumerate(SPLITS):
        ax = axes[i, 0]; ps = [m(r, s) for r in ROUNDS]
        ax.plot(range(3), [p['G_accuracy_verified'] for p in ps], 'o-', color=COLORS[0], label='Verifiable-subset accuracy')
        ax.fill_between(range(3), [p['G_accuracy_lower_all'] for p in ps], [p['G_accuracy_upper_all'] for p in ps], alpha=.2, color=COLORS[0], label='All-output identification bounds')
        ax.set_ylim(0, 1); ax.set_title(f'{s}: G correctness'); ax.legend(fontsize=7)
        ax = axes[i, 1]
        ax.plot(range(3), [p['unknown_fraction'] for p in ps], 'o-', color='#B14A55', label='Unknown / all')
        ax.plot(range(3), [p['truncated_fraction'] for p in ps], 's--', color='#555555', label='Truncated / all')
        ax.set_ylim(0, .15); ax.set_title(f'{s}: coverage exclusions'); ax.legend(fontsize=8)
        ax = axes[i, 2]
        for k, label, color in [('V_accuracy', 'V accuracy', COLORS[0]), ('balanced_accuracy', 'V balanced accuracy', COLORS[2]), ('all_accept_baseline_accuracy', 'Always accept baseline', '#777777')]:
            ax.plot(range(3), [p[k] for p in ps], 'o-', label=label, color=color)
        ax.set_ylim(0, 1); ax.set_title(f'{s}: V and baseline'); ax.legend(fontsize=8)
        for ax in axes[i]: ax.set_xticks(range(3), ROUNDS); ax.grid(alpha=.15)
    emit(fig, '10_coverage_and_baselines', 'Coverage, generator correctness and verifier baselines', 'Correctness shading: unknown-label identification bounds, NOT confidence intervals. Unknown labels remain unknown; truncation is a subset of unknown.')

    fig, axes = plt.subplots(2, 2, figsize=(11, 7))
    for i, s in enumerate(SPLITS):
        for j, metric in enumerate(['FNR', 'FPR']):
            ax = axes[i, j]
            for ri, (r, color) in enumerate(zip(ROUNDS, COLORS)):
                rs = [x for x in lengths if x['split'] == s and x['G_round'] == r]
                ys = [np.nan if x[metric] is None else x[metric] for x in rs]; xs = np.arange(4)+(ri-1)*.07
                ax.plot(xs, ys, 'o-', color=color, label=r)
                for x, y, entry in zip(xs, ys, rs):
                    if np.isfinite(y): ax.annotate(str(int(entry['correct' if metric == 'FNR' else 'wrong'])), (x, y), xytext=(0, 5+ri*6), textcoords='offset points', ha='center', fontsize=6, color=color)
            ax.set_xticks(range(4), [x[0] for x in LENGTHS]); ax.set_ylim(0, 1); ax.set_xlabel('Output tokens'); ax.set_title(f'{s}: {metric} by output length'); ax.legend(fontsize=8); ax.grid(alpha=.18)
    emit(fig, '11_length_stratified_errors', 'Exploratory length-stratified V0 errors', 'Small numbers above points = relevant true-class support. Missing denominator = NA. Strata can be sparse; length differences are not causal mechanisms.')

    fig, axes = plt.subplots(2, 2, figsize=(10, 9))
    cmap = plt.get_cmap('viridis').copy(); cmap.set_bad('#DDDDDD')
    for i, s in enumerate(SPLITS):
        qs = sorted({r['question_id'] for r in per_question if r['split'] == s})
        lookup = {(r['G_round'], r['question_id']): r for r in per_question if r['split'] == s}
        for j, metric in enumerate(['FNR', 'FPR']):
            values = np.array([[np.nan if lookup[(r, q)][metric] is None else lookup[(r, q)][metric] for r in ROUNDS] for q in qs])
            ax = axes[i, j]; im = ax.imshow(np.ma.masked_invalid(values), vmin=0, vmax=1, cmap=cmap, aspect='auto', interpolation='nearest')
            ax.set_xticks(range(3), ROUNDS); ticks = np.arange(0, len(qs), 25); ax.set_yticks(ticks, ticks+1)
            ax.set_ylabel('Question index, sorted ID'); ax.set_title(f'{s}: per-question {metric}')
    fig._drift_colorbar = (im, 'Conditional error rate')
    emit(fig, '12_per_question_error_heatmap', 'Question-level diagnostics across G rounds', 'Shared 0-1 color scale. Grey = no relevant true-class examples. At most 4 candidates/question: local proportions are noisy. IDs in per_question_metrics.csv.')

    fig, axes = plt.subplots(2, 2, figsize=(10, 7))
    vmax = max(max(t[k] for k in ['a', 'b', 'c', 'd']) for ss in result['transitions'].values() for t in ss.values())
    for i, s in enumerate(SPLITS):
        for j, r in enumerate(['G1', 'G2']):
            t = result['transitions'][s][r]; values = np.array([[t['a'], t['b']], [t['c'], t['d']]])
            ax = axes[i, j]; ax.imshow(values, vmin=0, vmax=vmax, cmap='Blues')
            for (y, x), v in np.ndenumerate(values): ax.text(x, y, f'{[["a", "b"], ["c", "d"]][y][x]} = {v}', ha='center', va='center', color='white' if v > vmax*.55 else '#173E52')
            ax.set_xticks([0, 1], ['V correct on new answer', 'V wrong on new answer']); ax.set_yticks([0, 1], ['V correct on G0 answer', 'V wrong on G0 answer'])
            ax.tick_params(labelsize=8); ax.set_title(f'{s}: G0 -> {r}; excluded unknown pairs={t["unknown_pairs"]}')
    emit(fig, '13_descriptive_judgement_transitions', 'Descriptive paired-question judgement transitions', 'Candidate index 0 only. Different G answers and truth can occur. These are NOT fixed-answer verifier-forgetting matrices.')

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))
    for ax, s in zip(axes, SPLITS):
        for r, color in zip(ROUNDS, COLORS):
            x = np.sort([z['output_tokens'] for z in data[f'{r}_{s}']]); y = np.arange(1, len(x)+1)/len(x)
            ax.plot(x, y, color=color, label=f'{r}: mean={m(r,s)["mean_output_tokens"]:.0f}, p95={m(r,s)["p95_output_tokens"]:.0f}')
        ax.set_title(s); ax.set_xlim(0, 4096); ax.set_ylim(0, 1); ax.set_xlabel('Output tokens'); ax.set_ylabel('Empirical cumulative fraction'); ax.legend(fontsize=8); ax.grid(alpha=.18)
    emit(fig, '14_output_length_cdf', 'Output length distributions, including unknown outputs', 'All candidates included. Shared native-context budget across rounds; no selective length reruns mixed into evaluation.')
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.2))
    for ax, s in zip(axes, SPLITS):
        values = np.array([[m(r, s)[metric] for r in ROUNDS] for metric in ['FNR', 'FPR', 'retention_all']])
        im = ax.imshow(values, vmin=0, vmax=1, cmap='viridis', aspect='auto')
        for (y, x), v in np.ndenumerate(values): ax.text(x, y, f'{v*100:.2f}%', ha='center', va='center', color='white' if v < .55 else 'black')
        ax.set_xticks(range(3), ROUNDS); ax.set_yticks(range(3), ['False rejection', 'False acceptance', 'Retention (all)']); ax.set_title(f'{s}: fixed V0')
    fig._drift_colorbar = (im, 'Rate, common scale')
    emit(fig, '15_v0_by_generator_heatmap', 'Fixed V0 x G rounds: error rates and retention', 'No Vrand or Vtarget checkpoints evaluated. All panels share the same 0-1 color scale; unknown excluded only from FNR/FPR.')
    pdf.close()

    # Deterministic representative cases include all four confusion cells and unknowns.
    cases = []; brief = []; case_lines = ['# 可复查案例（确定性抽样）', '', '每个轮次/题池按候选ID排序，每类最多2条。用于人工复查，不据此估计总体类型占比。V理由与G步骤未在本次新标注。', '']
    for key, rows in data.items():
        for z in rows:
            cell = 'unknown' if z['correct'] is None else 'TP' if z['correct'] and z['accepted'] else 'FN' if z['correct'] else 'FP' if z['accepted'] else 'TN'
            brief.append({k: z[k] for k in ['split', 'G_round', 'question_id', 'candidate_id', 'candidate_index', 'correct', 'accepted', 'score', 'output_tokens', 'truncated', 'verification_status', 'unknown_reason', 'gold_answer', 'extracted_answer']} | {'confusion_cell': cell, 'response_preview': z['text'][:1200]})
        for label, predicate in [('TP', lambda z: z['correct'] is True and z['accepted']), ('FN', lambda z: z['correct'] is True and not z['accepted']), ('FP', lambda z: z['correct'] is False and z['accepted']), ('TN', lambda z: z['correct'] is False and not z['accepted']), ('unknown', lambda z: z['correct'] is None)]:
            for z in sorted([z for z in rows if predicate(z)], key=lambda z: z['candidate_id'])[:2]:
                assert len(z['question']) < 32767 and len(z['text']) < 32767
                record = {k: z[k] for k in ['split', 'G_round', 'question_id', 'candidate_id', 'question', 'text', 'correct', 'accepted', 'score', 'gold_answer', 'extracted_answer', 'verification_status']}
                record.update({'category': label, 'manual_final_answer_review': '未审查', 'V_reason_review': '未提供/未审查', 'G_process_review': '未审查'}); cases.append(record)
                case_lines += [f'## {key} / {label} / {z["candidate_id"]}', z['question'], '', f'```text\n{z["text"]}\n```', f'来源答案={z["gold_answer"]}；独立标签={z["correct"]}；V接受={z["accepted"]}；代理分数={z["score"]:.6f}；核验状态={z["verification_status"]}。', '']
    (destination/'review_cases.md').write_text('\n'.join(case_lines)+'\n')
    transitions = [{'split': s, 'comparison': f'{r}-G0', **{k: v for k, v in t.items() if k != 'note'}} for s, rs in result['transitions'].items() for r, t in rs.items()]
    definitions = [{'metric': name, 'meaning': meaning} for name, meaning in [
        ('TP/FN/FP/TN', 'Y=1为独立最终答案正确，接受为预测正；未知排除'), ('FNR', 'FN/(TP+FN)'), ('FPR', 'FP/(FP+TN)'),
        ('precision_PPV', 'TP/(TP+FP)'), ('accepted_error_FDR', 'FP/(TP+FP)'), ('TPR', 'TP/(TP+FN)'), ('TNR', 'TN/(TN+FP)'),
        ('NPV', 'TN/(TN+FN)'), ('V_accuracy', '(TP+TN)/known'), ('balanced_accuracy', '(TPR+TNR)/2'), ('F1', '2TP/(2TP+FP+FN)'),
        ('MCC', '(TP*TN-FP*FN)/sqrt((TP+FP)(TP+FN)(TN+FP)(TN+FN))'), ('AUROC', 'P(score_correct>score_wrong)+0.5*P(tie)'),
        ('average_precision', 'sum over tied score thresholds of delta_recall * precision；非PR梯形面积'), ('Brier_proxy', 'mean((score-Y)^2), labelled subset'),
        ('ECE10_proxy', '10个等宽正类代理分数箱，sum n_bin/N * abs(mean_score - positive_fraction)'),
        ('retention_all', 'all accepted/all candidates，含未知'), ('G_accuracy_lower_all', 'correct/all；未知全错的识别下界，不是置信区间'),
        ('G_accuracy_upper_all', '(correct+unknown)/all；未知全对的识别上界，不是置信区间')]]
    workbook(destination/'v0_drift_review.xlsx', {'summary': summaries, 'metrics_ci95': long, 'paired_deltas': paired,
        'length_strata': lengths, 'per_question': per_question, 'complete_question_sensitivity': sensitivity,
        'transitions_candidate0': transitions, 'score_calibration_bins': calibration, 'metric_definitions': definitions,
        'candidates_brief': brief, 'cases': cases}, data)
    from .v0_drift_analysis import csv_write
    csv_write(destination/'score_calibration_bins.csv', calibration); csv_write(destination/'transitions_candidate0.csv', transitions)

    # Chinese narrative, with fixed definitions and explicit inferential limits.
    adjacent = next(d for d in paired if d['split'] == 'B' and d['comparison'] == 'G2-G1' and d['metric'] == 'FNR')
    unknown_delta = next(d for d in paired if d['split'] == 'B' and d['comparison'] == 'G2-G0' and d['metric'] == 'unknown_fraction')
    rows = ['# 固定 V0 跨生成器轮次可靠性诊断（第一版）', '',
        '**当前证据不支持“V0随G更新而明确、单向退化”的结论。** FPR约50%–55%，跨轮变化很小；B上G2的FNR较G0下降，但主要跨轮差值的95%区间仍跨0。更突出的事实是当前V0对自然错误的误接受偏高、实际SFT接受数据仍约一半错误。区间较宽，不能把未检出差异写成“已证明没有漂移”或“等价稳定”。', '',
        '## 数据与固定协议', '',
        '分析B200题、C96道有效题，每题每轮4候选，G0/G1/G2共3552条回答与同数V0评分。A只用于报告实际更新数据质量，不混入B/C判别成绩。C原100题中的4道坏题按题级排除；T保持封存。B是未来验证器更新池，C是开发池，本报告不称为最终独立测试。', '',
        '标签来自冻结的最终答案核验规则，不依赖被测V0；未知不当作错误。V0权重、提示、条件A/B似然评分与0.5阈值不变。B/C已缓存回答逐条重算统计，TP/FN/FP/TN与正式报告一致；不新增推理、不训练、不改阈值。图中的V版本只有V0，不虚构Vrand/Vtarget成绩。', '',
        '## 读图与指标定义', '',
        '正类Y=1为“独立核验最终答案正确”；V接受为预测正。混淆矩阵行是真值正确/错误，列是V接受/拒绝。', '',
        markdown_table(['指标', '定义与解释'], [['误拒率FNR', 'FN/(TP+FN)：正确回答中被V拒绝的比例'], ['误接受率FPR', 'FP/(FP+TN)：错误回答中被V接受的比例'],
            ['接受集错误率FDR', 'FP/(TP+FP)：有标签接受集中错误回答的比例，分母不同于FPR'], ['精确率PPV', 'TP/(TP+FP)：有标签接受集正确率=1-FDR'],
            ['召回/特异度', 'TPR=1-FNR；TNR=1-FPR'], ['平衡准确率', '(TPR+TNR)/2；全部接受基线为0.5'],
            ['AUROC / AP', '排序区分能力；AP为阈值阶梯加权的平均精确率，不是PR曲线梯形面积'], ['F1 / MCC / NPV', '完整表中给出；MCC综合四个混淆计数，NPV为拒绝集内错误答案比例'],
            ['保留率', '接受数/全部候选数，包含未知；可核验子集保留率另列'], ['Brier/ECE10', '把A/B归一化代理分数作二元正类分数的辅助校准诊断，不假设其已是可信概率']]), '',
        '指标定义参照[scikit-learn模型评价文档](https://scikit-learn.org/stable/modules/model_evaluation.html)；分数校准与判别能力是不同对象，参见[概率校准文档](https://scikit-learn.org/stable/modules/calibration.html)与[Guo等，ICML 2017](https://proceedings.mlr.press/v70/guo17a.html)。本次未拟合校准器。', '',
        '## 核心结果与95%区间', '',
        markdown_table(['题池/轮次', '可核验/未知', 'FNR [95%区间]', 'FPR [95%区间]', '平衡准确率', 'AUROC'],
            [[f'{s}/{r}', f'{int(m(r,s)["known"])}/{int(m(r,s)["unknown"])}', ci_text(m(r,s)['FNR'], metrics[f'{r}_{s}']['ci95']['FNR']),
              ci_text(m(r,s)['FPR'], metrics[f'{r}_{s}']['ci95']['FPR']), pct(m(r,s)['balanced_accuracy']), num(m(r,s)['AUROC'])] for s in SPLITS for r in ROUNDS]), '',
        '候选级指标按题聚类重采样2000次：每次抽取200个B题或96个C题，同题4候选一起进入，跨轮使用同一题目重采样索引。报告百分位95%点位区间；不是把所有候选当独立样本。区间仅反映本题池重采样不确定性，不包含重做G训练的随机性、参考答案错误和核验器系统误差。方法参见[bootstrap说明](https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.bootstrap.html)；本实现的重采样单位为整题。', '',
        '![混淆矩阵](figures/01_confusion_counts.png)', '',
        '![行归一化混淆矩阵](figures/02_confusion_row_normalized.png)', '',
        '![跨轮FNR/FPR](figures/03_fnr_fpr_cluster_ci.png)', '',
        '## 相对G0的有符号漂移', '',
        'ΔFNR=FNR(Gt)-FNR(G0)，ΔFPR同理；正数表示该项错误率升高，负数表示降低。这是V0面对不同G回答时的可靠性变化，既不是V0参数发生漂移，也不是输入分布距离。阈值没有逐轮重调。', '',
        markdown_table(['题池', '对比', '指标', '变化/百分点', '成对95%区间/百分点', '方向证据'],
            [[d['split'], d['comparison'], d['metric'], f'{d["delta"]*100:+.2f}', f'[{d["low"]*100:+.2f}, {d["high"]*100:+.2f}]', '点位区间不跨0' if d['CI_excludes_zero'] else '区间跨0']
             for d in paired if d['comparison'] in ['G1-G0', 'G2-G0'] and d['metric'] in ['FNR', 'FPR']]), '',
        '这些是探索性、未做多重比较校正的区间；不据单个边缘结果宣布机制。区间跨0不等于证明相同，未预注册等价界值，因此不做等价性结论。G2-G1及其他指标差值见Excel/CSV。', '',
        f'补充观察：B上G2相对G1的FNR变化为{adjacent["delta"]*100:+.2f}个百分点，其探索性点位区间为[{adjacent["low"]*100:+.2f},{adjacent["high"]*100:+.2f}]个百分点。这是相邻两轮的局部改善信号，未作多重比较校正，也不在C上重复出现；不能把它写成相对G0的确定改善或普遍跨轮规律。', '',
        '![成对漂移区间](figures/04_paired_drift_forest.png)', '',
        '## 接受池的实际风险', '',
        markdown_table(['题池/轮次', 'PPV正确率', 'FDR错误率（有标签接受集）', '全候选保留率', '接受集未知/总数', '全部接受集风险识别界'],
            [[f'{s}/{r}', pct(m(r,s)['precision_PPV']), pct(m(r,s)['accepted_error_FDR']), pct(m(r,s)['retention_all']),
              f'{int(m(r,s)["accepted_unknown"])}/{int(m(r,s)["accepted_total"])}', f'[{pct(m(r,s)["accepted_error_lower_all"])}, {pct(m(r,s)["accepted_error_upper_all"])}]'] for s in SPLITS for r in ROUNDS]), '',
        '风险识别界把接受集未知分别按全对/全错处理，不是抽样置信区间。FPR高表示错答案容易被放行，FDR还取决于正负样本组成；它们不能互相替代。ROC/PR与风险—保留率曲线仅描述冻结分数，不据本轮B/C真值改变实际训练筛选。', '',
        '![ROC](figures/05_roc_curves.png)', '', '![风险与保留率](figures/07_risk_coverage.png)', '',
        '## 排序能力与分数校准', '',
        markdown_table(['题池/轮次', 'AUROC [95%区间]', 'AP', 'Brier代理分数误差', 'ECE10代理分数误差', '平均代理分数/实际正确率'],
            [[f'{s}/{r}', ci_text(m(r,s)['AUROC'], metrics[f'{r}_{s}']['ci95']['AUROC'], percent=False), num(m(r,s)['average_precision']),
              num(m(r,s)['Brier_proxy']), num(m(r,s)['ECE10_proxy']), f'{num(m(r,s)["mean_proxy_score_known"])}/{num(m(r,s)["G_accuracy_verified"])}'] for s in SPLITS for r in ROUNDS]), '',
        'AUROC点估计约0.76–0.80，表明V0分数对正负答案有一定排序区分力；但当前0.5工作点仍放行约一半的错误。平均正类代理分数明显高于真实正确比例，ECE10约0.33–0.39。排序能力、固定工作点的误接受、分数校准应分别解释，不能把例如0.99的A/B归一化分数直接解读成99%的答案正确概率。AUROC相对G0的成对区间也跨0，不能仅凭点估计上升声称排序能力显著增强。', '',
        '![辅助分数校准](figures/09_proxy_score_reliability.png)', '',
        '## 覆盖、长度与基线敏感性', '',
        markdown_table(['题池/轮次', 'G正确率（可核验）', '完整输出正确率识别界', '未知率', '截断数', '输出均长', 'V总体准确率/全部接受基线'],
            [[f'{s}/{r}', pct(m(r,s)['G_accuracy_verified']), f'[{pct(m(r,s)["G_accuracy_lower_all"])}, {pct(m(r,s)["G_accuracy_upper_all"])}]',
              pct(m(r,s)['unknown_fraction']), int(m(r,s)['truncated']), f'{m(r,s)["mean_output_tokens"]:.1f}', f'{pct(m(r,s)["V_accuracy"])}/{pct(m(r,s)["all_accept_baseline_accuracy"])}'] for s in SPLITS for r in ROUNDS]), '',
        f'B上可核验G正确率从{pct(m("G0","B")["G_accuracy_verified"])}到{pct(m("G2","B")["G_accuracy_verified"])}，但完整输出正确率下界从{pct(m("G0","B")["G_accuracy_lower_all"])}到{pct(m("G2","B")["G_accuracy_lower_all"])}，未知从{int(m("G0","B")["unknown"])}条增至{int(m("G2","B")["unknown"])}条。因此不能只据可核验子集百分比宣称G2能力提高。C的G正确率也没有单调上升。', '',
        f'B上G2相对G0的未知比例变化为{unknown_delta["delta"]*100:+.2f}个百分点，成对95%点位区间为[{unknown_delta["low"]*100:+.2f},{unknown_delta["high"]*100:+.2f}]个百分点。覆盖变化值得单独关注；这也是为什么只比较已核验子集的正确率或FNR/FPR可能遗漏重要现象。该辅助比较同样未作多重校正。', '',
        '作为额外敏感性检查，仅保留三个轮次全部4候选均可核验的共同题目，结果如下。该子集按核验完整性筛选，可能偏向易解析题，只作辅助，不能替代主样本。', '',
        markdown_table(['题池/轮次', '共同完整题数', 'FNR', 'FPR'], [[f'{x["split"]}/{x["G_round"]}', x['common_fully_verifiable_questions'], pct(x['FNR']), pct(x['FPR'])] for x in sensitivity]), '',
        '例如C上全样本的G2-G0误拒率点差为-0.25个百分点，而共同完整题子集的点差约+1.23个百分点。方向会随分析子集改变，说明不能忽视标签覆盖与题目组成；这不是未知状态导致变化的因果证明。', '',
        '按固定长度层≤512、513–1024、1025–2048、>2048报告类内错误率与分母。长回答层可能只有极少正确例，零分母记NA而不填0。现有DAPO记录缺少冻结题型/难度及独立解题错误类型标注，本次不事后发明这些分层，不按每轮是否答错定义难度。', '',
        '![覆盖与基线](figures/10_coverage_and_baselines.png)', '', '![长度分层](figures/11_length_stratified_errors.png)', '',
        '## 描述性判断正确性转移', '',
        markdown_table(['题池/对比', 'a旧对新对', 'b旧对新错', 'c旧错新对', 'd旧错新错', '未知配对排除', '保持/退化/修复'],
            [[f'{s}/{r}-G0', t['a'], t['b'], t['c'], t['d'], t['unknown_pairs'], f'{pct(t["retained_judgement_correctness"])}/{pct(t["degraded_judgement_correctness"])}/{pct(t["repaired_judgement_correctness"])}'] for s, rs in result['transitions'].items() for r, t in rs.items()]), '',
        '按方案预定candidate_index=0、同题跨G配对。G回答和独立真值可能改变，上述“退化”只描述判断正确性从对到错，不能叫固定回答上的V遗忘。保持a/(a+b)、退化b/(a+b)、修复c/(c+d)，与数据保留率不同。目前Vrand/Vtarget尚未训练，因此不能交付它们在同批固定回答上的真实转移矩阵。', '',
        '![判断转移](figures/13_descriptive_judgement_transitions.png)', '',
        '## 实际训练质量与解释范围', '']
    for r in ['G1', 'G2']:
        q = json.loads((destination.parent/f'{r}_SFT_data_quality.json').read_text())['actual_SFT_quality']
        rows += [f'{r}实际SFT入选{q["selected"]}条：正确{q["correct"]}、错误{q["wrong"]}、未知{q["unknown"]}；已核验错误率{pct(q["verified_error_fraction"])}。独立标签仅在筛选落盘后统计，没有据真答案移除错例。', '']
    rows += ['本轮只有同一G的一条两次小规模SFT轨迹，未证明生成器稳定增强，因此不能据此检验“强G使V更难验证”的一般性因果命题。现有结果支持的描述是：V0有一定正负区分力，但高误接受是持续问题；这次弱反馈更新带来的主要错误率差值尚不足以确证单向可靠性漂移。后续V更新、GRPO、理由和过程分析需独立执行，不从本报告外推。', '',
        '## 文件索引', '',
        '- `v0_drift_review.xlsx`：主指标、2000次区间、成对差值、长度层、逐题指标、候选摘要和60条抽查案例。',
        '- `v0_drift_figures.pdf`：封面及15张完整图；`figures/`同时保存PNG、矢量PDF、SVG。',
        '- `metrics_summary.csv` / `metrics_ci95.csv` / `paired_deltas.csv`：可直接继续统计的数值表。',
        '- `metrics_and_inference.json`：完整指标、区间、转移与敏感性结果；`curves.json`：曲线数据。',
        '- `candidates_full.jsonl`：3552条B/C候选完整原文、独立标签和V0评分；`review_cases.md`：可读抽查案例。',
        '- `analysis_manifest.json`：分析参数、源文件哈希、实现哈希与原文件未更改的校验。', '',
        '复现命令：`scripts/pilot v0_drift_analysis --bootstrap 2000 --seed 20261022`。本次仅CPU统计与绘图，T不读取/评价。', '']
    report = '\n'.join(rows)+'\n'; (destination/'v0_drift_report.md').write_text(report)
    # Browser entry with tables and all plots; no external JS or network dependencies.
    html_rows = []
    for s in SPLITS:
        for r in ROUNDS:
            p = m(r, s); html_rows.append(f'<tr><td>{s}/{r}</td><td>{pct(p["FNR"])}</td><td>{pct(p["FPR"])}</td><td>{pct(p["accepted_error_FDR"])}</td><td>{num(p["AUROC"])}</td><td>{int(p["unknown"])}/{int(p["total"])}</td></tr>')
    gallery = ''.join(f'<section><h2>{html.escape(title)}</h2><img src="figures/{name}.png" loading="lazy"><p><a href="figures/{name}.pdf">PDF</a> · <a href="figures/{name}.svg">SVG</a></p></section>' for name, title in generated)
    (destination/'index.html').write_text(f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>V0跨轮可靠性诊断</title><style>body{{font:16px system-ui,sans-serif;max-width:1100px;margin:40px auto;padding:0 20px;color:#193245}}h1,h2{{line-height:1.3}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ccd7dd;padding:10px;text-align:right}}section{{margin:40px 0}}img{{max-width:100%;height:auto}}a{{color:#126b99}}p{{line-height:1.8}}</style><h1>固定V0跨G0/G1/G2可靠性诊断</h1><p>第一版：当前数据不支持明确的单向退化。FPR约50%–55%，相对G0的主要FNR/FPR差值区间跨0；不能据此证明等价稳定。B200题、C96题，每题4候选，T封存。</p><p><a href="v0_drift_report.md">完整中文报告</a> · <a href="v0_drift_review.xlsx">Excel分析表</a> · <a href="v0_drift_figures.pdf">全部图表PDF</a> · <a href="review_cases.md">抽查案例</a></p><table><tr><th>题池/G</th><th>误拒FNR</th><th>误接受FPR</th><th>接受集错误FDR</th><th>AUROC</th><th>未知/全部</th></tr>{''.join(html_rows)}</table><p>FNR=正确回答中被拒绝比例；FPR=错误回答中被接受比例；FDR=有标签接受集内错误比例。置信区间按题聚类bootstrap2000次，跨轮同题配对。未知不作为负例，0.5阈值不变。所有细节和限制见中文报告。</p>{gallery}</html>''')
